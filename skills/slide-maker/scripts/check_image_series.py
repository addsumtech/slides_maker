#!/usr/bin/env python3
"""The IMAGE SERIES gate — read from the saved FILE (the `+gen.<slot>` tags) and `series.json`, by both
delivery gates (render_deck --gate-check and codex_delivery_gate.py), so the runtimes cannot disagree.

Blocks: a series recorded with no readable/valid plan; a generated picture with no slot; a series with
NO slot placed through slot_picture; a generated person given a real-looking name, a role/quote, or
team / testimonial / customer framing on its slide with no 'fictional' label (the user's people rule,
2026-10-03 — conservative by design: better a miss than a false block). Notes: an unplaced slot; no
series-qc.json; an unresolved OFF-SERIES outlier.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

LABELS = ("fictional", "illustrative persona", "persona (illustrative)", "composite persona",
          "虚构", "示意人物", "虚构人物", "人物为虚构")
TEAM = re.compile(r"\b(our team|meet the team|founders?|testimonials?|what our (customers|clients) say|"
                  r"customer stor(y|ies)|case study)\b|我们的团队|团队介绍|创始团队|客户评价|用户评价|客户说|"
                  r"证言|用户心声", re.I)
ROLE = re.compile(r"\b(CEO|CTO|COO|CFO|founder|co-founder|head of|director|manager|lead|designer|engineer|"
                  r"researcher|professor|VP)\b", re.I)
# A NAME joined to a ROLE on ONE line by a separator ("Daniel Okafor, Head of Ceramics" / "张伟｜首席设计师").
# A bare Title Case heading ("Community Garden Program") is two capitalised words too — so a name alone,
# or a role word elsewhere on the slide, never blocks. Chinese names need a common surname first, because
# "城市菜园，社区负责人招募" has the same shape as a name + role.
EN_NAME = re.compile(r"\b[A-Z][a-z]{1,15}(?: [A-Z]\.)? [A-Z][a-zA-Z'-]{1,20}\b")
SEP = re.compile(r"^\s*[,|·—–-]")
ZH_SURNAMES = ("王李张刘陈杨黄赵吴周徐孙马朱胡郭何林罗高郑梁谢宋唐许韩冯邓曹彭曾肖田董袁潘蒋蔡余杜叶程苏魏吕"
               "丁任沈姚卢姜崔钟谭陆汪范金石廖贾夏韦方白邹孟熊秦邱江尹薛段雷侯龙史陶黎贺顾毛郝龚邵万钱严武戴莫孔汤")
ZH_NAME_ROLE = re.compile(r"(?<![一-鿿])[" + ZH_SURNAMES + r"][一-鿿]{1,2}\s*[，,｜|·—]\s*[^\n]{0,12}?"
                          r"(首席|总监|经理|负责人|设计师|工程师|研究员|教授|创始人|主任|老师)")
QUOTE_ATTR = re.compile(r"[\"“].{6,}[\"”]\s*[—–-]{1,2}\s*([A-Z][a-z]|[一-鿿])")


def _named(text):
    for line in re.split(r"[\n\x0b]", text):
        for m in EN_NAME.finditer(line):
            rest = line[m.end():]
            if SEP.match(rest) and ROLE.search(rest):
                return True
        if ZH_NAME_ROLE.search(line):
            return True
    return bool(TEAM.search(text)) or bool(QUOTE_ATTR.search(text))


def _near(root, name, depth=3):
    """The newest `name` at most `depth` folders below `root` — the documented layout writes the
    images, their manifest and the QC report to <deck>/assets/generated, two folders down."""
    hits = [p for d in range(depth + 1) for p in root.glob("/".join(["*"] * d + [name]))]
    return max(hits, key=lambda p: p.stat().st_mtime) if hits else None


def recorded_series(gates):
    for sec in ("design_plan", "design"):
        d = (gates or {}).get(sec) or {}
        if isinstance(d, dict) and d.get("imagery") == "series":
            return {"imagery": "series", "plan": d.get("image_series") or None}
    return None


def _slide_texts(slide):
    out = []
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
            out.append(sh.text_frame.text)
    return "\n".join(out)


def check(pptx, rec, deck_dir):
    import deckkit as dk
    import image_series as ims
    from pptx import Presentation
    findings, facts = [], {"slots": 0, "placed": 0, "generated": 0}
    if not rec or rec.get("imagery") != "series":
        return findings, facts
    if not rec.get("plan"):
        return [("block", "SERIES PLAN MISSING", "imagery is 'series' but design_plan.image_series names no "
                 "series.json — record the plan path")], facts
    pp = Path(rec["plan"])
    if not pp.is_absolute():
        pp = Path(deck_dir) / pp
    try:
        plan = ims.load(pp)
    except ValueError as e:
        return [("block", "SERIES PLAN MISSING", str(e))], facts
    probs = ims.check(plan)
    if probs:
        return [("block", "SERIES PLAN INVALID", "; ".join(probs[:4]) + (" …" if len(probs) > 4 else ""))], facts
    slots = {s["id"]: s for s in plan["slots"]}
    facts["slots"] = len(slots)
    placed = {}
    prs = Presentation(str(pptx))
    for n, slide in enumerate(prs.slides, 1):
        gens = []
        for sh in slide.shapes:
            sid = dk.generated_slot(sh)
            if sid is None:
                continue
            facts["generated"] += 1
            if sid not in slots:
                findings.append(("block", "UNPLANNED GENERATED IMAGE", "slide {}: a generated picture tagged "
                                 "{!r} has no slot in series.json — plan it (with its meaning line) or remove it"
                                 .format(n, sid)))
                continue
            placed[sid] = n
            gens.append(slots[sid])
        people = [s for s in gens if s.get("kind") in ("generic-person", "persona")]
        if people:
            txt = _slide_texts(slide)
            low = txt.lower()
            labels = [l.lower() for l in LABELS] + [str(s.get("persona_label", "")).lower() for s in people
                                                    if s.get("persona_label")]
            labelled = any(l and l in low for l in labels)
            named = _named(txt)
            unlabelled_persona = any(s.get("kind") == "persona" for s in people) and not labelled
            if (named and not labelled) or unlabelled_persona:
                findings.append(("block", "GENERATED PERSON NAMED", "slide {}: a GENERATED person ({}) sits "
                                 "beside a name/role/quote or team/testimonial framing with no visible "
                                 "'fictional' label — real people get a real photo; a persona is labelled "
                                 "on its slide".format(n, ", ".join(s["id"] for s in people))))
    facts["placed"] = len(placed)
    if not placed:
        findings.append(("block", "NO SLOT PLACED", "imagery is 'series' but no picture was placed through "
                         "image_series.slot_picture — place each slot with it so the deck carries its tags"))
    for sid in slots:
        if sid not in placed and placed:
            findings.append(("note", "SLOT NOT PLACED", "slot {!r} (slide {}) was planned but not placed"
                             .format(sid, slots[sid]["slide"])))
    qc = _near(pp.parent, "series-qc.json")
    if qc is None:
        man = _near(pp.parent, "image_prompt_manifest.json")
        findings.append(("note", "SERIES QC MISSING", "no series-qc.json — run: python3 scripts/image_series.py "
                         "qc {} --dir {}".format(pp, man.parent if man else "<the folder the images were "
                                                 "generated into>")))
    else:
        rep = json.loads(qc.read_text(encoding="utf-8"))
        ack = rep.get("acknowledged") or {}
        for sid in rep.get("outliers") or []:
            if sid not in ack:
                findings.append(("note", "SERIES QC OUTLIER", "slot {!r} is OFF-SERIES — regenerate it with "
                                 "--style-ref <key>, or record why in series-qc.json 'acknowledged'".format(sid)))
    return findings, facts
