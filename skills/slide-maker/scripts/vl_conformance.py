"""Build-time conformance checks of the visual-language contract, over any registered language (tests/test_vl_contract.py
runs them for all; scripts/check_vl_contract.py for the ones you name). Each returns a list of problems, empty = conforms.
Spec: docs/superpowers/specs/2026-10-15-visual-language-system-design.md §4.3."""
import contextlib
import io
import re
import tempfile
from pathlib import Path

import deckkit as dk
import visual_languages as vl
import vl_native as vn
from vl_contract import PAGES

SKILL = Path(__file__).resolve().parent.parent
REPO = SKILL.parent.parent     # the repo root holds the READMEs; an installed copy of the skill has none
_CJK = re.compile(r"[぀-ヿ㐀-鿿豈-﫿가-힯]")
_WORD = re.compile(r"[A-Za-z][A-Za-z'’\-]*")
# numerals a page may number itself with (page / point / section / sheet / figure numbers): 1-2 digits, CJK numerals
# one to ten, circled one to twenty — a longer number must be the caller's
_NUMERAL = re.compile(r"^(?:\d{1,2}|[一二三四五六七八九十]|[①-⑳])$")


def pages_missing(name, registry=None):
    reg = registry or vl.REGISTRY
    spec = reg[name]
    if spec["kind"] == "image-led":
        have = set(vl.LAYOUTS.get(name, {})) | {"points"}       # points: the shared image-led points page
    else:
        have = set(vn.COMPOSERS.get(name, {}))
    return ["{}: no composition for page {!r}".format(name, p) for p in PAGES if p not in have]


def _strings(o):
    if isinstance(o, str):
        yield o
    elif isinstance(o, dict):
        for v in o.values():
            yield from _strings(v)
    elif isinstance(o, (list, tuple)):
        for v in o:
            yield from _strings(v)


def allowed_tokens(name, caller_texts):
    """Upper-cased Latin words and single CJK characters the caller gave, plus the language's declared derived words."""
    tok = set()
    for t in caller_texts:
        tok |= {w.upper() for w in _WORD.findall(t)}
        tok |= set(_CJK.findall(t))
    for d in vl.REGISTRY[name]["derived"]:
        for words in d["words"].values():
            for w in words:
                tok |= {x.upper() for x in _WORD.findall(w)}
                tok |= set(_CJK.findall(w))
    return tok


def _sample_deck(name):
    """Build the language's sample deck (both grounds) and return (pptx paths, the caller's strings)."""
    td = Path(tempfile.mkdtemp(prefix="vlconf_"))
    paths = []
    for g in vl.VARIANTS[name]:
        with contextlib.redirect_stdout(io.StringIO()):
            paths.append(vl.build_sample(name, td, W=10.0, H=5.625, ground=g))
    if name in vl.NATIVE:
        texts = list(_strings(vl._SAMPLE_COPY_NATIVE[name]))
    else:
        texts = list(_strings(vl._SAMPLE_COPY["watercolour" if name == "storybook" else "photo"]))
    return paths, texts


def invented_words(name):
    from pptx import Presentation
    try:
        paths, texts = _sample_deck(name)
    except NotImplementedError as e:           # a language still being authored: report it, never a traceback
        return ["{}: its sample deck cannot be built yet — {}".format(name, e)]
    tok = allowed_tokens(name, texts)
    probs = []
    for p in paths:
        prs = Presentation(str(p))
        for i, s in enumerate(prs.slides, 1):
            for sh in s.shapes:
                if not sh.has_text_frame or sh.top is None or sh.top < 0:      # the off-page a11y title repeats a field
                    continue
                for w in _WORD.findall(sh.text_frame.text):
                    if w.upper() not in tok and not _NUMERAL.match(w):
                        probs.append("{} slide {}: {!r} is not the caller's and not declared derived".format(p.name, i, w))
                for c in _CJK.findall(sh.text_frame.text):
                    if c not in tok and not _NUMERAL.match(c):
                        probs.append("{} slide {}: {!r} is not the caller's and not declared derived".format(p.name, i, c))
    return sorted(set(probs))


def extras_not_refused(name):
    spec = vl.REGISTRY[name]
    probs = []
    for x, d in spec["extras"].items():
        if d["pages"] == "any":
            continue
        off = [p for p in PAGES if p not in d["pages"] and p != "image_text"][0]
        prs = dk.blank_deck(10.0, 5.625)
        with contextlib.redirect_stdout(io.StringIO()):
            k = vl.use(name, prs)
        fields = {"cover": dict(title="A title"), "section": dict(number="1", title="A section"),
                  "points": dict(title="Points", items=["One", "Two"]), "quote": dict(quote="A quote."),
                  "data": dict(number="3", label="a label"), "closing": dict(title="Thanks")}[off]
        try:
            getattr(k, off)(k.new_slide(), **fields, **{x: "value"})
            probs.append("{}: {}= on {} was accepted (it draws only on {})".format(name, x, off, d["pages"]))
        except NotImplementedError as e:
            probs.append("{}: {}= could not be tried on {} — {}".format(name, x, off, e))
        except (ValueError, TypeError) as e:
            if x not in str(e):
                probs.append("{}: {}= on {} was refused without naming it: {}".format(name, x, off, str(e)[:100]))
    return probs


def assets_missing(name):
    import json
    import sigs
    probs = []
    stems = [vl._sample_stem(name, g) for g in vl.VARIANTS[name]]
    man = json.loads((SKILL / "assets" / "vl" / "samples" / "manifest.json").read_text())
    for st in stems:
        if not (SKILL / "assets" / "vl" / "samples" / (st + ".jpg")).exists():
            probs.append("{}: no bundled sample {}.jpg".format(name, st))
        if st not in json.dumps(man):
            probs.append("{}: sample {} has no fingerprint in manifest.json".format(name, st))
    if not any(('"{}"'.format(name) in src) or ("'{}'".format(name) in src)
               for k, src in sigs.EXAMPLES.items() if k.startswith("vl_")):
        probs.append("{}: no sigs.py --example vl_* scaffold names it".format(name))
    ref = (SKILL / "references" / "visual-languages.md").read_text(encoding="utf-8")
    if "| `{}` |".format(name) not in ref:
        probs.append("{}: references/visual-languages.md has no table row for it".format(name))
    for readme in ("README.md", "README_CN.md"):
        if not (REPO / readme).is_file():       # an installed copy of the skill: no repo README to check
            continue
        if "samples/{}.jpg".format(name) not in (REPO / readme).read_text(encoding="utf-8"):
            probs.append("{}: {} gallery has no cell for it".format(name, readme))
    return probs


def suites_not_registry_driven():
    probs = []
    gen = (SKILL / "tests" / "test_native_generality.py").read_text(encoding="utf-8")
    if "for name in vl.LANGS" not in gen:
        probs.append("tests/test_native_generality.py no longer iterates vl.LANGS — a new language would skip the build matrix")
    aud = (SKILL / "tests" / "test_vl_audit_fixes.py").read_text(encoding="utf-8")
    if "vl.LANGS" not in aud:
        probs.append("tests/test_vl_audit_fixes.py no longer iterates vl.LANGS — a new language would skip the own-pages lint")
    return probs
