#!/usr/bin/env python3
"""vl_native2 — page compositions of the second set of NATIVE visual languages: starlit (星夜), broadsheet (报纸头版),
journal (学术期刊), tally (数据账本) and chalkboard (黑板报).

Same contract as vl_native, which imports this module from its foot (so any importer of vl_native gets these): each
page PLANS its text by measurement, draws its art clear of that text, then sets the text. Nothing here invents a word:
masthead / edition / inside / tags / running / authors / abstract / margin / total / doodle / ordered come from the
caller or nothing is drawn. Page numbers, figure numbers, the ■ end mark and the structural labels (Abstract, Inside,
Figure, Page, Note — in the page's own script) are derived, never typed in.
"""
from __future__ import annotations

import copy
import math
import re
from pathlib import Path

import deckkit as dk
import native_art as na
import visual_languages as vl
from vl_native import GROUNDS, _run_all, alt, ctx, display, fit_circle, points_of, register, text_of  # noqa: F401
import vl_native as _vn

HEADROOM = vl.HEADROOM   # words are measured in 97% of a column and drawn at its full width (visual_languages._flow)


_CLAUSE_FIELDS = ("title", "subtitle", "quote", "label", "line", "note", "body", "caption", "item_head", "item_line",
                  "abstract", "margin", "authors")


def _cjk_width(k, f_, t, sz):
    role = vl.TYPE[k.name][f_][1]
    return dk._natural_width_in([(t, vl._weight(k, f_, t))], sz, k.ea_face(role, t) or k.face(role))


def _clause_break(k, f_, t, w, start):
    """Chinese or Japanese that will wrap is packed CLAUSE by clause — each line as many whole clauses as fit — so it
    breaks only at its clause marks: "来访者握着螺丝刀，/ 志愿者只在旁边指导，/ 这就是…" (a square board's quote read
    "…只在旁边指 / 导，…") and never leaves one character alone ("…二氧化 / 碳", sample render 2026-10-08) — at the
    largest size, down to the field's floor, at which every clause fits a line (clean breaks beat a bigger size). Returns
    (text, size or None); a clause too long even then, or text with no clause mark, is left for the flow to wrap."""
    if f_ not in _CLAUSE_FIELDS or not t or "\n" in t or not dk._has_cjk(t):
        return t, None
    clauses = _vn._clauses(t)
    if len(clauses) < 2:
        return t, None
    base, _role, _b, _c, _i, floor = vl.TYPE[k.name][f_]
    sz = (start or {}).get(f_, base * ctx(k)[2])
    room = w - dk.TEXT_INSET_LR
    if _cjk_width(k, f_, t, sz) <= room:
        return t, None                             # one line: nothing to break
    widest = max(_cjk_width(k, f_, c, sz) for c in clauses)
    z = sz if widest <= room else sz * room / widest * 0.98
    if z < floor * ctx(k)[2]:
        return t, None
    lines = [""]
    for c in clauses:
        if lines[-1] and _cjk_width(k, f_, lines[-1] + c, z) > room:
            lines.append(c)
        else:
            lines[-1] += c
    return ("\n".join(lines), (z if z < sz else None)) if len(lines) > 1 else (t, None)


def _one_line_size(k, f_, t, w, start):
    """A short Chinese/Japanese head (no clause mark) a little too wide for its line is set a little smaller, on ONE
    line — wrapped, it left one character alone ("记下每次修 / 理"). None when it fits, or would need below 75% of its size
    or its floor (then it wraps, and a layout that cares refuses — chalkboard's row gives way to its grid)."""
    if f_ not in ("item_head", "label") or not t or "\n" in t or not dk._has_cjk(t) or len(t) > 14:
        return None
    base, _role, _b, _c, _i, floor = vl.TYPE[k.name][f_]
    s_ = ctx(k)[2]
    sz = (start or {}).get(f_, base * s_)
    room = w - dk.TEXT_INSET_LR
    nat = _cjk_width(k, f_, t, sz)
    if nat <= room:
        return None
    fit = sz * room / nat * 0.98
    return fit if fit >= max(0.75 * sz, floor * s_) else None


def flow(k, slide, page, col, items, **kw):
    """vl_native.flow with headroom: the words are MEASURED in 97% of the column and the boxes DRAWN at its full width.
    A renderer can set a line a hair wider than the measure — LibreOffice set "Three lines on the ledger" (11.39in of
    Arial Black, measured to fit 11.41in) on two lines, drawn over the first ledger row, and lint, reading the same
    measure, saw nothing. The rects report the drawn boxes; no box grows past its column."""
    x, y, w, h = col
    align = kw.get("align", "l")
    wm = w * HEADROOM                      # what the core measures in; it draws the boxes at w itself
    packed = []
    for f_, t in items:
        t2, z = _clause_break(k, f_, t, wm, kw.get("start"))
        packed.append((f_, t2))
        if z:
            kw = dict(kw, start=dict(kw.get("start") or {}, **{f_: z}))
    items = packed
    for f_, t in items:
        z = _one_line_size(k, f_, t, wm, kw.get("start"))
        if z:
            kw = dict(kw, start=dict(kw.get("start") or {}, **{f_: z}))
    rects, draw = _vn.flow(k, slide, page, col, items, **kw)
    out = dict(rects)
    dw = getattr(draw, "headroom", 0.0)

    texts_ = {t: f_ for f_, t in items}

    def go():
        from pptx.util import Emu
        n0 = len(slide.shapes)
        draw()
        for sh in list(slide.shapes)[n0:]:
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                # display text keeps the lines it was MEASURED in: drawn wider (the headroom below), a renderer re-wrapped
                # "Mending a lamp, / one volunteer at a / time" into a lone last word (the docs-only run, 2026-10-09)
                f_ = texts_.get(sh.text_frame.text)
                ex = getattr(draw, "exact", {}).get(f_)
                if ex:                       # set at the decided size rounded DOWN, never up into one more line
                    import math as _m
                    from pptx.util import Pt as _Pt
                    down = _m.floor(ex * 10) / 10.0
                    for p_ in sh.text_frame.paragraphs:
                        for r_ in p_.runs:
                            if r_.font.size is not None and r_.font.size.pt > down + 1e-6 and abs(r_.font.size.pt - round(ex, 1)) < 0.06:
                                r_.font.size = _Pt(down)
                role_ = vl.TYPE[k.name][f_][1] if f_ in vl.TYPE[k.name] else "body"
                face_ = (k.ea_face(role_, sh.text_frame.text) if dk._has_cjk(sh.text_frame.text) else None) or k.face(role_)
                # Chinese only with its real face: under a stand-in with no CJK glyphs (a Linux box without SimSun) the
                # imposed lines disagree with the lint's line model and read as overflow (the font simulation, 2026-10-09)
                if (f_ in _DISPLAY_FIELDS and len(sh.text_frame.paragraphs) == 1 and sh.text_frame.paragraphs[0].runs
                        and not (dk._has_cjk(sh.text_frame.text) and dk._font_substituted(face_))):
                    z = sh.text_frame.paragraphs[0].runs[0].font.size
                    if z is not None:
                        ls = vl._break_lines(k, f_, sh.text_frame.text, z.pt, sh.width / 914400.0 - dw)   # as measured
                        if len(ls) > 1:
                            _split_paragraph(sh.text_frame, ls)
    go.sizes = getattr(draw, "sizes", {})
    go.exact = getattr(draw, "exact", {})
    return out, go

# Japanese and Korean entries: 要旨 / 초록, 図 n, n面 / n면 confirmed against university thesis guides and newspaper
# pages (2026-10-08); 今号 / 이번 호 could not be confirmed, so the plain 目次 / 목차 (contents) and 주석 (annotation) stand.
LABELS = {
    "abstract": {"en": "Abstract", "zh": "摘要", "ja": "要旨", "ko": "초록"},
    "inside": {"en": "Inside", "zh": "本期", "ja": "目次", "ko": "목차"},
    "figure": {"en": "Figure {}", "zh": "图 {}", "ja": "図 {}", "ko": "그림 {}"},
    "page": {"en": "Page {}", "zh": "第 {} 版", "ja": "{} 面", "ko": "{}면"},
    "note": {"en": "Note", "zh": "注", "ja": "注", "ko": "주석"},
}


def lang_of(*texts, k=None, f=None):
    """The language a structural label is set in: the script of the label's own words (Hangul, kana, Han), else en.
    Kanji alone is Chinese OR Japanese, so a Han verdict asks the rest of the page (`f`) and the deck (`k`: its
    remembered words and every slide's text): kana or Hangul there wins (final review, 2026-10-09: a Japanese
    cover's sidebar read 本期, a figure 图 1)."""
    scr = dk.script_of("".join(t for t in texts if t))
    if scr == "han" and (k is not None or f):
        more = [v for v in (f or {}).values() if isinstance(v, str)]
        if k is not None:
            more += [v for v in k.memo.values() if isinstance(v, str)]
            more += [sh.text_frame.text for sl in k.prs.slides for sh in sl.shapes if getattr(sh, "has_text_frame", False)]
        wider = dk.script_of("".join(more))
        if wider in ("kana", "hangul"):
            scr = wider
    return {"han": "zh", "kana": "ja", "hangul": "ko"}.get(scr, "en")


def label(kind, *texts, k=None, f=None):
    return LABELS[kind][lang_of(*texts, k=k, f=f)]


def caps(t):
    """Latin set in capitals (kickers, meta lines); CJK has no case and is left as written."""
    return t.upper() if t and not dk._has_cjk(t) else t


def fields(f, names, capsed=()):
    """[(field, words)] for the fields of `names` the caller gave, in order; those in `capsed` set in capitals."""
    return [(n, caps(text_of(f, n)) if n in capsed else text_of(f, n)) for n in names if text_of(f, n)]


def plan_together(k, slide, page, specs):
    """Plan a page's like items — specs = [(rect, items, flow_kw)] — at ONE size per field: each alone first, then all
    again at the smallest size any of them needed. Planned alone, a Chinese chalk box whose body wrapped was set at
    14.8pt beside neighbours at 19pt (the docs-only run, 2026-10-09). Returns [(rects, draw)] in order."""
    first = [flow(k, slide, page, rect, items, **kw) for rect, items, kw in specs]
    shared = {}
    for _r, d in first:
        for f_, z in getattr(d, "sizes", {}).items():
            shared[f_] = min(shared.get(f_, z), z)
    out = []
    for rect, items, kw in specs:
        st = dict(kw.get("start") or {}, **{f_: shared[f_] for f_, _t in items if f_ in shared})
        out.append(flow(k, slide, page, rect, items, **dict(kw, start=st)))
    return out


_DISPLAY_FIELDS = ("title", "subtitle", "quote", "label", "line")


def _split_paragraph(tf, lines):
    """Rewrite a one-paragraph text frame as one paragraph per line in `lines` (its own words, in order), keeping each
    run's formatting and the line spacing. The lines are ONE block, so no paragraph spacing between them: dk.text's
    6pt space-after, copied to every line, grew a 5-line title 0.33in past its measured box (final review,
    2026-10-09). Returns False — frame untouched — when the lines do not tile the text."""
    from pptx.oxml.ns import qn as _qn
    from pptx.text.text import _Paragraph
    from pptx.util import Pt as _Pt
    p0 = tf.paragraphs[0]._p
    full = "".join(r.text for r in tf.paragraphs[0].runs)
    pos, spans = 0, []
    for ln in lines:
        ln = ln.strip()
        i = full.find(ln, pos) if ln else -1
        if i < 0 or full[pos:i].strip():
            return False
        spans.append((i, i + len(ln)))
        pos = i + len(ln)
    if full[pos:].strip():
        return False
    for a, b in spans:
        q = copy.deepcopy(p0)
        off = 0
        for r in list(q.iter(_qn("a:r"))):
            t = r.find(_qn("a:t"))
            txt = (t.text or "") if t is not None else ""
            ra, rb = off, off + len(txt)
            off = rb
            lo, hi = max(a, ra), min(b, rb)
            if lo >= hi:
                r.getparent().remove(r)
            else:
                t.text = txt[lo - ra:hi - ra]
        p0.addprevious(q)
        para = _Paragraph(q, tf)
        para.space_before, para.space_after = _Pt(0), _Pt(0)
    p0.getparent().remove(p0)
    return True


def stack(k, slide, page, x, w, y0, y1, groups, *, anchor="middle", align="l", **flow_kw):
    """Plan GROUPS of fields one under another — groups = [(items, gap_after_in), …], an empty group skipped — inside
    [y0, y1], the whole stack anchored top or middle. Two passes: measured at the top, then placed. The gaps give an
    underline or a rule its own room (flow's own gaps are for type). Returns (rects, draws)."""
    groups = [(items, gap) for items, gap in groups if items]
    if len(groups) > 1:
        # size every group TOGETHER first (in the room left after the gaps), then start each group at that size: planned
        # one by one, a long first group kept its full size and left the next none (a square board's closing, 2026-10-08)
        gaps = sum(g for _i, g in groups[:-1])
        _r0, d0 = flow(k, slide, page, (x, y0, w, max(0.1, y1 - y0 - gaps)), [it for items, _g in groups for it in items],
                       anchor="top", align=align, **flow_kw)
        flow_kw = dict(flow_kw, start=dict(flow_kw.get("start") or {}, **getattr(d0, "sizes", {})))

    def plan(top):
        rects, draws, y = {}, [], top
        for i, (items, gap) in enumerate(groups):
            r, d = flow(k, slide, page, (x, y, w, y1 - y), items, anchor="top", align=align, **flow_kw)
            rects.update(r)
            draws.append(d)
            y = max((v[1] + v[3] for v in r.values()), default=y) + (gap if i < len(groups) - 1 else 0.0)
        return rects, draws, y
    rects, draws, foot = plan(y0)
    if anchor == "middle" and groups:
        rects, draws, foot = plan(y0 + max(0.0, (y1 - foot) / 2.0))
    return rects, draws


def meet(a, b, pad=0.0):
    return a[0] < b[0] + b[2] + pad and b[0] < a[0] + a[2] + pad and a[1] < b[1] + b[3] + pad and b[1] < a[1] + a[3] + pad


def memo(k, f, name, fallback=None):
    """The caller's words for an extra that holds for the whole deck (masthead=, edition=, running=). Given on this
    page, they are remembered by REASSIGNING k.memo — compose() restores k.__dict__ shallowly when a layout fails, so
    an in-place update would survive the failed attempt. Else the remembered words, else `fallback`."""
    v = text_of(f, name)
    if v:
        k.memo = dict(k.memo, **{name: v})
    return k.memo.get(name) or fallback


def page_no(slide):
    prs = slide.part.package.presentation_part.presentation
    return [s.slide_id for s in prs.slides].index(slide.slide_id) + 1


def tags_of(f, n, lang):
    """tags=: one short label per point, from the caller. Absent → None (no tags drawn)."""
    v = f.get("tags")
    if v is None:
        return None
    if not isinstance(v, (list, tuple)) or len(v) != n or not all(str(t or "").strip() for t in v):
        raise ValueError("{}.points(): tags= takes one non-empty label per point ({} points), got {!r}".format(lang, n, v))
    return [str(t).strip() for t in v]


def inside_lines(v):
    """inside=: 1 to 4 short lines for a broadsheet cover's 'Inside' sidebar. Absent → None."""
    if v is None:
        return None
    if not isinstance(v, (list, tuple)) or not 1 <= len(v) <= 4 or not all(str(t or "").strip() for t in v):
        raise ValueError("broadsheet.cover(): inside= takes 1 to 4 non-empty lines, got {!r}".format(v))
    return [str(t).strip() for t in v]


def max_extra(size_pt, s, factor=2.3):
    """The most height (in) a raised initial can add to a body field whose base size is `size_pt` at scale `s`."""
    return (factor - 1.0) * size_pt * s * 1.2 / 72.0


def raise_initial(tb, factor=2.3, face=None):
    """Enlarge the first letter of `tb`'s first paragraph IN its line — a raised initial. pptx has no drop cap, and a
    two-box fake re-wraps differently in PowerPoint, Keynote and LibreOffice (the look-dev's "W" ran into its own
    word). Latin letters only; returns the height (in) the taller first line adds, which the box grows by."""
    from pptx.text.text import _Run
    from pptx.util import Emu, Pt
    p = tb.text_frame.paragraphs[0]
    if not p.runs:
        return 0.0
    r0 = p.runs[0]
    t = r0.text
    if not t or not t[0].isalpha() or dk._has_cjk(t[0]) or len(t) < 2:
        return 0.0
    size = r0.font.size.pt
    el = copy.deepcopy(r0._r)
    r0._r.addprevious(el)
    first = _Run(el, p)
    first.text, r0.text = t[0], t[1:]
    first.font.size, first.font.bold = Pt(size * factor), True
    if face:
        first.font.name = face
    added = (factor - 1.0) * size * 1.2 / 72.0
    tb.height = Emu(int(tb.height + added * 914400))
    dk._compose_tag(tb, flag="+initial")       # lint measures the first line apart (deckkit._ink_rect)
    return added


def end_mark(tb, color, face="Arial"):
    """Set the end-of-article ■ (planned into the words as " ■", so the measured height holds it) in its own
    run, in the accent. Returns whether a mark was found."""
    from pptx.text.text import _Run
    p = tb.text_frame.paragraphs[-1]
    if not p.runs or not p.runs[-1].text.endswith("■"):
        return False
    r = p.runs[-1]
    r.text = r.text[:-1]
    el = copy.deepcopy(r._r)
    r._r.addnext(el)
    m = _Run(el, p)
    m.text = "■"
    m.font.color.rgb = dk.RGBColor.from_string(color)
    m.font.name = face
    return True


def kit_for(name, slide):
    """A kit for an ORDINARY page: rs.ground() receives only the slide, so the ground use() set and the deck's
    remembered extras (vl._MEMO) are read back here."""
    prs = slide.part.package.presentation_part.presentation
    k = vl.Kit(name, prs, "both", None, None, None, vl._ACTIVE.get(name, "light"))
    k.memo = dict(vl._MEMO.get(name, {}))
    return k


def last_shape(slide, n0, text):
    """The text box drawn after the first n0 shapes whose words are `text` (to finish a raised initial or end mark)."""
    for sh in list(slide.shapes)[n0:][::-1]:
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.replace("\n", "").startswith(text[:12]):
            return sh
    return None


# ═══════════════════════════════════ starlit 星夜 ═══════════════════════════════════
def _st_sky(k, slide, keep, seed):
    na.starfield_png(slide, base=k.P["ground"], ink=k.P["ink"], glow=k.P["glow"], seed=seed, keep_clear=list(keep))


def _st_ground(k, slide):
    _st_sky(k, slide, (), seed=len(k.prs.slides))       # an ordinary page: the sky; its words are not known yet


GROUNDS["starlit"] = _st_ground


def _st_rule(k, slide, rects, a, b):
    """A short gold rule in the gap between field `a` and field `b` (stack() leaves that gap)."""
    if a not in rects or b not in rects:
        return lambda: None
    W, H, s, o = ctx(k)
    y = (rects[a][1] + rects[a][3] + rects[b][1]) / 2.0
    return lambda: na.seg(slide, W / 2 - 0.55 * s, y, W / 2 + 0.55 * s, y, k.P["text_accents"][0], w=1.0)


def _st_moon(k, slide, clear):
    """The crescent at the first corner clear of the words — top right, top left, bottom right; none when no corner
    is clear (it is the sky's, not the page's)."""
    W, H, s, o = ctx(k)
    d = 0.9 * s
    for cx, cy in ((0.86 * W, 0.16 * H), (0.14 * W, 0.16 * H), (0.86 * W, 0.80 * H)):
        r = (cx - 0.5 * d, cy - 0.6 * d, 1.21 * d, 1.1 * d)
        if not any(meet(r, c, 0.2 * s) for c in clear):
            na.crescent(slide, cx, cy, d, k.P["text_accents"][0], k.P["ground"])
            return


@register("starlit", "cover", alts=2)
def _st_cover(k, slide, f, image):
    W, H, s, o = ctx(k)
    x, w = (0.12 * W, 0.76 * W) if o == "land" else (0.08 * W, 0.84 * W)
    y0, y1 = (0.20 * H, 0.80 * H) if alt() == 0 else (0.07 * H, 0.92 * H)
    r, ds = stack(k, slide, "cover", x, w, y0, y1,
                  [(fields(f, ("kicker", "title"), ("kicker",)), 0.40 * s), (fields(f, ("subtitle",)), 0.0)], align="c")
    _st_sky(k, slide, r.values(), seed=1)
    _st_moon(k, slide, list(r.values()))
    _run_all([_st_rule(k, slide, r, "title", "subtitle")] + ds)
    return r


@register("starlit", "section", alts=2)
def _st_section(k, slide, f, image):
    W, H, s, o = ctx(k)
    num = text_of(f, "number")
    cx, cy = 0.5 * W, (0.32 if o == "land" else 0.26) * H
    d0 = fit_circle(k, cx, cy, (min(0.46 * H, 0.40 * W) if o == "land" else 0.62 * W) * (1.0 if alt() == 0 else 0.72))
    rects, draws = {}, []
    if num:
        r, d = flow(k, slide, "section", (cx - d0 * 0.42, cy - d0 * 0.3, d0 * 0.84, d0 * 0.6), [("number", num)],
                    anchor="middle", align="c", start={"number": 96 * s})
        rects.update(r); draws.append(d)
    top = cy + d0 / 2 + 0.1 * s if num else 0.12 * H
    r, d = flow(k, slide, "section", (0.10 * W, top, 0.80 * W, 0.92 * H - top),
                fields(f, ("kicker", "title"), ("kicker",)), anchor="top" if num else "middle", align="c")
    rects.update(r); draws.append(d)
    _st_sky(k, slide, rects.values(), seed=2)
    if num:
        na.radial_glow(slide, cx, cy, d0, k.P["glow"])
    _run_all(draws)
    return rects


@register("starlit", "image_text", alts=2)
def _st_image_text(k, slide, f, image):
    """The caller's picture in a round window with a gold rim, the words beside it (under it in portrait)."""
    W, H, s, o = ctx(k)
    if o == "land":
        d = min(0.70 * H, 0.40 * W) * (1.0 if alt() == 0 else 0.8)
        cx, cy = 0.07 * W + 0.09 * s + d / 2, 0.5 * H
        x0 = cx + d / 2 + 0.6 * s
        col = (x0, 0.10 * H, 0.93 * W - x0, 0.80 * H)
    else:
        d = min(0.80 * W, 0.40 * H) * (1.0 if alt() == 0 else 0.75)
        cx, cy = 0.5 * W, 0.06 * H + 0.09 * s + d / 2
        y0 = cy + d / 2 + 0.4 * s
        col = (0.08 * W, y0, 0.84 * W, 0.92 * H - y0)
    r, dr = flow(k, slide, "image_text", col, fields(f, ("kicker", "title", "body", "caption"), ("kicker",)),
                 anchor="middle" if o == "land" else "top")
    rim = d + 0.18 * s
    _st_sky(k, slide, list(r.values()) + [(cx - rim / 2, cy - rim / 2, rim, rim)], seed=3)
    vl._place_image(k, slide, image, (cx - d / 2, cy - d / 2, d, d),
                    vl.L_((0, 0, 1, 1), "frame", None, frame="ellipse"), "image_text")
    na.ring(slide, cx, cy, rim, k.P["text_accents"][0], w=1.25)
    dr()
    return r


@register("starlit", "quote", alts=2)
def _st_quote(k, slide, f, image):
    W, H, s, o = ctx(k)
    x, w = (0.20 * W, 0.60 * W) if (o == "land" and alt() == 0) else (0.08 * W, 0.84 * W)
    r, ds = stack(k, slide, "quote", x, w, 0.10 * H, 0.90 * H,
                  [([("mark", "“")] + fields(f, ("quote",)), 0.40 * s), (fields(f, ("attribution",), ("attribution",)), 0.0)],
                  align="c")
    _st_sky(k, slide, r.values(), seed=4)
    _run_all([_st_rule(k, slide, r, "quote", "attribution")] + ds)
    return r


@register("starlit", "data", alts=2)
def _st_data(k, slide, f, image):
    W, H, s, o = ctx(k)
    num = text_of(f, "number")
    cx, cy = 0.5 * W, (0.36 if o == "land" else 0.28) * H
    d0 = fit_circle(k, cx, cy, (min(0.60 * H, 0.46 * W) if o == "land" else 0.80 * W) * (1.0 if alt() == 0 else 0.75))
    rects, draws = {}, []
    if num:
        r, d = flow(k, slide, "data", (0.06 * W, cy - d0 * 0.32, 0.88 * W, d0 * 0.64), [("number", num)],
                    anchor="middle", align="c")
        rects.update(r); draws.append(d)
    top = cy + d0 / 2 + 0.05 * s if num else 0.30 * H
    r, d = flow(k, slide, "data", (0.10 * W, top, 0.80 * W, 0.90 * H - top), fields(f, ("label", "note")),
                anchor="top", align="c")
    rects.update(r); draws.append(d)
    _st_sky(k, slide, rects.values(), seed=5)
    if num:
        na.radial_glow(slide, cx, cy, d0, k.P["glow"])
    na.horizon_glow(slide, 0.86 * H, k.P["glow"], k.P["text_accents"][0])
    _run_all(draws)
    return rects


@register("starlit", "closing", alts=2)
def _st_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    x, w = (0.12 * W, 0.76 * W) if alt() == 0 else (0.07 * W, 0.86 * W)
    r, ds = stack(k, slide, "closing", x, w, (0.22 if alt() == 0 else 0.08) * H, 0.80 * H,
                  [(fields(f, ("title",)), 0.40 * s), (fields(f, ("line",)), 0.0)], align="c")
    _st_sky(k, slide, r.values(), seed=6)
    _st_moon(k, slide, list(r.values()))
    na.horizon_glow(slide, 0.86 * H, k.P["glow"], k.P["text_accents"][0])
    _run_all([_st_rule(k, slide, r, "title", "line")] + ds)
    return r


@register("starlit", "points", alts=3)
def _st_points(k, slide, f, image):
    """A constellation: one star per point, joined in order by one gold line. Across a landscape page the stars step
    high, low, high with their words under them; in portrait (and for long copy) they climb a zig-zag at the left
    with their words beside them; the last resort is a 2x2 ring. The line never reaches the words: from a high star it
    drops 0.75s over a whole column span, so at the column's edge it is still above y + 0.375s, and the words start at
    y + 0.42s."""
    W, H, s, o = ctx(k)
    pts = points_of(f.get("items"))
    n = len(pts)
    hh = ((0.20 if o == "land" else 0.16) if alt() == 0 else 0.32) * H      # long copy: the title gets more room
    r, d = flow(k, slide, "points", (0.07 * W, 0.07 * H, 0.86 * W, hh), fields(f, ("kicker", "title"), ("kicker",)))
    rects, draws, clear = dict(r), [d], list(r.values())
    top = max((v[1] + v[3] for v in r.values()), default=0.07 * H) + 0.3 * s
    bottom = H - 0.6

    def layout(t0):
        stars, cols = [], []
        if o == "land" and alt() == 0:
            span = 0.86 * W / n
            for i in range(n):
                x, y = 0.07 * W + span * (i + 0.5), t0 + (0.45 if i % 2 == 0 else 1.20) * s
                stars.append((x, y))
                cols.append((x - span / 2 + 0.12 * s, y + 0.42 * s, span - 0.24 * s, 0.92 * H - (y + 0.42 * s)))
            align = "c"
        elif alt() == 2 and n >= 3:
            # long copy on a short page: a 2x2 constellation read as a ring (1 → 2 across, down, 3 → 4 back), each star
            # at its cell's top-left, its words to the right of the star and under the row's line — so a line never
            # meets words: across-lines run above every label, the down-line runs left of the second label
            rows_ = -(-n // 2)
            cw_ = 0.86 * W / 2.0
            rh = (0.92 * H - t0) / rows_
            for c_, r_ in [(0, 0), (1, 0), (1, 1), (0, 1)][:n]:
                x, y = 0.07 * W + c_ * cw_ + 0.25 * s, t0 + r_ * rh + 0.25 * s
                stars.append((x, y))
                cols.append((x + 0.30 * s, y + 0.42 * s, cw_ - 0.75 * s, rh - 0.52 * s))
            align = "l"
        else:
            slot = (0.92 * H - t0) / n
            lx = 0.30 * W
            for i in range(n):
                y = t0 + slot * (i + 0.5)
                stars.append(((0.13 if i % 2 == 0 else 0.22) * W, y))
                cols.append((lx, y - min(0.30 * s, slot / 2), 0.92 * W - lx, slot - 0.06 * s))
            align = "l"
        planned = plan_together(k, slide, "points", [(col, [(x_, t) for x_, t in (("item_head", hd), ("item_line", ln)) if t],
                                                      {"anchor": "top", "align": align}) for (hd, ln), col in zip(pts, cols)])
        return stars, planned
    stars, planned = layout(top)
    if o == "land" and alt() == 0:          # the constellation and its words sit in the middle of the room under the title
        foot = max(v[1] + v[3] for tr, _d in planned for v in tr.values())
        shift = max(0.0, (bottom - foot) / 2.0)
        if shift > 0.05:
            stars, planned = layout(top + shift)
    for tr, td_ in planned:
        clear += list(tr.values())
        draws.append(td_)
    _st_sky(k, slide, clear + [(x - 0.35 * s, y - 0.35 * s, 0.7 * s, 0.7 * s) for x, y in stars], seed=7)
    for a, b in zip(stars, stars[1:]):
        na.seg(slide, a[0], a[1], b[0], b[1], k.P["text_accents"][0], w=0.9, alpha=0.55)
    for x, y in stars:
        na.radial_glow(slide, x, y, 1.0 * s, k.P["glow"])
        dk.decorative(na.disc(slide, x, y, 0.16 * s, k.P["glow"]), "a star of the constellation; its words sit beside it")
    _run_all(draws)
    return rects


# ═══════════════════════════════════ broadsheet 报纸头版 ═══════════════════════════════════
def bs_strip(k, slide, f, big=False):
    """The masthead strip every page carries: the paper's name (masthead=, else the cover's kicker, else none), a
    heavy rule, the caller's edition= at the left and the page number at the right, a hairline. Never invents a name,
    a date or an issue number. Draws at once and returns the y under it."""
    W, H, s, o = ctx(k)
    name = memo(k, f, "masthead", fallback=k.memo.get("_cover_kicker"))
    edition = memo(k, f, "edition")
    y = 0.04 * H
    draws = []
    if name:
        nh = (0.80 if big else 0.52) * s
        r, d = flow(k, slide, "masthead", (0.05 * W, y, 0.90 * W, nh), [("masthead", name)], anchor="middle", align="c",
                    start={"masthead": (44 if big else 26) * s})
        draws.append(d)
        y = max(y + nh, max(v[1] + v[3] for v in r.values())) + 0.04 * s
    rule_y, meta_y = y, y + 0.10 * s
    draws.append(lambda: na.seg(slide, 0.05 * W, rule_y, 0.95 * W, rule_y, k.P["ink"], w=2.5))
    if edition:
        r, d = flow(k, slide, "edition", (0.05 * W, meta_y, 0.62 * W, 0.30 * s), [("edition", caps(edition))])
        draws.append(d)
    pg = caps(label("page", name, edition, text_of(f, "title"), text_of(f, "quote"), k.memo.get("_title"), k=k, f=f).format(page_no(slide)))
    sz = vl.TYPE[k.name]["edition"][0] * s
    draws.append(lambda: dk.text(slide, 0.70 * W, meta_y, 0.25 * W, 0.30 * s,
                                 [k.runs(pg, sz, k.color("mute"), True, "meta")], align=dk.PP_ALIGN.RIGHT))
    hair = meta_y + 0.34 * s
    draws.append(lambda: na.seg(slide, 0.05 * W, hair, 0.95 * W, hair, k.P["ink"], w=0.6))
    _run_all(draws)
    return hair + 0.10 * s


def _with_initial(k, slide, draw, words):
    """Run a planned draw, then open the body it set on a raised initial (planned with room for it: max_extra)."""
    def go():
        n0 = len(slide.shapes)
        draw()
        tb = last_shape(slide, n0, words)
        if tb is not None:
            raise_initial(tb, face=k.face("display"))
    return go


@register("broadsheet", "cover", alts=2)
def _bs_cover(k, slide, f, image):
    """A front page: the headline across the page, a rule, the standfirst; the caller's inside= lines in a sidebar
    (under the standfirst in portrait). With no masthead= the cover's kicker names the paper and is not repeated."""
    W, H, s, o = ctx(k)
    kick, title = text_of(f, "kicker"), text_of(f, "title")
    as_name = bool(kick) and not text_of(f, "masthead") and not k.memo.get("masthead")
    if title:
        k.memo = dict(k.memo, _title=title)
    if as_name:
        k.memo = dict(k.memo, _cover_kicker=kick)
    top = bs_strip(k, slide, f, big=True) + 0.12 * s
    lines = inside_lines(f.get("inside"))
    head = ([("kicker", caps(kick))] if kick and not as_name else []) + ([("title", title)] if title else [])
    hh = ((0.50 if o == "land" else 0.36) if alt() == 0 else (0.62 if o == "land" else 0.48)) * H
    r, d = flow(k, slide, "cover", (0.05 * W, top, 0.90 * W, hh), head, anchor="top")
    rects, draws = dict(r), [d]
    yr = max((v[1] + v[3] for v in r.values()), default=top) + 0.16 * s
    art = [lambda: na.seg(slide, 0.05 * W, yr, 0.95 * W, yr, k.P["ink"], w=0.75)]
    side = bool(lines) and o == "land"
    sub, sb, foot = text_of(f, "subtitle"), yr, H - 0.6
    if sub:
        r, d = flow(k, slide, "cover", (0.05 * W, yr + 0.16 * s, (0.62 if side else 0.90) * W, foot - yr - 0.16 * s),
                    [("subtitle", sub)], anchor="top")
        rects.update(r); draws.append(d)
        sb = max(v[1] + v[3] for v in r.values())
    if lines:
        if side:
            ix, iy, iw = 0.71 * W, yr + 0.16 * s, 0.24 * W
            art.append(lambda: na.seg(slide, 0.685 * W, yr + 0.16 * s, 0.685 * W, foot, k.P["ink"], w=0.5))
        else:
            ix, iy, iw = 0.05 * W, sb + 0.3 * s, 0.90 * W
        r, d = flow(k, slide, "cover", (ix, iy, iw, foot - iy),
                    [("inside_h", caps(label("inside", *lines, k=k, f=f))), ("inside", "\n".join(lines))], anchor="top")
        draws.append(d)
    _run_all(art + draws)
    return rects


@register("broadsheet", "section", alts=2)
def _bs_section(k, slide, f, image):
    """A section front: the number in a black tab on a heavy rule, the section's title as its headline."""
    W, H, s, o = ctx(k)
    top = bs_strip(k, slide, f)
    num = text_of(f, "number")
    rects, art, draws = {}, [], []
    y = top + (0.6 if alt() == 0 else 0.3) * s
    if num:
        th = 1.10 * s
        tw = min(0.6 * W, max(1.1 * s, na.chip_width(num, 44 * s, k.face("numeral")) + 0.2 * s))
        r, d = flow(k, slide, "section", (0.05 * W + 0.1 * s, y + 0.10 * s, tw - 0.2 * s, th - 0.20 * s), [("number", num)],
                    anchor="middle", align="c", ink=k.P["ground"], accent=k.P["ground"], start={"number": 40 * s})
        rects.update(r); draws.append(d)
        art.append(lambda y=y: dk.box(slide, 0.05 * W, y, tw, th, fill=k.P["ink"]))
        y += th
    art.append(lambda y=y: na.seg(slide, 0.05 * W, y, 0.95 * W, y, k.P["ink"], w=3.0))
    r, d = flow(k, slide, "section", (0.05 * W, y + 0.3 * s, 0.90 * W, H - 0.6 - y - 0.3 * s),
                fields(f, ("kicker", "title"), ("kicker",)), anchor="top")
    rects.update(r); draws.append(d)
    _run_all(art + draws)
    return rects


@register("broadsheet", "image_text", alts=3)
def _bs_image_text(k, slide, f, image):
    """A news photograph with its caption under a hairline; the story's headline and body beside it (under it in
    portrait). Long copy: the words get more of the page — on a square board the story moves beside the photo."""
    W, H, s, o = ctx(k)
    top = bs_strip(k, slide, f) + 0.25 * s
    bottom = H - 0.6
    cap = text_of(f, "caption")
    rects, draws, art = {}, [], []
    if o == "land" or (alt() >= 1 and W >= 0.8 * H):
        img_x, img_w = 0.05 * W, (0.58, 0.46, 0.36)[alt()] * W
        cx0 = img_x + img_w + 0.35 * s
        col = (cx0, top, 0.95 * W - cx0, bottom - top)
        img_y1 = bottom
        if cap:
            cr_, cd = flow(k, slide, "image_text", (img_x, bottom - 0.9 * s, img_w, 0.9 * s), [("caption", cap)], anchor="bottom")
            img_y1 = cr_["caption"][1] - 0.22 * s
    else:
        img_x, img_w = 0.05 * W, 0.90 * W
        img_y1 = top + (0.42, 0.34, 0.26)[alt()] * H
        if cap:
            cr_, cd = flow(k, slide, "image_text", (img_x, img_y1 + 0.22 * s, img_w, 0.6 * s), [("caption", cap)], anchor="top")
        cy0 = (max(v[1] + v[3] for v in cr_.values()) if cap else img_y1) + 0.3 * s
        col = (0.05 * W, cy0, 0.90 * W, bottom - cy0)
    if cap:
        rects.update(cr_); draws.append(cd)
        hy = cr_["caption"][1] - 0.11 * s
        art.append(lambda: na.seg(slide, img_x, hy, img_x + img_w, hy, k.P["ink"], w=0.5))
    r, d = flow(k, slide, "image_text", col, fields(f, ("kicker", "title", "body"), ("kicker",)), anchor="top",
                start={"title": 40 * s})
    rects.update(r); draws.append(d)
    vl._place_image(k, slide, image, (img_x, top, img_w, img_y1 - top), vl.L_((0, 0, 1, 1), "frame", None), "image_text")
    _run_all(art + draws)
    return rects


@register("broadsheet", "quote", alts=2)
def _bs_quote(k, slide, f, image):
    """A pull quote between a heavy rule and a hairline, the quote mark hung in the margin, the source under it."""
    W, H, s, o = ctx(k)
    top = bs_strip(k, slide, f)
    x0, x1 = (0.12 * W, 0.88 * W) if (o == "land" and alt() == 0) else (0.06 * W, 0.94 * W)
    mw = 0.9 * s
    q, a = text_of(f, "quote"), text_of(f, "attribution")
    rects, draws = {}, []
    qy0, qy1 = top + 0.85 * s, H - 1.4 * s
    if q:
        r, d = flow(k, slide, "quote", (x0 + mw, qy0, x1 - x0 - mw, qy1 - qy0), [("quote", q)], anchor="middle")
        rects.update(r); draws.append(d)
        qt, qb = r["quote"][1], r["quote"][1] + r["quote"][3]
        _m, dm = flow(k, slide, "quote", (x0, qt - 0.12 * s, mw, 1.2 * s), [("mark", "“")], start={"mark": 80 * s})
        draws.append(dm)
    else:
        qt = qb = (qy0 + qy1) / 2.0
    art = [lambda: na.seg(slide, x0, qt - 0.3 * s, x1, qt - 0.3 * s, k.P["ink"], w=3.0),
           lambda: na.seg(slide, x0, qb + 0.25 * s, x1, qb + 0.25 * s, k.P["ink"], w=0.6)]
    if a:
        r, d = flow(k, slide, "quote", (x0 + mw, qb + 0.40 * s, x1 - x0 - mw, 0.8 * s), [("attribution", caps(a))])
        rects.update(r); draws.append(d)
    _run_all(art + draws)
    return rects


@register("broadsheet", "data", alts=2)
def _bs_data(k, slide, f, image):
    """By the numbers: the figure, a short rule, its label and note, in a boxed panel under a solid band — a header
    with no words, because the kit has none to give it."""
    W, H, s, o = ctx(k)
    top = bs_strip(k, slide, f)
    bw = ((0.42 if o == "land" else 0.80) if alt() == 0 else 0.86) * W
    bx, band, pad = (W - bw) / 2.0, 0.30 * s, 0.30 * s
    y_in, avail = top + 0.3 * s, (H - 0.6) - (top + 0.3 * s)
    items = fields(f, ("number", "label", "note"))

    def plan(y):
        return flow(k, slide, "data", (bx + pad, y + band + pad, bw - 2 * pad, avail - band - 2 * pad), items,
                    anchor="top", align="c")
    r, _unused = plan(y_in)
    bh = band + 2 * pad + (max(v[1] + v[3] for v in r.values()) - (y_in + band + pad)) + 0.1 * s
    by = y_in + max(0.0, (avail - bh) / 2.0)
    r, d = plan(by)
    art = [lambda: dk.box(slide, bx, by, bw, bh, fill=None, line=dk._as_rgb(k.P["ink"]), line_w=1.25),
           lambda: dk.decorative(dk.box(slide, bx, by, bw, band, fill=k.P["ink"]), "the header band of a numbers box")]
    if "number" in r and "label" in r:
        yl = (r["number"][1] + r["number"][3] + r["label"][1]) / 2.0
        art.append(lambda: na.seg(slide, W / 2 - 0.15 * bw, yl, W / 2 + 0.15 * bw, yl, k.P["ink"], w=0.6))
    _run_all(art + [d])
    return r


@register("broadsheet", "closing", alts=2)
def _bs_closing(k, slide, f, image):
    """The last story's end: its headline, its line, and the ■ end-of-article mark after the last words."""
    W, H, s, o = ctx(k)
    top = bs_strip(k, slide, f)
    t, ln = text_of(f, "title"), text_of(f, "line")
    items = [("title", t if ln or not t else t + " ■")] if t else []
    if ln:
        items.append(("line", ln + " ■"))
    col = (0.05 * W, top + 0.3 * s, (0.70 if o == "land" and alt() == 0 else 0.90) * W, H - 0.6 - top - 0.3 * s)
    r, d = flow(k, slide, "closing", col, items, anchor="middle")
    d()
    end_mark(list(slide.shapes)[-1], k.P["text_accents"][0])
    return r


@register("broadsheet", "points", alts=2)
def _bs_points(k, slide, f, image):
    """Newspaper columns: each point a column — the caller's tag over its head over its body, rules between; a Latin
    body opens on a raised initial when its column has room for the taller first line (and the one line the initial's
    width may push down); where it has not, the body is set plain — the words never move for an ornament. Portrait
    stacks the columns as rows; long copy gives the title more room."""
    W, H, s, o = ctx(k)
    top = bs_strip(k, slide, f)
    pts = points_of(f.get("items"))
    n = len(pts)
    tags = tags_of(f, n, "broadsheet")
    r, d = flow(k, slide, "points", (0.05 * W, top + 0.05 * s, 0.90 * W, (0.20 if alt() == 0 else 0.34) * H),
                fields(f, ("kicker", "title"), ("kicker",)),
                start={"title": 42 * s})                  # an inside page's head; the 80pt headline is the cover's
    rects, draws, art = dict(r), [d], []
    ctop = max((v[1] + v[3] for v in r.values()), default=top) + 0.30 * s
    bottom = H - 0.6
    gap = 0.30 * s
    if o == "land":                                   # columns across the page
        cw = (0.90 * W - (n - 1) * gap) / n
        boxes = [(0.05 * W + i * (cw + gap), ctop, cw, bottom - ctop) for i in range(n)]
        rules = [(b[0] - gap / 2, ctop, b[0] - gap / 2, bottom) for b in boxes[1:]]
    elif alt() == 1 and n >= 3:                       # a square or portrait page with long copy: two columns of rows
        cw, rows = (0.90 * W - gap) / 2, -(-n // 2)
        rh = (bottom - ctop - (rows - 1) * 0.15 * s) / rows
        boxes = [(0.05 * W + (i % 2) * (cw + gap), ctop + (i // 2) * (rh + 0.15 * s), cw, rh) for i in range(n)]
        rules = [(0.05 * W + cw + gap / 2, ctop, 0.05 * W + cw + gap / 2, bottom)] + \
                [(0.05 * W, ctop + r_ * (rh + 0.15 * s) - 0.08 * s, 0.95 * W, ctop + r_ * (rh + 0.15 * s) - 0.08 * s)
                 for r_ in range(1, rows)]
    else:                                             # portrait: the columns stacked as rows
        slot = (bottom - ctop) / n
        boxes = [(0.05 * W, ctop + i * slot, 0.90 * W, slot - 0.15 * s) for i in range(n)]
        rules = [(0.05 * W, b[1] - 0.08 * s, 0.95 * W, b[1] - 0.08 * s) for b in boxes[1:]]
    for x0, y0, x1, y1 in rules:
        art.append(lambda x0=x0, y0=y0, x1=x1, y1=y1: na.seg(slide, x0, y0, x1, y1, k.P["ink"], w=0.5))
    specs = [((x, y, w, h), ([("tag", caps(tags[i]))] if tags else []) + [("item_head", hd)] + ([("item_line", ln)] if ln else []),
              {"anchor": "top"}) for i, ((hd, ln), (x, y, w, h)) in enumerate(zip(pts, boxes))]
    for (hd, ln), (x, y, w, h), (tr, td_) in zip(pts, boxes, plan_together(k, slide, "points", specs)):
        initial = bool(ln) and ln[0].isalpha() and not dk._has_cjk(ln[0]) and "item_line" in tr
        if initial:                       # room in the column for the taller first line and the one line its width
            bx, by, bw, bh = tr["item_line"]  # may push down (lint measures a declared initial so: deckkit._ink_rect)
            pitch = td_.sizes["item_line"] * dk._LINT_LINE_H / 72.0
            initial = by + bh + 1.3 * pitch + pitch + 0.05 <= y + h
        draws.append(_with_initial(k, slide, td_, ln) if initial else td_)
    _run_all(art + draws)
    return rects


# ═══════════════════════════════════ journal 学术期刊 ═══════════════════════════════════
def jn_running_rect(k):
    W, H, s, o = ctx(k)
    return (0.07 * W, 0.035 * H, 0.66 * W, 0.32 * s)


def jn_running(k, slide, f):
    """The running head: the caller's running= (remembered for the deck), else the cover title when it fits one
    line at its floor, else no words — the page number at the right either way, a hairline under both. An explicit
    running= that does not fit is refused. Draws at once and returns the y under it."""
    W, H, s, o = ctx(k)
    # the caller's words, given here or remembered from an earlier page: refused when they cannot fit, never dropped
    explicit = text_of(f, "running") or k.memo.get("running")
    words = memo(k, f, "running", fallback=k.memo.get("_title"))
    y, h = 0.035 * H, 0.32 * s
    if words:
        try:
            _r, d = flow(k, slide, "running", jn_running_rect(k), [("running", words)])
            d()
        except vl.VLTextOverflow:
            if explicit:
                raise
    sz = vl.TYPE["journal"]["running"][0] * s
    dk.text(slide, 0.75 * W, y, 0.18 * W, h, [k.runs(str(page_no(slide)), sz, k.color("mute"), False, "numeral")],
            align=dk.PP_ALIGN.RIGHT)
    yl = y + h + 0.06 * s
    na.seg(slide, 0.07 * W, yl, 0.93 * W, yl, k.P["ink"], w=0.5)
    return yl + 0.25 * s


@register("journal", "cover", alts=2)
def _jn_cover(k, slide, f, image):
    """An article's first page: an accent bar at the left edge; kicker, title, subtitle and the caller's authors=;
    then a rule and the caller's abstract= under its label (in the abstract's own script). The block sits on the
    page's optical centre; never a running head on the cover."""
    W, H, s, o = ctx(k)
    t, ab = text_of(f, "title"), text_of(f, "abstract")
    if t:
        k.memo = dict(k.memo, _title=t)
    if text_of(f, "running"):          # the cover draws no running head, but its running= holds for every later page:
        flow(k, slide, "running", jn_running_rect(k), [("running", text_of(f, "running"))])   # planned, so a refusal lands HERE
    memo(k, f, "running")
    x, w = 0.10 * W, (0.78 if o == "land" else 0.82) * W
    lab_w = 0.15 * W if o == "land" else 0.0
    head = fields(f, ("kicker", "title", "subtitle", "authors"), ("kicker",))

    def plan(y0):
        out, ry = [], None
        r, d = flow(k, slide, "cover", (x, y0, w, H - 0.6 - y0), head, anchor="top")
        out.append((r, d))
        foot = max((v[1] + v[3] for v in r.values()), default=y0)
        if ab:
            ry = foot + 0.30 * s
            lab = caps(label("abstract", ab, k=k, f=f))
            if o == "land":
                lr, ld = flow(k, slide, "cover", (x, ry + 0.22 * s, lab_w - 0.2 * s, 0.5 * s), [("abstract_h", lab)])
                ar, ad = flow(k, slide, "cover", (x + lab_w, ry + 0.18 * s, w - lab_w, H - 0.6 - ry - 0.18 * s),
                              [("abstract", ab)], anchor="top")
            else:
                lr, ld = flow(k, slide, "cover", (x, ry + 0.2 * s, w, 0.4 * s), [("abstract_h", lab)])
                ay = max(v[1] + v[3] for v in lr.values()) + 0.1 * s
                ar, ad = flow(k, slide, "cover", (x, ay, w, H - 0.6 - ay), [("abstract", ab)], anchor="top")
            out += [(lr, ld), (ar, ad)]
            foot = max(v[1] + v[3] for v in ar.values())
        return out, foot, ry
    _o, foot, _ry = plan(0.08 * H)
    y0 = max(0.08 * H, (H - (foot - 0.08 * H)) / 2.0 - 0.1 * s) if alt() == 0 else 0.06 * H
    out, foot, ry = plan(y0)
    art = [lambda: dk.decorative(dk.box(slide, 0, 0, 0.035 * W, H, fill=k.P["accents"][0]), "the journal's edge bar")]
    if ry is not None:
        art.append(lambda: na.seg(slide, x, ry, x + w, ry, k.P["ink"], w=0.5))
    rects = {}
    for r, _d in out:
        rects.update(r)
    _run_all(art + [d for _r, d in out])
    return rects


@register("journal", "section", alts=2)
def _jn_section(k, slide, f, image):
    """A section heading: § and the caller's number (a numeral or roman numeral; other words are set as given) in the
    accent, then kicker and title, a rule under them."""
    W, H, s, o = ctx(k)
    top = jn_running(k, slide, f)
    num = text_of(f, "number")
    lab = ("§ " + num) if num and re.fullmatch(r"[0-9]+(\.[0-9]+)*|[IVXLCivxlc]+", num) else num
    items = ([("sec_no", lab)] if lab else []) + fields(f, ("kicker", "title"), ("kicker",))
    r, d = flow(k, slide, "section", (0.10 * W, top, 0.80 * W, H - 0.6 - top - (0.4 if alt() == 0 else 0.2) * s), items,
                anchor="middle")
    yb = max(v[1] + v[3] for v in r.values()) + 0.25 * s
    _run_all([lambda: na.seg(slide, 0.10 * W, yb, 0.90 * W, yb, k.P["ink"], w=0.6), d])
    return r


@register("journal", "image_text", alts=2)
def _jn_figure(k, slide, f, image):
    """The figure page: its title above; the caller's figure WHOLE (contain, never cropped); beside it the figure's
    label (kicker=, else Figure n counted over the deck's figure pages, in the page's own script), its caption (body=)
    and, under a hairline, the source line (caption=)."""
    W, H, s, o = ctx(k)
    top = jn_running(k, slide, f)
    n = k.memo.get("_fig", 0) + 1
    k.memo = dict(k.memo, _fig=n)
    lab = text_of(f, "kicker") or label("figure", text_of(f, "title"), text_of(f, "body"), k=k, f=f).format(n)
    r, d = flow(k, slide, "image_text", (0.07 * W, top, 0.86 * W, (0.16 if alt() == 0 else 0.28) * H), fields(f, ("title",)),
                anchor="top")
    rects, draws, art = dict(r), [d], []
    ty = max((v[1] + v[3] for v in r.values()), default=top) + 0.25 * s
    bottom = H - 0.6
    if o == "land":
        img = (0.07 * W, ty, (0.52 if alt() == 0 else 0.46) * W, bottom - ty)
        cx0 = img[0] + img[2] + 0.45 * s
        col = (cx0, ty, 0.93 * W - cx0, bottom - ty)
    else:
        img = (0.07 * W, ty, 0.86 * W, (bottom - ty) * (0.56 if alt() == 0 else 0.48))
        cy0 = img[1] + img[3] + 0.3 * s
        col = (0.07 * W, cy0, 0.86 * W, bottom - cy0)
    src = text_of(f, "caption")
    src_h = 0.9 * s if src else 0.0
    main = [("fig_label", caps(lab))] + fields(f, ("body",))
    r, d = flow(k, slide, "image_text", (col[0], col[1], col[2], col[3] - src_h), main, anchor="top")
    rects.update(r); draws.append(d)
    if src:
        yl = col[1] + col[3] - src_h
        sr, sd = flow(k, slide, "image_text", (col[0], yl + 0.15 * s, col[2], src_h - 0.15 * s), [("caption", src)])
        rects.update(sr); draws.append(sd)
        art.append(lambda: na.seg(slide, col[0], yl, col[0] + col[2], yl, k.P["ink"], w=0.4))
    path, alt_txt, slot = vl._resolve(k, image[0] if isinstance(image, (list, tuple)) else image)
    try:                                   # whole, and set at the top-left of its area like a printed figure — contained
        from PIL import Image as _Im       # and centred, a portrait photo floated with dead space on both sides
        iw, ih = _Im.open(path).size
        sc = min(img[2] / iw, img[3] / ih)
        img = (img[0], img[1], iw * sc, ih * sc)
    except Exception:
        pass
    pic = dk.picture(slide, path, *img, fit="contain", alt=alt_txt)
    if slot:
        dk._compose_tag(pic, gen=slot)
    _run_all(art + draws)
    return rects


@register("journal", "points", alts=3)
def _jn_points(k, slide, f, image):
    """Numbered findings (1, 2, 3 in the accent), each a bold head over its line; the caller's margin= beside them
    past an accent rule under its Note label (under the list in portrait or for long copy)."""
    W, H, s, o = ctx(k)
    top = jn_running(k, slide, f)
    pts = points_of(f.get("items"))
    n = len(pts)
    mg = text_of(f, "margin")
    # long copy: the title gets more room. 0.18H refused a kicker over a one-line Japanese title by 0.03in, and the
    # fallbacks (margin under the list) then had no room for three points (the non-Claude agent run, 2026-10-09)
    r, d = flow(k, slide, "points", (0.07 * W, top, 0.86 * W, (0.22 if alt() == 0 else 0.32) * H),
                fields(f, ("kicker", "title"), ("kicker",)))
    rects, draws, art = dict(r), [d], []
    ty = max((v[1] + v[3] for v in r.values()), default=top) + 0.35 * s
    bottom, list_bottom = H - 0.6, H - 0.6
    side = bool(mg) and o == "land" and alt() == 0
    lw = (0.62 if side else 0.86) * W
    if mg:
        if side:
            mx = 0.07 * W + lw + 0.5 * s
            mcol = (mx + 0.2 * s, ty, 0.93 * W - mx - 0.2 * s, bottom - ty)
            art.append(lambda: na.seg(slide, mx, ty, mx, bottom, k.P["text_accents"][0], w=0.75))
        else:                          # under the list: as tall as the note needs, up to 1.3in — not always 1.3in
            mw = 0.86 * W - 0.2 * s
            _m0, _d0 = flow(k, slide, "points", (0.07 * W + 0.2 * s, 0.0, mw, 1.3 * s),
                            [("margin_h", caps(label("note", mg, k=k, f=f))), ("margin", mg)], anchor="top")
            mh = min(1.3 * s, max(v[1] + v[3] for v in _m0.values()) + 0.05 * s)
            mcol = (0.07 * W + 0.2 * s, bottom - mh, mw, mh)
            art.append(lambda: na.seg(slide, 0.07 * W, bottom - mh, 0.07 * W, bottom, k.P["text_accents"][0], w=0.75))
            list_bottom = bottom - mh - 0.3 * s
        _mr, md = flow(k, slide, "points", mcol, [("margin_h", caps(label("note", mg, k=k, f=f))), ("margin", mg)], anchor="top")
        draws.append(md)
    ncol = 2 if (alt() == 2 and n >= 3) else 1        # long copy, last resort: the findings in two columns
    gap = 0.4 * s
    cw = (lw - (ncol - 1) * gap) / ncol
    rows = -(-n // ncol)
    slot = (list_bottom - ty) / rows
    nw = min(0.07 * W, 0.6 * s) if ncol == 2 else 0.07 * W
    cells = [(0.07 * W + (i % ncol) * (cw + gap), ty + (i // ncol) * slot) for i in range(n)]
    nums = plan_together(k, slide, "points", [((x, y, nw, slot), [("item_no", str(i + 1))], {"anchor": "top"})
                                              for i, (x, y) in enumerate(cells)])
    words = plan_together(k, slide, "points", [((x + nw, y, cw - nw, slot - 0.1 * s),
                                                [(x_, t) for x_, t in (("item_head", hd), ("item_line", ln)) if t], {"anchor": "top"})
                                               for (hd, ln), (x, y) in zip(pts, cells)])
    for (_nr, nd), (_tr, td_) in zip(nums, words):
        draws += [nd, td_]
    _run_all(art + draws)
    return rects


@register("journal", "quote", alts=2)
def _jn_quote(k, slide, f, image):
    """A block quotation, indented behind an accent rule, its source after an em dash."""
    W, H, s, o = ctx(k)
    top = jn_running(k, slide, f)
    a = text_of(f, "attribution")
    x = (0.16 if (o == "land" and alt() == 0) else 0.10) * W
    items = fields(f, ("quote",)) + ([("attribution", "— " + a)] if a else [])
    r, d = flow(k, slide, "quote", (x + 0.35 * s, top + 0.3 * s, 0.92 * W - x - 0.35 * s, H - 0.6 - top - 0.3 * s), items,
                anchor="middle")
    y0, y1 = min(v[1] for v in r.values()), max(v[1] + v[3] for v in r.values())
    _run_all([lambda: na.seg(slide, x, y0, x, y1, k.P["text_accents"][0], w=2.0), d])
    return r


@register("journal", "data", alts=2)
def _jn_data(k, slide, f, image):
    """The one number in a tinted panel, its label and note beside it (under it in portrait)."""
    W, H, s, o = ctx(k)
    top = jn_running(k, slide, f)
    num = text_of(f, "number")
    avail = H - 0.6 - (top + 0.3 * s)
    ph = min(avail, 0.62 * H)
    px, py, pw = 0.07 * W, top + 0.3 * s + (avail - ph) / 2.0, 0.86 * W
    rects, draws = {}, []
    if o == "land" and alt() == 0:
        nrect, col, anchor = (px + 0.4 * s, py, pw * 0.42, ph), (px + pw * 0.48, py + 0.3 * s, pw * 0.48, ph - 0.6 * s), "middle"
    else:
        nrect = (px + 0.3 * s, py + 0.2 * s, pw - 0.6 * s, ph * 0.45)
        col, anchor = (px + 0.3 * s, py + ph * 0.5, pw - 0.6 * s, ph * 0.46), "top"
    if num:
        r, d = flow(k, slide, "data", nrect, [("number", num)], anchor="middle", align="c")
        rects.update(r); draws.append(d)
    r, d = flow(k, slide, "data", col, fields(f, ("label", "note")), anchor=anchor)
    rects.update(r); draws.append(d)
    _run_all([lambda: dk.box(slide, px, py, pw, ph, fill=k.P["panel"])] + draws)
    return rects


@register("journal", "closing", alts=2)
def _jn_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    top = jn_running(k, slide, f)
    r, d = flow(k, slide, "closing", (0.10 * W, top, (0.70 if o == "land" and alt() == 0 else 0.80) * W, H - 0.6 - top),
                fields(f, ("title", "line")), anchor="middle")
    d()
    return r


# ═══════════════════════════════════ tally 数据账本 ═══════════════════════════════════
def _tl_ground(k, slide):
    na.grid_background(slide, base=k.P["ground"], ink=k.P["grid_ink"], step=0.25, major=4)


GROUNDS["tally"] = _tl_ground


def chip_size(k, text, max_w, field):
    """The size (pt) a pill of `text` takes to fit `max_w` — the field's size, shrinking toward its floor — or None."""
    W, H, s, o = ctx(k)
    base, role, _b, _c, _i, floor = vl.TYPE[k.name][field]
    sz, face = base * s, k.face(role)
    while na.chip_width(text, sz, face) > max_w and sz > floor * s + 1e-6:
        sz = max(floor * s, sz * 0.92)
    return sz if na.chip_width(text, sz, face) <= max_w else None


def tl_chip(k, slide, x, y, text, max_w, *, field="tag", fill=None):
    """A lime pill (or `fill`) for the caller's words, the darker or lighter ink by contrast. None when it cannot fit."""
    sz = chip_size(k, text, max_w, field)
    if sz is None:
        return None
    fill = fill or k.P["lime"]
    ink = k.P["chip_ink"] if vl._contrast(k.P["chip_ink"], fill) >= vl._contrast("FFFFFF", fill) else "FFFFFF"
    role = vl.TYPE[k.name][field][1]
    return na.chip(slide, x, y, text, size=sz, fill=fill, ink=ink, face=k.face(role), ea_face=k.ea_face(role, text))


def _tl_kicker(k, slide, f, x, y, max_w):
    """The kicker as a pill; a kicker too long for one is set as plain words — still the caller's, never dropped.
    Draws at once; returns its foot."""
    t = text_of(f, "kicker")
    if not t:
        return y
    c = tl_chip(k, slide, x, y, t, max_w, field="kicker")
    if c:
        return y + c[3]
    r, d = flow(k, slide, "kicker", (x, y, max_w, 1.0 * ctx(k)[2]), [("kicker", t)])
    d()
    return max(v[1] + v[3] for v in r.values())


def num_value(t):
    """A plain or percent number ("12", "1,250", "98.6%") as a float; None for anything else ("$4.2M", "3–5")."""
    m = re.fullmatch(r"\s*([0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(\.[0-9]+)?\s*%?\s*", t or "")
    return float((m.group(1) + (m.group(2) or "")).replace(",", "")) if m else None


def share_of(lang, num, total):
    """total= as a share: num/total for two plain or percent numbers of the same kind, num <= total; None when total
    is None. Anything else is refused by name — the bar must never show a fraction it cannot read (tally, interface)."""
    if total is None:
        return None
    v, t = num_value(num), num_value(total)
    if v is None or t is None:
        raise ValueError("{}.data(): total= draws a share bar, so number= and total= must both be plain, non-negative "
                         "numbers (12, 1,250, 98.6%) — got number={!r}, total={!r}".format(lang, num, total))
    if t <= 0:
        raise ValueError("{}.data(): total= must be greater than 0 — a share of nothing has no bar (got total={!r})"
                         .format(lang, total))
    if ("%" in (num or "")) != ("%" in total):
        raise ValueError("{}.data(): number= and total= must both be percentages or neither — got {!r} and "
                         "total={!r}".format(lang, num, total))
    if v > t:
        raise ValueError("{}.data(): the number {!r} is larger than total={!r} — a share cannot exceed its whole"
                         .format(lang, num, total))
    return v / t


@register("tally", "cover", alts=2)
def _tl_cover(k, slide, f, image):
    W, H, s, o = ctx(k)
    x, w = 0.07 * W, 0.86 * W
    y = _tl_kicker(k, slide, f, x, (0.16 if alt() == 0 else 0.07) * H, w) + 0.25 * s
    r, ds = stack(k, slide, "cover", x, (0.78 if o == "land" else 0.86) * W, y, H - 0.6,
                  [(fields(f, ("title",)), 0.70 * s), (fields(f, ("subtitle",)), 0.0)], anchor="top")
    art = []
    if "title" in r:
        yb = r["title"][1] + r["title"][3] + 0.30 * s
        art.append(lambda: dk.box(slide, x, yb, 2.4 * s, 0.09 * s, fill=k.P["text_accents"][0], round=True, r=0.045 * s))
    _run_all(art + ds)
    return r


@register("tally", "section", alts=2)
def _tl_section(k, slide, f, image):
    W, H, s, o = ctx(k)
    num = text_of(f, "number")
    rects, draws = {}, []
    if o == "land" and alt() == 0:
        nrect, col = (0.07 * W, 0.18 * H, 0.38 * W, 0.64 * H), (0.50 * W, 0.22 * H, 0.43 * W, 0.60 * H)
    else:
        nrect, col = (0.07 * W, 0.10 * H, 0.86 * W, 0.34 * H), (0.07 * W, 0.48 * H, 0.86 * W, 0.42 * H)
    if num:
        r, d = flow(k, slide, "section", nrect, [("number", num)], anchor="middle", start={"number": 150 * s})
        rects.update(r); draws.append(d)
    ky = _tl_kicker(k, slide, f, col[0], col[1], col[2]) + (0.2 * s if text_of(f, "kicker") else 0.0)
    r, d = flow(k, slide, "section", (col[0], ky, col[2], col[1] + col[3] - ky), fields(f, ("title",)), anchor="top")
    rects.update(r); draws.append(d)
    _run_all(draws)
    return rects


@register("tally", "image_text", alts=2)
def _tl_image_text(k, slide, f, image):
    """The caller's picture under a hairline frame; kicker pill, title, body and caption beside it."""
    W, H, s, o = ctx(k)
    if o == "land":
        img = (0.07 * W, 0.12 * H, (0.50 if alt() == 0 else 0.44) * W, 0.76 * H)
        cx0 = img[0] + img[2] + 0.5 * s
        col = (cx0, 0.12 * H, 0.93 * W - cx0, 0.76 * H)
    else:
        img = (0.07 * W, 0.07 * H, 0.86 * W, (0.40 if alt() == 0 else 0.32) * H)
        cy0 = img[1] + img[3] + 0.4 * s
        col = (0.07 * W, cy0, 0.86 * W, H - 0.6 - cy0)
    y = _tl_kicker(k, slide, f, col[0], col[1], col[2]) + (0.2 * s if text_of(f, "kicker") else 0.0)
    r, d = flow(k, slide, "image_text", (col[0], y, col[2], col[1] + col[3] - y), fields(f, ("title", "body", "caption")),
                anchor="top")
    vl._place_image(k, slide, image, img, vl.L_((0, 0, 1, 1), "frame", None), "image_text")
    frame = dk.box(slide, *img, fill=None, line=dk._as_rgb(k.P["ink"]), line_w=0.75)
    dk.decorative(frame, "a hairline frame around the picture")
    d()
    return r


@register("tally", "points", alts=3)
def _tl_points(k, slide, f, image):
    """Ledger rows: 01, 02 … in the accent, the point's head and line, the caller's tag in a pill at the right (under
    the words when the pills are wide); a rule above each row and one closing the ledger. A tag too long for its pill
    is refused by name, never dropped. Long copy: the title gets more room, then the ledger splits into two columns."""
    W, H, s, o = ctx(k)
    pts = points_of(f.get("items"))
    n = len(pts)
    tags = tags_of(f, n, "tally")
    x0, w0 = 0.07 * W, 0.86 * W
    y = _tl_kicker(k, slide, f, x0, 0.07 * H, w0) + 0.15 * s
    r, d = flow(k, slide, "points", (x0, y, w0, (0.18 if alt() == 0 else 0.30) * H), fields(f, ("title",)), anchor="top")
    rects, draws, art = dict(r), [d], []
    top = max((v[1] + v[3] for v in r.values()), default=y) + 0.40 * s
    bottom = H - 0.6
    ncol = 2 if (alt() == 2 and n >= 3) else 1
    gap = 0.4 * s
    w = (w0 - (ncol - 1) * gap) / ncol
    rows = -(-n // ncol)
    pitch = (bottom - top) / rows
    tag_sz = vl.TYPE["tally"]["tag"][0] * s
    tw = (max(na.chip_width(t, tag_sz, k.face("body")) for t in tags) + 0.3 * s) if tags else 0.0
    beside = bool(tags) and tw <= 0.30 * w
    nw = (1.7 if alt() == 0 else 1.1) * s
    rows_ = []
    for i, (hd, ln) in enumerate(pts):
        x, ry = x0 + (i % ncol) * (w + gap), top + (i // ncol) * pitch
        tcol = (x + nw + 0.2 * s, ry, w - nw - 0.2 * s - (tw if beside else 0.0), pitch - 0.24 * s)
        if tags and not beside:
            tcol = (tcol[0], tcol[1], tcol[2], tcol[3] - 0.45 * s)
        rows_.append((x, ry, tcol))
    nums = plan_together(k, slide, "points", [((x, ry, nw, pitch - 0.24 * s), [("item_no", "{:02d}".format(i + 1))],
                                               {"anchor": "middle", "start": {"item_no": (50 if alt() == 0 else 34) * s}})
                                              for i, (x, ry, _t) in enumerate(rows_)])
    words = plan_together(k, slide, "points", [(tcol, [(x_, t) for x_, t in (("item_head", hd), ("item_line", ln)) if t],
                                                {"anchor": "middle"}) for (hd, ln), (_x, _y, tcol) in zip(pts, rows_)])
    for i, ((x, ry, tcol), (_nr, nd), (_tr, td_)) in enumerate(zip(rows_, nums, words)):
        art.append(lambda x=x, ry=ry: na.seg(slide, x, ry - 0.12 * s, x + w, ry - 0.12 * s, k.P["ink"], w=0.6, alpha=0.5))
        draws += [nd, td_]
        if tags:
            mw = (tw - 0.3 * s) if beside else tcol[2]
            if chip_size(k, tags[i], mw, "tag") is None:
                raise vl.VLTextOverflow("tally.points(): the tag {!r} does not fit its pill even at the floor size — "
                                        "shorten it".format(tags[i][:40]))
            ch = tag_sz / 72.0 * 1.9
            cx_, cy_ = (x + w - tw + 0.3 * s, ry + (pitch - 0.24 * s - ch) / 2.0) if beside else (tcol[0], tcol[1] + tcol[3] + 0.1 * s)
            draws.append(lambda t=tags[i], cx_=cx_, cy_=cy_, mw=mw: tl_chip(k, slide, cx_, cy_, t, mw))
    yl = top + rows * pitch - 0.12 * s
    for c in range(ncol):
        xc = x0 + c * (w + gap)
        art.append(lambda xc=xc: na.seg(slide, xc, yl, xc + w, yl, k.P["ink"], w=0.6, alpha=0.5))
    _run_all(art + draws)
    return rects


@register("tally", "quote", alts=2)
def _tl_quote(k, slide, f, image):
    """The quote in heavy type behind a rounded accent bar; the source in a pill (as words when too long for one)."""
    W, H, s, o = ctx(k)
    x = (0.10 if alt() == 0 else 0.06) * W
    a = text_of(f, "attribution")
    col = (x + 0.45 * s, 0.14 * H, 0.90 * W - x - 0.45 * s, H - 0.6 - 0.14 * H - (0.9 * s if a else 0.0))
    r, d = flow(k, slide, "quote", col, fields(f, ("quote",)), anchor="middle")
    rects, draws = dict(r), [d]
    qt = min((v[1] for v in r.values()), default=col[1])
    qb = max((v[1] + v[3] for v in r.values()), default=col[1])
    art = [lambda: dk.box(slide, x, qt, 0.10 * s, max(0.2, qb - qt), fill=k.P["text_accents"][0], round=True, r=0.05 * s)]
    if a:
        ay = qb + 0.3 * s
        if chip_size(k, a, col[2], "attribution") is not None:
            draws.append(lambda: tl_chip(k, slide, col[0], ay, a, col[2], field="attribution"))
        else:
            r2, d2 = flow(k, slide, "quote", (col[0], ay, col[2], H - 0.6 - ay), [("attribution", a)])
            rects.update(r2); draws.append(d2)
    _run_all(art + draws)
    return rects


@register("tally", "data", alts=2)
def _tl_data(k, slide, f, image):
    """The giant number, its label and note — and, when the caller gives total=, a share bar: the number's share of
    the total with both numbers at its ends. A total the number cannot be read against is refused, never guessed."""
    W, H, s, o = ctx(k)
    num, total = text_of(f, "number"), text_of(f, "total")
    frac = share_of("tally", num, total)
    x, w = 0.07 * W, 0.86 * W
    bottom = H - 0.6 - (1.0 * s if frac is not None else 0.0)
    rects, draws, art = {}, [], []
    if o == "land" and alt() == 0:
        nrect, col, anchor = (x, 0.12 * H, 0.52 * W, bottom - 0.12 * H), (0.62 * W, 0.18 * H, 0.31 * W, bottom - 0.18 * H), "middle"
    else:
        hh = bottom - 0.08 * H
        nrect, col, anchor = (x, 0.08 * H, w, hh * 0.5), (x, 0.08 * H + hh * 0.52, w, hh * 0.48), "top"
    if num:
        r, d = flow(k, slide, "data", nrect, [("number", num)], anchor="middle")
        rects.update(r); draws.append(d)
    r, d = flow(k, slide, "data", col, fields(f, ("label", "note")), anchor=anchor)
    rects.update(r); draws.append(d)
    if frac is not None:
        by, sz = bottom + 0.25 * s, vl.TYPE["tally"]["tag"][0] * s
        art.append(lambda: na.share_bar(slide, x, by, w, frac, track=k.P["track"], fill=k.P["text_accents"][0], h=0.16 * s))
        art.append(lambda: dk.text(slide, x, by + 0.25 * s, w / 2, 0.35 * s, [k.runs(num, sz, k.color("mute"), True)]))
        art.append(lambda: dk.text(slide, x + w / 2, by + 0.25 * s, w / 2, 0.35 * s, [k.runs(total, sz, k.color("mute"), True)],
                                   align=dk.PP_ALIGN.RIGHT))
    _run_all(art + draws)
    return rects


@register("tally", "closing", alts=2)
def _tl_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    x = 0.07 * W
    col = (x, 0.20 * H, (0.78 if o == "land" else 0.86) * W, 0.60 * H) if alt() == 0 else (x, 0.08 * H, 0.86 * W, 0.78 * H)
    r, d = flow(k, slide, "closing", col, fields(f, ("title", "line")), anchor="middle")
    yb = max(v[1] + v[3] for v in r.values()) + 0.35 * s
    _run_all([d, lambda: dk.box(slide, x, yb, 2.4 * s, 0.09 * s, fill=k.P["text_accents"][0], round=True, r=0.045 * s)])
    return r


# ═══════════════════════════════════ chalkboard 黑板报 ═══════════════════════════════════
def _cb_ground(k, slide):
    na.board_frame(slide, k.P["wood"], k.P["chalk"])      # the frame inside the page, the ledge, a stick of chalk


GROUNDS["chalkboard"] = _cb_ground


def ink_w(k, field, text, size):
    """The width (in) one line of `text` takes at `size` in the field's face. Text with CJK in it is set in the
    East-Asian face, whose spaces and marks ("·") run wider than the Latin face measures — the sample's
    "科学课 · 第 3 讲" ran past its box — so it counts every CJK character an em and every other at least 0.6 em."""
    _b, role, bold, *_r = vl.TYPE[k.name][field]
    w = na.chip_width(text, size, k.face(role), bold=bool(bold)) - 1.2 * size / 72.0
    if dk._has_cjk(text):
        w = max(w, sum(1.0 if dk._has_cjk(ch) else 0.6 for ch in text) * size / 72.0)
    return w


def _one_line(rect, size):
    return rect[3] <= size * 1.2 / 72.0 * 1.5 + 0.06


def _cb_smudges(k, slide, keep, seed):
    """Up to three faint eraser smudges (rim-free radial fades) where no words are; fewer when the board is full."""
    import random as _random
    W, H, s, o = ctx(k)
    rnd = _random.Random(seed)
    placed = 0
    keep = [(x_, y_, w_, h_ + 0.35 * s) for x_, y_, w_, h_ in keep]   # and clear of a chalk underline under the words
    for _ in range(40):
        if placed == 3:
            break
        w = rnd.uniform(3.0, 4.5) * s
        x, y = rnd.uniform(0.3 * s, max(0.3 * s, W - w - 0.3 * s)), rnd.uniform(0.4 * s, H - 1.2 * s)
        r = (x, y, w, w * 0.22)
        if x + w > W - 0.2 * s or any(meet(r, c, 0.1 * s) for c in keep):
            continue
        keep = list(keep) + [r]                      # nor on each other
        e = dk.box(slide, x, y, w, w * 0.22, fill=k.P["ink"])
        e._element.spPr.find(dk.qn("a:prstGeom")).set("prst", "ellipse")
        na._radial(e, k.P["ink"], 0.06)
        dk.decorative(e, "an eraser smudge on the board; nothing reads from it")
        placed += 1


def _cb_doodle(k, slide, spec, cx, cy, size):
    """The caller's icon (doodle=, an icons.py spec) drawn in yellow chalk. Fetched while PLANNING, so an unknown
    icon raises (naming what it tried) before anything is drawn."""
    import icons as _ic
    import tempfile as _tf
    png = Path(_tf.gettempdir()) / "slide-maker-native-art" / "doodle_{}_{}.png".format(
        re.sub(r"[^A-Za-z0-9]+", "_", spec), k.P["text_accents"][0])
    png.parent.mkdir(parents=True, exist_ok=True)
    import os as _os
    tmp = str(png) + ".{}.tmp.png".format(_os.getpid())     # written whole, then renamed: two builds never read half
    _ic.icon_png(spec, tmp, color=k.P["text_accents"][0], px=320)
    _os.replace(tmp, str(png))

    def draw():
        pic = dk.icon(slide, str(png), cx - size / 2.0, cy - size / 2.0, size, alt=spec.split(":")[-1].replace("-", " "))
        dk.decorative(pic, "a chalk doodle beside the title; the title carries the meaning")
    return draw


@register("chalkboard", "cover", alts=2)
def _cb_cover(k, slide, f, image):
    """A lesson's first board: the kicker in a pink chalk box, the title with a double chalk underline, the subtitle;
    the caller's doodle= drawn in yellow chalk beside them — nothing drawn without it."""
    W, H, s, o = ctx(k)
    dd = text_of(f, "doodle")
    x, w = 0.08 * W, (0.60 if dd and o == "land" else 0.84) * W
    y0, y1 = (0.12 if alt() == 0 else 0.07) * H, (0.66 if dd and o != "land" else 0.86) * H
    r, ds = stack(k, slide, "cover", x, w, y0, y1,
                  [(fields(f, ("kicker", "title")), 0.40 * s), (fields(f, ("subtitle",)), 0.0)],
                  accent=k.P["text_accents"][1])
    art = []
    if "kicker" in r:
        kr = r["kicker"]
        sz = ds[0].sizes["kicker"]
        kw = min(kr[2], ink_w(k, "kicker", text_of(f, "kicker"), sz)) if _one_line(kr, sz) else kr[2]
        art.append(lambda: na.chalk_box(slide, kr[0] - 0.12 * s, kr[1] - 0.06 * s, kw + 0.24 * s, kr[3] + 0.12 * s,
                                        k.P["text_accents"][1], seed=3))
    if "title" in r:
        tr = r["title"]
        art.append(lambda: na.chalk_underline(slide, tr[0], tr[1] + tr[3] + 0.08 * s, min(tr[2], 0.55 * W),
                                              k.P["text_accents"][0], seed=5))
    if dd:
        size = 1.6 * s
        if o == "land":
            cx_, cy_ = 0.80 * W, 0.62 * H
        else:
            foot = max(v[1] + v[3] for v in r.values())
            cx_, cy_ = 0.70 * W, min(foot + 0.4 * s + size / 2, H - 0.9 * s - size / 2)
        art.append(_cb_doodle(k, slide, dd, cx_, cy_, size))
    _cb_smudges(k, slide, list(r.values()) + ([(cx_ - size / 2, cy_ - size / 2, size, size)] if dd else []), seed=1)
    _run_all(art + ds)
    return r


@register("chalkboard", "section", alts=2)
def _cb_section(k, slide, f, image):
    W, H, s, o = ctx(k)
    num = text_of(f, "number")
    rects, draws, art = {}, [], []
    if o == "land" and alt() == 0:
        cx, cy, R = 0.24 * W, 0.50 * H, min(0.30 * H, 0.17 * W)
        col = (0.44 * W, 0.16 * H, 0.48 * W, 0.68 * H)
    else:
        cx, cy = 0.50 * W, 0.24 * H
        R = min(0.16 * H, 0.30 * W) * (1.0 if alt() == 0 else 0.75)
        col = (0.08 * W, cy + R + 0.4 * s, 0.84 * W, H - 0.7 * s - (cy + R + 0.4 * s))
    R = fit_circle(k, cx, cy, 2 * R) / 2.0
    if num:
        r, d = flow(k, slide, "section", (cx - R * 0.75, cy - R * 0.55, 1.5 * R, 1.1 * R), [("number", num)],
                    anchor="middle", align="c", start={"number": 110 * s})
        rects.update(r); draws.append(d)
        art.append(lambda: na.chalk_ellipse(slide, cx, cy, R, R * 0.85, k.P["text_accents"][1], seed=9))
    r, ds = stack(k, slide, "section", col[0], col[2], col[1], col[1] + col[3], [(fields(f, ("kicker", "title")), 0.0)],
                  accent=k.P["text_accents"][1])
    rects.update(r); draws += ds
    if "title" in r:
        tr = r["title"]
        art.append(lambda: na.chalk_underline(slide, tr[0], tr[1] + tr[3] + 0.08 * s, min(tr[2], 0.4 * W),
                                              k.P["text_accents"][0], seed=11))
    _cb_smudges(k, slide, list(rects.values()) + [(cx - R, cy - R, 2 * R, 2 * R)], seed=2)
    _run_all(art + draws)
    return rects


@register("chalkboard", "image_text", alts=2)
def _cb_image_text(k, slide, f, image):
    """The caller's picture pinned to the board by four chalk corner marks, the words beside it (under it in
    portrait)."""
    W, H, s, o = ctx(k)
    if o == "land":
        img = (0.08 * W, 0.14 * H, (0.46 if alt() == 0 else 0.40) * W, 0.70 * H)
        cx0 = img[0] + img[2] + 0.6 * s
        col = (cx0, 0.12 * H, 0.92 * W - cx0, 0.76 * H)
    else:
        img = (0.10 * W, 0.08 * H, 0.80 * W, (0.40 if alt() == 0 else 0.32) * H)
        cy0 = img[1] + img[3] + 0.5 * s
        col = (0.08 * W, cy0, 0.84 * W, H - 0.7 * s - cy0)
    r, ds = stack(k, slide, "image_text", col[0], col[2], col[1], col[1] + col[3],
                  [(fields(f, ("kicker", "title", "body", "caption")), 0.0)], anchor="middle" if o == "land" else "top",
                  accent=k.P["text_accents"][1])
    vl._place_image(k, slide, image, img, vl.L_((0, 0, 1, 1), "frame", None), "image_text")
    L = 0.35 * s
    x0, y0, x1, y1 = img[0] - 0.08 * s, img[1] - 0.08 * s, img[0] + img[2] + 0.08 * s, img[1] + img[3] + 0.08 * s
    for i, (ax, ay, dx, dy) in enumerate(((x0, y0, 1, 1), (x1, y0, -1, 1), (x1, y1, -1, -1), (x0, y1, 1, -1))):
        for sh_ in na.chalk_path(slide, [(ax + dx * L, ay), (ax, ay), (ax, ay + dy * L)], k.P["ink"], w=2.2, seed=20 + i, passes=1):
            dk.overlap_intent(sh_, "a chalk corner mark holds the picture to the board: it sits on the picture's corner")
    _run_all(ds)
    return r


@register("chalkboard", "points", alts=3)
def _cb_points(k, slide, f, image):
    """Chalk boxes, one per point, each sized to its measured words, a circled number at its corner; arrows between
    them only when the caller says the points happen in order (ordered=True) — numbers alone read as a list. Across
    a landscape page the boxes form a row; otherwise they stack; long copy gets a 2x2 grid, then a roomier title."""
    W, H, s, o = ctx(k)
    pts = points_of(f.get("items"))
    n = len(pts)
    ordered = f.get("ordered", False)
    if not isinstance(ordered, bool):
        raise TypeError("chalkboard.points(): ordered= is True or False, got {!r}".format(ordered))
    r, d = flow(k, slide, "points", (0.08 * W, 0.08 * H, 0.84 * W, (0.26 if alt() == 2 else 0.18) * H),
                fields(f, ("kicker", "title")), accent=k.P["text_accents"][1])
    rects, draws, art = dict(r), [d], []
    if "title" in r:
        tr = r["title"]
        art.append(lambda: na.chalk_underline(slide, tr[0], tr[1] + tr[3] + 0.08 * s, min(tr[2], 0.36 * W),
                                              k.P["text_accents"][0], seed=2))
    top = max((v[1] + v[3] for v in r.values()), default=0.08 * H) + 0.50 * s
    bottom = H - 0.75 * s
    cols = [k.P["text_accents"][0], k.P["text_accents"][1], k.P["accents"][2], k.P["ink"]]
    # the designed layout: a row across a landscape page, a stack in portrait; long copy: a 2x2 grid, then the
    # grid under a roomier title (a roomier title alone only lets it stay large and squeeze the boxes)
    ncol = (n if o == "land" else 1) if alt() == 0 else (2 if n >= 3 else 1)
    nrow = -(-n // ncol)
    gap, pad, numd = ((0.55, 0.30, 0.64) if alt() == 0 else (0.35, 0.20, 0.50))
    gap, pad, numd = gap * s, pad * s, numd * s
    bw = (0.84 * W - (ncol - 1) * gap) / ncol
    above = ncol == n and n > 1                       # one row: the number sits above the words, else beside them

    head_start = {}
    if ncol == n and n > 1:
        # a row of narrow boxes only when every head still reads as a head — a Chinese one on one line, a Latin one in
        # at most two — at one shared size no smaller than 75% of the field's; else the grid takes over (stress render:
        # "共享工 / 具", "Write / every / repair / down")
        base_h, role_h, bold_h = vl.TYPE[k.name]["item_head"][:3]
        tw_ = (bw - 2 * pad) * HEADROOM - dk.TEXT_INSET_LR

        def head_lines(hd, z_):
            face_ = k.ea_face(role_h, hd) or k.face(role_h)
            return dk._measure_lines([(hd, bool(bold_h))], z_, tw_, font=face_)
        z_ = base_h * s
        while any(head_lines(hd, z_) > (1 if dk._has_cjk(hd) else 2) for hd, _ln in pts):
            z_ *= 0.95
            if z_ < 0.75 * base_h * s:
                raise vl.VLTextOverflow("chalkboard.points(): {} boxes in a row are too narrow for their heads — the "
                                        "grid takes over".format(n))
        head_start = {"item_head": z_}

    def plan_box(i, y, h, start=None):
        hd, ln = pts[i]
        x = 0.08 * W + (i % ncol) * (bw + gap)
        if above:
            tcol = (x + pad, y + pad + numd + 0.15 * s, bw - 2 * pad, h - (2 * pad + numd + 0.15 * s))
        else:
            tcol = (x + pad + numd + 0.25 * s, y + pad, bw - 2 * pad - numd - 0.25 * s, h - 2 * pad)
        return flow(k, slide, "points", tcol, [(x_, t) for x_, t in (("item_head", hd), ("item_line", ln)) if t], anchor="top",
                    start=dict(head_start, **(start or {})) or None)
    # pass 1: each row's height from its boxes' own words; pass 2: the rows centred in the room under the title. A
    # grid's narrow boxes try the number beside the words first, then above them (a square board, 2026-10-08)
    slot = (bottom - top - (nrow - 1) * gap) / nrow
    opts = [above] if (ncol == 1 or above) else [False, True]
    for j, above in enumerate(opts):
        try:
            # every box at ONE size per field — the smallest any box needed alone (a Chinese box's body was smaller
            # than its neighbours', 2026-10-09) — then each row as tall as its tallest box at that size
            shared = {}
            for i in range(n):
                for f_, z in getattr(plan_box(i, top, slot)[1], "sizes", {}).items():
                    shared[f_] = min(shared.get(f_, z), z)
            row_h = []
            for rw in range(nrow):
                hs = [max(v[1] + v[3] for v in plan_box(i, top, slot, shared)[0].values()) - top + pad
                      for i in range(rw * ncol, min(n, (rw + 1) * ncol))]
                row_h.append(max(max(hs), numd + 2 * pad))
            break
        except vl.VLTextOverflow:
            if j == len(opts) - 1:
                raise
    y = top + max(0.0, (bottom - top - sum(row_h) - (nrow - 1) * gap) / 2.0)
    boxes = []
    for rw in range(nrow):
        for i in range(rw * ncol, min(n, (rw + 1) * ncol)):
            boxes.append((0.08 * W + (i % ncol) * (bw + gap), y, bw, row_h[rw]))
        y += row_h[rw] + gap
    for i, (x, y, w, h_) in enumerate(boxes):
        c = cols[i % len(cols)]
        _tr, td_ = plan_box(i, y, h_ + 0.01, shared)
        art.append(lambda x=x, y=y, w=w, h_=h_, c=c, i=i: na.chalk_box(slide, x, y, w, h_, c, seed=20 + i))
        ncx, ncy = x + pad + numd / 2.0, y + pad + numd / 2.0
        art.append(lambda ncx=ncx, ncy=ncy, c=c, i=i: na.chalk_ellipse(slide, ncx, ncy, numd / 2.0, numd / 2.0, c, seed=30 + i))
        _nr, nd = flow(k, slide, "points", (ncx - numd / 2.0, ncy - numd / 2.0, numd, numd), [("item_no", str(i + 1))],
                       anchor="middle", align="c", ink=c, accent=c)
        draws += [nd, td_]
        if ordered and i < n - 1:
            nx, ny, nw_, nh = boxes[i + 1]
            if ny == y:                               # the next box is beside this one: across the gap
                ay = y + min(h_, nh) / 2.0
                art.append(lambda x=x, w=w, ay=ay, i=i: na.chalk_arrow(slide, x + w + 0.10 * s, ay, x + w + gap - 0.10 * s, ay,
                                                                     k.P["ink"], seed=40 + i))
            elif ncol == 1:                           # stacked: straight down the gap
                ax = x + w / 2.0
                art.append(lambda ax=ax, y=y, h_=h_, i=i: na.chalk_arrow(slide, ax, y + h_ + 0.06 * s, ax, y + h_ + gap - 0.06 * s,
                                                                       k.P["ink"], seed=40 + i))
            else:                                     # the grid's next row: through the central gap, corner to corner
                art.append(lambda x=x, y=y, h_=h_, nx=nx, nw_=nw_, ny=ny, i=i: na.chalk_arrow(
                    slide, x - 0.05 * s, y + h_ + 0.06 * s, nx + nw_ + 0.05 * s, ny - 0.06 * s, k.P["ink"], seed=40 + i))
    _cb_smudges(k, slide, list(rects.values()) + boxes, seed=3)
    _run_all(art + draws)
    return rects


@register("chalkboard", "quote", alts=2)
def _cb_quote(k, slide, f, image):
    W, H, s, o = ctx(k)
    x = max(0.17 * W, 1.6 * s) if alt() == 0 else max(0.10 * W, 1.45 * s)
    a = text_of(f, "attribution")
    r, ds = stack(k, slide, "quote", x, 0.90 * W - x, 0.16 * H, H - 0.75 * s,
                  [(fields(f, ("quote",)), 0.45 * s), ([("attribution", "— " + a)] if a else [], 0.0)],
                  accent=k.P["text_accents"][1])
    art = []
    if "quote" in r:
        qr = r["quote"]
        _m, md = flow(k, slide, "quote", (x - 1.35 * s, qr[1] - 0.25 * s, 1.2 * s, 1.4 * s), [("mark", "“")],
                      start={"mark": 110 * s})
        ds.append(md)
        art.append(lambda: na.chalk_underline(slide, qr[0], qr[1] + qr[3] + 0.08 * s, min(qr[2], 0.22 * W),
                                              k.P["text_accents"][0], seed=8))
    _cb_smudges(k, slide, list(r.values()) + [(x - 1.35 * s, 0.0, 1.2 * s, H)], seed=4)
    _run_all(art + ds)
    return r


@register("chalkboard", "data", alts=2)
def _cb_data(k, slide, f, image):
    """The number circled in pink chalk, a chalk arrow from it to its label and note."""
    W, H, s, o = ctx(k)
    num = text_of(f, "number")
    rects, draws, art = {}, [], []
    if o == "land" and alt() == 0:
        cx, cy, R = 0.28 * W, 0.48 * H, min(0.30 * H, 0.18 * W)
        col = (0.56 * W, 0.22 * H, 0.36 * W, 0.56 * H)
    else:
        cx, cy = 0.5 * W, 0.26 * H
        R = min(0.17 * H, 0.32 * W) * (1.0 if alt() == 0 else 0.8)
        col = (0.10 * W, cy + R + 0.7 * s, 0.80 * W, H - 0.75 * s - (cy + R + 0.7 * s))
    R = fit_circle(k, cx, cy, 2 * R) / 2.0
    side = o == "land" and alt() == 0
    if num:
        # beside the words the figure keeps to its circle; centred above them it may be as wide as the page allows,
        # and the chalk ring widens to hold it ("1,250,000" was refused on a square board)
        nrect = (cx - 1.05 * R, cy - 0.7 * R, 2.1 * R, 1.4 * R) if side else (0.08 * W, cy - 0.7 * R, 0.84 * W, 1.4 * R)
        r, d = flow(k, slide, "data", nrect, [("number", num)], anchor="middle", align="c")
        rects.update(r); draws.append(d)
        rx = 1.15 * R
        if not side:
            rx = min(0.48 * W, max(rx, ink_w(k, "number", num, d.sizes.get("number", 48)) / 2.0 + 0.35 * s))
        art.append(lambda rx=rx: na.chalk_ellipse(slide, cx, cy, rx, 0.92 * R, k.P["text_accents"][1], seed=9, turns=1.12))
    r, ds = stack(k, slide, "data", col[0], col[2], col[1], col[1] + col[3], [(fields(f, ("label", "note")), 0.0)],
                  anchor="middle" if o == "land" else "top")
    rects.update(r); draws += ds
    if num and r:
        ty = min(v[1] for v in r.values())
        if o == "land" and alt() == 0:
            art.append(lambda: na.chalk_arrow(slide, cx + 1.2 * R, cy - 0.1 * R, col[0] - 0.2 * s, ty + 0.25 * s,
                                              k.P["ink"], seed=60))
        else:
            art.append(lambda: na.chalk_arrow(slide, cx, cy + 0.98 * R, cx, ty - 0.12 * s, k.P["ink"], seed=60))
    _cb_smudges(k, slide, list(rects.values()) + [(cx - 1.2 * R, cy - R, 2.4 * R, 2 * R)], seed=5)
    _run_all(art + draws)
    return rects


@register("chalkboard", "closing", alts=2)
def _cb_closing(k, slide, f, image):
    W, H, s, o = ctx(k)
    dd = text_of(f, "doodle")
    w = (0.62 if dd and o == "land" else 0.84) * W
    r, ds = stack(k, slide, "closing", 0.08 * W, w, (0.14 if alt() == 0 else 0.07) * H, (0.68 if dd and o != "land" else 0.86) * H,
                  [(fields(f, ("title",)), 0.45 * s), (fields(f, ("line",)), 0.0)])
    art = []
    if "title" in r:
        tr = r["title"]
        art.append(lambda: na.chalk_underline(slide, tr[0], tr[1] + tr[3] + 0.08 * s, min(tr[2], 0.5 * W),
                                              k.P["text_accents"][0], seed=12))
    if dd:
        size = 1.6 * s
        if o == "land":
            cx_, cy_ = 0.82 * W, 0.50 * H
        else:
            foot = max(v[1] + v[3] for v in r.values())
            cx_, cy_ = 0.70 * W, min(foot + 0.4 * s + size / 2, H - 0.9 * s - size / 2)
        art.append(_cb_doodle(k, slide, dd, cx_, cy_, size))
    _cb_smudges(k, slide, list(r.values()) + ([(cx_ - size / 2, cy_ - size / 2, size, size)] if dd else []), seed=6)
    _run_all(art + ds)
    return r
