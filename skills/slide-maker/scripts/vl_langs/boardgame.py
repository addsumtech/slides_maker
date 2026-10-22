"""boardgame — the Board game language (棋盘游戏): its ground painter, its card and its seven pages.

A classic property board, generic — no brand, no game's words. The cover is the whole board: a ring of squares with
coloured bands round the page, the title on the logo banner in the board's centre, a pair of dice, a pawn on the
start corner; a section is one big corner square carrying the caller's number, the board's two edges running off from
it; the points are squares on the track (ordered=True: the pawn on the first, hops square to square, the squares
numbered); a quote is printed on a drawn card from the deck; the key figure sits on the square the pawn landed on,
the dice beside it; the closing is the finish corner, the pawn arriving. Every page PLANS its words first
(vl_native2.flow / stack: measured, never under a floor), draws its art clear of them, then sets them. The only words
it writes itself are the point numbers of an ordered track (the manifest's derived words)."""
from __future__ import annotations

import math

import deckkit as dk
import native_art as na
import visual_languages as vl
from vl_native import _run_all, alt, ctx, points_of, register, text_of
import vl_native2 as _v2   # names looked up at call time: vl_native2 may still be importing when this module loads

NAME = "boardgame"

_LAYER = "a board game's pieces lie on each other by design: a pawn stands on its square, dice tumble onto the board"


# ─────────────────────────────────────────── drawing primitives ───────────────────────────────────────────
def _decor(sh, reason, overlap=None):
    if sh is None:
        return None
    dk.decorative(sh, reason)
    if overlap:
        dk.overlap_intent(sh, overlap if isinstance(overlap, str) else _LAYER)
    return sh


def _poly(slide, pts, fill, reason, *, line=None, line_w=0.75, alpha=None, overlap=None):
    return _decor(na.poly(slide, pts, fill=fill, line=line, line_w=line_w, alpha=alpha), reason, overlap)


def _mpoly(slide, polys, fill, reason, *, alpha=None, overlap=None):
    """Many polygons as ONE shape, each clipped to the page first (multipoly does not clip)."""
    W, H = na.page_size(slide)
    polys = [q for q in (na.clip_to_page(p, W, H) for p in polys) if len(q) >= 3]
    return _decor(na.multipoly(slide, polys, fill, alpha=alpha), reason, overlap)


def _clip_seg(x0, y0, x1, y1, W, H):
    """Liang–Barsky: the part of a segment on the page, or None."""
    dx, dy = x1 - x0, y1 - y0
    t0, t1 = 0.0, 1.0
    for p, q in ((-dx, x0), (dx, W - x0), (-dy, y0), (dy, H - y0)):
        if abs(p) < 1e-12:
            if q < 0:
                return None
            continue
        t = q / p
        if p < 0:
            t0 = max(t0, t)
        else:
            t1 = min(t1, t)
        if t0 > t1:
            return None
    return (x0 + t0 * dx, y0 + t0 * dy), (x0 + t1 * dx, y0 + t1 * dy)


def _segs(slide, segs, color, reason, *, w=0.75, alpha=None, overlap=None):
    """Pen strokes (the board's printed lines) as ONE shape, each cut to the page."""
    W, H = na.page_size(slide)
    cut = [c for c in (_clip_seg(a[0], a[1], b[0], b[1], W, H) for a, b in segs) if c and
           math.hypot(c[1][0] - c[0][0], c[1][1] - c[0][1]) > 0.005]
    return _decor(na.strokes(slide, cut, color, w=w, alpha=alpha), reason, overlap)


def _rect(x, y, w, h):
    return [(x, y), (x + w, y), (x + w, y + h), (x, y + h)]


def _rrect(x, y, w, h, r, n=6, top_only=False):
    """A rounded rectangle's outline (screen y down); top_only rounds the top two corners."""
    r = max(0.0, min(r, w / 2.0, h / 2.0))
    if r <= 1e-4:
        return _rect(x, y, w, h)
    pts = []
    corners = ((x + w - r, y + r, -90.0), (x + w - r, y + h - r, 0.0), (x + r, y + h - r, 90.0), (x + r, y + r, 180.0))
    for i, (cx, cy, a0) in enumerate(corners):
        if top_only and i in (1, 2):
            pts.append((x + w, y + h) if i == 1 else (x, y + h))
            continue
        for j in range(n + 1):
            a = math.radians(a0 + 90.0 * j / n)
            pts.append((cx + r * math.cos(a), cy + r * math.sin(a)))
    return pts


def _rot(pts, cx, cy, ang):
    ca, sa = math.cos(ang), math.sin(ang)
    return [(cx + (px - cx) * ca - (py - cy) * sa, cy + (px - cx) * sa + (py - cy) * ca) for px, py in pts]


def _circle(cx, cy, r, n=28, ry=None):
    ry = r if ry is None else ry
    return [(cx + r * math.cos(2 * math.pi * i / n), cy + ry * math.sin(2 * math.pi * i / n)) for i in range(n)]


def _recolour(slide, draw, color):
    """Run a planned draw, then set every run it drew in `color` (text set on the logo banner)."""
    n0 = len(slide.shapes)
    draw()
    shapes = list(slide.shapes)[n0:]
    for sh in shapes:
        if getattr(sh, "has_text_frame", False):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = dk.RGBColor.from_string(color)
    return shapes


def _lint_foot(sh):
    """Where the delivery lint reckons a text frame's words end: its top, 0.04in, then its own line estimate x size x
    1.25 (1.4 for CJK) — lower than the glyphs' box when it counts a line more (a Chinese quote whose 。 it wraps
    alone), and a banner or card round the words must clear it."""
    bb = dk._bbox_in(sh)
    r = dk._ink_rect(sh, bb) if bb else None
    if not r:
        return 0.0
    tf = sh.text_frame
    size, lines = r[2][0], r[2][1]
    try:
        import lint_deck as _ld
        paras = [[(r_.text, r_.font.size.pt if r_.font.size else size) for r_ in p_.runs] for p_ in tf.paragraphs
                 if p_.runs]
        faces = {}
        for p_ in tf.paragraphs:
            for r_ in p_.runs:
                if r_.font.name and r_.text:
                    faces[(r_.font.name, bool(r_.font.bold))] = faces.get((r_.font.name, bool(r_.font.bold)), 0) + \
                        len(r_.text)
        face, bold = max(faces, key=faces.get) if faces else (None, False)
        lines = max(lines, _ld._est_lines(paras, bb[2], face, bold))
        size = max([size] + [z for pr in paras for _t, z in pr])
    except Exception:          # the lint's own model is a refinement; the glyph box above still stands
        pass
    return bb[1] + 0.04 + lines * size / 72.0 * (1.4 if dk._has_cjk(tf.text) else 1.25)


def _ink_box(sh):
    bb = dk._bbox_in(sh)
    r = dk._ink_rect(sh, bb) if bb else None
    return r[0] if r else bb


# ─────────────────────────────────────────── the pieces ───────────────────────────────────────────
_PIPS = {1: ((0, 0),), 2: ((-1, -1), (1, 1)), 3: ((-1, -1), (0, 0), (1, 1)), 4: ((-1, -1), (1, -1), (-1, 1), (1, 1)),
         5: ((-1, -1), (1, -1), (0, 0), (-1, 1), (1, 1)), 6: ((-1, -1), (1, -1), (-1, 0), (1, 0), (-1, 1), (1, 1))}


def _die(k, slide, cx, cy, size, face, ang=0.0):
    """A die seen from above at (cx, cy): its rounded top face turned by `ang`, its thickness showing below-right,
    the pips of `face`."""
    top = _rot(_rrect(cx - size / 2.0, cy - size / 2.0, size, size, size * 0.20), cx, cy, ang)
    d = size * 0.13
    side = _poly(slide, [(px + d * 0.45, py + d) for px, py in top], k.P["die_side"],
                 "the side of a drawn die, its thickness: a game piece, nothing to read", overlap=True)
    if side is not None:
        na.soft_shadow(side, blur=0.06, dist=0.03, alpha=0.30, color=k.P["shadow"])
    _poly(slide, top, k.P["die"], "the top face of a drawn die: a game piece, nothing to read", line=k.P["die_side"],
          line_w=0.5, overlap=True)
    off, pr = size * 0.26, size * 0.085
    pips = [_circle(*_rot([(cx + dx * off, cy + dy * off)], cx, cy, ang)[0], pr, 14) for dx, dy in _PIPS[face]]
    _mpoly(slide, pips, k.P["pip"], "the pips on a drawn die: a game piece, they count nothing on the page",
           overlap=True)


_PAWN = ((0.000, 0.40), (0.030, 0.42), (0.085, 0.42), (0.115, 0.38), (0.150, 0.31), (0.230, 0.245), (0.360, 0.19),
         (0.500, 0.155), (0.575, 0.150), (0.600, 0.215), (0.640, 0.228), (0.675, 0.195), (0.700, 0.115))


def _gradient(sh, stops, angle=0.0):
    """Swap a shape's solid fill for a linear gradient: stops = [(pos 0..1, colour)], `angle` in degrees (0: left to
    right). The pawn's roundness in ONE shape — a separate shade and highlight were 'sibling cards' to the lint."""
    from lxml import etree
    from pptx.oxml.ns import qn
    sp = sh._element.spPr
    sf = sp.find(qn("a:solidFill"))
    if sf is None:
        return sh
    g = etree.fromstring(
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1"><a:gsLst>{}'
        '</a:gsLst><a:lin ang="{}" scaled="0"/></a:gradFill>'.format(
            "".join('<a:gs pos="{}"><a:srgbClr val="{}"/></a:gs>'.format(int(p * 100000), na.hexstr(c))
                    for p, c in stops), na.angle(angle)))
    sf.addprevious(g)
    sp.remove(sf)
    return sh


def _pawn(k, slide, cx, base_y, h, *, fill=None):
    """A pawn standing with its base's centre on (cx, base_y), `h` tall: its shadow on the board, then the body (base,
    bell, collar) and the round head as ONE shape, lit from the left and shaded on the right, a glint on its head."""
    c, cd, cl = fill or k.P["pawn"], k.P["pawn_dark"], k.P["pawn_light"]
    _poly(slide, _circle(cx + 0.06 * h, base_y + 0.005 * h, 0.50 * h, 32, 0.085 * h), k.P["shadow"],
          "the pawn's shadow on the board; nothing to read", alpha=0.22, overlap=True)
    right = [(cx + hw * h, base_y - t * h) for t, hw in _PAWN]
    left = [(cx - hw * h, base_y - t * h) for t, hw in reversed(_PAWN)]
    hy = base_y - 0.835 * h
    piece = _mpoly(slide, [right + left, _circle(cx, hy, 0.165 * h, 32)], c,
                   "a drawn pawn: a game piece, nothing to read", overlap=True)
    if piece is not None:
        _gradient(piece, [(0.0, cl), (0.30, c), (0.62, c), (1.0, cd)])
    _poly(slide, _circle(cx - 0.055 * h, hy - 0.055 * h, 0.045 * h, 16), cl, "light catching the pawn's head",
          alpha=0.85, overlap=True)


def _arrow(k, slide, x, y, w, h, color, *, point="left"):
    """The start corner's big arrow inside (x, y, w, h), pointing along the track."""
    hw = min(w * 0.42, h * 1.0)
    if point == "left":
        pts = [(x, y + h / 2), (x + hw, y), (x + hw, y + h * 0.30), (x + w, y + h * 0.30), (x + w, y + h * 0.70),
               (x + hw, y + h * 0.70), (x + hw, y + h)]
    else:
        pts = [(x + w, y + h / 2), (x + w - hw, y), (x + w - hw, y + h * 0.30), (x, y + h * 0.30), (x, y + h * 0.70),
               (x + w - hw, y + h * 0.70), (x + w - hw, y + h)]
    _poly(slide, pts, color, "the start corner's arrow: the way round the board, nothing to read", overlap=True)


def _star(k, slide, cx, cy, r, color):
    pts = [(cx + (r if i % 2 == 0 else r * 0.42) * math.cos(-math.pi / 2 + math.pi * i / 5),
            cy + (r if i % 2 == 0 else r * 0.42) * math.sin(-math.pi / 2 + math.pi * i / 5)) for i in range(10)]
    _poly(slide, pts, color, "a star printed on a corner of the board; ornament, nothing to read", line=k.P["line"],
          line_w=0.5, overlap=True)


def _hop(k, slide, p0, p1, lift, *, side=None):
    """A pawn's hop from p0 to p1: a dotted arc lifted by `lift` (towards `side`, a unit normal; default up) and an
    arrowhead where it lands."""
    W, H, s, o = ctx(k)
    (x0, y0), (x1, y1) = p0, p1
    nx, ny = side or (0.0, -1.0)
    c = ((x0 + x1) / 2.0 + nx * lift * 2.0, (y0 + y1) / 2.0 + ny * lift * 2.0)
    pts = []
    for i in range(25):
        t = i / 24.0
        u = 1 - t
        pts.append((u * u * x0 + 2 * u * t * c[0] + t * t * x1, u * u * y0 + 2 * u * t * c[1] + t * t * y1))
    ln = na.polyline(slide, pts[:-2], k.P["hop"], w=1.6 * s)
    if ln is not None:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
        _decor(ln, "the pawn's hop from square to square: the order the caller gave, drawn", overlap=True)
    a = math.atan2(y1 - pts[-4][1], x1 - pts[-4][0])
    L, Wd = 0.21 * s, 0.075 * s
    ca, sa = math.cos(a), math.sin(a)
    head = [(x1, y1), (x1 - ca * L - sa * Wd, y1 - sa * L + ca * Wd), (x1 - ca * L + sa * Wd, y1 - sa * L - ca * Wd)]
    _poly(slide, head, k.P["hop"], "the arrowhead where the pawn's hop lands", overlap=True)


def _card_back(k, slide, x, y, w, h, ang):
    """A face-down card from the deck at (x, y, w, h) turned by `ang`: its coloured back, an inner frame, a lozenge."""
    W, H, s, o = ctx(k)
    cx, cy = x + w / 2.0, y + h / 2.0
    r = 0.14 * s
    body = _poly(slide, _rot(_rrect(x, y, w, h, r), cx, cy, ang), k.P["card_back"],
                 "a face-down card from the deck under the one drawn; a game piece, nothing to read", overlap=True)
    if body is not None:
        na.soft_shadow(body, blur=0.10, dist=0.04, alpha=0.25, color=k.P["shadow"])
    m = 0.14 * s
    fr = _rot(_rrect(x + m, y + m, w - 2 * m, h - 2 * m, r * 0.6), cx, cy, ang)
    _decor(na.polyline(slide, fr + fr[:1], k.P["card_back_ink"], w=0.9), "the printed frame on the card's back",
           overlap=True)
    dm = min(w, h) * 0.16
    _poly(slide, _rot([(cx, cy - dm), (cx + dm * 0.7, cy), (cx, cy + dm), (cx - dm * 0.7, cy)], cx, cy, ang),
          k.P["card_back_ink"], "the lozenge printed on the card's back", alpha=0.85, overlap=True)


def _card_face(k, slide, x, y, w, h, band):
    """A face-up card: its paper, a thin edge, a coloured header band across its top."""
    W, H, s, o = ctx(k)
    r = 0.14 * s
    body = _poly(slide, _rrect(x, y, w, h, r), k.P["panel"],
                 "the card the words are printed on; the words carry the meaning", line=k.P["line"], line_w=0.6)
    if body is not None:
        na.soft_shadow(body, blur=0.12, dist=0.05, alpha=0.28, color=k.P["shadow"])
    hb = 0.40 * s
    _poly(slide, _rrect(x, y, w, hb, r, top_only=True), band, "the card's coloured header band; it only marks the card",
          line=k.P["line"], line_w=0.6)
    return hb


def _square(k, slide, x, y, w, h, band, side, bd, *, reason="a square of the board the words are printed on"):
    """One square of the track at (x, y, w, h): its tile and printed edge, its colour band on `side` (None: plain)."""
    _poly(slide, _rect(x, y, w, h), k.P["square"], reason, line=k.P["line"], line_w=1.0)
    if band:
        br = {"top": (x, y, w, bd), "bottom": (x, y + h - bd, w, bd), "left": (x, y, bd, h),
              "right": (x + w - bd, y, bd, h)}[side]
        _poly(slide, _rect(*br), band, "the square's colour band: the property group it belongs to, nothing to read",
              line=k.P["line"], line_w=1.0)


def _badge(k, slide, page, cx, cy, d, text):
    """PLAN a numbered token (a disc in the card's colour) at (cx, cy) on a square's band: (rects, draw)."""
    W, H, s, o = ctx(k)
    nr, nd = _v2.flow(k, slide, page, (cx - d * 0.62, cy - d * 0.5, d * 1.24, d), [("item_no", text)], anchor="middle",
                      align="c")

    def draw():
        _poly(slide, _circle(cx, cy, d / 2.0, 32), k.P["panel"], "the token carrying the square's number",
              line=k.P["line"], line_w=0.9)
        nd()
    return nr, draw


# ─────────────────────────────────────────── the board's track ───────────────────────────────────────────
_PT_BANDS = (4, 1, 5, 6, 3, 2, 7, 0)    # the band colours the points' squares take, in turn: red, sky, yellow, green…


def _ring_pattern(total, bands, start=0):
    """The band colour of each square in playing order: the property groups in turn (two or three squares of a
    colour), a plain square after each group — the rhythm of a real board, never one colour for a whole side."""
    out, i = [], start
    sizes = (2, 3, 3, 3, 3, 3, 3, 2)
    while len(out) < total:
        j = i % len(bands)
        out += [bands[j]] * sizes[j % len(sizes)] + [None]
        i += 1
    return out[:total]


def _band_rect(rect, side, bd):
    x, y, w, h = rect
    return {"top": (x, y, w, bd), "bottom": (x, y + h - bd, w, bd), "left": (x, y, bd, h),
            "right": (x + w - bd, y, bd, h)}[side]


def _band_edge(band, side):
    """The printed line between a band and the rest of its square."""
    bx, by, bw, bh = band
    return {"top": ((bx, by + bh), (bx + bw, by + bh)), "bottom": ((bx, by), (bx + bw, by)),
            "left": ((bx + bw, by), (bx + bw, by + bh)), "right": ((bx, by), (bx, by + bh))}[side]


def _track(k, slide, cells, *, emblems=True, bd=None, why="the board's track of squares; nothing to read in it"):
    """Squares with no words on them, as few shapes as possible: cells = [(rect, band side, colour or None)] — every
    tile in ONE shape, each colour's bands in one shape, every printed line in one; a plain square (no colour) gets a
    small mark, a face-down card or a ring, as a board's special squares do."""
    W, H, s, o = ctx(k)
    tiles, bands, segs, plain = [], {}, [], []
    for rect, side, col in cells:
        x, y, w, h = rect
        tiles.append(_rect(x, y, w, h))
        segs += [((x, y), (x + w, y)), ((x + w, y), (x + w, y + h)), ((x, y + h), (x + w, y + h)), ((x, y), (x, y + h))]
        if col:
            b = _band_rect(rect, side, bd or 0.25 * (h if side in ("top", "bottom") else w))
            # one shape per colour PER SIDE (a colour's bands on two sides spanned the board's centre, and the lint
            # read that box as a card round the words); on an upright side one per square — runs of two and three
            # facing each other across the board read to the lint as a row of cards of uneven heights
            key = (col, side) if side in ("top", "bottom") else (col, side, len(bands))
            bands.setdefault(key, []).append(_rect(*b))
            segs.append(_band_edge(b, side))
        elif col is None and side:
            plain.append((rect, side))
    _mpoly(slide, tiles, k.P["square"], why)
    for key, polys in bands.items():
        _mpoly(slide, polys, key[0], "the colour bands of one property group on the board; nothing to read in them")
    _segs(slide, segs, k.P["line"], "the board's printed lines round its squares; nothing to read in them",
          w=0.9 * max(0.8, s), alpha=0.85 if not _dark(k) else 0.60)
    if not emblems:
        return
    for j, ((rx, ry, rw, rh), side) in enumerate(plain):
        cx, cy = rx + rw / 2.0, ry + rh / 2.0
        u = min(rw, rh, 0.85 * s)
        if j % 2 == 0:                       # a small face-down card, laid across the square
            cw, ch = (u * 0.40, u * 0.56) if side in ("top", "bottom") else (u * 0.56, u * 0.40)
            _poly(slide, _rrect(cx - cw / 2, cy - ch / 2, cw, ch, u * 0.07), k.P["card_back"],
                  "a small card printed on a plain square of the board; nothing to read", alpha=0.9, overlap=True)
        else:                                # a ring
            ring = _circle(cx, cy, u * 0.21, 28)
            _decor(na.polyline(slide, ring + ring[:1], k.P["mute"], w=1.3 * max(0.8, s)),
                   "a ring printed on a plain square of the board; nothing to read", overlap=True)


def _ring(k, slide, x0, y0, x1, y1, D, *, sq=0.80):
    """The board's ring of squares inside (x0, y0)-(x1, y1), `D` deep: the four corner squares, the property squares
    with their colour bands facing the centre, the plain squares with a small mark. Returns the corners' rects."""
    nh = max(3, int(round((x1 - x0 - 2 * D) / (sq * D))))
    nv = max(2, int(round((y1 - y0 - 2 * D) / (sq * D))))
    wh, wv = (x1 - x0 - 2 * D) / nh, (y1 - y0 - 2 * D) / nv
    corners = {"br": (x1 - D, y1 - D, D, D), "bl": (x0, y1 - D, D, D), "tl": (x0, y0, D, D), "tr": (x1 - D, y0, D, D)}
    cells = [(r_, "", "") for r_ in corners.values()]
    order = [((x1 - D - (i + 1) * wh, y1 - D, wh, D), "top") for i in range(nh)] + \
            [((x0, y1 - D - (i + 1) * wv, D, wv), "right") for i in range(nv)] + \
            [((x0 + D + i * wh, y0, wh, D), "bottom") for i in range(nh)] + \
            [((x1 - D, y0 + D + i * wv, D, wv), "left") for i in range(nv)]
    cells += [(r_, side, col) for (r_, side), col in zip(order, _ring_pattern(len(order), k.P["bands"]))]
    _track(k, slide, cells)
    return corners


def _edge_run(k, slide, x, y, D, step, n, direction, band_side, start=0, *, bd=None):
    """A straight run of `n` squares of the board, `D` deep and `step` wide, from (x, y) going `direction` ("right",
    "left", "down"); the bands on `band_side`. Returns the cells (for the caller's own marks)."""
    cells = []
    pattern = _ring_pattern(n, k.P["bands"], start)
    for i in range(n):
        if direction == "right":
            r_ = (x + i * step, y, step, D)
        elif direction == "left":
            r_ = (x - (i + 1) * step, y, step, D)
        else:
            r_ = (x, y + i * step, D, step)
        cells.append((r_, band_side, pattern[i]))
    _track(k, slide, cells, bd=bd)
    return cells


def _dark(k):
    import vl_contract as vc
    return vc.is_dark(k.P)


def _text_boxes(slide, n0):
    return [b for b in (_ink_box(sh) for sh in list(slide.shapes)[n0:]
                        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()) if b]


def _dice_pair(k, slide, cx, cy, d, faces=(3, 5)):
    """Two dice just thrown, turned the same amount either way: equal outlines, so the lint never reads them as two
    cards of a row sized unevenly."""
    _die(k, slide, cx - 0.56 * d, cy + 0.10 * d, d, faces[0], -0.26)
    _die(k, slide, cx + 0.62 * d, cy - 0.12 * d, d, faces[1], 0.26)


def _place_dice(k, slide, area, spots, d, taken, pad):
    """The pair of dice at the first of `spots` (centres) whose box stays in `area` and clear of `taken`."""
    ax, ay, aw, ah = area
    for cx, cy in spots:
        box = (cx - 1.25 * d, cy - 0.80 * d, 2.55 * d, 1.65 * d)
        if box[0] >= ax and box[1] >= ay and box[0] + box[2] <= ax + aw and box[1] + box[3] <= ay + ah \
                and not any(_v2.meet(box, t_, pad) for t_ in taken):
            _dice_pair(k, slide, cx, cy, d)
            return box
    return None


# ─────────────────────────────────────────── the pages ───────────────────────────────────────────
@register(NAME, "cover", alts=3)
def _bg_cover(k, slide, f, image):
    """The whole board: the ring of squares round the page; in its centre the kicker, the title on the logo banner and
    the subtitle, a pair of dice beside them; the pawn on the start corner, its arrow pointing the way round."""
    W, H, s, o = ctx(k)
    e = 0.16 * s
    land = o == "land"                       # a square or portrait board gives its narrower centre more room
    D = ((0.92, 0.72, 0.58) if land else (0.70, 0.60, 0.50))[alt()] * s
    m = ((0.40, 0.30, 0.24) if land else (0.26, 0.22, 0.18))[alt()] * s
    px, py = (0.42 if land else 0.30) * s, 0.22 * s
    x0, y0 = e + D + m, e + D + m
    cw, ch = W - 2 * x0, H - 2 * y0
    groups = [(_v2.fields(f, ("kicker",), ("kicker",)), 0.22 * s + py),
              (_v2.fields(f, ("title",)), py + 0.30 * s),
              (_v2.fields(f, ("subtitle",)), 0.0)]
    live = [g for g in groups if g[0]]
    ti = next((i for i, g in enumerate(live) if g[0][0][0] == "title"), None)
    r, ds = _v2.stack(k, slide, "cover", x0 + px, cw - 2 * px, y0 + py, y0 + ch - py, groups, anchor="middle",
                      align="c", **({} if land else {"start": {"title": 44 * s}}))
    corners = _ring(k, slide, e, e, W - e, H - e, D)
    sx, sy, _w, _h = corners["br"]               # the start corner: its arrow, the pawn standing on it
    _arrow(k, slide, sx + 0.10 * D, sy + 0.64 * D, 0.80 * D, 0.26 * D, k.P["arrow"])
    _pawn(k, slide, sx + 0.50 * D, sy + 0.60 * D, 0.74 * D)
    tlx, tly, _w, _h = corners["tl"]
    _star(k, slide, tlx + D / 2, tly + D / 2, 0.30 * D, k.P["bands"][5])
    trx, try_, _w, _h = corners["tr"]
    _die(k, slide, trx + D / 2, try_ + D / 2, 0.44 * D, 5, 0.25)
    blx, bly, _w, _h = corners["bl"]
    for j, (dx, dy, a_) in enumerate(((0.26, 0.18, -0.16), (0.34, 0.24, 0.10))):
        _poly(slide, _rot(_rrect(blx + dx * D, bly + dy * D, 0.36 * D, 0.52 * D, 0.05 * D), blx + D / 2, bly + D / 2,
                          a_), (k.P["card_back"], k.P["bands"][3])[j],
              "a stack of cards printed on a corner of the board; nothing to read", line=k.P["line"], line_w=0.5,
              overlap=True)
    n0 = len(slide.shapes)
    for i, d in enumerate(ds):
        if i != ti:
            d()
            continue
        shapes = _recolour(slide, d, k.P["banner_ink"])
        boxes = [b for b in (_ink_box(sh_) for sh_ in shapes if getattr(sh_, "has_text_frame", False)) if b]
        if not boxes:
            continue
        bx0, by0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
        bx1, by1 = max(b[0] + b[2] for b in boxes), max(b[1] + b[3] for b in boxes)
        by1 = max([by1] + [_lint_foot(sh_) for sh_ in shapes if getattr(sh_, "has_text_frame", False)])
        bw = min(cw, max(bx1 - bx0 + 2 * px, 0.50 * cw))
        cxm = r["title"][0] + r["title"][2] / 2.0
        bx, by, bh = cxm - bw / 2.0, by0 - py, by1 - by0 + 2 * py
        b = _poly(slide, _rect(bx, by, bw, bh), k.P["banner"], "the logo banner the title is printed on; the title "
                  "carries the meaning", line=k.P["banner_edge"], line_w=1.6 * max(0.8, s))
        if b is None:
            continue
        na.soft_shadow(b, blur=0.0, dist=0.07 * s, alpha=0.9, color=k.P["banner_edge"], direction=45.0)
        q = 0.07 * s
        fr = _rect(bx + q, by + q, bw - 2 * q, bh - 2 * q)
        ln = _decor(na.polyline(slide, fr + fr[:1], k.P["banner_ink"], w=0.8, alpha=0.6),
                    "the hairline frame printed inside the logo banner")
        shapes[0]._element.addprevious(b._element)
        if ln is not None:
            shapes[0]._element.addprevious(ln._element)
        r["banner"] = (bx, by, bw, bh)
    taken = _text_boxes(slide, n0) + ([r["banner"]] if "banner" in r else [])
    dsz = (0.66, 0.54, 0.46)[alt()] * s
    area = (e + D + 0.12 * s, e + D + 0.12 * s, W - 2 * (e + D + 0.12 * s), H - 2 * (e + D + 0.12 * s))
    ax, ay, aw, ah = area
    spots = [(ax + aw - 1.35 * dsz, ay + ah - 0.92 * dsz), (ax + 1.35 * dsz, ay + ah - 0.92 * dsz),
             (ax + aw - 1.35 * dsz, ay + 0.92 * dsz), (ax + 1.35 * dsz, ay + 0.92 * dsz)]
    _place_dice(k, slide, area, spots, dsz, taken, 0.18 * s)
    r.pop("banner", None)
    return r


@register(NAME, "section", alts=3)
def _bg_section(k, slide, f, image):
    """A corner of the board seen close: one big corner square at the top left carrying the section number, the
    board's two edges running off the page from it; the kicker and the title in the board's centre."""
    W, H, s, o = ctx(k)
    e = 0.16 * s
    C = ((0.34, 0.29, 0.24) if o == "land" else (0.27, 0.23, 0.20))[alt()] * min(W, H)
    step = 0.62 * C
    num = text_of(f, "number")
    m = 0.50 * s
    tx0 = e + C + m
    r, ds = _v2.stack(k, slide, "section", tx0, W - tx0 - 0.55 * s, e + C + 0.35 * s, H - 0.50 * s,
                      [(_v2.fields(f, ("kicker",), ("kicker",)), 0.15 * s), (_v2.fields(f, ("title",)), 0.0)],
                      anchor="middle")
    _edge_run(k, slide, e + C, e, C, step, int(math.ceil((W - e - C) / step)), "right", "bottom", start=4)
    _edge_run(k, slide, e, e + C, C, step, int(math.ceil((H - e - C) / step)), "down", "right", start=2)
    _poly(slide, _rect(e, e, C, C), k.P["banner"], "the corner square the section number is printed on",
          line=k.P["line"], line_w=1.2 * max(0.8, s))
    q = 0.09 * C
    fr = _rect(e + q, e + q, C - 2 * q, C - 2 * q)
    _decor(na.polyline(slide, fr + fr[:1], k.P["banner_ink"], w=1.0, alpha=0.55),
           "the frame printed inside the corner square")
    if num:
        pad = 0.18 * C
        nr, nd = _v2.flow(k, slide, "section", (e + pad, e + pad, C - 2 * pad, C - 2 * pad), [("number", num)],
                          anchor="middle", align="c", accent=k.P["banner_ink"],
                          start={"number": (120 if alt() == 0 else 96) * s})
        r.update(nr)
        ds.append(_centred(slide, nd))
    _run_all(ds)
    return r


@register(NAME, "image_text", alts=3)
def _bg_image_text(k, slide, f, image):
    """The caller's picture as a card drawn from the deck — a card with a coloured header band, the picture on it, the
    caption under the picture; two face-down cards of the deck under it. The kicker, title and body beside it (under
    it on a portrait or square page)."""
    W, H, s, o = ctx(k)
    a = alt()
    cap = text_of(f, "caption")
    hb, m = 0.40 * s, 0.20 * s
    land = o == "land"
    if land:
        cx0, cw = 0.08 * W, (0.42, 0.37, 0.32)[a] * W
        pw = cw - 2 * m
        ph = min(pw * 0.74, (0.54 if cap else 0.62) * H)
    else:
        cw = (0.78, 0.74, 0.70)[a] * W
        cx0 = (W - cw) / 2.0
        pw = cw - 2 * m
        ph = min(pw * 0.66, (0.30, 0.24, 0.19)[a] * H)
    cap_h = 0.0
    if cap:
        cap = _unwidow(k, "caption", cap, vl.TYPE[NAME]["caption"][0] * s, pw)
        cr0, _cd0 = _v2.flow(k, slide, "image_text", (0, 0, pw, 1.1 * s), [("caption", cap)])
        cap_h = max(v[1] + v[3] for v in cr0.values()) + 0.10 * s
    card_h = hb + m + ph + (0.10 * s + cap_h if cap else 0.0) + m
    cy0 = (H - card_h) / 2.0 if land else 0.07 * H
    img = (cx0 + m, cy0 + hb + m, pw, ph)
    r, ds = {}, []
    if cap:
        cr, cd = _v2.flow(k, slide, "image_text", (img[0], img[1] + ph + 0.10 * s, pw, cap_h + 0.05 * s),
                          [("caption", cap)])
        r.update(cr)
        ds.append(cd)
    if land:
        tx = cx0 + cw + 0.80 * s
        col = (tx, 0.12 * H, W - 0.07 * W - tx, 0.76 * H)
    else:
        ty = cy0 + card_h + 0.45 * s
        col = (0.10 * W, ty, 0.80 * W, H - 0.55 * s - ty)
    tr, tds = _v2.stack(k, slide, "image_text", col[0], col[2], col[1], col[1] + col[3],
                        [(_v2.fields(f, ("kicker",), ("kicker",)), 0.12 * s), (_v2.fields(f, ("title",)), 0.18 * s),
                         (_fx(k, f, ("body",), col[2]), 0.0)], anchor="middle")
    r.update(tr)
    for dx, dy in ((0.24 * s, -0.16 * s), (0.11 * s, -0.06 * s)):     # one turn for both: equal outlines (a row of
        _card_back(k, slide, cx0 + dx, cy0 + dy, cw, card_h, 0.045)    # uneven 'cards' to the lint otherwise)
    _card_face(k, slide, cx0, cy0, cw, card_h, k.P["bands"][6])
    vl._place_image(k, slide, image, img, vl.L_((0, 0, 1, 1), "frame", None), "image_text")
    _run_all(tds + ds)
    return r


def _ordered(f):
    v = f.get("ordered")
    if v not in (None, True, False):
        raise ValueError("boardgame.points(): ordered= takes True or False, got {!r}".format(v))
    return bool(v)


def _unwidow(k, field, text, size, w):
    """Latin words a little too many for their last line ("…and that is the whole / evening."): set in the same number
    of lines at a narrower, balanced measure, as explicit lines — the caller's words, only broken elsewhere. Judged at
    the width the box is DRAWN at (what the renderer wraps to), balanced within the width it is MEASURED in. CJK keeps
    its clause breaks (vl_native2.flow)."""
    if not text or "\n" in text or dk._has_cjk(text):
        return text
    m = w * vl.HEADROOM
    drawn = vl._break_lines(k, field, text, size, w)
    if len(drawn) < 2 or not (vl._widowed(k, field, text, size, w) or vl._widowed(k, field, text, size, m)):
        return text
    for f_ in (1.0, 0.94, 0.88, 0.82, 0.76, 0.70, 0.64):
        ls2 = vl._break_lines(k, field, text, size, m * f_)
        if len(ls2) > len(drawn):
            break
        if len(ls2) == len(drawn) and not vl._widowed(k, field, text, size, m * f_) \
                and vl._lines_hold(k, field, text, ls2, size, m) and " ".join(ls2).split() == text.split():
            return "\n".join(l_.strip() for l_ in ls2)
    return text


def _fx(k, f, names, w):
    """vl_native2.fields with the smaller fields (an attribution, a note, a caption, a body, a label, a line) balanced
    away from a lone last word at their design size; a title's widow is the flow's own business."""
    W, H, s, o = ctx(k)
    out = []
    for n_, t in _v2.fields(f, names):
        if n_ not in ("title", "quote", "kicker", "mark", "number") and n_ in vl.TYPE[NAME]:
            t = _unwidow(k, n_, t, vl.TYPE[NAME][n_][0] * s, w)
        out.append((n_, t))
    return out


def _balanced_items(k, items, probe, width):
    """The points' words with their widows balanced away, at the sizes the probe decided."""
    sizes = {f_: z for _pr, d_ in probe for f_, z in getattr(d_, "sizes", {}).items()}
    out = []
    for it in items:
        out.append([(f_, _unwidow(k, f_, t, sizes.get(f_, vl.TYPE[NAME][f_][0]), width)) for f_, t in it])
    return out, sizes


@register(NAME, "points", alts=4)
def _bg_points(k, slide, f, image):
    """A stretch of the board's track: each point a square with its colour band, its head and its line. ordered=True:
    the squares are numbered on tokens, the pawn stands on the first and dotted hops lead square to square. A wide page
    sets the squares in a row across it, the track running on off the page at both ends; four squares too narrow for
    a row make a two-by-two block (a loop 1 2 / 4 3 when ordered); a portrait or square page stacks them as the board's
    side edge. Long copy: a smaller title, then less padding, then the block."""
    W, H, s, o = ctx(k)
    pts = points_of(f.get("items"))
    n = len(pts)
    ordered = _ordered(f)
    a = alt()
    land = o == "land"
    mx = 0.07 * W if land else 0.08 * W
    w_all = W - 2 * mx
    row_ok = n <= 3 or w_all / n >= 2.35 * s
    mode = ("row" if row_ok and a < 3 else "grid") if land else ("col" if a < 3 or n < 3 else "grid")
    r, ds = _v2.stack(k, slide, "points", mx, w_all, 0.07 * H, (0.30, 0.33, 0.36, 0.36)[a] * H,
                      [(_v2.fields(f, ("kicker",), ("kicker",)), 0.08 * s), (_v2.fields(f, ("title",)), 0.0)],
                      anchor="top", start=None if a == 0 else {"title": (36 if a == 1 else 28) * s})
    head_foot = max((v[1] + v[3] for v in r.values()), default=0.07 * H)
    pawn_h = (0.80 if a < 2 else 0.66) * s if ordered else 0.0
    top = head_foot + (0.34 if a < 2 else 0.24) * s + (pawn_h * 0.72 if ordered else 0.0)
    bottom = H - (0.42 if land else 0.50) * s
    pad = (0.20, 0.18, 0.13, 0.13)[a] * s
    tok = 0.40 * s
    items = [[("item_head", hd)] + ([("item_line", ln)] if ln else []) for hd, ln in pts]
    nos = [str(i + 1) for i in range(n)]
    cols = [k.P["bands"][_PT_BANDS[i % len(_PT_BANDS)]] for i in range(n)]
    art, marks = [], []
    if mode in ("row", "grid"):
        bd = (tok + 0.12 * s) if ordered else 0.28 * s
        ncol = 2 if mode == "grid" else n
        nrow = -(-n // ncol)
        loop = mode == "grid" and ordered
        w0 = w_all - (0.45 * s if loop else 0.0)               # a loop's hops run down its right side…
        bot = bottom - (0.40 * s if loop and n == 4 else 0.0)   # …and back under it
        sw = w0 / ncol if mode == "grid" else min(w0 / ncol, 3.0 * s)
        x0 = mx if mode == "grid" else (W - n * sw) / 2.0      # a short row sits mid-track, plain squares either side
        room = (bot - top) / nrow - bd - 2 * pad
        if room < 0.45 * s:
            raise vl.VLTextOverflow("boardgame.points(): the track has {:.2f}in for its words".format(room))
        probe = _v2.plan_together(k, slide, "points", [((0, 0, sw - 2 * pad, room), it, {"anchor": "top", "align": "c"})
                                                       for it in items])
        items, sizes = _balanced_items(k, items, probe, sw - 2 * pad)
        words = [max((v[1] + v[3] for v in pr.values()), default=0.0) for pr, _d in probe]
        need = [h_ + bd + 2 * pad for h_ in words]
        rows_h = [max(need[i] for i in range(rw * ncol, min(n, rw * ncol + ncol))) for rw in range(nrow)]
        if sum(rows_h) > bot - top + 1e-6:
            raise vl.VLTextOverflow("boardgame.points(): the squares need {:.2f}in, the page has {:.2f}in".format(
                sum(rows_h), bot - top))
        if mode == "row":                    # a board's square is about as tall as it is wide
            rows_h = [max(rows_h[0], min((bot - top) * 0.88, sw * 1.0))]
        y0 = top + (bot - top - sum(rows_h)) * (0.62 if mode == "row" else 0.45)
        order = [0, 1, 3, 2] if loop else list(range(4))        # ordered in a block: a loop, 1 2 / 4 3
        cells = []
        for i in range(n):
            ci, ri = order[i] % ncol, order[i] // ncol
            cells.append((x0 + ci * sw, y0 + sum(rows_h[:ri]), sw, rows_h[ri]))
        # the words of a row's squares start level (their heads aligned), the row's tallest centred in its square
        tall = [max(words[i] for i in range(n) if order[i] // ncol == order[j] // ncol) for j in range(n)]
        rects = [(cx + pad, cy + bd + pad + (ch - bd - 2 * pad - tall[i]) / 2.0, sw - 2 * pad, tall[i] + 0.02 * s)
                 for i, (cx, cy, _cw, ch) in enumerate(cells)]
        planned = _v2.plan_together(k, slide, "points", [(rc, it, {"anchor": "top", "align": "c", "start": sizes})
                                                         for rc, it in zip(rects, items)])
        if mode == "row":                    # the track runs on off the page at both ends
            th = rows_h[0]
            nl, nr_ = int(math.ceil(x0 / sw - 1e-6)), int(math.ceil((W - x0 - n * sw) / sw - 1e-6))
            art.append(lambda th=th, nl=nl: _edge_run(k, slide, x0, y0, th, sw, nl, "left", "top", start=7, bd=bd))
            art.append(lambda th=th, nr_=nr_: _edge_run(k, slide, x0 + n * sw, y0, th, sw, nr_, "right", "top", start=3,
                                                        bd=bd))
        for i, ((cx, cy, cw_, ch), (pr, pd)) in enumerate(zip(cells, planned)):
            r.update(pr)
            ds.append(pd)
            art.append(lambda cx=cx, cy=cy, cw_=cw_, ch=ch, i=i: _square(k, slide, cx, cy, cw_, ch, cols[i], "top", bd))
            if ordered:
                nr, nd = _badge(k, slide, "points", cx + pad + tok / 2.0, cy + bd / 2.0, tok, nos[i])
                r.update(nr)
                marks.append(nd)
        if ordered:
            c1 = cells[0]
            marks.append(lambda: _pawn(k, slide, c1[0] + c1[2] * 0.62, c1[1] + bd * 0.60, pawn_h))
            hops = []
            for i in range(n - 1):
                (ax, ay, aw, ah), (bx, by, bw_, bh_) = cells[i], cells[i + 1]
                if abs(ay - by) < 1e-6 and bx > ax:          # along the top, rightwards
                    hops.append(((ax + 0.80 * aw, ay - 0.05 * s), (bx + 0.36 * bw_, by - 0.05 * s), None))
                elif abs(ax - bx) < 1e-6:                    # down the right side
                    hops.append(((ax + aw + 0.05 * s, ay + 0.60 * ah), (bx + bw_ + 0.05 * s, by + 0.40 * bh_),
                                 (1.0, 0.0)))
                else:                                        # back along the foot, leftwards
                    hops.append(((ax + 0.36 * aw, ay + ah + 0.05 * s), (bx + 0.64 * bw_, by + bh_ + 0.05 * s),
                                 (0.0, 1.0)))
            for p0, p1, side in hops:
                marks.append(lambda p0=p0, p1=p1, side=side: _hop(k, slide, p0, p1, 0.20 * s, side=side))
    else:
        # the board's side edge: the squares stacked, each band on the left
        bw = (tok + 0.20 * s) if ordered else 0.30 * s
        x0 = mx + (0.42 * s if ordered else 0.0)
        w0 = W - mx - x0
        vpad = (0.15, 0.14, 0.11, 0.11)[a] * s                  # a row's words need less air above and below
        room = bottom - top - 2 * vpad * n
        probe = _v2.plan_together(k, slide, "points", [((0, 0, w0 - bw - 2 * pad, max(0.3 * s, room / n)), it,
                                                       {"anchor": "top"}) for it in items])
        items, sizes = _balanced_items(k, items, probe, w0 - bw - 2 * pad)
        orig = [max(max((v[1] + v[3] for v in pr.values()), default=0.0) + 2 * vpad, (0.95 if ordered else 0.70) * s)
                for pr, _d in probe]
        if sum(orig) > bottom - top + 1e-6:
            raise vl.VLTextOverflow("boardgame.points(): the squares need {:.2f}in, the page has {:.2f}in".format(
                sum(orig), bottom - top))
        even = min((bottom - top) * 0.86 / n, 2.10 * s)         # the side edge's squares, as even as the words allow
        need = [max(h_, even) for h_ in orig]
        while sum(need) > bottom - top + 1e-6 and even > min(orig):
            even *= 0.92
            need = [max(h_, even) for h_ in orig]
        if sum(need) > bottom - top + 1e-6:
            need = orig
        y = top + max(0.0, bottom - top - sum(need)) * 0.30
        cells = []
        for h_ in need:
            cells.append((x0, y, w0, h_))
            y += h_
        rects = [(cx + bw + pad, cy + vpad, cw_ - bw - 2 * pad, ch - 2 * vpad) for cx, cy, cw_, ch in cells]
        planned = _v2.plan_together(k, slide, "points", [(rc, it, {"anchor": "middle", "start": sizes})
                                                         for rc, it in zip(rects, items)])
        for i, ((cx, cy, cw_, ch), (pr, pd)) in enumerate(zip(cells, planned)):
            r.update(pr)
            ds.append(pd)
            art.append(lambda cx=cx, cy=cy, cw_=cw_, ch=ch, i=i:
                       _square(k, slide, cx, cy, cw_, ch, cols[i], "left", bw))
            if ordered:
                nr, nd = _badge(k, slide, "points", cx + bw / 2.0, cy + ch / 2.0, tok, nos[i])
                r.update(nr)
                marks.append(nd)
        if ordered:
            c1 = cells[0]
            marks.append(lambda: _pawn(k, slide, c1[0] + c1[2] - 0.75 * s, c1[1] + 0.10 * s, pawn_h))
            for i in range(n - 1):
                p0 = (cells[i][0] - 0.05 * s, cells[i][1] + cells[i][3] * 0.62)
                p1 = (cells[i + 1][0] - 0.05 * s, cells[i + 1][1] + cells[i + 1][3] * 0.40)
                marks.append(lambda p0=p0, p1=p1: _hop(k, slide, p0, p1, 0.16 * s, side=(-1.0, 0.0)))
    _run_all(art + marks + ds)
    return r


@register(NAME, "quote", alts=3)
def _bg_quote(k, slide, f, image):
    """The quote printed on a card drawn from the deck: a card with a coloured header band, the quote on it, the
    attribution under a short rule; the deck's face-down cards under it."""
    W, H, s, o = ctx(k)
    a = alt()
    land = o == "land"
    cw = (min(0.58 * W, 8.0 * s), min(0.70 * W, 9.4 * s), 0.84 * W)[a] if land else (0.80, 0.86, 0.90)[a] * W
    pad = (0.50, 0.42, 0.34)[a] * s
    hb = 0.40 * s
    mv = 0.50 * s
    x = (W - cw) / 2.0 - 0.10 * s          # the deck's cards peek out to the right: the pair sits centred
    has_a = bool(text_of(f, "attribution"))
    groups = [(_v2.fields(f, ("mark", "quote")), 0.36 * s), (_fx(k, f, ("attribution",), cw - 2 * pad), 0.0)]
    room = H - 2 * mv - hb - 2 * pad
    r0, d0 = _v2.stack(k, slide, "quote", x + pad, cw - 2 * pad, 0.0, room, groups, anchor="top", align="c")
    qz = {f_: z for d_ in d0 for f_, z in getattr(d_, "sizes", {}).items()}.get("quote")
    q = groups[0][0]
    if qz and q and q[-1][0] == "quote":     # "…and that is the whole / evening.": the quote balanced at its own size
        bal = _unwidow(k, "quote", q[-1][1], qz, cw - 2 * pad)
        if bal != q[-1][1]:
            groups = [(q[:-1] + [("quote", bal)], groups[0][1])] + groups[1:]
            r0, d0 = _v2.stack(k, slide, "quote", x + pad, cw - 2 * pad, 0.0, room, groups, anchor="top", align="c")
    used = max((v[1] + v[3] for v in r0.values()), default=0.0)
    card_h = hb + pad + used + pad
    y = (H - card_h) / 2.0
    # planned again in the SAME room, only lower: re-planned in a room cut to the words, a short quote that GREW in the
    # roomy probe was set smaller and left the card a blank inch under its attribution
    r, ds = _v2.stack(k, slide, "quote", x + pad, cw - 2 * pad, y + hb + pad, y + hb + pad + room, groups, anchor="top",
                      align="c")
    n0 = len(slide.shapes)
    _run_all(ds)
    texts = [sh for sh in list(slide.shapes)[n0:]
             if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
    near = (lambda f_: next((sh for sh in texts if f_ in r and abs(sh.top / 914400.0 - r[f_][1]) < 0.02), None))
    qsh, ash = near("quote"), near("attribution")
    if qsh is not None and ash is not None:
        # the flow sizes a long quote's box by the lint's line model, often a line or two longer than the lines it
        # sets: on a card that air sat between the quote and its attribution (a portrait page, 1.5in). The
        # attribution is drawn up to the quote's foot as the delivery lint reckons it, plus the planned gap
        # — and pushed DOWN to it as well: on other face metrics (GitHub's Linux runner) the quote's foot fell below
        # the planned attribution, and the rule set halfway between them crossed it (felt/4:3/zh, 2026-10-22)
        ib = _ink_box(qsh)
        want = max(_lint_foot(qsh), ib[1] + ib[3]) + 0.36 * s
        at = ash.top / 914400.0
        if abs(at - want) > 0.01:
            ash.top = int(round(want * 914400))
            r["attribution"] = (r["attribution"][0], want) + tuple(r["attribution"][2:])
    # the card reaches past the words by its padding as the delivery lint reckons them too (lines x size x 1.25/1.4);
    # cut to them it is centred again, the words moving with it
    foot = max(max(_lint_foot(sh), _ink_box(sh)[1] + _ink_box(sh)[3]) for sh in texts) if texts else y + card_h - pad
    new_h = foot + pad - y
    if new_h < card_h - 0.01:
        dy = (card_h - new_h) / 2.0
        for sh in texts:
            sh.top = int(round(sh.top + dy * 914400))
        r = {f_: (v[0], v[1] + dy) + tuple(v[2:]) for f_, v in r.items()}
        y += dy
    card_h = new_h
    n1 = len(slide.shapes)
    for dx, dy_ in ((0.26 * s, -0.17 * s), (0.12 * s, -0.06 * s)):     # one turn for both: equal outlines
        _card_back(k, slide, x + dx, y + dy_, cw, card_h, 0.045)
    _card_face(k, slide, x, y, cw, card_h, k.P["bands"][3])
    if has_a and qsh is not None and ash is not None:
        ib = _ink_box(qsh)
        yy = (max(_lint_foot(qsh), ib[1] + ib[3]) + ash.top / 914400.0) / 2.0
        _segs(slide, [((x + cw / 2.0 - 0.45 * s, yy), (x + cw / 2.0 + 0.45 * s, yy))], k.P["line"],
              "the short rule printed between the quote and its attribution", w=1.0)
    _behind(slide, n1, texts[0] if texts else None)
    return r


def _centred(slide, draw):
    """A planned draw whose text boxes are set middle-anchored: a figure shrunk to its square's width keeps the box
    planned for it, and set at the box's top it sat high in the square (1,250,000 on a 10in page)."""
    def go():
        n0 = len(slide.shapes)
        draw()
        for sh in list(slide.shapes)[n0:]:
            if getattr(sh, "has_text_frame", False):
                sh.text_frame.vertical_anchor = dk.MSO_ANCHOR.MIDDLE
    go.sizes = getattr(draw, "sizes", {})
    return go


def _behind(slide, n0, anchor):
    """Move the shapes drawn after the first n0 to just under `anchor` (the art sized to the words it carries)."""
    if anchor is None:
        return
    for sh in list(slide.shapes)[n0:]:
        anchor._element.addprevious(sh._element)


@register(NAME, "data", alts=3)
def _bg_data(k, slide, f, image):
    """The key figure on the square the pawn landed on: a big square of the board with its colour band, the number
    printed in it, the pawn standing on its band and the dice that brought it there tumbling at its foot; the label and
    the note beside it (under it, centred, on a portrait or square page). The square hugs the figure."""
    W, H, s, o = ctx(k)
    a = alt()
    land = o == "land"
    num = text_of(f, "number")
    bd, pad, foot = 0.50 * s, 0.32 * s, 0.22 * s
    pawn_h = 0.85 * s
    if land:
        ty, th = 0.20 * H, 0.64 * H
        tmin, tmax = 0.28 * W, (0.46, 0.40, 0.34)[a] * W
    else:
        ty = 0.05 * H + pawn_h * 0.75
        th = (0.40, 0.33, 0.27)[a] * H
        tmin, tmax = 0.50 * W, 0.80 * W
    r, ds = {}, []
    tw = tmin
    size = None
    if num:
        _nr, nd0 = _v2.flow(k, slide, "data", (0, 0, tmax - 2 * pad, th - bd - 2 * pad - foot), [("number", num)],
                            anchor="middle", align="c")
        size = getattr(nd0, "sizes", {}).get("number", vl.TYPE[NAME]["number"][0] * s)
        tw = min(tmax, max(tmin, _v2.ink_w(k, "number", num, size) + 2 * pad + 0.35 * s))
    tx = 0.08 * W if land else (W - tw) / 2.0
    if num:
        # figures have no descenders: centred by its line box a big number sits high in its square — lowered by a
        # tenth of its size (within the foot's room) it sits on the square's optical centre
        dz = min(0.10 * size / 72.0, foot)
        nr, nd = _v2.flow(k, slide, "data", (tx + pad, ty + bd + pad + dz, tw - 2 * pad, th - bd - 2 * pad - foot),
                          [("number", num)], anchor="middle", align="c", start={"number": size})
        r.update(nr)
        ds.append(_centred(slide, nd))
    if land:
        cx0 = tx + tw + 0.90 * s
        tr, td = _v2.flow(k, slide, "data", (cx0, 0.16 * H, W - 0.07 * W - cx0, 0.70 * H),
                          _fx(k, f, ("label", "note"), W - 0.07 * W - cx0), anchor="middle")
    else:
        cy0 = ty + th + 0.50 * s
        tr, td = _v2.flow(k, slide, "data", (0.10 * W, cy0, 0.80 * W, H - 0.55 * s - cy0),
                          _fx(k, f, ("label", "note"), 0.80 * W), anchor="top", align="c")
    r.update(tr)
    ds.append(td)
    _square(k, slide, tx, ty, tw, th, k.P["bands"][1], "top", bd, reason="the square of the board the key figure is "
            "printed on; the figure carries the meaning")
    _pawn(k, slide, tx + tw - 0.70 * s, ty + bd * 0.62, pawn_h)
    _dice_pair(k, slide, tx + tw - 0.62 * s, ty + th - 0.02 * s, 0.46 * s, faces=(4, 2))
    _run_all(ds)
    return r


def _checks(k, slide, x, y, w, n):
    """A chequered flag pattern filling the square (x, y, w): the finish."""
    u = w / n
    _poly(slide, _rect(x, y, w, w), k.P["check_b"], "the finish square's chequer, its light squares",
          line=k.P["line"], line_w=1.0)
    dark = [_rect(x + i * u, y + j * u, u, u) for i in range(n) for j in range(n) if (i + j) % 2 == 0]
    _mpoly(slide, dark, k.P["check_a"], "the finish square's chequer: the end of the game, nothing to read")


@register(NAME, "closing", alts=3)
def _bg_closing(k, slide, f, image):
    """The finish: the track along the foot of the page running into the chequered finish corner, the pawn arriving
    on it after its last hops; the title and its line in the board's centre above."""
    W, H, s, o = ctx(k)
    a = alt()
    land = o == "land"
    e = 0.16 * s
    C = ((0.30, 0.26, 0.22) if land else (0.30, 0.26, 0.23))[a] * min(W, H)
    step = 0.62 * C
    pawn_h = 0.74 * C
    cx0, cy0 = W - e - C, H - e - C
    rise = pawn_h - 0.60 * C + 0.35 * s
    x, w = (0.10, 0.08, 0.07)[a] * W, (0.80, 0.84, 0.86)[a] * W
    r, ds = _v2.stack(k, slide, "closing", x, w, 0.10 * H, cy0 - max(rise, 0.0) - 0.40 * s,
                      [(_v2.fields(f, ("title",)), 0.22 * s), (_fx(k, f, ("line",), w), 0.0)], anchor="middle",
                      align="c")
    n = int(math.ceil(cx0 / step))
    cells = _edge_run(k, slide, cx0, cy0, C, step, n, "left", "top", start=5)
    _checks(k, slide, cx0 + 0.10 * C, cy0 + 0.10 * C, 0.80 * C, 5)
    _segs(slide, [((cx0, cy0), (cx0 + C, cy0)), ((cx0 + C, cy0), (cx0 + C, cy0 + C)), ((cx0, cy0 + C),
                  (cx0 + C, cy0 + C)), ((cx0, cy0), (cx0, cy0 + C))], k.P["line"], "the finish corner's printed edge",
          w=0.9 * max(0.8, s))
    hops = [(cells[1][0][0] + 0.5 * step, cy0 - 0.05 * s), (cells[0][0][0] + 0.5 * step, cy0 - 0.05 * s),
            (cx0 + 0.24 * C, cy0 - 0.05 * s)]
    for p0, p1 in zip(hops, hops[1:]):
        _hop(k, slide, p0, p1, 0.30 * s)
    _pawn(k, slide, cx0 + 0.56 * C, cy0 + 0.60 * C, pawn_h)
    _run_all(ds)
    return r


# ─────────────────────────────────────────── ordinary pages ───────────────────────────────────────────
def _ground_boardgame(slide, role, index):
    """An ordinary page (a chart, an agenda) on the board: a run of the board's squares along the top edge, their
    bands facing down into the page; the content rect under it, in REFERENCE inches."""
    K = vl.rs._K[0]
    W, H = vl.rs._canvas(slide)
    k = _v2.kit_for(NAME, slide)
    Wr, Hr, s, o = ctx(k)
    D = 0.50 * s
    step = 0.80 * D
    _edge_run(k, slide, -0.37 * step, 0.16 * s, D, step, int(math.ceil(Wr / step)) + 1, "right", "bottom",
              start=int(index) % 8)
    top = (0.16 * s + D) / K + 0.35
    return (0.6, top, W - 1.2, H - top - 0.55)


_card_boardgame = vl._card_for(NAME)
