"""scoreboard — the Scoreboard language (体育记分牌): its ground painter, its card and its seven pages.

A stadium under floodlights: chalk pitch markings on mowed turf (stripes a shade off the ground), the title lit on
the scoreboard above a pitch seen from the stand, the section number on the back of a jersey, the picture on the big
screen with its caption in the ticker, the points as a lineup beside their formation on a half pitch, the quote on
a fans' banner hung from the stand rail, the key number lit on an LED board, and on the last page the whole pitch
from above with the last words on the centre spot. Every page PLANS its words first (vl_native2.flow / stack:
measured, never under a floor), draws its art clear of them — a pitch line that would run through words is cut
around them — then sets them. The only words it writes itself are the jersey numbers 1 2 3 4 on the lineup page
(the manifest's derived words)."""
from __future__ import annotations

import math

import deckkit as dk
import native_art as na
import visual_languages as vl
from vl_native import _run_all, alt, ctx, points_of, register, text_of
import vl_native2 as _v2   # names looked up at call time: vl_native2 may still be importing when this module loads
from pptx.oxml.ns import qn

NAME = "scoreboard"
U = na.U
PL, PW = 105.0, 68.0          # the pitch, in metres: its length and its width


# ─────────────────────────────────────────── drawing helpers ───────────────────────────────────────────
def _m(k):
    W, H, s, o = ctx(k)
    return 0.55 * s


def _decl(sh, reason, over=None):
    if sh is None:
        return None
    dk.decorative(sh, reason)
    if over:
        dk.overlap_intent(sh, over)
    return sh


def _alpha_line(sh, alpha):
    if sh is not None and alpha is not None:
        clr = sh._element.spPr.find(qn("a:ln")).find(qn("a:solidFill"))[0]
        clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(round(alpha * 100000)))}))
    return sh


def _paths(slide, runs, color, *, w=1.0, alpha=None, dash=None, cap="rnd"):
    """Open polylines [[(x, y), …], …] (inches) as ONE outline shape — a pitch's whole line work is one object."""
    runs = [r_ for r_ in runs if len(r_) >= 2]
    if not runs:
        return None
    xs = [p[0] for r_ in runs for p in r_]
    ys = [p[1] for r_ in runs for p in r_]
    x0, y0 = min(xs), min(ys)
    bw, bh = max(max(xs) - x0, 0.01), max(max(ys) - y0, 0.01)
    u = lambda p: na.orn._pt((p[0] - x0) / bw * U, (p[1] - y0) / bh * U)          # noqa: E731
    d = "".join("<a:moveTo>{}</a:moveTo>".format(u(r_[0])) + "".join("<a:lnTo>{}</a:lnTo>".format(u(q)) for q in r_[1:])
                for r_ in runs)
    sh = na._custom(slide, x0, y0, bw, bh, '<a:path w="{u}" h="{u}" fill="none">{d}</a:path>'.format(u=U, d=d),
                    line=na.hexstr(color), line_w=w)
    ln = sh._element.spPr.find(qn("a:ln"))
    ln.set("cap", cap)
    _alpha_line(sh, alpha)
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    return sh


def _simplify(pts, eps=0.004):
    """Ramer–Douglas–Peucker: the polyline with its collinear points dropped (a densified straight line is two)."""
    if len(pts) < 3:
        return pts
    (x0, y0), (x1, y1) = pts[0], pts[-1]
    dx, dy = x1 - x0, y1 - y0
    n = math.hypot(dx, dy)
    far, idx = -1.0, 0
    for i in range(1, len(pts) - 1):
        # a CLOSED run (its ends meet: a touchline loop, a whole circle) is measured from its start point — measured
        # from the zero-length chord, every point read as on it and the loop collapsed to nothing
        dd = (abs(dy * (pts[i][0] - x0) - dx * (pts[i][1] - y0)) / n if n > 1e-9
              else math.hypot(pts[i][0] - x0, pts[i][1] - y0))
        if dd > far:
            far, idx = dd, i
    if far <= eps:
        return [pts[0], pts[-1]]
    return _simplify(pts[:idx + 1], eps)[:-1] + _simplify(pts[idx:], eps)


def _dense(pts, step):
    out = [pts[0]]
    for a, b in zip(pts, pts[1:]):
        n = max(1, int(math.ceil(math.hypot(b[0] - a[0], b[1] - a[1]) / step)))
        out += [(a[0] + (b[0] - a[0]) * i / n, a[1] + (b[1] - a[1]) * i / n) for i in range(1, n + 1)]
    return out


def _cut(lines, box, keep, step=0.02):
    """The lines (polylines, inches) kept inside `box` (x, y, w, h) and OUT of every `keep` rect: a pitch line that
    would run through words or a drawn object stops at its edge and resumes past it. Returns the runs left."""
    bx, by, bw, bh = box

    def ok(p):
        if not (bx - 1e-6 <= p[0] <= bx + bw + 1e-6 and by - 1e-6 <= p[1] <= by + bh + 1e-6):
            return False
        return not any(r_[0] < p[0] < r_[0] + r_[2] and r_[1] < p[1] < r_[1] + r_[3] for r_ in keep)
    runs = []
    for ln in lines:
        cur = []
        for p in _dense(ln, step):
            if ok(p):
                cur.append(p)
            else:
                if len(cur) >= 2:
                    runs.append(_simplify(cur))
                cur = []
        if len(cur) >= 2:
            runs.append(_simplify(cur))
    return runs


def _grow(r_, pad):
    return (r_[0] - pad, r_[1] - pad, r_[2] + 2 * pad, r_[3] + 2 * pad)


def _clip_convex(poly, clip):
    """Sutherland–Hodgman: `poly` cut to the CONVEX polygon `clip` (counter-clockwise or clockwise)."""
    def side(a, b, p):
        return (b[0] - a[0]) * (p[1] - a[1]) - (b[1] - a[1]) * (p[0] - a[0])
    sign = 1.0 if sum(side(clip[i - 1], clip[i], clip[(i + 1) % len(clip)]) for i in range(len(clip))) > 0 else -1.0
    out = list(poly)
    for i in range(len(clip)):
        a, b = clip[i - 1], clip[i]
        inp, out = out, []
        for j in range(len(inp)):
            p, q = inp[j - 1], inp[j]
            pin, qin = sign * side(a, b, p) >= 0, sign * side(a, b, q) >= 0
            if pin != qin:
                dx, dy = q[0] - p[0], q[1] - p[1]
                ex, ey = b[0] - a[0], b[1] - a[1]
                den = ex * dy - ey * dx
                t = ((p[0] - a[0]) * ey - (p[1] - a[1]) * ex) / den if den else 0.0
                out.append((p[0] + t * dx, p[1] + t * dy))
            if qin:
                out.append(q)
        if not out:
            break
    return out


def _led(k, slide, rect, *, pitch=None, color=None, reason="the scoreboard's unlit LED dots; texture, nothing to read"):
    """An LED matrix filling `rect`: rows of round dots (a dotted line per row, one shape) — the lamps of a board."""
    W, H, s, o = ctx(k)
    x, y, w, h = rect
    p = pitch or 0.12 * s
    d = p / 2.0                                  # sysDot: a dot as long as the line is wide, then a gap as long
    n = int((h - d) / p)
    if n < 1 or w < 4 * p:
        return None
    y0 = y + (h - n * p) / 2.0
    cols = int((w - d) / p)
    x0 = x + (w - cols * p) / 2.0 + d / 2.0
    x1 = x0 + cols * p
    runs = [[(x0, y0 + i * p), (x1, y0 + i * p)] for i in range(n + 1)]
    sh = _paths(slide, runs, color or k.P["led_off"], w=d * 72.0, dash="sysDot")
    return _decl(sh, reason, "the LED dots lie on the board's display, under its words: layered by design")


def _bands(k, slide, horizontal=None):
    """The mowed turf: stripes a shade off the ground across the whole page (one shape) — across the pitch's length,
    so upright stripes on a landscape page and level ones on a portrait page. An odd count, so the shape spans the
    whole page edge to edge and everything on the page sits inside it."""
    W, H, s, o = ctx(k)
    if horizontal is None:
        horizontal = o != "land"
    n = 15
    polys = []
    for i in range(0, n, 2):
        if horizontal:
            a, b = H * i / n, H * (i + 1) / n
            polys.append([(0.0, a), (W, a), (W, b), (0.0, b)])
        else:
            a, b = W * i / n, W * (i + 1) / n
            polys.append([(a, 0.0), (b, 0.0), (b, H), (a, H)])
    sh = na.multipoly(slide, polys, k.P["band"])
    return _decl(sh, "the mowed stripes of the turf: the ground's texture, nothing to read",
                 "the turf's stripes lie under everything on the page: the ground, layered by design")


# ─────────────────────────────────────────── the pitch ───────────────────────────────────────────
def _arc(cx, cy, r, a0, a1, step=4.0):
    n = max(3, int(abs(a1 - a0) / step))
    return [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * i / n)),
             cy + r * math.sin(math.radians(a0 + (a1 - a0) * i / n))) for i in range(n + 1)]


def _markings(half=False, goals=True):
    """A pitch's markings as polylines in metres — x along its length (0 to 105), y across it (0 to 68) — and its
    spots. `half`: the attacking half only (x from the halfway line to the goal line at 105)."""
    lines, spots = [], []
    x0 = PL / 2.0 if half else 0.0
    lines.append([(x0, 0.0), (PL, 0.0), (PL, PW), (x0, PW), (x0, 0.0)])
    if not half:
        lines.append([(PL / 2.0, 0.0), (PL / 2.0, PW)])
    lines.append(_arc(PL / 2.0, PW / 2.0, 9.15, -90.0, 90.0) if half else _arc(PL / 2.0, PW / 2.0, 9.15, 0.0, 360.0))
    spots.append((PL / 2.0, PW / 2.0))
    a = math.degrees(math.acos(5.5 / 9.15))
    for side in ((1,) if half else (0, 1)):
        mx = (lambda x: x) if side == 0 else (lambda x: PL - x)
        lines.append([(mx(0.0), 13.84), (mx(16.5), 13.84), (mx(16.5), 54.16), (mx(0.0), 54.16)])
        lines.append([(mx(0.0), 24.84), (mx(5.5), 24.84), (mx(5.5), 43.16), (mx(0.0), 43.16)])
        base = 0.0 if side == 0 else 180.0
        lines.append(_arc(mx(11.0), PW / 2.0, 9.15, base - a, base + a))
        spots.append((mx(11.0), PW / 2.0))
        if goals:
            lines.append([(mx(0.0), 30.34), (mx(-1.9), 30.34), (mx(-1.9), 37.66), (mx(0.0), 37.66)])
        for cy, a0 in ((0.0, 0.0 if side == 0 else 90.0), (PW, 270.0 if side == 0 else 180.0)):
            lines.append(_arc(mx(0.0), cy, 1.6, a0, a0 + 90.0, step=15.0))
    return lines, spots


def _map(lines, fn, step_m=1.5):
    """Map pitch polylines through `fn` (metres → inches), densified first so a projected circle stays round."""
    out = []
    for ln in lines:
        pts = []
        for a, b in zip(ln, ln[1:]):
            n = max(1, int(math.hypot(b[0] - a[0], b[1] - a[1]) / step_m))
            pts += [(a[0] + (b[0] - a[0]) * i / n, a[1] + (b[1] - a[1]) * i / n) for i in range(n)]
        pts.append(ln[-1])
        out.append([fn(*p) for p in pts])
    return out


def _top_down(rect, upright):
    """Metres → inches for a pitch seen from above in `rect`: its length across the page, or down it when upright."""
    rx, ry, rw, rh = rect
    if upright:
        return lambda x, y: (rx + y / PW * rw, ry + x / PL * rh)
    return lambda x, y: (rx + x / PL * rw, ry + y / PW * rh)


def _perspective(cx, y_far, y_near, w_far, w_near, across):
    """Metres → inches for a pitch seen from the stand: the axis `across` ("length" or "width") lies across the page,
    the other runs away from the viewer — its far line at y_far, w_far wide; its near line at y_near, w_near wide."""
    zf = w_near / w_far
    yh = (zf * y_far - y_near) / (zf - 1.0)

    def f(x, y):
        if across == "length":
            a, t = (x - PL / 2.0) / PL, y / PW
        else:
            a, t = (y - PW / 2.0) / PW, x / PL
        z = zf + t * (1.0 - zf)
        return (cx + a * w_near / z, yh + (y_near - yh) / z)
    return f


def _near_line(y_far, w_far, zf, across, rho):
    """Where the near line of a pitch seen from the stand falls so that, halfway into it, a metre running away from the
    viewer looks `rho` of a metre running across — a circle on the centre spot reads as a wide ellipse, never a tall
    one (set by the page's height alone, a portrait pitch was stretched 3x in depth: its centre circle stood upright)."""
    zm = (zf + 1.0) / 2.0
    w_near = w_far * zf
    deep, wide = (PW, PL) if across == "length" else (PL, PW)
    return y_far + rho * zm * deep * w_near / (zf * wide)


def _chalk(k, slide, lines, box, keep, *, w=1.4, alpha=None, reason="the pitch's chalk markings; they draw the "
                                                                    "pitch the page is set on, nothing to read"):
    W, H, s, o = ctx(k)
    runs = _cut(lines, box, keep, step=0.015 * max(1.0, s))
    if not runs:
        return None
    xs = [p[0] for r_ in runs for p in r_]
    ys = [p[1] for r_ in runs for p in r_]
    # line work that is all one thin stroke (a halfway line cut round the words) is drawn run by run: as one shape its
    # box spans the gap it leaves for the words, and reads as a rule drawn through them
    groups = [[r_] for r_ in runs] if min(max(xs) - min(xs), max(ys) - min(ys)) < 0.12 else [runs]
    out = None
    for g in groups:
        out = _decl(_paths(slide, g, k.P["chalk"], w=w * s, alpha=_chalk_alpha(k) if alpha is None else alpha), reason)
    return out


def _chalk_alpha(k):
    return 0.42 if k.ground == "light" else 0.78


def _spots(k, slide, pts, d, keep, alpha=None):
    """The pitch's painted spots (centre, penalty) — those clear of the words — as one shape."""
    pts = [p for p in pts if not any(r_[0] - d < p[0] < r_[0] + r_[2] + d and r_[1] - d < p[1] < r_[1] + r_[3] + d
                                     for r_ in keep)]
    W, H, s, o = ctx(k)
    for px, py in pts:                 # one disc each: spots in a row as ONE shape read as a thin rule to the lint
        if not (d < px < W - d and d < py < H - d):
            continue
        poly = [(px + d / 2.0 * math.cos(2 * math.pi * i / 16), py + d / 2.0 * math.sin(2 * math.pi * i / 16))
                for i in range(16)]
        sh = na.multipoly(slide, [poly], k.P["chalk"], alpha=_chalk_alpha(k) if alpha is None else alpha)
        _decl(sh, "one of the pitch's painted spots (the centre spot, a penalty spot); nothing to read")


def _ink_rects(k, rects, draws, texts, align, pad):
    """Where each planned field's words are, padded: a field on one line is cut to its words (centred, left or right
    as set), a wrapped one keeps its box. The rects a pitch line is cut around."""
    W, H, s, o = ctx(k)
    sizes = {}
    for d_ in draws:
        sizes.update(getattr(d_, "sizes", {}))
    out = []
    for f_, (x, y, w, h) in rects.items():
        t, z = texts.get(f_), sizes.get(f_)
        if t and z and "\n" not in t and h <= z * 1.25 / 72.0 * 1.45 + 0.08:
            iw = min(w, _v2.ink_w(k, f_, t, z) + 0.12 * s)
            x = x + (w - iw) / 2.0 if align == "c" else (x + w - iw if align == "r" else x)
            w = iw
        out.append(_grow((x, y, w, h), pad))
    return out


def _block(k, slide, page, x, w, y0, y1, groups, align="l", start=None):
    """PLAN `groups` (vl_native2.stack's) from the top of [y0, y1]: (rects, draws, height). Planned again at another top
    with the same room (`_block(…, y0 + d, y1 + d, …)`), the same sizes and lines come back — so a stack measured
    once can be moved to centre a picture-and-words group on the page."""
    kw = {"anchor": "top", "align": align}
    if start:
        kw["start"] = start
    r, ds = _v2.stack(k, slide, page, x, w, y0, y1, groups, **kw)
    h = (max(v[1] + v[3] for v in r.values()) - y0) if r else 0.0
    return r, ds, h


# ─────────────────────────────────────────── stadium furniture ───────────────────────────────────────────
def _board(k, slide, rect, *, fr=None, led=True, lamps=True):
    """A scoreboard at `rect`: its frame (lit along the top), the display inside and the display's unlit LED dots.
    Returns the display's rect."""
    W, H, s, o = ctx(k)
    x, y, w, h = rect
    fr = 0.11 * s if fr is None else fr
    frame = dk.box(slide, x, y, w, h, fill=k.P["bezel"], line=dk._as_rgb(k.P["bezel_hi"]), line_w=1.0)
    dk.decorative(frame, "the scoreboard's frame; what the board shows carries the meaning")
    disp = (x + fr, y + fr, w - 2 * fr, h - 2 * fr)
    panel = dk.box(slide, *disp, fill=k.P["panel"])
    dk.decorative(panel, "the scoreboard's display, the ground its words are lit on")
    if led:
        _led(k, slide, _grow(disp, -0.10 * s))
    if lamps and w > 2.4 * s:                       # two small lamps at the frame's top corners
        dd = min(0.075 * s, fr * 0.65)
        _spot_lamps(k, slide, [(x + fr * 0.5 + 0.5 * dd + 0.02 * s, y + fr * 0.5),
                               (x + w - fr * 0.5 - 0.5 * dd - 0.02 * s, y + fr * 0.5)], dd)
    return disp


def _spot_lamps(k, slide, pts, d):
    polys = [[(px + d / 2.0 * math.cos(2 * math.pi * i / 12), py + d / 2.0 * math.sin(2 * math.pi * i / 12))
              for i in range(12)] for px, py in pts]
    sh = na.multipoly(slide, polys, k.P["accents"][0])
    return _decl(sh, "the indicator lamps on the scoreboard's frame; nothing to read",
                 "the lamps are set into the board's frame: layered by design")


def _legs(k, slide, x0, x1, y_top, y_bot, w):
    """Two posts under a board (or a screen) from its foot down to `y_bot`."""
    W, H, s, o = ctx(k)
    if y_bot - y_top < 0.08 * s:
        return
    polys = [[(cx - w / 2.0, y_top), (cx + w / 2.0, y_top), (cx + w / 2.0, y_bot), (cx - w / 2.0, y_bot)]
             for cx in (x0, x1)]
    braces = []
    if y_bot - y_top > 0.9 * s:                      # a cross brace between the posts
        yb = y_top + (y_bot - y_top) * 0.45
        braces = [[(x0, yb - 0.025 * s), (x1, yb - 0.025 * s), (x1, yb + 0.025 * s), (x0, yb + 0.025 * s)]]
    sh = na.multipoly(slide, polys + braces, k.P["post"])
    _decl(sh, "the posts that hold the board up; part of the stadium, nothing to read",
          "the posts run up behind the board they hold: layered by design")


def _floodlight(k, slide, cx, top, foot, bank_w):
    """A floodlight mast: a lamp bank (four by three lamps in a frame) on a tapering mast, a soft halo round the bank."""
    W, H, s, o = ctx(k)
    bh = bank_w * 0.62
    glow = na.radial_glow(slide, cx, top + bh / 2.0, bank_w * 2.6, k.P["glow"], alpha=0.30)
    dk.overlap_intent(glow, "the floodlight's halo lies behind its lamp bank and mast: layered by design")
    mast = na.poly(slide, [(cx - 0.035 * s, top + bh), (cx + 0.035 * s, top + bh), (cx + 0.075 * s, foot),
                           (cx - 0.075 * s, foot)], fill=k.P["post"])
    _decl(mast, "the floodlight's mast; part of the stadium, nothing to read",
          "the mast stands behind its lamp bank and the halo: layered by design")
    bank = dk.box(slide, cx - bank_w / 2.0, top, bank_w, bh, fill=k.P["post"], line=dk._as_rgb(k.P["bezel_hi"]),
                  line_w=0.75)
    _decl(bank, "the floodlight's lamp bank; part of the stadium, nothing to read",
          "the lamp bank sits on its mast in its own halo: layered by design")
    lw, lh = bank_w / 4.0, bh / 3.0
    rr = 0.36 * min(lw, lh)
    lamps = [[(cx - bank_w / 2.0 + (i + 0.5) * lw + rr * math.cos(2 * math.pi * a / 14),
               top + (j + 0.5) * lh + rr * math.sin(2 * math.pi * a / 14)) for a in range(14)]
             for i in range(4) for j in range(3)]
    sh = na.multipoly(slide, lamps, k.P["lamp"])
    _decl(sh, "the floodlight's lamps; part of the stadium, nothing to read",
          "the lamps are set in their bank: layered by design")


def _shirt_outline(x, y, w, h):
    """A jersey seen from the back: shoulders, short sleeves, a straight body; the shallow back neckline."""
    body = [(0.36, 0.0), (0.18, 0.055), (0.0, 0.30), (0.13, 0.445), (0.235, 0.36), (0.225, 1.0), (0.775, 1.0),
            (0.765, 0.36), (0.87, 0.445), (1.0, 0.30), (0.82, 0.055), (0.64, 0.0)]
    neck = [(0.64 - 0.28 * i / 8.0, 0.045 * math.sin(math.pi * i / 8.0)) for i in range(1, 8)]
    return [(x + u * w, y + v * h) for u, v in body + neck]


def _shirt(k, slide, x, y, w, h, *, fill=None, trim=None, detail=True, reason=None):
    """A jersey at (x, y, w, h): its body, the collar and two stripes round each sleeve's cuff."""
    W, H, s, o = ctx(k)
    fill = fill or k.P["jersey"]
    trim = trim or k.P["jersey_trim"]
    sh = na.poly(slide, _shirt_outline(x, y, w, h), fill=fill)
    _decl(sh, reason or "the jersey the number is printed on; the number carries the meaning")
    if not detail:
        return sh
    P = lambda u, v: (x + u * w, y + v * h)                                              # noqa: E731
    collar = [P(0.36, 0.0)] + [P(0.36 + 0.28 * i / 10.0, 0.045 * math.sin(math.pi * i / 10.0)) for i in range(1, 10)] \
        + [P(0.64, 0.0), P(0.655, 0.035)] \
        + [P(0.64 - 0.28 * i / 10.0, 0.045 * math.sin(math.pi * i / 10.0) + 0.045) for i in range(0, 11)] \
        + [P(0.345, 0.035)]
    parts = [[collar]]
    for left in (True, False):
        polys = []
        mir = (lambda u: u) if left else (lambda u: 1.0 - u)
        a, b = P(mir(0.0), 0.30), P(mir(0.13), 0.445)                    # the cuff's edge
        sleeve = [a, P(mir(0.18), 0.055), P(mir(0.235), 0.36), b]
        ex, ey = b[0] - a[0], b[1] - a[1]
        ln = math.hypot(ex, ey) or 1e-9
        nx, ny = ey / ln, -ex / ln
        if left is False:
            nx, ny = -nx, -ny
        cx_, cy_ = (sleeve[1][0] + sleeve[2][0]) / 2.0, (sleeve[1][1] + sleeve[2][1]) / 2.0
        if (cx_ - a[0]) * nx + (cy_ - a[1]) * ny < 0:                 # point the normal INTO the sleeve
            nx, ny = -nx, -ny
        for o0, o1 in ((0.05, 0.085), (0.115, 0.15)):
            q = [(a[0] - ex * 0.3 + nx * o0 * w, a[1] - ey * 0.3 + ny * o0 * w),
                 (b[0] + ex * 0.3 + nx * o0 * w, b[1] + ey * 0.3 + ny * o0 * w),
                 (b[0] + ex * 0.3 + nx * o1 * w, b[1] + ey * 0.3 + ny * o1 * w),
                 (a[0] - ex * 0.3 + nx * o1 * w, a[1] - ey * 0.3 + ny * o1 * w)]
            c = _clip_convex(q, sleeve)
            if len(c) >= 3:
                polys.append(c)
        parts.append(polys)
    # the collar and each cuff's stripes are shapes of their own: one shape spanning all three held the number's box
    # inside it and read as a card the number ran out of (TEXT PADDING, the delivery lint)
    for polys in parts:
        tr = na.multipoly(slide, polys, trim)
        _decl(tr, "the jersey's collar or the stripes round a cuff; ornament, nothing to read",
              "the collar and the cuff stripes are printed on the jersey: layered by design")
    return sh


def _hanger(k, slide, cx, neck_y, w, rail_half):
    """A coat hanger's hook rising from the jersey's neck and curling over a short rail — a shirt hung in the dressing
    room. Returns the rail's height."""
    W, H, s, o = ctx(k)
    r = 0.12 * s
    yc = neck_y - 0.40 * s                       # the hook's curl: centred on (cx + r, yc)
    rail_y = yc - r + 0.06 * s                   # the rail runs through the top of the curl
    rl = na.polyline(slide, [(max(0.05, cx - rail_half), rail_y), (min(W - 0.05, cx + rail_half), rail_y)],
                     k.P["post"], w=6.0 * s)
    _decl(rl, "the dressing-room rail the jersey hangs on; nothing to read",
          "the hanger's hook curls over the rail: layered by design")
    pts = [(cx, neck_y + 0.04 * s), (cx, yc)] + \
        [(cx + r + r * math.cos(math.radians(a)), yc + r * math.sin(math.radians(a))) for a in range(195, 416, 15)]
    hk = na.polyline(slide, pts, k.P["bezel_hi"], w=2.6 * s)
    _decl(hk, "the coat hanger's hook the jersey hangs from; nothing to read",
          "the hook rises from inside the jersey's collar and over the rail: layered by design")
    return rail_y


# ─────────────────────────────────────────── the pages ───────────────────────────────────────────
@register(NAME, "cover", alts=3)
def _sb_cover(k, slide, f, image):
    """The stadium: the kicker, title and subtitle lit on the scoreboard; under it the pitch seen from the stand —
    its stripes and chalk lines running away to the far touchline the board stands over — and on a landscape page a
    floodlight mast at each side. Long copy: a wider board and less pitch, then the board alone."""
    W, H, s, o = ctx(k)
    m = _m(k)
    land = o == "land"
    pitch = alt() < 2
    if not pitch:
        _bands(k, slide)
    if land:
        bw = (0.62, 0.74, 0.84)[alt()] * W
        top = 0.08 * H
        y_max = (0.66, 0.74)[alt()] * H if pitch else H - m
    else:
        bw = (0.86, 0.90, 0.90)[alt()] * W
        top = 0.06 * H
        y_max = (0.58, 0.68)[alt()] * H if pitch else H - m
    bx = (W - bw) / 2.0
    fr, pad = 0.11 * s, 0.36 * s
    ix, iw = bx + fr + pad, bw - 2 * (fr + pad)
    kick = text_of(f, "kicker")
    groups = [(_v2.fields(f, ("kicker",), ("kicker",)), 0.34 * s), (_v2.fields(f, ("title",)), 0.20 * s),
              (_v2.fields(f, ("subtitle",)), 0.0)]

    def plan(y0):
        return _v2.stack(k, slide, "cover", ix, iw, y0, y_max - fr - pad, groups, anchor="top", align="c")
    r, ds = plan(top + fr + pad)
    foot = max(v[1] + v[3] for v in r.values())
    bh = foot + pad + fr - top
    if not pitch:                                    # the board alone: centred on the page
        top = max(m * 0.5, (H - bh) / 2.0)
        y_max = top + bh + 0.01 * s
        r, ds = plan(top + fr + pad)
    board = (bx, top, bw, bh)
    art = []
    if pitch:
        y_far = top + bh + 0.42 * s
        if land:
            w_far = min(0.92 * W, max(bw * 1.12, 0.74 * W))
            proj = _perspective(W / 2.0, y_far, _near_line(y_far, w_far, 1.80, "length", 0.42), w_far, w_far * 1.80,
                                "length")
        else:
            w_far = 0.66 * W
            proj = _perspective(W / 2.0, y_far, _near_line(y_far, w_far, 2.4, "width", 0.50), w_far, w_far * 2.4,
                                "width")
        art.append(lambda: _persp_pitch(k, slide, proj, "length" if land else "width", y_far))
        leg_w = 0.10 * s
        art.append(lambda: _legs(k, slide, bx + 0.22 * bw, bx + 0.78 * bw, top + bh - 0.05 * s, y_far + 0.05 * s, leg_w))
        if land and bx - m > 1.0 * s:
            bank = min(0.95 * s, (bx - m) * 0.62)
            for cx in (bx / 2.0 + 0.10 * s, W - bx / 2.0 - 0.10 * s):
                art.append(lambda cx=cx, bank=bank: _floodlight(k, slide, cx, top + 0.10 * s, y_far + 0.25 * s, bank))
    rule_y = None
    if kick and "kicker" in r and "title" in r:
        rule_y = (r["kicker"][1] + r["kicker"][3] + r["title"][1]) / 2.0
    _run_all(art)
    _board(k, slide, board)
    if rule_y is not None:
        rl = na.strokes(slide, [((bx + fr + 0.25 * s, rule_y), (bx + bw - fr - 0.25 * s, rule_y))], k.P["bezel_hi"],
                        w=1.0)
        _decl(rl, "the rule under the board's heading; it only divides the heading from the title",
              "the rule is printed on the board's display: layered by design")
    _run_all(ds)
    return r


def _persp_pitch(k, slide, proj, across, y_far):
    """The pitch seen from the stand: the mowed stripes in perspective (one shape) and the chalk markings."""
    W, H, s, o = ctx(k)
    polys = []
    # the stripes cross the pitch's LENGTH in both views, 7.5 m apart; the run-off past the near line is mowed too
    lo, hi = (-4.0, PL + 4.0) if across == "length" else (-4.0, PL + 16.0)
    y_lo, y_hi = (-3.0, PW + 9.0) if across == "length" else (-3.0, PW + 3.0)
    n = max(2, int(round((hi - lo) / 7.5)))
    for i in range(0, n, 2):
        x0, x1 = lo + (hi - lo) * i / n, lo + (hi - lo) * (i + 1) / n
        quad = [proj(x0, y_lo), proj(x1, y_lo), proj(x1, y_hi), proj(x0, y_hi)]
        pts = na.clip_to_page(quad, W, H)
        if len(pts) >= 3:
            polys.append(pts)
    sh = na.multipoly(slide, polys, k.P["band"]) if polys else None
    _decl(sh, "the mowed stripes of the pitch seen from the stand; texture, nothing to read",
          "the turf lies under the board's posts and the floodlights: layered by design")
    lines, spots = _markings()
    box = (0.0, 0.0, W, H)
    _chalk(k, slide, _map(lines, proj), box, [], w=1.5,
           reason="the pitch's chalk markings seen from the stand; they draw the stadium, nothing to read")
    _spots(k, slide, [proj(*p) for p in spots], 0.07 * s, [])


@register(NAME, "section", alts=3)
def _sb_section(k, slide, f, image):
    """The section number on the back of a jersey hung in the dressing room; the kicker and title beside it (under it
    on a portrait or square page). A number too wide for the jersey — or none — and the page is the kick-off: the
    number, kicker and title on the centre spot, the halfway line and the centre circle cut round them."""
    W, H, s, o = ctx(k)
    m = _m(k)
    num = text_of(f, "number")
    land = o == "land"
    _bands(k, slide)
    if not num or alt() == 2:
        return _kickoff(k, slide, f, "section", [("number", num)] if num else [])
    r = {}
    if land:
        sw = min((0.33 if alt() == 0 else 0.27) * W, 0.66 * H)
        sh_ = 0.90 * sw
        sx = m + 0.45 * s
        sy = (H - sh_) / 2.0 + 0.30 * s
        col = (sx + sw + 0.85 * s, 0.14 * H, W - m - (sx + sw + 0.85 * s), 0.72 * H)
    else:
        sw = min((0.60 if alt() == 0 else 0.48) * W, (0.36 if alt() == 0 else 0.29) * H)
        sh_ = 0.90 * sw
        sx = (W - sw) / 2.0
        sy = m + 0.62 * s
        col = (0.09 * W, sy + sh_ + 0.50 * s, 0.82 * W, H - m - (sy + sh_ + 0.50 * s))
        # the jersey and the words as ONE group at the page's optical centre: the words measured first
        groups = [(_v2.fields(f, ("kicker",), ("kicker",)), 0.16 * s), (_v2.fields(f, ("title",)), 0.0)]
        _r, _d, th = _block(k, slide, "section", col[0], col[2], col[1], col[1] + col[3], groups, align="c")
        d = max(0.0, (H - m - (col[1] + th)) * 0.45)
        sy += d
        col = (col[0], col[1] + d, col[2], col[3])
    torso = (sx + 0.25 * sw, sy + 0.16 * sh_, 0.50 * sw, 0.74 * sh_)
    nr, nd = _v2.flow(k, slide, "section", torso, [("number", num)], anchor="middle", align="c",
                      accent=k.P["jersey_ink"], start={"number": min(170 * s, torso[3] * 72.0 / 1.25)})
    size = getattr(nd, "sizes", {}).get("number", 0.0)
    if size < 0.42 * min(170 * s, torso[3] * 72.0 / 1.25):     # a long number shrunk to a sliver: the kick-off reads it
        raise vl.VLTextOverflow("scoreboard.section(): the number {!r} sets at {:.0f}pt on the jersey — trying the "
                                "kick-off layout".format(num, size))
    r.update(nr)
    if land:
        tr, tds = _v2.stack(k, slide, "section", col[0], col[2], col[1], col[1] + col[3],
                            [(_v2.fields(f, ("kicker",), ("kicker",)), 0.16 * s), (_v2.fields(f, ("title",)), 0.0)],
                            anchor="middle", align="l")
    else:
        tr, tds, _h = _block(k, slide, "section", col[0], col[2], col[1], col[1] + col[3], groups, align="c")
    r.update(tr)
    _hanger(k, slide, sx + sw / 2.0, sy + 0.01 * sh_, sw, sw * 0.62)
    _shirt(k, slide, sx, sy, sw, sh_)
    _run_all([nd] + tds)
    return r


def _kickoff(k, slide, f, page, head, groups=None):
    """The kick-off: `head` (the number, when there is one), the kicker and the title — or `groups` — centred on the
    centre spot of a pitch seen from above; the halfway line and the centre circle cut round the words."""
    W, H, s, o = ctx(k)
    m = _m(k)
    land = o == "land"
    x, w = (0.16 * W, 0.68 * W) if land else (0.09 * W, 0.82 * W)
    groups = groups or [(head, 0.10 * s), (_v2.fields(f, ("kicker",), ("kicker",)), 0.14 * s),
                        (_v2.fields(f, ("title",)), 0.0)]
    r, ds = _v2.stack(k, slide, page, x, w, 0.10 * H, H - m, groups, anchor="middle", align="c",
                      start={"number": 120 * s} if head else None)
    texts = {f_: t for g, _gap in groups for f_, t in g}
    keep = _ink_rects(k, r, ds, texts, "c", 0.16 * s)
    q = 0.32 * s
    cx, cy = W / 2.0, (min(v[1] for v in keep) + max(v[1] + v[3] for v in keep)) / 2.0
    # the centre circle holds the words; words too big for any circle on the page keep only the halfway line (a
    # circle cut round them left stray arcs at the words' corners)
    need = max(math.hypot(v[0] + dx - cx, v[1] + dy - cy) for v in keep for dx in (0.0, v[2]) for dy in (0.0, v[3]))
    R = max(1.0 * s, need + 0.12 * s)
    lines = [_arc(cx, cy, R, 0.0, 360.0, step=2.0)] if R <= 0.46 * min(W, H) and cy - R > q and cy + R < H - q else []
    lines.append([(cx, q), (cx, H - q)] if land else [(q, cy), (W - q, cy)])
    _chalk(k, slide, lines, (q, q, W - 2 * q, H - 2 * q), keep, w=1.6)
    _run_all(ds)
    return r


@register(NAME, "image_text", alts=3)
def _sb_image_text(k, slide, f, image):
    """The caller's picture on the stadium's big screen — a heavy frame on two posts — the caption running in the
    screen's ticker under the picture; the kicker, title and body beside it (under it on a portrait or square
    page)."""
    W, H, s, o = ctx(k)
    m = _m(k)
    land = o == "land"
    _bands(k, slide)
    cap = text_of(f, "caption")
    fr = 0.13 * s
    if land:
        pw = (0.48, 0.42, 0.36)[alt()] * W
        ph = min(pw * 9.0 / 16.0, 0.56 * H)
        px = m + 0.25 * s + fr
    else:
        pw = (0.84, 0.76, 0.66)[alt()] * W - 2 * fr
        ph = min(pw * 9.0 / 16.0, (0.30, 0.26, 0.21)[alt()] * H)
        px = (W - pw) / 2.0
    r, ds = {}, []
    tick_h = 0.0
    tpad, bpad = max(0.10 * s, 0.09), max(0.12 * s, 0.14)
    if cap:
        cr, cd = _v2.flow(k, slide, "image_text", (px + 0.16 * s, 0.0, pw - 0.32 * s, 1.4 * s), [("caption", cap)],
                          anchor="top", ink=k.P["ink"])
        tick_h = max(v[1] + v[3] for v in cr.values()) + tpad + bpad    # the lint asks ~0.1in under the words
    screen_h = fr + ph + tick_h + fr
    legs = (0.75 if land else 0.40) * s
    tgroups = [(_v2.fields(f, ("kicker",), ("kicker",)), 0.12 * s), (_v2.fields(f, ("title",)), 0.18 * s),
               (_v2.fields(f, ("body",)), 0.0)]
    if land:
        py = (H - (screen_h + legs)) / 2.0 + fr
    else:
        py = m * 0.6 + fr
        top_ = py - fr + screen_h + legs + 0.40 * s
        _r, _d, th = _block(k, slide, "image_text", 0.09 * W, 0.82 * W, top_, H - m, tgroups)
        py += max(0.0, (H - m - (top_ + th)) * 0.45)       # the screen and the words as one group, centred
    if cap:
        cr, cd = _v2.flow(k, slide, "image_text", (px + 0.16 * s, py + ph + tpad, pw - 0.32 * s, tick_h - tpad),
                          [("caption", cap)], anchor="top", ink=k.P["ink"])
        r.update(cr)
        ds.append(cd)
    sx, sy, sw_, sh_ = px - fr, py - fr, pw + 2 * fr, screen_h
    foot = sy + sh_ + legs
    if land:
        x0 = sx + sw_ + 0.85 * s
        col = (x0, 0.14 * H, W - m - x0, 0.72 * H)
        tr, tds = _v2.stack(k, slide, "image_text", col[0], col[2], col[1], col[1] + col[3], tgroups, anchor="middle")
    else:
        d = foot + 0.40 * s - top_
        tr, tds, _h = _block(k, slide, "image_text", 0.09 * W, 0.82 * W, top_ + d, H - m + d, tgroups)
    r.update(tr)
    _legs(k, slide, sx + 0.24 * sw_, sx + 0.76 * sw_, sy + sh_ - 0.05 * s, foot, 0.11 * s)
    gl = na.strokes(slide, [((max(0.1 * s, sx - 0.35 * s), foot), (min(W - 0.1 * s, sx + sw_ + 0.35 * s), foot))],
                    k.P["chalk"], w=1.4 * s, alpha=_chalk_alpha(k))
    _decl(gl, "the touchline the big screen stands behind; nothing to read")
    frame = dk.box(slide, sx, sy, sw_, sh_, fill=k.P["bezel"], line=dk._as_rgb(k.P["bezel_hi"]), line_w=1.0)
    dk.decorative(frame, "the big screen's frame; the picture on it carries the meaning")
    if cap:
        tick = dk.box(slide, px, py + ph, pw, tick_h, fill=k.P["panel"])
        dk.decorative(tick, "the screen's ticker, the strip the caption runs in")
        bar = dk.box(slide, px, py + ph, 0.09 * s, tick_h, fill=k.P["accents"][0])
        _decl(bar, "the ticker's lit edge; it only marks where the caption runs",
              "the lit edge is part of the ticker strip: layered by design")
    dd = min(0.07 * s, fr * 0.6)
    _spot_lamps(k, slide, [(sx + fr * 0.5 + 0.03 * s, sy + fr * 0.5), (sx + sw_ - fr * 0.5 - 0.03 * s, sy + fr * 0.5)], dd)
    vl._place_image(k, slide, image, (px, py, pw, ph), vl.L_((0, 0, 1, 1), "frame", None), "image_text")
    _run_all(ds + tds)
    return r


def _jersey_mark(k, slide, page, x, y, w, num, start=None):
    """PLAN a small jersey carrying `num` at (x, y), `w` wide: returns (rects, draw)."""
    W, H, s, o = ctx(k)
    h = 0.95 * w
    kw = {"anchor": "middle", "align": "c", "ink": k.P["jersey_ink"]}
    if start:
        kw["start"] = start
    nr, nd = _v2.flow(k, slide, page, (x, y + 0.14 * h, w, 0.80 * h), [("item_no", num)], **kw)

    def draw():
        _shirt(k, slide, x, y, w, h, detail=False, reason="a jersey carrying the point's number; the number carries "
                                                          "the meaning")
        nd()
    return nr, draw


@register(NAME, "points", alts=4)
def _sb_points(k, slide, f, image):
    """The lineup: each point a row — its jersey number on a small shirt, its head and line — beside the formation
    on a half pitch, where the same shirts take their places. Long copy: a smaller pitch, then the list alone, then
    two columns of rows."""
    W, H, s, o = ctx(k)
    m = _m(k)
    land = o == "land"
    _bands(k, slide)
    pts = points_of(f.get("items"))
    n = len(pts)
    nos = [str(i + 1) for i in range(n)]
    pitch = alt() < 2
    x0, w0 = m + 0.1 * s, W - 2 * m - 0.2 * s
    r, ds = _v2.stack(k, slide, "points", x0, w0, m, m + (0.24 if alt() == 0 else 0.30) * H,
                      [(_v2.fields(f, ("kicker",), ("kicker",)), 0.08 * s), (_v2.fields(f, ("title",)), 0.0)],
                      anchor="top", start=None if alt() < 2 else {"title": 30 * s})
    top = max((v[1] + v[3] for v in r.values()), default=m) + (0.40 if alt() < 3 else 0.28) * s
    bottom = H - m
    art = []
    # the half pitch: upright, the halfway line at its top and the goal at its foot (68 wide, 52.5 deep)
    side = W >= 0.95 * H                 # a landscape or square page: the pitch beside the list
    if pitch:
        if side:
            frac = (0.40 if alt() == 0 else 0.32) if land else (0.36 if alt() == 0 else 0.30)
            ph_ = min(bottom - top, frac * W / (PW / (PL / 2.0)))
            pw_ = ph_ * PW / (PL / 2.0)
            prect = (x0, top + (bottom - top - ph_) / 2.0, pw_, ph_)
            lx, ly0, lw = x0 + pw_ + 0.75 * s, top, W - m - 0.1 * s - (x0 + pw_ + 0.75 * s)
        else:
            pw_ = min(w0 * (0.70 if alt() == 0 else 0.56), (bottom - top) * (0.40 if alt() == 0 else 0.30) * PW / (PL / 2.0))
            ph_ = pw_ * (PL / 2.0) / PW
            prect = ((W - pw_) / 2.0, top, pw_, ph_)
            lx, ly0, lw = x0, top + ph_ + 0.45 * s, w0
    else:
        lx, ly0, lw = x0, top, w0
    cols = 2 if (alt() == 3 and n >= 3) else 1
    cgap = 0.5 * s
    cw = (lw - (cols - 1) * cgap) / cols
    nw = 0.60 * s                       # the jersey mark's width
    if pitch and side and cw - nw - 0.30 * s < 2.9 * s:
        # beside the pitch the list's column is too narrow to set its lines (a square page): the list alone instead
        raise vl.VLTextOverflow("scoreboard.points(): the list beside the pitch is {:.2f}in wide — trying the list "
                                "alone".format(cw - nw - 0.30 * s))
    gap = 0.22 * s
    rows = -(-n // cols)
    items = [[("item_head", hd)] + ([("item_line", ln)] if ln else []) for hd, ln in pts]
    share = (bottom - ly0 - gap * (rows - 1)) / rows
    probe = _v2.plan_together(k, slide, "points", [((0, 0, cw - nw - 0.30 * s, max(0.3 * s, share - 0.1 * s)), it,
                                                    {"anchor": "top"}) for it in items])
    hs = [max(max((v[1] + v[3] for v in pr.values()), default=0.0) + 0.06 * s, 0.95 * nw + 0.08 * s)
          for pr, _d in probe]
    row_h = [max(hs[i] for i in range(rw * cols, min(n, rw * cols + cols))) for rw in range(rows)]
    block = sum(row_h) + gap * (rows - 1)
    if block > bottom - ly0 + 1e-6:
        raise vl.VLTextOverflow("scoreboard.points(): the lineup needs {:.2f}in, the page has {:.2f}in".format(
            block, bottom - ly0))
    if pitch and side:
        ly = ly0 + max(0.0, (bottom - ly0 - block) / 2.0)
    elif pitch:                          # the pitch and the list under it move together, as one group
        d = max(0.0, (bottom - ly0 - block) * 0.40)
        prect = (prect[0], prect[1] + d, prect[2], prect[3])
        ly = ly0 + d
    else:
        ly = ly0 + max(0.0, (bottom - ly0 - block) * 0.35)
    ys = []
    for rw in range(rows):
        ys.append(ly)
        ly += row_h[rw] + gap
    sizes = {f_: z for _pr, d_ in probe for f_, z in getattr(d_, "sizes", {}).items()}     # as measured (airmail)
    if alt() < 3 and sizes.get("item_head", 1e9) < 0.70 * vl.TYPE[NAME]["item_head"][0] * s:
        # the lineup's heads squeezed to under 70% of their size beside the pitch: a roomier layout sets them larger
        raise vl.VLTextOverflow("scoreboard.points(): the heads set at {:.0f}pt in this layout — trying a roomier "
                                "one".format(sizes["item_head"]))
    boxes = [(lx + (i % cols) * (cw + cgap), ys[i // cols], row_h[i // cols]) for i in range(n)]
    planned = _v2.plan_together(k, slide, "points", [((bx + nw + 0.30 * s, by + 0.03 * s, cw - nw - 0.30 * s, bh), it,
                                                      {"anchor": "top", "start": sizes})
                                                     for (bx, by, bh), it in zip(boxes, items)])
    marks = []
    for i, ((bx, by, bh), (pr, pd)) in enumerate(zip(boxes, planned)):
        r.update(pr)
        ds.append(pd)
        mr, md = _jersey_mark(k, slide, "points", bx, by + 0.04 * s, nw, nos[i])
        r.update(mr)
        marks.append(md)
    # the rules between rows (in each column)
    segs = []
    for i, (bx, by, bh) in enumerate(boxes):
        if i // cols > 0:
            yl = by - gap / 2.0
            segs.append(((bx, yl), (bx + cw, yl)))
    if segs:
        art.append(lambda: _decl(na.strokes(slide, segs, k.P["chalk"], w=0.75, alpha=0.30),
                                 "the chalk rule between two rows of the lineup; it only divides them"))
    if pitch:
        art.append(lambda: _formation(k, slide, prect, nos, nw * (0.92 if land else 0.85)))
    _run_all(art + marks + ds)
    return r


_FORMATION = {2: ((0.30, 0.42), (0.70, 0.42)),
              3: ((0.50, 0.31), (0.22, 0.47), (0.78, 0.47)),
              4: ((0.50, 0.29), (0.19, 0.40), (0.81, 0.40), (0.50, 0.51))}


def _formation(k, slide, rect, nos, mw):
    """The half pitch seen from above in `rect` (the halfway line at its top, the goal at its foot) with a jersey for
    each point in its place: two side by side, three in a triangle, four in a diamond."""
    W, H, s, o = ctx(k)
    rx, ry, rw, rh = rect
    fn = lambda x, y: (rx + y / PW * rw, ry + (x - PL / 2.0) / (PL / 2.0) * rh)        # noqa: E731
    lines, spots = _markings(half=True)
    mw = min(mw, 0.19 * rh, 0.15 * rw)        # a small pitch takes smaller shirts: four never touch
    mh = 0.95 * mw
    places = [(rx + u * rw - mw / 2.0, ry + v * rh - mh / 2.0) for u, v in _FORMATION[len(nos)]]
    keep = [_grow((px, py, mw, mh), 0.08 * s) for px, py in places]
    g = 0.25 * s
    _chalk(k, slide, _map(lines, fn), (rx - g, ry - g, rw + 2 * g, rh + 2 * g), keep, w=1.3,
           reason="the half pitch the lineup's formation is drawn on; nothing to read in its lines")
    _spots(k, slide, [fn(*p) for p in spots], 0.06 * s, keep)
    for (px, py), no in zip(places, nos):
        nr, nd = _jersey_mark(k, slide, "points", px, py, mw, no)
        nd()


@register(NAME, "quote", alts=2)
def _sb_quote(k, slide, f, image):
    """A fans' banner hung from the stand's rail: the quote painted on it between two stripes, the attribution under
    it; the stand's rows of seats above the rail."""
    W, H, s, o = ctx(k)
    m = _m(k)
    land = o == "land"
    _bands(k, slide)
    att = text_of(f, "attribution")
    # the words painted on the banner: the quote, or — given only an attribution — the attribution itself, so a word
    # the caller gave is never dropped (an attribution alone crashed the page, audit A2, 2026-10-22)
    qf = ("quote", text_of(f, "quote")) if text_of(f, "quote") else ("attribution", att)
    if qf[0] == "attribution":
        att = None
    bx, bw = ((0.09 if alt() == 0 else 0.06) * W, (0.82 if alt() == 0 else 0.88) * W) if land else (0.06 * W, 0.88 * W)
    padx, pady = 0.50 * s, 0.46 * s
    tie = 0.34 * s
    stand = (0.16 if land else 0.12) * H if alt() == 0 else 0.09 * H
    y_room = H - m - stand - tie
    r, ds = {}, []
    ah = 0.0
    if att:
        ar, ad = _v2.flow(k, slide, "quote", (bx, 0.0, bw, 1.6 * s), [("attribution", att)], anchor="top", align="r")
        ah = max(v[1] + v[3] for v in ar.values()) + 0.30 * s
    qr, qd = _v2.flow(k, slide, "quote", (bx + padx + 0.03 * (bw - 2 * padx), 0.0, 0.94 * (bw - 2 * padx),
                                          y_room - ah - 2 * pady), [qf],
                      anchor="top", align="c", ink=k.P["banner_ink"])
    qh = max(v[1] + v[3] for v in qr.values())
    bh = qh + 2 * pady
    qsz = getattr(qd, "sizes", {}).get(qf[0])
    if qsz and qh <= qsz * 1.25 / 72.0 * 1.45 + 0.08:        # one line: the banner is cut to it, centred
        # slack past the measured words: cut to them, a Chinese line the lint models a little wider wrapped its '。'
        cut = max(0.42 * W, 1.15 * _v2.ink_w(k, qf[0], qf[1], qsz) + 2 * padx + 1.8 * s)
        if cut < bw:
            bx, bw = bx + (bw - cut) / 2.0, cut
            if att:                                            # the attribution, re-measured under the cut banner
                ar, ad = _v2.flow(k, slide, "quote", (bx, 0.0, bw, 1.6 * s), [("attribution", att)], anchor="top",
                                  align="r")
                ah = max(v[1] + v[3] for v in ar.values()) + 0.30 * s
    block = tie + bh + ah
    # the banner hangs at the page's optical centre; the stand above it grows to meet it
    rail_y = min(H - m - block, max(stand, (H - block) * (0.46 if land else 0.52)))
    by = rail_y + tie
    # measured in 94% of the banner's inner width and drawn across all of it: a Chinese line measured to fill its box
    # was modelled a hair wider by the delivery lint, which wrapped its '。' onto a line of its own (4:3 sweep)
    iw = bw - 2 * padx
    qr, qd = _v2.flow(k, slide, "quote", (bx + padx + 0.03 * iw, by + pady, 0.94 * iw, qh + 0.02 * s),
                      [qf], anchor="top", align="c", ink=k.P["banner_ink"],
                      start=dict(getattr(qd, "sizes", {})) or None)
    r.update({f_: (bx + padx, v[1], iw, v[3]) for f_, v in qr.items()})
    ds.append(_widened(slide, qd, bx + padx, iw))
    if att:
        ar, ad = _v2.flow(k, slide, "quote", (bx, by + bh + 0.30 * s, bw, ah), [("attribution", att)], anchor="top",
                          align="r")
        r.update(ar)
        ds.append(ad)
    _stand(k, slide, rail_y)
    _banner(k, slide, bx, by, bw, bh, rail_y)
    _run_all(ds)
    return r


def _widened(slide, draw, x, w):
    """`draw`, with the text boxes it sets widened to (x, w) about their centres — words measured in a narrower column
    keep that slack in the box they are set in."""
    def go():
        n0 = len(slide.shapes)
        draw()
        for sh in list(slide.shapes)[n0:]:
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                sh.left, sh.width = dk.Inches(x), dk.Inches(w)
    go.sizes = getattr(draw, "sizes", {})
    return go


def _stand(k, slide, rail_y):
    """The stand above the rail: its terrace, rows of seats (staggered, smaller toward the back) and the front rail."""
    W, H, s, o = ctx(k)
    terrace = dk.box(slide, 0.0, 0.0, W, rail_y, fill=k.P["panel"])
    _decl(terrace, "the stand's terrace above the pitch; the banner hanging from it carries the words")
    y = rail_y - 0.24 * s
    sw_ = 0.065 * s                  # a seat's height: the dashed line's width (a dash is four widths, a gap three)
    row = 0
    while y > 0.10 * s and row < 60:
        period = 7.0 * sw_
        x0 = -period + (period / 2.0 if row % 2 else 0.0)
        sh = _paths(slide, [[(max(0.0, x0), y), (W, y)]], k.P["seat"], w=sw_ * 72.0, dash="dash")
        _decl(sh, "a row of seats in the stand; texture, nothing to read",
              "the seats are set on the stand's terrace: layered by design")
        y -= sw_ * 3.2
        sw_ = max(0.030 * s, sw_ * 0.95)
        row += 1
    rail = dk.box(slide, 0.0, rail_y - 0.035 * s, W, 0.07 * s, fill=k.P["bezel_hi"])
    _decl(rail, "the stand's front rail the banner hangs from; nothing to read",
          "the banner's ties hang from the rail: layered by design")


def _banner(k, slide, x, y, w, h, rail_y):
    """A cloth banner at (x, y, w, h): its body (the foot a little wavy), a stripe along its top and its foot, and the
    ties up to the rail."""
    W, H, s, o = ctx(k)
    amp = 0.035 * s
    foot = [(x + w - w * i / 40.0, y + h + amp * math.sin(2 * math.pi * 3.0 * i / 40.0)) for i in range(41)]
    body = na.poly(slide, [(x, y), (x + w, y)] + foot, fill=k.P["banner"])
    _decl(body, "the fans' banner the quote is painted on; the words carry the meaning")
    na.soft_shadow(body, blur=0.12, dist=0.05, alpha=0.35, color="000000")
    t1, t2 = 0.13 * s, 0.05 * s
    stripes = [[(x, y + t1), (x + w, y + t1), (x + w, y + t1 + t2), (x, y + t1 + t2)],
               [(x, y + h - t1 - t2 - amp), (x + w, y + h - t1 - t2 - amp), (x + w, y + h - t1 - amp),
                (x, y + h - t1 - amp)]]
    st = na.multipoly(slide, stripes, k.P["banner_trim"])
    _decl(st, "the stripes painted along the banner's top and foot; ornament, nothing to read",
          "the stripes are painted on the banner: layered by design")
    ties = [((x + w * u, rail_y), (x + w * u, y + 0.04 * s)) for u in (0.03, 0.35, 0.65, 0.97)]
    tl = na.strokes(slide, ties, k.P["bezel_hi"], w=1.6 * s)
    _decl(tl, "the cords that tie the banner to the rail; nothing to read",
          "the cords run from the rail onto the banner's top edge: layered by design")


@register(NAME, "data", alts=3)
def _sb_data(k, slide, f, image):
    """The number lit on the scoreboard's LED display, the board on its posts; the label and the note beside it (under
    it on a portrait or square page)."""
    W, H, s, o = ctx(k)
    m = _m(k)
    land = o == "land"
    _bands(k, slide)
    num = text_of(f, "number")
    if not num:                      # no figure to light: the label and the note on the centre spot, no empty board
        return _kickoff(k, slide, f, "data", [], groups=[(_v2.fields(f, ("label",)), 0.14 * s),
                                                          (_v2.fields(f, ("note",)), 0.0)])
    fr, pad = 0.13 * s, 0.30 * s
    if land:
        bw = (0.50, 0.44, 0.38)[alt()] * W
        bh = min(0.56 * H, bw * 0.70)
        bx, by = m + 0.2 * s, (H - bh - 0.9 * s) / 2.0
        col = (bx + bw + 0.85 * s, 0.14 * H, W - m - (bx + bw + 0.85 * s), 0.72 * H)
        leg_foot = min(H - 0.3 * s, by + bh + 0.9 * s)
    else:
        bw = (0.84, 0.80, 0.72)[alt()] * W
        bh = min((0.34, 0.28, 0.24)[alt()] * H, bw * 0.62)
        bx, by = (W - bw) / 2.0, m
        leg_foot = by + bh + 0.42 * s
        col = (0.09 * W, leg_foot + 0.35 * s, 0.82 * W, H - m - leg_foot - 0.35 * s)
        # the board and the words as ONE group at the page's optical centre: the words measured first
        r0, _d0 = _v2.flow(k, slide, "data", col, _v2.fields(f, ("label", "note")), anchor="top", align="c")
        th = (max(v[1] + v[3] for v in r0.values()) - col[1]) if r0 else 0.0
        d = max(0.0, (H - m - (col[1] + th)) * 0.45)
        by, leg_foot, col = by + d, leg_foot + d, (col[0], col[1] + d, col[2], col[3])
    disp = (bx + fr, by + fr, bw - 2 * fr, bh - 2 * fr)
    r, ds = {}, []
    if num:
        nr, nd = _v2.flow(k, slide, "data", _grow(disp, -pad), [("number", num)], anchor="middle", align="c",
                          start={"number": min(215 * s, (disp[3] - 2 * pad) * 72.0 / 1.22)})
        r.update(nr)
        ds.append(nd)
    tr, td = _v2.flow(k, slide, "data", col, _v2.fields(f, ("label", "note")), anchor="middle" if land else "top",
                      align="l" if land else "c")
    r.update(tr)
    ds.append(td)
    _legs(k, slide, bx + 0.25 * bw, bx + 0.75 * bw, by + bh - 0.05 * s, leg_foot, 0.12 * s)
    _board(k, slide, (bx, by, bw, bh), fr=fr)
    _run_all(ds)
    return r


@register(NAME, "closing", alts=2)
def _sb_closing(k, slide, f, image):
    """Full time: the whole pitch from above, every marking in chalk; the title and its line on the centre spot, the
    halfway line and the centre circle cut round them (the circle drawn wide enough to hold short words)."""
    W, H, s, o = ctx(k)
    m = _m(k)
    land = o == "land"
    _bands(k, slide)
    q = 0.42 * s
    aw, ah = W - 2 * q, H - 2 * q
    long_, short_ = (aw, ah) if land else (ah, aw)
    ratio = max(1.30, min(1.62, long_ / short_))
    if long_ / short_ > ratio:
        long_ = short_ * ratio
    else:
        short_ = long_ / ratio
    pw_, ph_ = (long_, short_) if land else (short_, long_)
    rect = ((W - pw_) / 2.0, (H - ph_) / 2.0, pw_, ph_)
    if land:
        cw = (0.58 if alt() == 0 else 0.64) * pw_
        y0, y1 = rect[1] + 0.10 * ph_, rect[1] + 0.90 * ph_
    else:
        cw = (0.84 if alt() == 0 else 0.90) * pw_
        y0, y1 = rect[1] + 0.22 * ph_, rect[1] + 0.78 * ph_
    y1 = min(y1, H - m)
    r, ds = _v2.stack(k, slide, "closing", (W - cw) / 2.0, cw, y0, y1,
                      [(_v2.fields(f, ("title",)), 0.20 * s), (_v2.fields(f, ("line",)), 0.0)], anchor="middle",
                      align="c")
    texts = {"title": text_of(f, "title"), "line": text_of(f, "line")}
    keep = _ink_rects(k, r, ds, texts, "c", 0.18 * s)
    fn = _top_down(rect, not land)
    lines, spots = _markings()
    # the centre circle: wide enough to hold short words; long words keep the pitch's own circle, cut round them
    cx, cy = W / 2.0, H / 2.0
    R0 = 9.15 / PW * short_
    bx0, by0 = min(v[0] for v in keep), min(v[1] for v in keep)
    bx1, by1 = max(v[0] + v[2] for v in keep), max(v[1] + v[3] for v in keep)
    need = math.hypot(max(abs(bx0 - cx), abs(bx1 - cx)), max(abs(by0 - cy), abs(by1 - cy))) + 0.05 * s
    R = need if need <= 0.44 * short_ else R0
    R = max(R, R0)
    mapped = _map([ln for ln in lines if not _is_centre_circle(ln)], fn)
    mapped.append(_arc(cx, cy, R, 0.0, 360.0, step=2.0))
    g = 0.30 * s
    _chalk(k, slide, mapped, (rect[0] - g, rect[1] - g, rect[2] + 2 * g, rect[3] + 2 * g), keep, w=1.5)
    _spots(k, slide, [fn(*p) for p in spots], 0.08 * s, keep)
    _run_all(ds)
    return r


def _is_centre_circle(ln):
    return len(ln) > 40 and abs(ln[0][0] - (PL / 2.0 + 9.15)) < 1e-6 and abs(ln[0][1] - PW / 2.0) < 1e-6


# ─────────────────────────────────────────── ordinary pages ───────────────────────────────────────────
def _ground_scoreboard(slide, role, index):
    """An ordinary page (a chart, an agenda) on the pitch: the mowed stripes, the touchline round the page with its
    corner arcs; the content rect inside the touchline, in REFERENCE inches."""
    W, H = vl.rs._canvas(slide)
    k = _v2.kit_for(NAME, slide)
    Wr, Hr, s, o = ctx(k)
    _bands(k, slide)
    q, c = 0.30 * s, 0.30 * s
    lines = [[(q, q), (Wr - q, q), (Wr - q, Hr - q), (q, Hr - q), (q, q)],
             _arc(q, q, c, 0.0, 90.0, 15.0), _arc(Wr - q, q, c, 90.0, 180.0, 15.0),
             _arc(Wr - q, Hr - q, c, 180.0, 270.0, 15.0), _arc(q, Hr - q, c, 270.0, 360.0, 15.0)]
    _chalk(k, slide, lines, (0.0, 0.0, Wr, Hr), [], w=1.2, alpha=_chalk_alpha(k) * 0.8,
           reason="the touchline round the page and its corner arcs: the pitch the page is set on")
    return (0.75, 0.70, W - 1.5, H - 1.4)


_card_scoreboard = vl._card_for(NAME)
