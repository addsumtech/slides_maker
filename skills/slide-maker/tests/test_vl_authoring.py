"""Phase 4 — the authoring kit and any-agent use (spec §7): a sixteenth language is started by one command and named
by the checks until it is finished, without breaking the fifteen; the docs and the scripts' printed output are enough
for an agent that is not Claude. Prints `N passed, 0 failed`."""
import contextlib, copy, io, os, pprint, re, shutil, subprocess, sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import vl_contract as vc, vl_specs

ok, bad = [], []
def check(cond, why):
    (ok if cond else bad).append(why)
WORK = Path(tempfile.mkdtemp(prefix="vlauth_"))
SHIPPED = vl_specs.load()


def _write_spec(d, spec):
    (d / (spec["name"] + ".py")).write_text("SPEC = " + pprint.pformat(spec, width=118, sort_dicts=False) + "\n",
                                            encoding="utf-8")


def mosaic_spec():
    """tally's manifest as a sixteenth language: its own name, order, painters, and a ground/accent far enough from
    every shipped language's that the manifest validates on its own."""
    s = copy.deepcopy(SHIPPED["tally"])
    s["name"], s["order"] = "mosaic", max(v["order"] for v in SHIPPED.values()) + 1
    s["painters"] = {"ground": "_ground_mosaic", "card": "_card_mosaic"}
    for g, (ground, accent) in zip(s["grounds"], (("FAF6FC", "6A1B9A"), ("1E0B2E", "F06292"))):
        s["grounds"][g]["palette"]["ground"] = ground
        s["grounds"][g]["palette"]["accents"] = [accent] + list(s["grounds"][g]["palette"]["accents"][1:])
    s["look"]["palette"] = s["grounds"]["light"]["palette"]
    return s


def specs_dir_with(*extra):
    d = Path(tempfile.mkdtemp(prefix="vlspecs_", dir=WORK))
    for f in (ROOT / "scripts" / "vl_specs").glob("*.py"):
        if not f.name.startswith("_"):
            shutil.copy(f, d / f.name)
    for s in extra:
        _write_spec(d, s)
    return d


def run(code_or_args, env_extra=None, script=False):
    env = dict(os.environ, **(env_extra or {}))
    cmd = [sys.executable] + (code_or_args if script else ["-c", "import sys; sys.path.insert(0, 'scripts'); " + code_or_args])
    return subprocess.run(cmd, cwd=ROOT, env=env, capture_output=True, text=True, timeout=600)


# ── Task 1: a language's code is discovered; a manifest whose painters are not written breaks only that language ──
_m = mosaic_spec()
check(not [p for p in vc.validate(_m)], "the mosaic fixture manifest validates on its own ({})".format(vc.validate(_m)[:1]))
check(not [p for p in vc.distinct_pairs(dict(SHIPPED, mosaic=_m)) if "mosaic" in p],
      "the mosaic fixture is distinct from every shipped language")
D16 = specs_dir_with(_m)
ENV16 = {"VL_SPECS_DIR": str(D16)}
r = run("import deckkit as dk, visual_languages as vl; prs = dk.blank_deck(13.333, 7.5); k = vl.use('tally', prs); "
        "k.cover(k.new_slide(), title='Fix it together'); print('ok')", ENV16)
check(r.returncode == 0 and "ok" in r.stdout, "a 16th manifest with no code leaves tally working ({})".format(r.stderr[-160:]))
r = run("import deckkit as dk, visual_languages as vl; prs = dk.blank_deck(13.333, 7.5); vl.use('mosaic', prs)", ENV16)
check(r.returncode != 0 and "vl_langs/mosaic.py" in r.stderr and "_ground_mosaic" in r.stderr and "KeyError" not in r.stderr,
      "using the unready language is refused by name, naming the file to write ({})".format(r.stderr[-160:]))
r = run([str(ROOT / "scripts" / "check_vl_contract.py"), "mosaic"], ENV16, script=True)
check(r.returncode == 1 and "_ground_mosaic" in r.stdout and "Traceback" not in r.stdout + r.stderr,
      "check_vl_contract names the unwritten painters instead of crashing ({} {})".format(r.returncode, (r.stdout + r.stderr)[-160:]))
_bad = copy.deepcopy(_m); _bad["painters"] = {"ground": "ground", "card": "_card_mosaic"}
check(any("painters" in p and "_ground_mosaic" in p for p in vc.validate(_bad)),
      "a painter not named _ground_<name> / _card_<name> is refused by the manifest check")

# ── Task 2: an unfilled field and a stub page are named, with the file to edit ──
_s = mosaic_spec()
_s["display"]["zh"] = "FILL: the language's Chinese name"
_s["grounds"]["night"]["palette"]["accents"][0] = "FILL: one accent that reads on the night ground"
_p = vc.validate(_s)
check(any("display.zh" in x and "Chinese name" in x for x in _p), "an unfilled display.zh is named by its path and hint ({})".format(_p[:2]))
check(any("grounds.night.palette.accents[0]" in x for x in _p), "an unfilled list item is named by its path")
check(_p and all("not filled" in x for x in _p), "while fields are unfilled, only they are reported (not as bad hexes)")
check(vc.unfilled(mosaic_spec()) == [], "a filled manifest has nothing unfilled")
import vl_native as vn, vl_conformance as vcf
vn.stub("stubcheck", "cover")
_pm = vcf.pages_missing("stubcheck", registry={"stubcheck": {"kind": "drawn"}})
check(any("cover" in x and "still a stub" in x and "scripts/vl_langs/stubcheck.py" in x for x in _pm),
      "a stub page is a missing page, naming the file to compose it in ({})".format(_pm[:1]))
check(len(_pm) == 7, "the six unregistered pages are missing too ({})".format(len(_pm)))
try:
    vn.COMPOSERS["stubcheck"]["cover"](None, None, {}, None); _raised = None
except NotImplementedError as e:
    _raised = str(e)
check(_raised and "stubcheck" in _raised and "cover" in _raised, "a stub raises, naming its language and page")

# ── Task 3: new_visual_language.py starts a language; check_vl_contract names everything still to do ──
GEN = str(ROOT / "scripts" / "new_visual_language.py")
CHK = str(ROOT / "scripts" / "check_vl_contract.py")
SD, LD = specs_dir_with(), Path(tempfile.mkdtemp(prefix="vllangs_", dir=WORK))
ENVG = {"VL_SPECS_DIR": str(SD), "VL_LANGS_DIR": str(LD)}
r = run([GEN, "mosaic", "--like", "tally", "--specs-dir", str(SD), "--langs-dir", str(LD)], script=True)
check(r.returncode == 0 and (SD / "mosaic.py").is_file() and (LD / "mosaic.py").is_file()
      and re.search(r"NEXT: .*python3 \S*check_vl_contract\.py mosaic", r.stdout),
      "the generator writes the manifest and the module and prints the next command ({})".format((r.stdout + r.stderr)[-200:]))
_gen = vl_specs.load(SD, drafts=True).get("mosaic") if (SD / "mosaic.py").is_file() else None
check("mosaic" not in vl_specs.load(SD) and "mosaic" in vl_specs.DRAFTS,
      "an ordinary load leaves the draft out and names it in DRAFTS")
check(bool(_gen) and _gen["display"]["zh"].startswith("FILL:") and _gen["painters"] == {"ground": "_ground_mosaic", "card": "_card_mosaic"}
      and _gen["look"]["palette"] is _gen["grounds"]["light"]["palette"],
      "the written manifest loads, its Chinese hint intact, its light palette one object")
r = run([CHK, "mosaic"], ENVG, script=True)
_out = r.stdout + r.stderr
for _path in ("display.en", "display.zh", "rationale", "offer.en[0]", "grounds.light.palette.ground", "grounds.light.label",
              "look.fonts.both.display", "samples.copy"):
    check("mosaic: {} is not filled".format(_path) in _out, "check_vl_contract names the unfilled {} ({})".format(_path, _out[-160:]))
check(r.returncode == 1 and "Traceback" not in _out, "an untouched draft is reported, not crashed on")
r = run("import deckkit as dk, visual_languages as vl; prs = dk.blank_deck(13.333, 7.5); k = vl.use('tally', prs); "
        "k.cover(k.new_slide(), title='Fix it together'); print('ok')", ENVG)
check(r.returncode == 0 and "ok" in r.stdout, "an untouched draft leaves tally working ({})".format(r.stderr[-160:]))
for _args, _why, _word in ((["tally"], "an existing language", "already"), (["Pebble"], "an uppercase name", "lowercase"),
                           (["my-lang"], "a hyphen", "lowercase"), (["9lives"], "a leading digit", "lowercase"),
                           (["mosaic2", "--like", "editorial"], "an image-led model", "image-led"),
                           (["mosaic3", "--like", "nosuch"], "an unknown model", "nosuch")):
    r = run([GEN] + _args + ["--specs-dir", str(SD), "--langs-dir", str(LD)], script=True)
    # the message must say WHY — python's own exit 2 for a missing script passed this check vacuously at first
    check(r.returncode == 2 and _word in (r.stdout + r.stderr) and "new_visual_language" in (r.stdout + r.stderr)
          and (_args[0] == "tally" or not (SD / (_args[0] + ".py")).exists()),
          "the generator refuses {}, saying why ({} {})".format(_why, r.returncode, (r.stdout + r.stderr)[-120:]))
# end to end: decide every FILL (here: the values of the mosaic fixture) …
_want = mosaic_spec()


def _at(o, path):
    for part in re.findall(r"[^.\[\]]+|\[\d+\]", path):
        o = o[int(part[1:-1])] if part.startswith("[") else o[part]
    return o


def _fill(o, path=""):
    if isinstance(o, dict):
        return {k: _fill(v, "{}.{}".format(path, k) if path else k) for k, v in o.items()}
    if isinstance(o, list):
        return [_fill(v, "{}[{}]".format(path, i)) for i, v in enumerate(o)]
    if isinstance(o, tuple):
        return tuple(_fill(v, "{}[{}]".format(path, i)) for i, v in enumerate(o))
    return copy.deepcopy(_at(_want, path)) if isinstance(o, str) and o.startswith("FILL:") else o
if _gen:
    _filled = _fill(_gen)
    for _k in ("extras", "derived", "samples"):
        _filled[_k] = copy.deepcopy(_want[_k])
    _filled["look"]["palette"] = _filled["grounds"]["light"]["palette"]
    _write_spec(SD, _filled)
r = run([CHK, "mosaic"], ENVG, script=True)
_out = r.stdout + r.stderr
check(r.returncode == 1 and _out.count("is still a stub") == 7 and "scripts/vl_langs/mosaic.py" in _out,
      "with the manifest decided, the seven stub pages are named ({})".format(_out[-200:]))
# … then compose (here: borrow tally's compositions) — what remains is the shipping list, each item named
_mod = (LD / "mosaic.py").read_text(encoding="utf-8") if (LD / "mosaic.py").is_file() else ""
check("vn.stub(" in _mod, "the module drafts one stub per page")
(LD / "mosaic.py").write_text(_mod + "\n\nfor _p in PAGES:      # composed: here, tally's compositions\n"
                              "    vn.register('mosaic', _p, vn.ALTS[('tally', _p)])(vn.COMPOSERS['tally'][_p])\n",
                              encoding="utf-8")
r = run([CHK, "mosaic"], ENVG, script=True)
_fails = [l for l in (r.stdout + r.stderr).splitlines() if l.strip().startswith("FAIL")]
check(_fails and all(any(w in l for w in ("bundled sample", "fingerprint", "sigs.py --example", "table row", "gallery"))
                     for l in _fails),
      "composed, only the shipping steps remain — samples, scaffold, docs row, gallery ({})".format(_fails[:3]))

# ── Task 4: --describe prints a language's card and an example that runs as printed ──
VLPY = str(ROOT / "scripts" / "visual_languages.py")
from pptx import Presentation as _P
for _n, _sp in SHIPPED.items():
    r = run([VLPY, "--describe", _n], script=True)
    _card = r.stdout
    check(r.returncode == 0 and all(h in _card for h in ("grounds:", "pages:", "extras:", "derived words:",
                                                          "offered for:", "fonts:", "\nexample (")),
          "--describe {}: every section of the card is printed ({})".format(_n, (r.stdout + r.stderr)[-120:]))
    check(all(x + "=" in _card for x in _sp["extras"]), "--describe {}: every extra is listed by name".format(_n))
    _code = _card.split("\nexample (", 1)[-1].split("\n", 1)[-1] if "\nexample (" in _card else ""
    _wd = Path(tempfile.mkdtemp(prefix="desc_", dir=WORK))
    rr = subprocess.run([sys.executable, "-c", _code], cwd=_wd, capture_output=True, text=True, timeout=600)
    _decks = list(_wd.glob("*.pptx"))
    check(rr.returncode == 0 and len(_decks) == 1 and len(_P(str(_decks[0])).slides) == 4,
          "--describe {}: its example runs as printed and saves a 4-page deck ({})".format(_n, rr.stderr[-160:]))
r = run([VLPY, "--describe", "nosuch"], script=True)
check(r.returncode == 2 and "tally" in r.stderr, "--describe of an unknown name lists the registered languages")
r = run([VLPY, "--describe", "mosaic"], ENV16, script=True)
check(r.returncode == 0 and "not ready" in r.stdout and "vl_langs/mosaic.py" in r.stdout and "\nexample (" not in r.stdout,
      "--describe of a language with no code says what to write, and prints no example ({})".format((r.stdout + r.stderr)[-160:]))
SDd, LDd = specs_dir_with(), Path(tempfile.mkdtemp(prefix="vllangs_", dir=WORK))
run([GEN, "pebble", "--specs-dir", str(SDd), "--langs-dir", str(LDd)], script=True)
r = run([VLPY, "--describe", "pebble"], {"VL_SPECS_DIR": str(SDd), "VL_LANGS_DIR": str(LDd)}, script=True)
check(r.returncode == 0 and "draft" in r.stdout and "check_vl_contract.py pebble" in r.stdout,
      "--describe of a draft says it is a draft and how to see what to decide ({})".format((r.stdout + r.stderr)[-160:]))

# ── Task 5: sigs.py places by CHARACTER offsets (M1); --in / --ground are checked when typed (M8) ──
import sigs
sigs.EXAMPLES["probe_cjk"] = ('s = dk.add_slide(prs)\n'
                              't = "中文标题"; dk.text(s, 0.8, 0.6, 6.0, 1.0, [[(t, 24, dk.RGBColor(0x22, 0x22, 0x22), True, False)]])\n')
try:
    _placed = sigs.placed_example("probe_cjk", "tally"); _err = None
except Exception as e:
    _placed, _err = "", "{}: {}".format(type(e).__name__, e)
try:
    compile(_placed, "<placed>", "exec"); _compiles = True
except SyntaxError:
    _compiles = False
check(_err is None and _compiles and "dk.text(s, x, y, w, h, [[(t," in _placed and '"中文标题"' in _placed,
      "a scaffold with Chinese before its call is placed with its arguments intact ({})".format(_err or _placed[-140:]))
del sigs.EXAMPLES["probe_cjk"]
SIGS = str(ROOT / "scripts" / "sigs.py")
for _args, _word in ((["timeline", "--example", "--ground", "night"], "--ground goes with --in"),
                     (["timeline", "--example", "--in", "nosuch"], "tally"),
                     (["timeline", "--example", "--in", "tally", "--ground", "nosuch"], "night")):
    r = run([SIGS] + _args, script=True)
    check(r.returncode == 2 and _word in r.stderr, "sigs.py {} is refused when typed, saying why ({} {})".format(
        " ".join(_args), r.returncode, (r.stdout + r.stderr)[-120:]))

# ── Task 6: a gate's pass summary never sits beside a fault of its own section (friction 2) ──
def _fixture_deck(d, n=3, name="d.pptx"):
    code = ("import sys; sys.path.insert(0, {!r}); import deckkit as dk; prs = dk.blank_deck(13.333, 7.5)\n"
            "for t in {!r}:\n    s = dk.add_slide(prs); dk.text(s, 1, 1, 6, 1, [[(t, 28, dk.RGBColor(0x22, 0x22, 0x22), True, False)]])\n"
            "prs.save({!r})").format(str(ROOT / "scripts"), ["Page {}".format(i + 1) for i in range(n)], name)
    subprocess.run([sys.executable, "-c", code], cwd=d, check=True, capture_output=True)
    return d / name


GD = Path(tempfile.mkdtemp(prefix="gates_", dir=WORK))
_deck = _fixture_deck(GD)
run([str(ROOT / "scripts" / "deck_gates.py"), "init", str(GD), "--slides", "3"], script=True)
r = run([str(ROOT / "scripts" / "render_deck.py"), str(_deck), "--gate-check"], script=True)
_gc = r.stdout + r.stderr
check("hand-off gate(s) failed" in _gc, "the template record fails the gate check (fixture sanity)")
check(not ("[gates] blind read:" in _gc and "`blind_read`" in _gc),
      "blind read is not reported answered beside its own 'answers is missing'")
check(not ("all accounted for" in _gc and re.search(r"\[\d+/\d+\] taste\b", _gc)),
      "the taste ledger is not reported accounted for beside a taste fault")
check(not re.search(r"\[gates\] (build shape|audience brief): .*<", _gc),
      "no pass summary prints a template placeholder as its result ({})".format(
          re.findall(r"\[gates\] (?:build shape|audience brief): .*", _gc)[:2]))

# ── Task 7: the waiver the GROUND REPEAT message names exists — "deliberate repeat: <why>", history checks only ──
import json as _json
import check_register_pixels as crp
from PIL import Image as _Img


def _rp_deck(tag, ground, accents, palette_line, rows):
    d = Path(tempfile.mkdtemp(prefix="rp_" + tag + "_", dir=WORK))
    (d / "render").mkdir()
    for i in range(3):
        im = _Img.new("RGB", (480, 270), ground)
        px = im.load()
        for j, a in enumerate(accents):
            for y in range(20 + j * 40, 36 + j * 40):
                for x in range(20, 260):
                    px[x, y] = a
        im.save(d / "render" / ("slide%02d.png" % (i + 1)))
    (d / ".deck-gates.json").write_text(_json.dumps({"design_plan": {"palette": palette_line}}), encoding="utf-8")
    tp = d / "taste.md"
    tp.write_text("## LOOK HISTORY\n| date | deck | look | canvas | motif |\n|---|---|---|---|---|\n" + "\n".join(
        "| 2026-01-0%d | %s | x | %s | y |" % (i + 1, nm, " ".join("#%02X%02X%02X" % c for c in cols))
        for i, (nm, cols) in enumerate(rows)) + "\n", encoding="utf-8")
    return d, str(tp)


_W, _B = (0xFF, 0xFF, 0xFF), (0x24, 0x38, 0xF0)
_hist = [("white1", [(0xFF, 0xFF, 0xFF)]), ("white2", [(0xFA, 0xFA, 0xFA)])]
for _tag, _line, _want_repeat in (("plain", "#FFFFFF / #2438F0", True),
                                  ("named", "#FFFFFF / #2438F0 — deliberate repeat: the user named the light tally ground", False),
                                  ("empty", "#FFFFFF / #2438F0 — deliberate repeat:", True)):
    _d, _tp = _rp_deck(_tag, _W, [_B], _line, _hist)
    _probs, _facts = crp.check(_d, taste=_tp)
    _codes = {c for c, _m in _probs}
    check(("GROUND REPEAT" in _codes) == _want_repeat, "repeat waiver ({}): GROUND REPEAT {} ({})".format(
        _tag, "held" if _want_repeat else "waived", sorted(_codes)))
    if _tag == "plain":
        _msg = next((m for c, m in _probs if c == "GROUND REPEAT"), "")
        check("deliberate repeat: <why>" in _msg, "the GROUND REPEAT message names the exact words of its waiver")
    if _tag == "named":
        check("user named" in str(_facts.get("repeat", "")), "the waiver's reason is recorded in the facts")
_d, _tp = _rp_deck("hue", _W, [_B], "#FFFFFF / #C4122E — deliberate repeat: the user named the light ground", _hist)
_codes = {c for c, _m in crp.check(_d, taste=_tp)[0]}
check(_codes & {"DECLARED HUES ABSENT", "DECLARED PALETTE ABSENT"},
      "the repeat waiver never silences a declared colour that is not on the pages ({})".format(sorted(_codes)))

# ── Task 8: one deck's render/ is not another's (friction 4) ──
import hashlib as _hl
RENDER = str(ROOT / "scripts" / "render_deck.py")
RD = Path(tempfile.mkdtemp(prefix="twodecks_", dir=WORK))
_a, _b = _fixture_deck(RD, 2, "a.pptx"), _fixture_deck(RD, 2, "b.pptx")


def _png_hash():
    p = RD / "render" / "slide01.png"
    return _hl.md5(p.read_bytes()).hexdigest() if p.is_file() else None


r = run([RENDER, str(_a)], script=True)
check(r.returncode == 0 and _png_hash(), "deck a renders into render/ ({})".format((r.stdout + r.stderr)[-120:]))
_h = _png_hash()
r = run([RENDER, str(_b)], script=True)
_o = r.stdout + r.stderr
check(r.returncode == 1 and "a.pptx" in _o and "folder of its own" in _o and _png_hash() == _h,
      "a second deck into the same render/ is refused, naming the first and the way out; a's pages kept ({})".format(_o[-200:]))
r = run([RENDER, str(_a)], script=True)
check(r.returncode == 0, "the same deck renders again into its own render/")
r = run([RENDER, str(_b), str(RD / "render-b")], script=True)
check(r.returncode == 0 and (RD / "render-b" / "slide01.png").is_file(), "deck b renders into the out dir it is given")
_cp = RD / "render" / ".render-cache.json"
_c = _json.loads(_cp.read_text(encoding="utf-8")) if _cp.is_file() else {}
check(_c.get("deck") == "a.pptx", "the render cache records which deck it holds ({})".format(_c.get("deck")))
_c.pop("deck", None); _cp.write_text(_json.dumps(_c), encoding="utf-8")
r = run([RENDER, str(_b)], script=True)
check(r.returncode == 0, "a cache written before the deck key existed is not refused")
run([str(ROOT / "scripts" / "deck_gates.py"), "init", str(RD), "--slides", "2"], script=True)
r = run([RENDER, str(_a), "--gate-check"], script=True)
check("render/ holds b.pptx's pages" in r.stdout + r.stderr,
      "the gate check names renders made from another deck ({})".format((r.stdout + r.stderr)[-160:]))

# ── Task 9: the docs an agent reads say what the scripts really do ──
_gd = Path(tempfile.mkdtemp(prefix="gatesd_", dir=WORK))
r = run([VLPY, "--gates", "tally", "--deck", str(_gd)], script=True)
check(r.returncode == 0 and "light — the default" in r.stdout and "--ground night" in r.stdout,
      "--gates without --ground says it recorded the default ground, and how to pick the other ({})".format(r.stdout[:160]))
r = run([VLPY, "--gates", "tally", "--deck", str(_gd), "--ground", "night"], script=True)
check(r.returncode == 0 and "deliberate repeat: the user named the night ground" in r.stdout,
      "--gates with a named ground prints the repeat waiver to use when the user chose it")

GUIDE = (ROOT / "references" / "visual-language-authoring.md").read_text(encoding="utf-8") \
    if (ROOT / "references" / "visual-language-authoring.md").is_file() else ""
check(bool(GUIDE), "references/visual-language-authoring.md exists")
for _scr in sorted(set(re.findall(r"scripts/([a-z_0-9]+\.py)", GUIDE))):
    check((ROOT / "scripts" / _scr).is_file(), "the guide's scripts/{} exists".format(_scr))
for _scr, _flags in re.findall(r"scripts/([a-z_0-9]+\.py)((?: [^`\n]*?)?)`", GUIDE):
    for _fl in re.findall(r"(--[a-z][a-z-]+)", _flags):
        _src = (ROOT / "scripts" / _scr).read_text(encoding="utf-8") if (ROOT / "scripts" / _scr).is_file() else ""
        check('"{}"'.format(_fl) in _src, "the guide's {} {} is a real flag of that script".format(_scr, _fl))
for _ref in sorted(set(re.findall(r"references/([a-z_\-]+\.md)", GUIDE))):
    check((ROOT / "references" / _ref).is_file(), "the guide's references/{} exists".format(_ref))
for _doc in ("SKILL.md", "references/visual-languages.md"):
    check("visual-language-authoring.md" in (ROOT / _doc).read_text(encoding="utf-8"), "{} points to the authoring guide".format(_doc))
# the guide's example composition runs, in the language Task 3 composed
_ex = re.search(r"```python\n(.*?)```", GUIDE, re.S)
check(bool(_ex), "the guide carries a composition example")
if _ex and (LD / "mosaic.py").is_file():
    (LD / "mosaic.py").write_text((LD / "mosaic.py").read_text(encoding="utf-8") + "\n\n" + _ex.group(1), encoding="utf-8")
    r = run("import deckkit as dk, visual_languages as vl; prs = dk.blank_deck(13.333, 7.5); k = vl.use('mosaic', prs); "
            "k.cover(k.new_slide(), kicker='Q3 review', title='The lending library, quarter three'); "
            "fs = [f for f in dk.lint_layout(prs, verbose=False) if f[1] == 'CRITICAL']; print('criticals', len(fs), fs[:1])", ENVG)
    check(r.returncode == 0 and "criticals 0" in r.stdout,
          "the guide's cover example builds a lint-clean cover ({})".format((r.stdout + r.stderr)[-200:]))

# ── Task 11: what the two docs-only runs hit (EN broadsheet, ZH chalkboard — 2026-10-16) ──
import audience_brief as _ab
_brief = {"who": "社区里想借工具、但从没来过图书馆的居民", "decisions": [
    {"decision": "怎么借", "needs": "登记要带什么证件，在哪里登记"},
    {"decision": "什么时候去", "needs": "开放时间和每次能借多久"},
    {"decision": "借哪一件工具最合适", "needs": "有哪些工具可借，适合做什么"}]}
_f = _ab.faults(_brief)
check(not any("is empty" in x for x in _f if "decisions[0]" in x or "decisions[1]" in x),
      "a short Chinese decision is not called empty ({})".format([x for x in _f if "decisions[" in x][:1]))
check(any("decisions[0]" in x and "short" in x for x in _f), "a too-short decision is called too short, with the bar")
r = run([str(ROOT / "scripts" / "arc_divergence.py"), "--help"], script=True)
check("path" in r.stdout.lower() and "json" in r.stdout.lower(), "arc_divergence --help says it takes a PATH to a JSON file")
import check_image_provenance as _cip
_msg = " ".join(m for _c, m in _cip.check_rows(["not a token"])) if hasattr(_cip, "check_rows") else ""
_src = (ROOT / "scripts" / "check_image_provenance.py").read_text(encoding="utf-8")
check("the whole field as the string `n/a" in _src, "the BAD TOKEN message names the n/a form a deck with no picture uses")
_vld = (ROOT / "references" / "visual-languages.md").read_text(encoding="utf-8")
for _needle, _why in (("interview.picks", "the auto checkpoint's picks table"), ("declare_delivery", "how to declare the delivery mode"),
                      ("provenance.waived", "how to record a deck whose facts are all the user's"),
                      ("build_<deck stem>.py", "the build script's name the STYLE APPLIED check looks for"),
                      ("inline_ran", "the critic waiver's extra field on a host with no dispatch"),
                      ('"waived_category": "conservative"', "the composition record of a conservative deck")):
    check(_needle in _vld, "visual-languages.md says {} ({})".format(_why, _needle))

# ── Task 12: the whole-branch review's findings, each pinned ──
# (I1) the render refusal's remedy must be one the gates accept: they read render/ beside the deck, so the way out is
# a folder of its own — "render into render-b" left deck b failing its gates for ever
RD2 = Path(tempfile.mkdtemp(prefix="twodecks2_", dir=WORK))
_a2, _b2 = _fixture_deck(RD2, 2, "a.pptx"), _fixture_deck(RD2, 2, "b.pptx")
run([RENDER, str(_a2)], script=True)
r = run([RENDER, str(_b2)], script=True)
_o = r.stdout + r.stderr
check(r.returncode == 1 and "folder of its own" in _o and "render-b" not in _o,
      "the refusal sends the second deck to a folder of its own, not to render-b ({})".format(_o[-160:]))
run([str(ROOT / "scripts" / "deck_gates.py"), "init", str(RD2), "--slides", "2"], script=True)
r = run([RENDER, str(_b2), "--gate-check"], script=True)
_o = r.stdout + r.stderr
check("render/ holds a.pptx's pages" in _o and "folder of its own" in _o and "render-b" not in _o,
      "the gate's remedy is the same one ({})".format(_o[-160:]))
(RD2 / "a.pptx").rename(RD2 / "c.pptx")             # the deck that owned render/ was renamed: a plain re-render fixes it
r = run([RENDER, str(RD2 / "c.pptx"), "--gate-check"], script=True)
_o = r.stdout + r.stderr
check("render/ holds a.pptx's pages" in _o and "re-render" in _o and "folder of its own" not in _o,
      "renders of a deck no longer here are fixed by a plain re-render, and the gate says so ({})".format(_o[-160:]))
# (I2) the waiver's own template is not a waiver
for _txt in ("#FFF / #2438F0 — deliberate repeat: <why>", "#FFF — deliberate repeat: ....", "#FFF — deliberate repeat:  "):
    check(crp._repeat_waiver({"palette": _txt}) is None, "{!r} is not a waiver".format(_txt[-28:]))
check(crp._repeat_waiver({"palette": "#FFF — deliberate repeat: 用户点名要这个底色"}), "a Chinese reason is a waiver")
# (I3) importing vl_native (or vl_native2) BEFORE visual_languages leaves a ready language ready
for _first in ("vl_native", "vl_native2"):
    r = run("import {}; import deckkit as dk, visual_languages as vl; prs = dk.blank_deck(13.333, 7.5); "
            "k = vl.use('mosaic', prs); k.cover(k.new_slide(), title='Fix it together'); print('ok')".format(_first), ENVG)
    check(r.returncode == 0 and "ok" in r.stdout, "importing {} first leaves mosaic usable ({})".format(_first, r.stderr[-160:]))
# a module that raises is named as raising, not as unwritten
_bad = Path(tempfile.mkdtemp(prefix="vllangs_bad_", dir=WORK))
(_bad / "mosaic.py").write_text("raise RuntimeError('half written')\n", encoding="utf-8")
r = run("import deckkit as dk, visual_languages as vl; prs = dk.blank_deck(13.333, 7.5); vl.use('mosaic', prs)",
        {"VL_SPECS_DIR": str(D16), "VL_LANGS_DIR": str(_bad)})
check("half written" in r.stderr and "not written yet" not in r.stderr,
      "a module that raises is reported with its error ({})".format(r.stderr[-160:]))
# (minor → fixed) a name whose painters collide with this module's own functions (_card_for, _card_geom …)
r = run([GEN, "for", "--specs-dir", str(SD), "--langs-dir", str(LD)], script=True)
check(r.returncode == 2 and "collide" in r.stderr, "the generator refuses a name whose painters collide ({})".format(r.stderr[-120:]))

for line in ok:
    print("  ok   " + line)
for line in bad:
    print("  FAIL " + line)
print("\n{} passed, {} failed".format(len(ok), len(bad)))
sys.exit(1 if bad else 0)
