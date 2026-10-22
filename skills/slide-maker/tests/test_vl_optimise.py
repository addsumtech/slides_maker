"""Phase 3 — the languages' own pages, measured (spec §6): quote scale, title fit, rules at their content, designed dark
grounds, the smallest tier, distinctiveness, and the per-language fixes. Prints `N passed, 0 failed`."""
import contextlib, io, os, re, sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import deckkit as dk, visual_languages as vl, lint_deck
import vl_native2 as vls

ok, bad = [], []
def check(cond, why):
    (ok if cond else bad).append(why)
WORK = Path(tempfile.mkdtemp(prefix="vlopt_")); os.chdir(WORK)
# a size or a rendered gap measured on the languages' REAL faces: the ubuntu CI runner has none of them (its stand-ins
# are wider, and its CJK stand-in has no glyphs for the layout to measure), so those checks record a SKIP there — printed
# at the end — instead of failing on the stand-in's widths (as test_native_languages_p4 does; 2026-10-16)
import vl_specs as _vls
_FACES = {f for sp in _vls.load().values() for v in (sp["look"]["fonts"].get("mac") or sp["look"]["fonts"]["both"]).values()
          for f in ([v] if isinstance(v, str) else [])} | {"Songti SC", "Hiragino Sans GB"}
FACES_HERE = sys.platform == "darwin" and not any(dk._font_substituted(f) for f in _FACES)
skipped = []


def check_mac(cond, why):
    if FACES_HERE:
        check(cond, why)
    else:
        skipped.append(why)


def _deck(L, g=None):
    prs = dk.blank_deck(13.333, 7.5)
    with contextlib.redirect_stdout(io.StringIO()):
        k = vl.use(L, prs, ground=g) if g else vl.use(L, prs)
    return prs, k


def _hard(prs, tag):
    p = WORK / (tag + ".pptx"); prs.save(str(p)); buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
        lint_deck.lint(str(p), static_ok=True)
    return [l.strip()[:150] for l in buf.getvalue().splitlines()
            if re.match(r"\s+slide \d+: [A-Z]", l) and "[warn]" not in l and "[stats]" not in l]


def _size_of(slide, words):
    """The largest run size of the text shape holding `words` (the quote's own runs)."""
    best = 0.0
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False) and words.lower() in sh.text_frame.text.replace("\n", " ").lower():
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.text.strip():
                        best = max(best, r.font.size.pt)
    return best


# ── Task 1 (F1): a quote page scales with the quote; the attribution is a readable line, not a caption ──
SHORT, LONG = "Fix it together.", " ".join(["Every repair we made that evening was a small argument for keeping"
                                            " things, and the room kept making it until the last lamp came on."] * 3)
for L in vl.LANGS:
    base = vl.TYPE[L]["quote"][0]
    prs, k = _deck(L)
    s1 = k.new_slide()
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            k.quote(s1, quote=SHORT, attribution="Repair café volunteer")
    except Exception as e:
        check(False, "{}: a short quote builds ({}: {})".format(L, type(e).__name__, str(e)[:80])); continue
    z = _size_of(s1, "Fix it")                  # poster stacks a quote one line per shape
    # ink sets a CJK couplet vertically; poster stacks its quote as display type already sized to the page (27% of it)
    check(z >= 1.3 * base - 1e-6 or L in ("ink", "poster"), "{}: a 3-word quote is set larger than the language's quote size "
          "({:.1f}pt vs {}pt)".format(L, z, base))
    za = _size_of(s1, "repair caf")             # some languages set it in caps
    check(za >= 16 - 1e-6, "{}: the attribution is >= 16pt ({:.1f}pt)".format(L, za))
    s2 = k.new_slide()
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            k.quote(s2, quote=LONG, attribution="Repair café volunteer")
        refused = None
    except vl.VLTextOverflow as e:
        refused = str(e)
    judged = ("slide 1:",) if refused else ("slide 1:", "slide 2:")     # a refused quote leaves its page empty, loudly
    found = [f for f in _hard(prs, "q_" + L) if f.startswith(judged)]
    check(not found, "{}: the quote pages are lint-clean ({})".format(L, found[:1]))
dk.set_theme(None)

# ── Task 2 (F2): a cover title wraps before it shrinks — a phrase per line only while that costs <= 15% of the size ──
TITLES = [("short", "Fix it together", 0.85), ("comma", "Bring it broken, take it home working", 0.85),
          ("cjk", "带着坏东西来，带着好东西走", 0.85)]
for L in vl.LANGS:
    base = vl.TYPE[L]["title"][0]
    for tag, t, need in TITLES:
        prs, k = _deck(L)
        sl = k.new_slide()
        try:
            pic = str(vl.ASSETS / "photo/hall-repair.jpg") if L in vl.IMAGE_LED else None   # the narrow column
            with contextlib.redirect_stdout(io.StringIO()):
                k.cover(sl, title=t, kicker="Repair café", **({"image": pic} if pic else {}))
        except Exception as e:
            check(False, "{} cover ({}) builds ({}: {})".format(L, tag, type(e).__name__, str(e)[:70])); continue
        words = t.split()[0] if tag != "cjk" else t[:2]
        z = _size_of(sl, words)
        # poster's title size is the CEILING of a display stack sized to the page, not a target
        # ink sets a CJK title vertically in the approved two-column scroll (bounded by the column's height)
        check_mac(z >= need * base - 1e-6 or L == "poster" or (L == "ink" and tag == "cjk"), "{} cover title ({}) holds >= {:.0%} of {}pt ({:.1f}pt)".format(L, tag, need, base, z))
dk.set_theme(None)


def _drawn_lines(k, field, sh):
    """The lines text shape `sh` draws: each paragraph, wrapped at the shape's measure the way the renderer wraps it."""
    paras = ["".join(r.text for r in p.runs).strip() for p in sh.text_frame.paragraphs]
    paras = [p for p in paras if p]
    size = max(r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size)
    box = sh.width / 914400.0
    return [l_.strip() for p in paras for l_ in vl._break_lines(k, field, p, size, box * vl.HEADROOM)]


# ── Task 2 (F2, sheet review 2026-10-16): a cover title that wraps never leaves its last word alone when the SAME line
# count can end on two — collage's "a repair / night" stood alone because the widest line was the one that had to take it,
# and a narrower measure cannot move a word INTO the widest line ──
for L in vl.LANGS:
    for t in ("Why every street needs a repair night", "Bring it broken, take it home working, and bring a neighbour"):
        prs, k = _deck(L)
        sl = k.new_slide()
        pic = str(vl.ASSETS / "photo/hall-repair.jpg") if L in vl.IMAGE_LED else None
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                k.cover(sl, title=t, kicker="Repair café", **({"image": pic} if pic else {}))
        except Exception as e:
            check(False, "{} cover ({!r}) builds ({}: {})".format(L, t[:12], type(e).__name__, str(e)[:70])); continue
        head = " ".join(t.split()[:2]).lower()
        sh = max((x for x in sl.shapes if getattr(x, "has_text_frame", False)
                  and head in x.text_frame.text.replace("\n", " ").lower()),
                 key=lambda x: max((r.font.size.pt for p in x.text_frame.paragraphs for r in p.runs if r.font.size), default=0))
        ls = _drawn_lines(k, "title", sh)
        # poster stacks a headline one word a line on purpose ("THREE / MOVES."), which is a stack, not a widow
        lone = len(ls) > 1 and len(ls[-1].split()) == 1 and not (L == "poster" and len(ls[-2].split()) == 1)
        check(not lone, "{} cover title {!r} does not end on a lone word ({})".format(L, t[:12], " / ".join(ls)))
dk.set_theme(None)
# …and a balanced measure is one the lint's line model also holds in the same lines: collage's 4.23in measure broke
# in four lines by the glyphs and five by the measure lint reads, so the height guard threw it away and the widow came back
prs, k = _deck("collage")
_t = "Why every street needs a repair night"
_n = len(vl._break_lines(k, "title", _t, 62.0, 4.785))
_wb = vl._balanced_width(k, "title", _t, 62.0, 4.785)
_face, _bold = k.face(vl.TYPE["collage"]["title"][1]), vl._weight(k, "title", _t)
check(_wb is None or dk._measure_lines([(_t, _bold)], 62.0, max(0.2, _wb - vl._INSET), font=_face) <= _n,
      "a balanced measure keeps the line count by the lint's model too ({} at {})".format(
          None if _wb is None else dk._measure_lines([(_t, _bold)], 62.0, max(0.2, _wb - vl._INSET), font=_face), _wb))
dk.set_theme(None)

# ── Task 5 (D): chalkboard: the quote's chalk underline runs the whole LAST line, to its measured ink width ──
def _quote_shape(slide, words):
    """The quote's own text shape (the largest-set one holding `words`)."""
    cand = [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False)
            and words.lower() in sh.text_frame.text.replace("\n", " ").lower()]
    return max(cand, key=lambda sh: max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs
                                         if r.font.size), default=0))


def _last_line_ink(k, field, sh):
    """(left edge, ink width) in inches of the last line drawn in text shape `sh`: its last paragraph, wrapped at the
    shape's width when it is a paragraph the renderer wraps."""
    paras = [p for p in sh.text_frame.paragraphs if "".join(r.text for r in p.runs).strip()]
    size = max(r.font.size.pt for p in paras for r in p.runs if r.font.size)
    box = sh.width / 914400.0
    text = "".join(r.text for r in paras[-1].runs).strip()
    last = vl._break_lines(k, field, text, size, box * vl.HEADROOM)[-1].strip()
    return sh.left / 914400.0 + 0.028, vl._line_width(k, field, last, size)


def _underline_span(slide, below):
    """(left, right) in inches of the chalk strokes (freeforms) drawn below y=`below`."""
    xs = [(sh.left / 914400.0, (sh.left + sh.width) / 914400.0) for sh in slide.shapes
          if sh.shape_type == 5 and sh.top / 914400.0 >= below]
    return (min(a for a, _ in xs), max(b for _, b in xs)) if xs else None


CB_QUOTES = [("short", "Less, but better."),
             ("mid", "Nature does not hurry, yet everything is accomplished."),
             ("long", " ".join(["Every repair we made that evening was a small argument for keeping things, and the "
                                "room kept making it until the last lamp came on."] * 2))]
for g in ("light", "slate"):
    for tag, qt in CB_QUOTES:
        prs, k = _deck("chalkboard", g)
        sl = k.new_slide()
        with contextlib.redirect_stdout(io.StringIO()):
            k.quote(sl, quote=qt, attribution="Lao Tzu")
        sh = _quote_shape(sl, qt.split()[0])
        left, ink = _last_line_ink(k, "quote", sh)
        span = _underline_span(sl, (sh.top + sh.height) / 914400.0 - 0.12)
        check(span is not None, "chalkboard/{} {} quote: a chalk underline is drawn".format(g, tag))
        if span:
            cover = (span[1] - span[0]) / ink
            check(0.97 <= cover <= 1.04 and abs(span[0] - left) <= 0.06,
                  "chalkboard/{} {} quote: the underline runs the last line's ink width ({:.2f}in of {:.2f}in = {:.0%}, "
                  "starting {:+.2f}in from the line)".format(g, tag, span[1] - span[0], ink, cover, span[0] - left))
dk.set_theme(None)

# ── Task 5 (D): interface: data / quote tail / points / window bar / dark bezel — measured on the built shapes ──
def _in(v):
    return v / 914400.0


def _texts(slide):
    return [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]


def _by_text(slide, t):
    return [sh for sh in _texts(slide) if sh.text_frame.text.strip() == t]


def _top_size(sh):
    return max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0.0)


def _fill_hex(sh):
    try:
        return str(sh.fill.fore_color.rgb)
    except Exception:
        return None


def _line_hex(sh):
    try:
        return str(sh.line.color.rgb)
    except Exception:
        return None


def _cross(a, b):
    """True when the two shapes' rects overlap."""
    return (a.left < b.left + b.width and b.left < a.left + a.width and a.top < b.top + b.height
            and b.top < a.top + a.height)


UI_PICK = lambda L_: _deck("interface", L_)      # noqa: E731

# data: the label under the number; the bar names the fill's end (the number) and its right end (the total)
for g, num, tot in (("light", "12", "40"), ("dark", "12", "40"), ("light", "39", "40"), ("light", "40", "40")):
    prs, k = UI_PICK(g)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.data(sl, number=num, label="of {} teams switched".format(tot), total=tot)
    big = max(_by_text(sl, num), key=_top_size)
    lab = _by_text(sl, "of {} teams switched".format(tot))[0]
    tag = "interface/{} data {} of {}".format(g, num, tot)
    if (num, tot) == ("12", "40"):
        check(_in(lab.top) >= _in(big.top) + _in(big.height) - 0.02,
              "{}: the label sits BELOW the number (label top {:.2f}in, number bottom {:.2f}in)".format(
                  tag, _in(lab.top), _in(big.top) + _in(big.height)))
    bars = sorted([sh for sh in sl.shapes if sh.shape_type == 1 and abs(_in(sh.height) - 0.16) < 0.01 and _in(sh.width) > 2],
                  key=lambda sh: -sh.width)
    check(len(bars) >= 1, "{}: a progress bar is drawn".format(tag))
    if not bars:
        continue
    trk, fill = bars[0], bars[-1]
    ends = [(num, _in(fill.left + fill.width)), (tot, _in(trk.left + trk.width))]
    under = [sh for sh in _texts(sl) if _top_size(sh) < 40 and sh.top >= trk.top + trk.height - 9144
             and sh.text_frame.text.strip() in (num, tot)]           # the labels printed under the bar
    check(len(under) == 2, "{}: both ends of the bar are labelled ({} labels under it)".format(tag, len(under)))
    if (num, tot) == ("12", "40") and len(under) == 2:             # a share well clear of both edges
        for words, edge in ends:
            lb = next(sh for sh in under if sh.text_frame.text.strip() == words)
            right = _in(lb.left + lb.width) - 0.028
            check(abs(right - edge) <= 0.08, "{}: '{}' labels its end of the bar ({:.2f}in vs the bar's {:.2f}in)".format(
                tag, words, right, edge))
    if len(under) == 2:
        check(not _cross(under[0], under[1]), "{}: the two bar labels do not overlap".format(tag))
dk.set_theme(None)

# quote: the bubble has a real tail (a triangle protruding from its edge), its outline >= 1.5:1 against card and ground
for g in ("light", "dark"):
    prs, k = UI_PICK(g)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.quote(sl, quote="We stopped asking where the request lived. It was just there.", attribution="Priya, operations lead")
    panel, ground = k.P["panel"], k.P["ground"]
    body = max([sh for sh in sl.shapes if sh.shape_type == 1 and _fill_hex(sh) == panel], key=lambda sh: sh.width * sh.height)
    tails = [sh for sh in sl.shapes if sh.shape_type == 5 and _line_hex(sh)]
    check(len(tails) >= 1, "interface/{} quote: the bubble draws its tail".format(g))
    if tails:
        t = tails[0]
        out = _in(body.left) - _in(t.left)
        c1, c2 = vl._contrast(_line_hex(t), panel), vl._contrast(_line_hex(t), ground)
        check(out >= 0.15 and _in(t.height) >= 0.2,
              "interface/{} quote: the tail protrudes {:.2f}in from the bubble's edge, {:.2f}in tall".format(g, out, _in(t.height)))
        check(c1 >= 1.5 and c2 >= 1.5,
              "interface/{} quote: the tail's edge reads {:.2f}:1 against the card and {:.2f}:1 against the ground".format(g, c1, c2))
        cb = vl._contrast(_line_hex(body), panel)
        check(cb >= 1.5 and _line_hex(body) == _line_hex(t),
              "interface/{} quote: the bubble's outline matches the tail's ({:.2f}:1 against the card)".format(g, cb))
dk.set_theme(None)

# points: with 4 items a head is never set smaller than its 20pt badge numeral
P4 = [("Approvals first", "One gate before anything ships."), ("One source of truth", "Every team reads the same list."),
      ("A weekly digest", "Monday morning, in your inbox."), ("Quiet hours", "Nothing pings after six.")]
for g in ("light", "dark"):
    for tag, items, kw, size in (("16:9", P4, {}, (13.333, 7.5)), ("16:9 heads only", [(h, None) for h, _l in P4], {}, (13.333, 7.5)),
                                 ("4:3", P4, {}, (10.0, 7.5)), ("with tags", P4, {"tags": ["New", "On", "Beta", "Off"]}, (13.333, 7.5))):
        prs = dk.blank_deck(*size)
        with contextlib.redirect_stdout(io.StringIO()):
            k = vl.use("interface", prs, ground=g)
            sl = k.new_slide()
            k.points(sl, title="Four settings", items=items, **kw)
        heads = [_top_size(sh) for sh in _texts(sl) if sh.text_frame.text.strip() in [h for h, _l in P4]]
        badge = [_top_size(sh) for sh in _texts(sl) if sh.text_frame.text.strip() in ("1", "2", "3", "4")]
        check(len(heads) == 4 and min(heads) >= 20 - 1e-6 and min(heads) >= max(badge or [0]) - 1e-6,
              "interface/{} points x4 ({}): heads hold >= 20pt and the badge numeral ({:.1f}pt vs {:.1f}pt)".format(
                  g, tag, min(heads or [0]), max(badge or [0])))
dk.set_theme(None)

# …and long copy that cannot hold 20pt in the first layout moves to a roomier one (two columns) instead of shrinking
LONG4 = [(h + " and then some more words to make it long", ln + " " + ln) for h, ln in
         (("Approvals first, always, before shipping", "One gate before anything ships to a customer, reviewed by two people."),
          ("One source of truth for every team", "Every team reads the same list and nobody keeps a private spreadsheet."),
          ("A weekly digest on Monday morning", "Monday morning, in your inbox, with what changed and who changed it."),
          ("Quiet hours after six in the evening", "Nothing pings after six unless somebody marks the request as urgent."))]
for n in (3, 4):
    prs, k = UI_PICK("light")
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.points(sl, title="Settings", items=LONG4[:n])
    heads = [_top_size(sh) for sh in _texts(sl) if any(sh.text_frame.text.strip().startswith(h[:18]) for h, _l in LONG4[:n])]
    check(len(heads) == n, "interface points x{} with long copy: every head is drawn".format(n))
    check_mac(len(heads) == n and min(heads) >= 20 - 1e-6,
          "interface points x{} with long copy: the heads hold >= 20pt in a roomier layout ({:.1f}pt)".format(n, min(heads or [0])))
dk.set_theme(None)

# window: no crumb= and no status= -> no header strip (the window stays); given, the strip is drawn
def _window_of(slide):
    shapes = [sh for sh in slide.shapes if sh.shape_type == 1]      # a window's drop edge is drawn first, same size
    return max(reversed(shapes), key=lambda sh: sh.width * sh.height, default=None)


def _strip(slide):
    """The window's header hairline: as wide as the window itself (a row hairline starts inside its padding)."""
    win = _window_of(slide)
    return [sh for sh in slide.shapes if sh.shape_type == 1 and 0.008 < _in(sh.height) < 0.016
            and abs(sh.left - win.left) < 9144 and abs(sh.width - win.width) < 9144]


for page, kw in (("cover", dict(kicker="Product tour", title="One board for every request")),
                 ("section", dict(number="02", kicker="Part two", title="Setting it up")),
                 ("points", dict(title="Three settings", items=["Approvals first", "One source of truth", "A weekly digest"]))):
    for extra, want in (({}, False), ({"crumb": "Launch › Setup"}, True), ({"status": ("Live", "on")}, True)):
        prs, k = UI_PICK("light")
        sl = k.new_slide()
        with contextlib.redirect_stdout(io.StringIO()):
            getattr(k, page)(sl, **kw, **extra)
        win = _window_of(sl)
        check(bool(_strip(sl)) == want and win is not None and _in(win.width) > 8,
              "interface {} {}: the window {} its header strip".format(page, sorted(extra) or "(no crumb, no status)",
                                                                         "keeps" if want else "drops"))
dk.set_theme(None)

# image_text: the phone bezel on the dark ground is dark with a hairline, not the near-white ink; light is unchanged
for g in ("light", "dark"):
    prs, k = UI_PICK(g)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.image_text(sl, kicker="Inside the app", title="Requests land in one inbox",
                     body="Every request from every channel arrives in the same list.", caption="The inbox view",
                     image=str(vl.ASSETS / "photo/hall-repair.jpg"))
    bez = max([sh for sh in sl.shapes if sh.shape_type == 1], key=lambda sh: sh.width * sh.height)
    fh = _fill_hex(bez)
    share = _in(bez.width) * _in(bez.height) / (13.333 * 7.5)
    if g == "dark":
        check(vl._lum(fh) < 0.06 and _line_hex(bez) is not None,
              "interface/dark image_text: the bezel is dark ({} luminance {:.2f}, {:.0%} of the page) with a hairline ({})".format(
                  fh, vl._lum(fh), share, _line_hex(bez)))
    else:
        check(fh == k.P["ink"] and _line_hex(bez) is None,
              "interface/light image_text: the bezel is the ink colour, unchanged ({})".format(fh))
dk.set_theme(None)

# ── Task 5 (D): wayfinding: quote takes line=; data board rows; ordered strip map labels and route ──
WF_PIC = lambda g: _deck("wayfinding", g)      # noqa: E731


def _signs(slide):
    """The enamel sign panels (rounded autoshapes filled with the language's sign colour)."""
    return [sh for sh in slide.shapes if sh.shape_type == 1 and _fill_hex(sh) is not None]


# quote: line= draws the route stub and the line's roundel in that line's colour; without line= nothing is added
for g in ("light", "night"):
    for lv, ci in ((("B", "blue"), 1), (("M10", "green"), 2), ("A", 0)):
        prs, k = WF_PIC(g)
        sl = k.new_slide()
        with contextlib.redirect_stdout(io.StringIO()):
            k.quote(sl, quote="The best signs are the ones you never notice reading.", attribution="A station designer", line=lv)
        code = lv if isinstance(lv, str) else lv[0]
        col = k.P["lines"][ci]
        sign = next(sh for sh in _signs(sl) if _fill_hex(sh) == k.P["sign"])
        stubs = [sh for sh in sl.shapes if sh.shape_type == 9 and _line_hex(sh) == col]
        rds = [sh for sh in _by_text(sl, code)]
        tag = "wayfinding/{} quote line={!r}".format(g, lv)
        check(bool(stubs) and min(sh.left for sh in stubs) <= 9144 and
              max(sh.left + sh.width for sh in stubs) >= sign.left - 91440,
              "{}: a route stub in the line's colour runs from the page edge to the sign".format(tag))
        check(len(rds) == 1 and any(_fill_hex(sh) == col and _cross(sh, rds[0]) for sh in sl.shapes if sh.shape_type == 1),
              "{}: the roundel '{}' is drawn on the line's colour".format(tag, code))
    prs, k = WF_PIC(g)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.quote(sl, quote="The best signs are the ones you never notice reading.", attribution="A station designer")
    check(not [sh for sh in sl.shapes if sh.shape_type == 9], "wayfinding/{} quote without line=: no route is drawn".format(g))
# …and the page with the line holds on other canvases and with a long code, lint-clean (the first roundel straddled the
# sign's corner and overlapped it)
for size in ((13.333, 7.5), (10.0, 7.5), (7.5, 10.0), (7.5, 7.5)):
    for lv in (("B", "blue"), ("M10", "green"), ("一号线", "red")):
        prs = dk.blank_deck(*size)
        with contextlib.redirect_stdout(io.StringIO()):
            k = vl.use("wayfinding", prs)
            sl = k.new_slide()
            k.quote(sl, quote="The best signs are the ones you never notice reading.", attribution="A station designer", line=lv)
        found = [f for f in _hard(prs, "wfq_{}x{}_{}".format(size[0], size[1], lv[0])) if f.startswith("slide 1:")]
        stub = [sh for sh in sl.shapes if sh.shape_type == 9 and _line_hex(sh) == k.P["lines"][{"blue": 1, "green": 2, "red": 0}[lv[1]]]]
        check(not found and stub, "wayfinding quote line={} on {}x{}: lint-clean with its route ({})".format(lv[0], size[0], size[1],
                                                                                                            found[:1]))
dk.set_theme(None)

# data: the board's rows under the figure are bold ~26pt Courier, not 20pt regular beside a 96pt figure
for g in ("light", "night"):
    prs, k = WF_PIC(g)
    sl = k.new_slide()
    rows = [("North loop", "93.1%"), ("South loop", "88.0%"), ("Harbour", "91.7%"), ("Airport", "79.5%")]
    with contextlib.redirect_stdout(io.StringIO()):
        k.data(sl, number="96.4%", label="Harbour line on time", board=rows, note="Source: the operator's weekly report")
    for words in [w for r_ in rows for w in r_]:
        sh = _by_text(sl, words)[0]
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
        check(runs[0].font.bold is True and _top_size(sh) >= 25.5 - 1e-6,
              "wayfinding/{} data board row '{}' is bold at >= 26pt ({} {:.1f}pt)".format(g, words, runs[0].font.bold, _top_size(sh)))
    found = [f for f in _hard(prs, "wfdata_" + g) if f.startswith("slide 1:")]
    check(not found, "wayfinding/{} data board with 4 rows and a note is lint-clean ({})".format(g, found[:1]))
dk.set_theme(None)

# ordered points (the strip map): every label is left-aligned; the last wraps inside its own box on the page; the route
# sits under the sign (<= 1in from the sign to the first label), not at mid-page
def _strip_page(g, items):
    prs, k = WF_PIC(g)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.points(sl, title="Stops", ordered=True, items=items)
    sign = min((sh for sh in _signs(sl) if _fill_hex(sh) == k.P["sign"]), key=lambda sh: sh.top)
    return prs, sl, sign


STRIPS = [("heads 4", ["Sign up", "First project", "Invite a teammate", "Upgrade"]),
          ("with lines 4", [("Sign up", "Takes a minute."), ("First project", "Pick a template."),
                            ("Invite a teammate", "By email."), ("Upgrade", "When you need more seats.")]),
          ("long last 4", ["Sign up", "First project", "Invite a teammate", "Upgrade to the plan with unlimited seats"]),
          ("3", ["Sign up", "First project", "Upgrade"]), ("2", ["Sign up", "Upgrade"])]
for g in ("light", "night"):
    for tag, items in STRIPS:
        prs, sl, sign = _strip_page(g, items)
        heads = [h if isinstance(h, str) else h[0] for h in items]
        labels = [sh for sh in _texts(sl) if sh.text_frame.text.split("\n")[0].strip() in heads]
        aligns = {p.alignment for sh in labels for p in sh.text_frame.paragraphs}
        check(len(labels) == len(items) and dk.PP_ALIGN.RIGHT not in aligns,
              "wayfinding/{} strip map ({}): every label is left-aligned ({} labels)".format(g, tag, len(labels)))
        check(all(sh.left >= 0 and sh.left + sh.width <= 13.333 * 914400 for sh in labels),
              "wayfinding/{} strip map ({}): every label's box is on the page".format(g, tag))
        gap = _in(min(sh.top for sh in labels)) - _in(sign.top + sign.height)
        check(gap <= 1.0 + 1e-6, "wayfinding/{} strip map ({}): the first label is {:.2f}in under the sign".format(g, tag, gap))
        found = [f for f in _hard(prs, "wfstrip_{}_{}".format(g, tag.replace(" ", "_"))) if f.startswith("slide 1:")]
        check(not found, "wayfinding/{} strip map ({}) is lint-clean ({})".format(g, tag, found[:1]))


# ── Task 5 (B): measuring helpers — the pptx's own geometry, and the rendered picture where one can be made ──
import shutil, subprocess
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

CAN_RENDER = bool(shutil.which("soffice") and shutil.which("pdftoppm"))
try:
    import numpy as np
    from PIL import Image
except ImportError:                                  # no picture to measure: the geometry checks still run
    CAN_RENDER = False


def _inch(v):
    return Emu(v).inches


def _rules(slide):
    """The slide's straight rules: [(x0, y0, x1, y1)] in inches (LINE shapes)."""
    return [tuple(_inch(v) for v in (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height))
            for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.LINE]


def _words(slide):
    """(text, x0, y0, x1, y1, max_size_pt) of every text shape on the page proper."""
    out = []
    for sh in slide.shapes:
        if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip() and sh.top >= 0:
            z = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size and r.text.strip()),
                    default=0.0)
            out.append((sh.text_frame.text.replace("\n", " ").replace("\v", " "), _inch(sh.left), _inch(sh.top),
                        _inch(sh.left + sh.width), _inch(sh.top + sh.height), z))
    return out


_PNG = {}


def _render(prs, tag, dpi=100):
    """The deck's pages as PIL images (cached), or None when no renderer is installed. A renderer that is installed but
    returns nothing (another soffice holding the profile) is retried on its own profile, then FAILS the run — a measured
    check must not skip itself without a word."""
    if not CAN_RENDER:
        return None
    if tag not in _PNG:
        pp = WORK / (tag + ".pptx"); prs.save(str(pp))
        pages = []
        for attempt in range(3):
            subprocess.run(["soffice", "-env:UserInstallation=file://" + str(WORK / ("lo_profile%d" % attempt)), "--headless",
                            "--convert-to", "pdf", "--outdir", str(WORK), str(pp)], capture_output=True, timeout=240)
            subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(WORK / (tag + ".pdf")), str(WORK / (tag + "_pg"))],
                           capture_output=True, timeout=240)
            pages = [Image.open(f).convert("RGB") for f in sorted(WORK.glob(tag + "_pg*.png"))]
            if pages:
                break
        check(len(pages) == len(prs.slides), "the renderer drew all {} pages of {} (got {})".format(len(prs.slides), tag, len(pages)))
        _PNG[tag] = pages
    return _PNG[tag] or None


def _ink(img, box, dpi=100, thr=70):
    """The bounding box (inches) of what is drawn inside `box` (x0, y0, x1, y1) — pixels more than `thr` (summed
    channel difference) away from the box's median colour — or None when the box is empty."""
    a = np.asarray(img).astype(int)
    x0, y0, x1, y1 = [max(0, int(round(v * dpi))) for v in box]
    r = a[y0:y1, x0:x1]
    if r.size == 0:
        return None
    hit = np.abs(r - np.median(r.reshape(-1, 3), axis=0)).sum(axis=2) > thr
    ys, xs = np.where(hit)
    if not len(ys):
        return None
    return ((xs.min() + x0) / dpi, (ys.min() + y0) / dpi, (xs.max() + 1 + x0) / dpi, (ys.max() + 1 + y0) / dpi)


# ── Task 5 (B): ink: a points page's rules end at their column's text (+0.2in), not at a fixed 4.8in ──
INK_CASES = {
    "cjk-short": dict(title="两道工序", items=[("洗盏", "器净，心先静"), ("候汤", "水沸，如蟹眼")], seal="序"),
    "cjk-long": dict(title="四道工序", items=[("洗盏", "器净，心先静，盏中不留一滴旧水"), ("候汤", "水沸如蟹眼，听松风而不急不躁"),
                                              ("分茶", "浅斟留七分，余下三分是人情"), ("奉茶", "双手捧盏，视线低垂，敬客人先")], seal="序"),
    "latin-short": dict(title="Two steps", items=[("Rinse", "Clean the cup"), ("Heat", "Water just off the boil")], seal="T"),
    "latin-long": dict(title="Four steps of tea", items=[
        ("Rinse the cup", "Clean the cup until nothing of the old water is left behind in it"),
        ("Heat the water", "Let it come just off the boil and listen for the sound of pine wind"),
        ("Pour thin", "Pour seven parts of ten and leave the rest for company"),
        ("Serve", "Hold the cup in both hands with your eyes lowered and offer it first to the guest")], seal="T"),
}
INK_GAP = (0.10, 0.32)         # a rule ends 0.2in (+-) past the last of its words
for L_, g_ in (("ink", "light"), ("ink", "night")):
    prs, k = _deck(L_, g_)
    for tag, kw in INK_CASES.items():
        s_ = k.new_slide()
        with contextlib.redirect_stdout(io.StringIO()):
            k.points(s_, **kw)
    pages = list(prs.slides)
    imgs = _render(prs, "inkrules_" + g_)
    for (tag, kw), s_, pg in zip(INK_CASES.items(), pages, range(len(pages))):
        rules = _rules(s_)
        words = [w_ for w_ in _words(s_) if w_[0] != kw["seal"] and w_[0] != "a11y"]
        vert = [r_ for r_ in rules if abs(r_[0] - r_[2]) < 1e-3]
        horiz = [r_ for r_ in rules if abs(r_[1] - r_[3]) < 1e-3]
        check(vert or horiz, "ink/{} {}: the points page draws its rules".format(g_, tag))
        if vert:                                              # the vertical scroll: a rule is the left edge of its column
            vert.sort()
            for j, r_ in enumerate(vert):
                right_edge = vert[j + 1][0] if j + 1 < len(vert) else 13.333
                col = [w_ for w_ in words if w_[1] >= r_[0] - 0.01 and w_[3] <= right_edge + 0.01 and w_[2] > 0.3]
                low = max(w_[4] for w_ in col)                # the planned text's bottom
                # geometry: the rule is cut to the words — it used to run to a fixed 0.76H whatever the copy was
                check(r_[3] - low <= 0.32, "ink/{} {}: rule {} ends {:.2f}in below its column's text box (<= 0.32in)".format(
                    g_, tag, j, r_[3] - low))
                if imgs:                                      # the picture: the last ink of the column vs the rule's foot
                    ib = _ink(imgs[pg], (r_[0] + 0.06, 0.35, right_edge - 0.06, low + 0.02))
                    gap = r_[3] - ib[3] if ib else None
                    check_mac(gap is not None and INK_GAP[0] <= gap <= INK_GAP[1],
                          "ink/{} {}: rule {} ends {} past the column's last ink (want {:.2f}-{:.2f}in)".format(
                              g_, tag, j, "{:.2f}in".format(gap) if gap is not None else "?", *INK_GAP))
        elif horiz:                                           # the horizontal rows: a rule hangs under ITS row's words
            horiz.sort(key=lambda r_: r_[1])
            for j, r_ in enumerate(horiz):
                above = [w_ for w_ in words if w_[4] <= r_[1] + 0.01 and w_[2] > (horiz[j - 1][1] if j else 0.3)
                         and w_[1] >= r_[0] - 0.01]
                low = max(w_[4] for w_ in above)
                check(0.0 <= r_[1] - low <= 0.32, "ink/{} {}: row rule {} hangs {:.2f}in under its row's text (<= 0.32in)".format(
                    g_, tag, j, r_[1] - low))
dk.set_theme(None)


def _ink_runs(img, box, dpi=100, gap=0.10, thr=70):
    """The horizontal bands of ink in `box`: [(y0, y1)] in inches, bands closer than `gap` merged."""
    a = np.asarray(img).astype(int)
    x0, y0, x1, y1 = [max(0, int(round(v * dpi))) for v in box]
    r = a[y0:y1, x0:x1]
    hit = (np.abs(r - np.median(r.reshape(-1, 3), axis=0)).sum(axis=2) > thr).any(axis=1)
    out, start, last = [], None, None
    for i, h in enumerate(hit):
        if h:
            if start is None:
                start = i
            elif i - last > gap * dpi:
                out.append((start, last)); start = i
            last = i
    if start is not None:
        out.append((start, last))
    return [((a_ + y0) / dpi, (b_ + 1 + y0) / dpi) for a_, b_ in out]


# ── Task 5 (B): poster: the quote's attribution sits ~0.4in under the quote's last line (it sat 1.12in under) ──
PQ = {"one line": "Make it.", "three lines": "A room is a promise you can walk into.",
      "five lines": "A room is a promise you can walk into, and a street is a long hallway of rooms that nobody locked, so "
                    "we keep the doors open and the benches warm for whoever is next in line to mend something."}
for g_ in ("light", "paper"):
    prs, k = _deck("poster", g_)
    for tag, q_ in PQ.items():
        with contextlib.redirect_stdout(io.StringIO()):
            k.quote(k.new_slide(), quote=q_, attribution="From the manifesto")
    imgs = _render(prs, "pq_" + g_)
    for (tag, q_), s_, pg in zip(PQ.items(), prs.slides, range(3)):
        wd = _words(s_)
        qb = next(w_ for w_ in wd if w_[0].upper().startswith(q_[:8].upper()))
        ab = next(w_ for w_ in wd if w_[0].startswith("From the"))
        if imgs:
            runs = _ink_runs(imgs[pg], (qb[1], qb[2], qb[3], ab[4] + 0.05))
            gap = runs[-1][0] - runs[-2][1] if len(runs) >= 2 else None
            check(gap is not None and 0.30 <= gap <= 0.55,
                  "poster/{} quote ({}): the attribution's ink starts {} under the quote's last line (want 0.30-0.55in)".format(
                      g_, tag, "{:.2f}in".format(gap) if gap is not None else "?"))
        else:                                        # no renderer: the planned boxes — the attribution right under the quote's
            check(-0.1 <= ab[2] - qb[4] <= 0.3, "poster/{} quote ({}): the attribution's box starts {:.2f}in from the quote's".format(
                g_, tag, ab[2] - qb[4]))
        check(not [f for f in _hard(prs, "pqh_" + g_ + tag.replace(" ", "")) if f.startswith("slide {}:".format(pg + 1))],
              "poster/{} quote ({}): the page stays lint-clean".format(g_, tag))
dk.set_theme(None)


# ── Task 5 (B): poster: the points staircase fills the page (the bottom-left 28% was empty) and still climbs with the order ──
PP = {"3 points": ("Three moves.", ["Share the tools", "Open the door", "Keep it local"]),
      "2 points": ("Two moves.", ["Share the tools", "Open the door"]),
      "4 points": ("Four moves for a street", ["Share the tools", "Open the door", "Keep it local", "Mend it together"]),
      "one-line title": ("Moves", ["Share the tools", "Open the door", "Keep it local"])}
for g_ in ("light", "paper"):
    prs, k = _deck("poster", g_)
    for tag, (t_, its) in PP.items():
        with contextlib.redirect_stdout(io.StringIO()):
            k.points(k.new_slide(), title=t_, items=its)
    for pg, ((tag, (t_, its)), s_) in enumerate(zip(PP.items(), prs.slides)):
        bars = sorted([(_inch(sh.left), _inch(sh.top), _inch(sh.width), _inch(sh.height)) for sh in s_.shapes
                       if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and _inch(sh.width) > 1.0], key=lambda b_: b_[0])
        check(len(bars) == len(its), "poster/{} points ({}): one bar a point ({} bars)".format(g_, tag, len(bars)))
        if len(bars) != len(its):
            continue
        # the staircase still encodes order: each step is taller than the one before and stands to its right
        check(all(b2[3] > b1[3] + 0.2 and b2[0] > b1[0] + b1[2] - 1e-6 and abs((b2[1] + b2[3]) - (b1[1] + b1[3])) < 0.01
                  for b1, b2 in zip(bars, bars[1:])),
              "poster/{} points ({}): the steps rise left to right on one baseline ({})".format(
                  g_, tag, ["{:.2f}".format(b_[3]) for b_ in bars]))
        # the bottom-left quarter of the page, under the title: how much of it carries ink (bars)
        tw = [w_ for w_ in _words(s_) if w_[0].upper().startswith(t_.split()[0].upper())][0]
        reg = (0.0, tw[4], 6.667, 7.5)              # below the title's box, left half
        area = (reg[2] - reg[0]) * (reg[3] - reg[1])
        cover = sum(max(0.0, min(b_[0] + b_[2], reg[2]) - max(b_[0], reg[0])) * max(0.0, min(b_[1] + b_[3], reg[3]) - max(b_[1], reg[1]))
                    for b_ in bars)
        check(cover / area >= 0.35, "poster/{} points ({}): the bottom-left under the title is {:.0%} filled (want >= 35%)".format(
            g_, tag, cover / area))
        check(not [f for f in _hard(prs, "ppt_" + g_ + str(pg)) if f.startswith("slide {}:".format(pg + 1))],
              "poster/{} points ({}): the page stays lint-clean".format(g_, tag))
dk.set_theme(None)


from pptx.enum.shapes import MSO_SHAPE


def _fill_of(sh):
    try:
        return str(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
    except Exception:
        return None


def _cards(slide, min_w=1.5):
    """The paper cards of a cutpaper page: rounded rectangles >= min_w wide -> [(x, y, w, h, fill, rotation)]."""
    return [(_inch(sh.left), _inch(sh.top), _inch(sh.width), _inch(sh.height), _fill_of(sh), sh.rotation)
            for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
            and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE and _inch(sh.width) >= min_w]


# ── Task 5 (B): cutpaper: the cover's card is as tall as its words + padding + the hill tuck (it was a fixed 0.47H) ──
CCV = {"one word": dict(kicker="A paper-cut science story", title="Seeds"),
       "sample": dict(kicker="A paper-cut science story", title="How seeds travel"),
       "long": dict(kicker="A paper-cut science story", title="Why every forest on this hillside began as one small seed",
                    subtitle="Wings, water and fur carry them further than you would think."),
       "cjk": dict(kicker="剪纸科学故事", title="种子如何旅行")}
for g_ in ("light", "night"):
    prs, k = _deck("cutpaper", g_)
    for tag, kw in CCV.items():
        with contextlib.redirect_stdout(io.StringIO()):
            k.cover(k.new_slide(), **kw)
    imgs = _render(prs, "ccv_" + g_)
    for pg, ((tag, kw), s_) in enumerate(zip(CCV.items(), prs.slides)):
        cards = sorted(_cards(s_), key=lambda c_: -c_[2] * c_[3])
        check(bool(cards), "cutpaper/{} cover ({}): the title card is drawn".format(g_, tag))
        if not cards:
            continue
        cx_, cy_, cw_, ch_ = cards[0][:4]
        inside = [w_ for w_ in _words(s_) if w_[1] >= cx_ - 0.01 and w_[3] <= cx_ + cw_ + 0.01 and w_[2] >= cy_ - 0.01
                  and w_[4] <= cy_ + ch_ + 0.01]
        low = max(w_[4] for w_ in inside)
        # the card runs 0.4in of padding + the 0.3in the near hill tucks it under past its words, and no further
        check(0.55 <= cy_ + ch_ - low <= 0.80, "cutpaper/{} cover ({}): the card ends {:.2f}in under its words "
              "(want 0.55-0.80: padding + tuck)".format(g_, tag, cy_ + ch_ - low))
        if imgs:                                     # the picture: the card as it SHOWS, above the near hill's crest
            im = imgs[pg]
            px_ = lambda x, y: im.getpixel((int(round(x * 100)), int(round(y * 100))))        # noqa: E731
            col = px_(cx_ + cw_ - 0.25, cy_ + 0.2)
            shown = 99.0
            for i_ in range(12):                      # per column, from below the card up to its own colour: the crest
                x = cx_ + 0.3 + (cw_ - 0.6) * i_ / 11.0
                y = min(cy_ + ch_, im.size[1] / 100.0 - 0.05)
                while y > cy_ + 0.3 and sum(abs(a_ - b_) for a_, b_ in zip(px_(x, y), col)) >= 40:
                    y -= 0.01
                shown = min(shown, y)
            ib = _ink(im, (cx_ + 0.3, cy_ + 0.05, cx_ + cw_ - 0.3, shown - 0.02))
            blank = shown - ib[3] if ib else None
            check(blank is not None and 0.20 <= blank <= 0.75, "cutpaper/{} cover ({}): {} of bare card shows under the last "
                  "ink (want 0.20-0.75in)".format(g_, tag, "{:.2f}in".format(blank) if blank is not None else "?"))
        check(not [f for f in _hard(prs, "ccvh_" + g_ + str(pg)) if f.startswith("slide {}:".format(pg + 1))],
              "cutpaper/{} cover ({}): the page stays lint-clean".format(g_, tag))
dk.set_theme(None)


# ── Task 5 (B): cutpaper: the points cards are as tall as their content + padding, all one height (22-30% was empty) ──
CPT = {"3 short": ("Three ways a seed gets around", [("Wind", "Wings and parachutes drift far."),
                                                      ("Water", "Some seeds float to a new shore."),
                                                      ("Animals", "Hooks hitch a ride on fur.")], None),
       "2 short": ("Two ways", [("Wind", "Drift far."), ("Water", "Float to a shore.")], None),
       "icons": ("Three ways a seed gets around", [("Wind", "Wings and parachutes drift far."),
                                                   ("Water", "Some seeds float to a new shore."),
                                                   ("Animals", "Hooks hitch a ride on fur.")],
                 ["lucide:wind", "lucide:droplets", "lucide:paw-print"]),
       "4 long": ("Four ways a seed gets around the hillside", [
           ("Wind carries the light ones", "Wings and parachutes let a seed drift for a whole afternoon before it lands."),
           ("Water floats the heavy ones", "Some seeds are buoyant enough to ride a stream down to a brand new shore."),
           ("Animals hitch a ride", "Hooks and burrs catch on fur and stay there until they are scratched loose."),
           ("People carry them too", "Every pocket and boot sole on the path is a small seed vehicle.")], None),
       "cjk": ("种子旅行的三种方式", [("风", "种子随风飘得很远。"), ("水", "种子漂到新的岸边。"),
                                     ("动物", "种子挂在毛上搭车。")], None)}
for g_ in ("light", "night"):
    prs, k = _deck("cutpaper", g_)
    for tag, (t_, its, ic_) in CPT.items():
        with contextlib.redirect_stdout(io.StringIO()):
            k.points(k.new_slide(), title=t_, items=its, **({"icons": ic_} if ic_ else {}))
    for pg, ((tag, (t_, its, ic_)), s_) in enumerate(zip(CPT.items(), prs.slides)):
        cards = [c_ for c_ in _cards(s_) if c_[3] >= 1.0]
        check(len(cards) == len(its), "cutpaper/{} points ({}): one card a point ({})".format(g_, tag, len(cards)))
        if not cards:
            continue
        hs = [c_[3] for c_ in cards]
        check(max(hs) - min(hs) < 0.01, "cutpaper/{} points ({}): the cards share one height ({})".format(
            g_, tag, ["{:.2f}".format(h_) for h_ in hs]))
        wd = _words(s_)
        discs = [(_inch(sh.left), _inch(sh.top), _inch(sh.left + sh.width), _inch(sh.top + sh.height)) for sh in s_.shapes
                 if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.OVAL and _inch(sh.width) < 1.5]
        blanks = []
        for ci, (cx_, cy_, cw_, ch_, _f, _r) in enumerate(cards):
            inside = [w_[4] for w_ in wd if w_[1] >= cx_ - 0.01 and w_[3] <= cx_ + cw_ + 0.01 and w_[2] >= cy_ - 0.01
                      and w_[4] <= cy_ + ch_ + 0.05]
            inside += [d_[3] for d_ in discs if d_[0] >= cx_ - 0.01 and d_[2] <= cx_ + cw_ + 0.01 and d_[1] >= cy_ - 0.01]
            blanks.append((cy_ + ch_ - max(inside)) / ch_)
        # one shared height means the card with the fewest words keeps a line's worth more paper than the fullest one
        check(0.04 <= min(blanks) and min(blanks) <= 0.15 and max(blanks) <= 0.22,
              "cutpaper/{} points ({}): the cards are {} empty under their content (fullest 4-15%, emptiest <= 22%)".format(
                  g_, tag, "/".join("{:.0%}".format(b_) for b_ in blanks)))
        check(not [f for f in _hard(prs, "cptl_" + g_ + str(pg)) if f.startswith("slide {}:".format(pg + 1))],
              "cutpaper/{} points ({}): the page stays lint-clean (no UNEVEN CARD HEIGHTS)".format(g_, tag))
dk.set_theme(None)


# ── Task 5 (B): cutpaper night: the cards are the cream paper of the back sheet, not pure white on navy ──
CREAM = "F1E7C8"
for g_, want in (("light", "FFFFFF"), ("night", CREAM)):
    prs, k = _deck("cutpaper", g_)
    with contextlib.redirect_stdout(io.StringIO()):
        k.cover(k.new_slide(), kicker="A paper-cut science story", title="How seeds travel")
        k.points(k.new_slide(), title="Three ways a seed gets around", items=[("Wind", "Drift far."), ("Water", "Float."),
                                                                          ("Animals", "Hitch a ride.")])
        k.quote(k.new_slide(), quote="Every forest began as one small seed.", attribution="A paper-cut science story")
        k.closing(k.new_slide(), title="Plant one", line="and watch it travel")
    for pg, (page, s_) in enumerate(zip(("cover", "points", "quote", "closing"), prs.slides)):
        cards = [c_ for c_ in _cards(s_) if abs(c_[5]) < 0.01]                  # the page's own cards (the back sheets tilt)
        check(bool(cards), "cutpaper/{} {}: the cards are drawn".format(g_, page))
        check(all(c_[4] == want for c_ in cards), "cutpaper/{} {}: every card is {} ({})".format(
            g_, page, want, sorted({c_[4] for c_ in cards})))
        wd = [sh for sh in s_.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip() and sh.top >= 0]
        for c_ in cards:
            for sh in wd:
                x_, y_ = _inch(sh.left), _inch(sh.top)
                if not (c_[0] - 0.01 <= x_ and x_ + _inch(sh.width) <= c_[0] + c_[2] + 0.01 and c_[1] - 0.01 <= y_ <= c_[1] + c_[3]):
                    continue
                for p_ in sh.text_frame.paragraphs:
                    for r_ in p_.runs:
                        if r_.text.strip() and r_.font.color and r_.font.color.type and str(r_.font.color.rgb) != "FFFFFF":
                            cr_ = dk.contrast_ratio(str(r_.font.color.rgb), want)
                            check(cr_ >= 4.5, "cutpaper/{} {}: {!r} ({}) reads on its card at {:.2f}:1 (want >= 4.5)".format(
                                g_, page, r_.text[:14], r_.font.color.rgb, cr_))
dk.set_theme(None)
pal_n = vl.REGISTRY["cutpaper"]["grounds"]["night"]["palette"]
check(pal_n["paper"].upper() == CREAM and dk.contrast_ratio(pal_n["paper"], pal_n["ground"]) < 14.0,
      "cutpaper night: the cards' paper is the cream of the back sheet ({})".format(pal_n.get("paper")))
check(all(dk.contrast_ratio(pal_n[key], pal_n["paper"]) >= 4.5 for key in ("card_ink", "card_mute", "card_accent")),
      "cutpaper night: card_ink, card_mute and card_accent read on the paper at >= 4.5:1 ({})".format(
          ["{:.2f}".format(dk.contrast_ratio(pal_n[key], pal_n["paper"])) for key in ("card_ink", "card_mute", "card_accent")]))
check(vl.REGISTRY["cutpaper"]["grounds"]["light"]["palette"]["paper"].upper() == "FFFFFF", "cutpaper day: the paper stays white")
# `panel` is the ordinary page's card (it carries the light default inks): the night keeps it navy
check(dk.contrast_ratio(pal_n["ink"], pal_n["panel"]) >= 4.5, "cutpaper night: the ordinary page's panel still carries the default ink")

# ── Task 5 (B): drafting points: the leader's head outranks the line under it (it was 11pt Courier under a 15pt line) ──
DPT = {"3 short": ("Three layers, one frame", [("Floor", "One continuous plate."), ("Walls", "Panels that slide on a track."),
                                                ("Roof", "A single span, no columns.")]),
       "2 short": ("Two layers", [("Floor", "One continuous plate."), ("Walls", "Panels that slide.")]),
       "4 long": ("Four layers, one frame", [
           ("Floor", "One continuous plate poured in a single day and left to cure for a week."),
           ("Walls", "Panels that slide on a track and lock at any of the nine positions."),
           ("Roof", "A single span with no columns, lifted whole by four cranes."),
           ("Skin", "A rain screen that can be replaced one panel at a time.")])}
for g_ in ("light", "cyanotype"):
    prs, k = _deck("drafting", g_)
    for tag, (t_, its) in DPT.items():
        with contextlib.redirect_stdout(io.StringIO()):
            k.points(k.new_slide(), title=t_, items=its)
    for pg, ((tag, (t_, its)), s_) in enumerate(zip(DPT.items(), prs.slides)):
        shapes = [sh for sh in s_.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip() and sh.top >= 0]
        for head, line in its:
            hs = [sh for sh in shapes if sh.text_frame.text.strip() == head]
            ls = [sh for sh in shapes if sh.text_frame.text.strip().startswith(line[:12])]
            if not hs or not ls:
                check(False, "drafting/{} points ({}): {} and its line are drawn".format(g_, tag, head)); continue
            hr, lr = hs[0].text_frame.paragraphs[0].runs[0], ls[0].text_frame.paragraphs[0].runs[0]
            # ordinary copy: 16pt+ and clear of the line; copy so long that every field is squeezed toward its floor:
            # still bold serif and larger than the line under it
            roomy = tag != "4 long"
            check(hr.font.size.pt >= (16 if roomy else 11) and hr.font.size.pt >= lr.font.size.pt + (2 if roomy else 0.5)
                  and bool(hr.font.bold) and hr.font.name != "Courier New",
                  "drafting/{} points ({}): {!r} is {:.0f}pt {} {} over a {:.0f}pt line (want {}, bold serif, larger than the line)".format(
                      g_, tag, head, hr.font.size.pt, "bold" if hr.font.bold else "regular", hr.font.name, lr.font.size.pt,
                      ">= 16pt" if roomy else "squeezed copy still ranks head over line"))
        check(not [f for f in _hard(prs, "dptl_" + g_ + str(pg)) if f.startswith("slide {}:".format(pg + 1))],
              "drafting/{} points ({}): the page stays lint-clean".format(g_, tag))
dk.set_theme(None)


# ── Task 5 (B): drafting cover and quote: a ghost drawing fills the blank right of the words, clear of every word ──
def _ghosts(slide):
    """The faint drawing: the outline-only freeform shapes declared decorative -> [(x0, y0, x1, y1)]."""
    return [(_inch(sh.left), _inch(sh.top), _inch(sh.left + sh.width), _inch(sh.top + sh.height)) for sh in slide.shapes
            if sh.shape_type == MSO_SHAPE_TYPE.FREEFORM and "+decor" in sh.name.split(":")[0] and "drafting" in sh.name]


for g_ in ("light", "cyanotype"):
    prs, k = _deck("drafting", g_)
    with contextlib.redirect_stdout(io.StringIO()):
        k.cover(k.new_slide(), kicker="Schematic 01 · a modular reading room", title="A room built layer by layer.",
                subtitle="Floor, walls and roof drawn as one frame, each layer free to move.")
        k.quote(k.new_slide(), quote="“Draw the quiet first, then the walls.”", attribution="Design principle")
        k.quote(k.new_slide(), quote="Draw the quiet first, then the walls, then the roof, then the light that will fall "
                                     "across the floor at four in the afternoon.", attribution="Design principle")
    imgs = _render(prs, "dgh_" + g_)
    for pg, (page, s_) in enumerate(zip(("cover", "quote", "long quote"), prs.slides)):
        gh = _ghosts(s_)
        check(len(gh) >= 1, "drafting/{} {}: a ghost drawing is declared decorative ({} shapes)".format(g_, page, len(gh)))
        if not gh:
            continue
        gx0, gy0, gx1, gy1 = min(g[0] for g in gh), min(g[1] for g in gh), max(g[2] for g in gh), max(g[3] for g in gh)
        wd = _words(s_)
        over = [w_[0][:18] for w_ in wd if w_[1] < gx1 + 0.02 and gx0 - 0.02 < w_[3] and w_[2] < gy1 + 0.02 and gy0 - 0.02 < w_[4]]
        check(not over, "drafting/{} {}: the ghost drawing ({:.1f}-{:.1f}in across) touches no word box ({})".format(
            g_, page, gx0, gx1, over))
        check(gx1 - gx0 >= 2.2 and gy1 - gy0 >= 2.2 and gx0 > 6.0,
              "drafting/{} {}: the ghost fills the right of the sheet ({:.1f}x{:.1f}in from x={:.1f})".format(
                  g_, page, gx1 - gx0, gy1 - gy0, gx0))
        if imgs:                                   # it SHOWS: a faint line, but there, and nothing else is in that corner
            ib = _ink(imgs[pg], (gx0 - 0.1, gy0 - 0.1, gx1 + 0.1, gy1 + 0.1), thr=45)
            check(ib is not None and (ib[2] - ib[0]) >= 0.6 * (gx1 - gx0), "drafting/{} {}: the ghost drawing is visible".format(g_, page))
        check(not [f for f in _hard(prs, "dghl_" + g_ + str(pg)) if f.startswith("slide {}:".format(pg + 1))],
              "drafting/{} {}: the page stays lint-clean".format(g_, page))

# ── Task 5 (C): shared helpers for the four language fixes (starlit, broadsheet, journal, tally) ──
from pptx.util import Emu as _Emu
W13 = 13.333


def _text_shape(slide, words):
    """The drawn text box holding `words` (not the accessibility title placeholder, which repeats the page's words)."""
    for sh in slide.shapes:
        if (getattr(sh, "has_text_frame", False) and not sh.is_placeholder
                and words.lower() in sh.text_frame.text.replace("\n", " ").lower()):
            return sh
    return None


def _rect(sh):
    return (sh.left / 914400.0, sh.top / 914400.0, sh.width / 914400.0, sh.height / 914400.0)


def _page(L, g, page, **fields):
    prs, k = _deck(L, g)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        getattr(k, page)(sl, **fields)
    return prs, k, sl


# ── Task 5 (C): starlit: the quote box is wide enough to hold the quote's size ──
_SQ = " ".join(("Every repair we made that evening was a small argument for keeping things, and the room kept making it "
                "until the last lamp came on. We stayed late, and nobody wanted to be the one who turned the lights off "
                "at last.").split()[:38])
_prs, _k, _sl = _page("starlit", "light", "quote", quote=_SQ, attribution="Repair café volunteer")
_qw = _rect(_text_shape(_sl, "“"))[2]          # the quote mark's box is the column itself (the quote's own box narrows to balance its lines)
check(_qw >= 0.74 * W13, "starlit quote: the quote's column is >= 0.74 of the page wide ({:.2f}in = {:.0%}; was 8.00in = 60%)".format(_qw, _qw / W13))
_qz = _size_of(_sl, "Every repair")
check(_qz >= vl.TYPE["starlit"]["quote"][0] - 1e-6, "starlit quote: a 38-word quote holds the language's {}pt in the wider box ({:.1f}pt; "
      "was 33.0pt)".format(vl.TYPE["starlit"]["quote"][0], _qz))
_prs, _k, _sl = _page("starlit", "light", "quote", quote="We measured the year in conversations, not in launches.", attribution="x")
check(_size_of(_sl, "We measured") >= 1.3 * vl.TYPE["starlit"]["quote"][0], "starlit quote: the 55-character sample is still set larger than the base size "
      "({:.1f}pt)".format(_size_of(_sl, "We measured")))

# ── Task 5 (C): starlit: the numeral reads on its glow (>= 4.5:1 over the glow's centre, both grounds) ──
import native_art as _na                       # noqa: E402


def _glow_centre(slide, ground_hex):
    """The colour at the centre of the page's radial glow, blended over the ground the way a renderer blends it."""
    from lxml import etree
    for sh in slide.shapes:
        x = etree.tostring(sh._element).decode()
        m = re.search(r'<a:gs pos="0"><a:srgbClr val="(\w{6})"><a:alpha val="(\d+)"/>', x)
        if m and '<a:path path="circle">' in x:
            c, a = m.group(1), int(m.group(2)) / 100000.0
            return "".join("{:02X}".format(round((1 - a) * int(ground_hex[i:i + 2], 16) + a * int(c[i:i + 2], 16))) for i in (0, 2, 4)), a
    return None, None


for _g in ("light", "dawn"):
    for _pg, _fl in (("data", dict(number="12", label="conversations that changed our plan", note="Each one is archived.")),
                     ("section", dict(number="3", title="The numbers", kicker="Part three"))):
        _prs, _k, _sl = _page("starlit", _g, _pg, **_fl)
        _num = _text_shape(_sl, _fl["number"])
        _ink = str(_num.text_frame.paragraphs[0].runs[0].font.color.rgb)
        _c, _a = _glow_centre(_sl, vl._hex(_k.P["ground"]))
        check(_c is not None and vl._contrast(_ink, _c) >= 4.5,
              "starlit/{} {}: the numeral {} is >= 4.5:1 over the glow's centre #{} ({}; was ~3.0:1 at alpha 0.42)".format(
                  _g, _pg, _ink, _c, "{:.2f}:1 at alpha {:.2f}".format(vl._contrast(_ink, _c), _a) if _c else "no glow"))
        check(_c is not None and vl._contrast(_c, vl._hex(_k.P["ground"])) >= 1.5,
              "starlit/{} {}: the glow is still a glow ({:.2f}:1 over the ground at its centre)".format(
                  _g, _pg, vl._contrast(_c, vl._hex(_k.P["ground"])) if _c else 0))
dk.set_theme(None)

# ── Task 5 (C): broadsheet: column rules end at the words, the Inside rule at its list, the figure is set at the measure ──
from pptx.enum.shapes import MSO_SHAPE_TYPE as _MST        # noqa: E402


def _lines_of(slide, vertical):
    """[(x, y0, y1)] of the page's vertical rules (or [(y, x0, x1)] of its horizontal ones), in inches."""
    out = []
    for sh in slide.shapes:
        if sh.shape_type != _MST.LINE:
            continue
        x, y, w, h = _rect(sh)
        if vertical and w < 0.01 and h > 0.3:
            out.append((x, y, y + h))
        if not vertical and h < 0.01 and w > 0.3:
            out.append((y, x, x + w))
    return out


def _text_eq(slide, text):
    """The lowest drawn text box whose words are exactly `text` (a numeral "1" is also the running head's page number)."""
    hits = [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False) and not sh.is_placeholder
            and sh.text_frame.text.strip() == text]
    return max(hits, key=lambda sh: sh.top) if hits else None


_BS_PT = [("Bring it broken", "Anything that switches on, unzips or wobbles is welcome."),
          ("Fix it together", "A volunteer sits beside you, but your hands do the work."),
          ("Take it home", "What could not be fixed today goes on a list for next time.")]
_BS_MORE = " Then there is more to say about it, because it is the part people remember best." * 3
for _tag, _its in (("short copy", _BS_PT), ("long copy", [(a, b + _BS_MORE) for a, b in _BS_PT]),
                   ("very long copy", [(a, b + _BS_MORE * 2) for a, b in _BS_PT])):
    _prs, _k, _sl = _page("broadsheet", "light", "points", title="How it works", tags=["Idea", "Evening", "Hands"], items=_its)
    _rules = _lines_of(_sl, True)
    _top = min(r[1] for r in _rules)
    _foot = max((sh.top + sh.height) / 914400.0 for sh in _sl.shapes
                if getattr(sh, "has_text_frame", False) and not sh.is_placeholder and sh.top / 914400.0 >= _top - 0.01)
    _end = max(r[2] for r in _rules)
    check(len(_rules) == 2 and all(abs(r[2] - _end) < 0.01 for r in _rules),
          "broadsheet points ({}): two column rules, level".format(_tag))
    if _tag == "short copy":
        check(0.15 <= _end - _foot <= 0.30, "broadsheet points (short copy): the rules end 0.2in under the words ({:.2f}in under; "
              "was {:.2f}in, to the page's bottom)".format(_end - _foot, 6.90 - _foot))
    if _tag != "short copy":
        check(min(_foot + 0.15, 6.90) <= _end <= 6.91,
              "broadsheet points ({}): the rules still reach the words and stop at the page's bottom (rule to {:.2f}, "
              "words to {:.2f}, bottom 6.90)".format(_tag, _end, _foot))
    if _tag != "very long copy":         # (the very long one trips lint's ORPHANED PUNCTUATION in the unchanged column text too)
        check(not _hard(_prs, "bsp_" + _tag.replace(" ", "_")), "broadsheet points ({}): lint-clean".format(_tag))

_prs, _k, _sl = _page("broadsheet", "light", "cover", title="Every street needs a night for fixing things",
                      subtitle="How a repair evening works, and why it spreads.", masthead="The Repair Weekly", edition="No. 12",
                      inside=["The idea", "The evening", "The numbers"])
_vr = [r for r in _lines_of(_sl, True) if r[0] > 0.5 * W13]
_lst = _text_eq(_sl, "The idea\nThe evening\nThe numbers")
_lb = (_lst.top + _lst.height) / 914400.0
check(len(_vr) == 1 and abs(_vr[0][2] - _lb) <= 0.06,
      "broadsheet cover: the Inside rule ends at its list ({:.2f}in vs the list's {:.2f}in; was 6.90in, 0.75in past it)".format(
          _vr[0][2] if _vr else -1, _lb))

for _g in ("light", "salmon"):
    for _num, _lab, _note in (("3", "conversations that changed our plan", "Each one is written up in the archive."),
                              ("1,250,000", "visitors since the first night", ""), ("三成", "的物品得到修复", "每一件都记录在档案里。")):
        _fl = dict(number=_num, label=_lab, **({"note": _note} if _note else {}))
        _prs, _k, _sl = _page("broadsheet", _g, "data", **_fl)
        _nsh = _text_eq(_sl, _num)
        _z = _size_of(_sl, _num) if _nsh is not None else 0
        _ink = vl._hex(_k.P["ink"]).upper()
        _solid = [sh for sh in _sl.shapes if sh.shape_type == _MST.AUTO_SHAPE and sh.fill.type == 1
                  and str(sh.fill.fore_color.rgb).upper() == _ink]
        _hr = _lines_of(_sl, False)
        _span = max((r[2] - r[1] for r in _hr if r[0] > 1.2), default=0)
        _need = 180 if len(_num) > 6 else 200            # a nine-character figure is bounded by the measure itself (191pt)
        check(_nsh is not None and not _solid and _span >= 0.88 * W13 - 0.01, "broadsheet/{} data {!r}: no black cap ({} solid "
              "ink boxes), and the block's rules run the measure ({:.1f}in)".format(_g, _num, len(_solid), _span))
        check_mac(_nsh is not None and _z >= _need - 1e-6 and not _solid and _span >= 0.88 * W13 - 0.01,
              "broadsheet/{} data {!r}: the figure is >= {}pt ({:.0f}pt; was 150pt), no black cap ({} solid ink boxes), and the "
              "block's rules run the measure ({:.1f}in = {:.0%})".format(_g, _num, _need, _z, len(_solid), _span, _span / W13))
    _prs, _k, _sl = _page("broadsheet", _g, "data", number="3", label="conversations that changed our plan", note="Each one is archived.")
    _ns, _ls = _text_eq(_sl, "3"), _text_shape(_sl, "conversations")
    _cr = [r for r in _lines_of(_sl, True) if r[0] > 1.0]
    _ln = _rect(_ls)
    check(len(_cr) == 1 and _rect(_ns)[0] + _rect(_ns)[2] <= _cr[0][0] + 1e-6 <= _ln[0] + 1e-6 and _cr[0][1] <= _ln[1] and _cr[0][2] >= _ln[1] + _ln[3],
          "broadsheet/{} data: the label and note stand in a column ruled off from the figure (rule x={:.2f}in between the figure "
          "and the words, spanning them)".format(_g, _cr[0][0] if _cr else -1))
for _W, _H in ((13.333, 7.5), (7.5, 7.5), (7.5, 13.333), (10.0, 7.5)):
    _prs = dk.blank_deck(_W, _H)
    with contextlib.redirect_stdout(io.StringIO()):
        _kk = vl.use("broadsheet", _prs, ground="light")
        for _num, _lab, _note in (("3", "conversations that changed our plan", "Each one is written up in the archive."),
                                  ("98.6%", "of repairs held after a month", ""), ("三成", "的物品得到修复", "每一件都记录在档案里。")):
            _kk.data(_kk.new_slide(), number=_num, label=_lab, masthead="The Repair Weekly", **({"note": _note} if _note else {}))
    check(not _hard(_prs, "bsd_{}x{}".format(int(_W), int(_H))), "broadsheet data {}x{}: the three figures are lint-clean".format(_W, _H))
dk.set_theme(None)

# ── Task 5 (C): journal: the margin note's rule ends at its words, numerals stand on their heads' baseline, the green figure sits on a mat ──
_JN_ITEMS = [("Three methods", "Same data, same metric."), ("Four accelerations", "From 2x to 8x."),
             ("Five metrics", "Scored the same way.")]
for _tag, _note in (("two-line note", "All runs on one scanner."),
                    ("long note", "All runs on one scanner, with the same coil, the same protocol and the same reader. " * 9)):
    for _g in ("light", "green"):
        _prs, _k, _sl = _page("journal", _g, "points", title="What we compared", items=_JN_ITEMS, margin=_note)
        _accent = vl._hex(_k.P["text_accents"][0]).upper()
        _rule = [sh for sh in _sl.shapes if sh.shape_type == _MST.LINE and _rect(sh)[2] < 0.01 and _rect(sh)[3] > 0.3
                 and str(sh.line.color.rgb).upper() == _accent]
        _nt = _text_shape(_sl, _note[:12])
        _lst = _text_shape(_sl, "Scored the same way")
        _foot_n, _foot_l = _rect(_nt)[1] + _rect(_nt)[3], _rect(_lst)[1] + _rect(_lst)[3]
        _end = _rect(_rule[0])[1] + _rect(_rule[0])[3] if _rule else -1
        check(len(_rule) == 1 and abs(_end - max(_foot_n, _foot_l)) <= 0.06,
              "journal/{} margin ({}): the note's rule ends at the lower of the note and the list ({:.2f}in; note {:.2f}, list {:.2f}; "
              "was 6.90in)".format(_g, _tag, _end, _foot_n, _foot_l))
        check(_end <= 6.90 + 0.01 and (_tag != "two-line note" or _end <= 6.9 - 0.5),
              "journal/{} margin ({}): the rule is not the page's height ({:.2f}in of 6.90)".format(_g, _tag, _end))
        check(not _hard(_prs, "jnm_{}_{}".format(_g, _tag.split("-")[0].split()[0])), "journal/{} margin ({}): lint-clean".format(_g, _tag))


def _ink_bottom(img, rect, dpi, bg):
    """The lowest row (in) of ink inside `rect` (in) of a rendered page: pixels that differ from the ground."""
    from PIL import Image, ImageChops
    x, y, w, h = [int(round(v * dpi)) for v in rect]
    crop = img.crop((x, y, x + w, y + h)).convert("RGB")
    d = ImageChops.difference(crop, Image.new("RGB", crop.size, bg)).convert("L").point(lambda v: 255 if v > 40 else 0)
    box = d.getbbox()
    return None if box is None else (y + box[3]) / dpi


def _render_png(prs, tag, dpi=200):
    """The first page of `prs` rendered by LibreOffice + pdftoppm (None when either is missing)."""
    import shutil, subprocess
    from PIL import Image
    if not (shutil.which("soffice") and shutil.which("pdftoppm")):
        return None
    f = WORK / (tag + ".pptx")
    prs.save(str(f))
    prof = (WORK / "lo_profile").as_uri()
    subprocess.run(["soffice", "-env:UserInstallation=" + prof, "--headless", "--convert-to", "pdf", "--outdir", str(WORK), str(f)],
                   capture_output=True, timeout=240)
    pdf = WORK / (tag + ".pdf")
    if not pdf.exists():
        return None
    subprocess.run(["pdftoppm", "-r", str(dpi), "-f", "1", "-l", "1", "-png", str(pdf), str(WORK / tag)], capture_output=True, timeout=120)
    png = sorted(WORK.glob(tag + "-*.png"))
    return Image.open(png[0]).convert("RGB") if png else None


# the numerals and the heads: baselines from the face metrics, then — where the faces and a renderer exist — measured on a render
for _g in ("light", "green"):
    _prs, _k, _sl = _page("journal", _g, "points", title="What we compared", items=_JN_ITEMS, margin="All runs on one scanner.")
    _gaps = []
    for _i, (_hd, _ln) in enumerate(_JN_ITEMS):
        _nb, _hb = _text_eq(_sl, str(_i + 1)), _text_shape(_sl, _hd)
        _zn = max(r.font.size.pt for p_ in _nb.text_frame.paragraphs for r in p_.runs)
        _zh = max(r.font.size.pt for p_ in _hb.text_frame.paragraphs for r in p_.runs)
        _gaps.append((_rect(_nb)[1] + vls.first_baseline(_k, "item_no", str(_i + 1), _zn)) - (_rect(_hb)[1] + vls.first_baseline(_k, "item_head", _hd, _zh)))
    check(max(abs(g_) for g_ in _gaps) <= 0.01, "journal/{} points: each numeral's first baseline is its head's, by the faces' metrics "
          "(max gap {:.3f}in; was 0.12in)".format(_g, max(abs(g_) for g_ in _gaps)))
    if dk._font_substituted("Georgia") or dk._font_substituted("Times New Roman"):
        continue                                       # no real faces here: the metrics above are all there is to measure
    _img = _render_png(_prs, "jn_base_" + _g)
    if _img is None:
        continue
    _bg = tuple(int(vl._hex(_k.P["ground"])[i:i + 2], 16) for i in (0, 2, 4))
    _rows = []
    for _i, (_hd, _ln) in enumerate(_JN_ITEMS):
        _nb, _hb = _text_eq(_sl, str(_i + 1)), _text_shape(_sl, _hd)
        _a, _b = _ink_bottom(_img, _rect(_nb), 200, _bg), _ink_bottom(_img, _rect(_hb), 200, _bg)
        _rows.append(None if _a is None or _b is None else _a - _b)
    check(all(r_ is not None for r_ in _rows) and max(abs(r_) for r_ in _rows) <= 0.025,
          "journal/{} points: on a render, the numeral's ink stands on its head's baseline (bottoms differ by {} in; was ~0.12in)".format(
              _g, ", ".join("{:+.3f}".format(r_) for r_ in _rows if r_ is not None)))


def _pdf_first_origins(prs, tag):
    """The baseline (in, from the page top) of the first text span on each page of `prs`, read from LibreOffice's PDF
    (the glyphs' own origins); None when soffice or PyMuPDF is missing."""
    import shutil, subprocess
    try:
        import fitz
    except ImportError:
        return None
    if not shutil.which("soffice"):
        return None
    f = WORK / (tag + ".pptx")
    prs.save(str(f))
    subprocess.run(["soffice", "-env:UserInstallation=" + (WORK / "lo_profile").as_uri(), "--headless", "--convert-to", "pdf",
                    "--outdir", str(WORK), str(f)], capture_output=True, timeout=600)
    pdf = WORK / (tag + ".pdf")
    if not pdf.exists():
        return None
    out = []
    for page in fitz.open(str(pdf)):
        sp = [s_ for b_ in page.get_text("rawdict")["blocks"] for l_ in b_.get("lines", []) for s_ in l_["spans"]]
        out.append(sp[0]["origin"][1] / 72.0 if sp else None)
    return out


# first_baseline() is where the RENDER sets the first baseline, in every face (bug 7). It modelled the face's own ascent
# plus line gap; LibreOffice lays out a PPTX line FONT-INDEPENDENTLY — 1.2 em a line, its baseline 1.0 em under the
# line's top, proportional spacing over 100% adding its extra above — measured on 16 faces from Courier New (ascent
# 0.83 em) to Zapfino (1.88 em), 40–200pt, spacing 0.9–1.5. The face model put a 150pt Arial Black figure's baseline
# 0.21in low (pixel's sprite hung under its digits) and a Hiragino Sans GB line's 0.41 em low.
_BL_FACES = ("Arial Black", "Impact", "Arial", "Georgia", "Times New Roman", "Trebuchet MS", "Courier New")
_bl_prs = dk.blank_deck(13.333, 7.5)
with contextlib.redirect_stdout(io.StringIO()):
    _bl_k = vl.use("journal", _bl_prs)
_bl_cases = [(f_, z_, t_) for f_ in _BL_FACES for z_ in (40, 80, 150, 200) for t_ in ("3608", "Hx")]
_bl_cases += [(f_, z_, "中文36") for f_ in ("Songti SC", "Hiragino Sans GB") for z_ in (40, 150)]
_bl_model = []
for _f, _z, _t in _bl_cases:
    _bl_k.face = lambda role, *_a, _f=_f: _f                 # the kit's face for every role: this case's face
    _bl_k.ea_face = lambda role, *_a, _f=_f: _f
    _cjk = dk._has_cjk(_t)
    _used = _f if _cjk else dk.numeral_run_face(_t, _f)      # the face the kit's own run draws (lining figures)
    _bold = vl._weight(_bl_k, "item_head", _t)
    _bl_model.append(0.5 + vls.first_baseline(_bl_k, "item_head", _t, _z))
    dk.text(_bl_prs.slides.add_slide(_bl_prs.slide_layouts[6]), 0.5, 0.5, 12.0, 6.5,
            [[(_t, _z, dk.DEEP, _bold, False, _used, _used if _cjk else None)]], space_after=0)
_bl_seen = _pdf_first_origins(_bl_prs, "first_baseline")
if _bl_seen is None:
    skipped.append("first_baseline against a LibreOffice render (soffice or PyMuPDF missing)")
else:
    _bl_err = {}
    for (_f, _z, _t), _m, _o in zip(_bl_cases, _bl_model, _bl_seen):
        _e = abs(_m - _o) if _o is not None else 9.0
        _bl_err[_f] = max(_bl_err.get(_f, 0.0), _e)
    check_mac(max(_bl_err.values()) <= 0.01,
              "first_baseline() is where LibreOffice sets the first baseline, 40–200pt, every face (worst error per "
              "face: {})".format(", ".join("{} {:.3f}in".format(f_, e_) for f_, e_ in _bl_err.items())))

_FIG = str(vl.ASSETS / "figure" / "illustrative-curves.jpg")
for _g in ("light", "green"):
    _prs, _k, _sl = _page("journal", _g, "image_text", title="Error grows with acceleration", body="Error against acceleration.",
                          caption="Source: an illustrative figure", image=_FIG)
    _pic = [sh for sh in _sl.shapes if sh.shape_type == _MST.PICTURE][0]
    _pr = _rect(_pic)
    _mats = [sh for sh in _sl.shapes if sh.shape_type == _MST.AUTO_SHAPE and sh.fill.type == 1
             and str(sh.fill.fore_color.rgb).upper() == vl._hex(_k.P["ink"]).upper()
             and _rect(sh)[0] <= _pr[0] and _rect(sh)[1] <= _pr[1]
             and _rect(sh)[0] + _rect(sh)[2] >= _pr[0] + _pr[2] and _rect(sh)[1] + _rect(sh)[3] >= _pr[1] + _pr[3]]
    from PIL import Image as _Im2
    _iw, _ih = _Im2.open(_FIG).size
    _whole = abs((_pr[2] / _pr[3]) - (_iw / _ih)) < 0.01 and not (_pic.crop_left or _pic.crop_right or _pic.crop_top or _pic.crop_bottom)
    if _g == "green":
        _m = _rect(_mats[0]) if _mats else None
        _edges = [] if not _m else [_pr[0] - _m[0], _pr[1] - _m[1], _m[0] + _m[2] - _pr[0] - _pr[2], _m[1] + _m[3] - _pr[1] - _pr[3]]
        check(len(_mats) == 1 and all(abs(e_ - 0.15) <= 0.01 for e_ in _edges) and _whole,
              "journal/green figure: the whole picture sits on a 0.15in cream mat ({})".format(", ".join("{:.3f}".format(e_) for e_ in _edges)))
        check(_m is not None and _m[0] >= 0.05 and _m[1] + _m[3] <= 7.5 - 0.55, "journal/green figure: the mat stays on the page, above the footer band")
        check(not _hard(_prs, "jnfig_green"), "journal/green figure: lint-clean")
    else:
        check(not _mats and _whole, "journal/light figure: unchanged — no mat, the whole picture ({} mats)".format(len(_mats)))
dk.set_theme(None)

# ── Task 5 (C): tally: the pills share one size per page and stand on the column's right margin ──
def _pills(slide, k):
    """[(text, size_pt, right_edge_in)] of the page's pills: the lime rounded boxes and the words set on them."""
    lime = vl._hex(k.P["lime"]).upper()
    out = []
    for sh in slide.shapes:
        if sh.shape_type == _MST.AUTO_SHAPE and sh.fill.type == 1 and str(sh.fill.fore_color.rgb).upper() == lime:
            x, y, w, h = _rect(sh)
            for t in slide.shapes:
                if getattr(t, "has_text_frame", False) and t.text_frame.text.strip() and abs(_rect(t)[0] - x) < 0.005 \
                        and abs(_rect(t)[1] - y) < 0.005:
                    out.append((t.text_frame.text, t.text_frame.paragraphs[0].runs[0].font.size.pt, x + w))
    return out


_TL_PTS = [("New members joined", "Mostly through word of mouth"), ("Tools went out on loan", "Drills and ladders led the list"),
           ("Items came back repaired", "Fixed at the monthly evening")]
_TL_LONG = "A tag that is much too long for its pill and has to shrink to fit its column"
for _g in ("light", "night"):
    for _tag, _tags, _cv in (("3 tags", ["Members", "Loans", "Repairs"], (13.333, 7.5)),
                             ("CJK tags", ["成员", "借出", "修复"], (13.333, 7.5)),
                             ("a long tag on a square page", ["Short", _TL_LONG, "Mid tag"], (7.5, 7.5)),
                             ("a long tag on a portrait page", ["Short", _TL_LONG, "Mid tag"], (7.5, 13.333))):
        _prs = dk.blank_deck(*_cv)
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                _k = vl.use("tally", _prs, ground=_g)
                _k.points(_k.new_slide(), title="Where the quarter went", tags=_tags, items=_TL_PTS)
        except vl.VLTextOverflow as _e:
            # refusing loudly is the contract when a tag cannot fit even at its floor; on the real face the long tag
            # fits, and only a wider STAND-IN (CI's Linux DejaVu) may refuse it — the suite stopped dead there before
            _face = _k.face(vl.TYPE["tally"]["tag"][1])
            check(dk._font_substituted(_face), "tally/{} points ({}): builds on its real face ({})".format(_g, _tag, str(_e)[:70]))
            dk.set_theme(None)
            continue
        _pl = _pills(_prs.slides[0], _k)
        _sizes = sorted({round(z_, 2) for _t, z_, _r in _pl})
        _rights = [r_ for _t, _z, r_ in _pl]
        check(len(_pl) == 3 and len(_sizes) == 1,
              "tally/{} points ({}): the pills share one size ({}; was 11.0/12.0 beside the words, 9–12 for a long tag)".format(
                  _g, _tag, "/".join("{:g}pt".format(z_) for z_ in _sizes)))
        if "long tag" not in _tag:                   # (under their words, a long tag's pills hang left, flush to the words)
            _edge = 0.93 * _cv[0]
            check(len(_pl) == 3 and all(abs(r_ - _edge) <= 0.005 for r_ in _rights),
                  "tally/{} points ({}): every pill's right edge is the margin at {:.2f}in ({}; was ragged by up to 0.16in)".format(
                      _g, _tag, _edge, ", ".join("{:.3f}".format(r_) for r_ in _rights)))
        check(not _hard(_prs, "tlp_{}_{}".format(_g, _tag.replace(" ", "_"))), "tally/{} points ({}): lint-clean".format(_g, _tag))

# ── shared helpers for the Task 5 (A) checks (image-led languages) ──
def _deck_at(L, g, W, H):
    prs = dk.blank_deck(W, H)
    with contextlib.redirect_stdout(io.StringIO()):
        k = vl.use(L, prs, ground=g)
    return prs, k


def _text_shapes(slide):
    return [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]


def _run_of(sh):
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size and r.text.strip():
                return r
    return None


# ── Task 5 (A): editorial: a points item number never wraps ──
# "01" in Arial Black 30.5pt measured 0.5651in in a box whose inner width was 0.5644in, so LibreOffice set "0" over "1".
# The number's box is now its measured width + >= 10% (measured by both widths the kit uses), at 2, 3 and 4 points.
import display_type as _dt
_PTS = [("Your broken thing", "Any size that fits through the door."), ("An hour", "Most repairs take less."),
        ("Curiosity", "Volunteers explain every step."), ("A cup of tea", "There is always a kettle on.")]
for L in ("editorial", "soft", "collage"):
    for g in vl.VARIANTS[L]:
        for n in (2, 3, 4):
            for (W, H, pic) in ((13.333, 7.5, str(vl.ASSETS / "photo/tools-tray.jpg")), (13.333, 7.5, None),
                                (7.5, 10.0, str(vl.ASSETS / "photo/tools-tray.jpg"))):
                prs, k = _deck_at(L, g, W, H)
                sl = k.new_slide()
                try:
                    with contextlib.redirect_stdout(io.StringIO()):
                        k.points(sl, kicker="A repair café", title="Things to bring", items=_PTS[:n], image=pic)
                except Exception as e:
                    check(False, "{}/{}: points x{} builds ({}: {})".format(L, g, n, type(e).__name__, str(e)[:70])); continue
                nums = [sh for sh in _text_shapes(sl) if re.fullmatch(r"\d\d", sh.text_frame.text.strip())]
                worst = 9.0
                for sh in nums:
                    r = _run_of(sh)
                    face, bold = r.font.name, bool(r.font.bold)
                    need = max(dk._natural_width_in([(r.text, bold)], r.font.size.pt, face),
                               _dt._glyph_width(r.text, r.font.size.pt, face, bold) or 0.0)
                    worst = min(worst, (sh.width / 914400.0 - dk.TEXT_INSET_LR) / need)
                check(len(nums) == n and worst >= 1.10 - 1e-6,
                      "{}/{} {} points x{}{}: every item number's box holds its digits with >= 10% to spare "
                      "(tightest {:.3f}x, {} numbers)".format(L, g, "landscape" if W > H else "portrait", n,
                                                              "" if pic else " (no picture)", worst, len(nums)))
dk.set_theme(None)


# ── Task 5 (A): editorial: the full-bleed cover is cropped for the part of it that is VISIBLE, biased to the top ──
# The picture was cover-fitted to the whole 13.33in page while the panel covers x >= 7.2in: a 853x1280 portrait kept
# 37.4% of its height in a 7.2in-wide window, and the crop (centred) cut the heads at the top edge.
from PIL import Image as _PIL


def _visible_source(slide):
    """(x0, y0, x1, y1) as fractions of the SOURCE picture that the viewer sees on `slide`'s first picture — the
    picture's frame and crop, minus the opaque panel box that follows it — and the pixel aspect of that window."""
    W, H = slide.part.package.presentation_part.presentation.slide_width / 914400.0, \
        slide.part.package.presentation_part.presentation.slide_height / 914400.0
    shapes = list(slide.shapes)
    pi = next(i for i, sh in enumerate(shapes) if sh.shape_type == 13)
    pic = shapes[pi]
    panel = next(sh for sh in shapes[pi + 1:] if sh.shape_type == 1)
    px, py, pw, ph = (v / 914400.0 for v in (panel.left, panel.top, panel.width, panel.height))
    if pw < 0.9 * W:                                    # a side band
        vis = (px + pw, 0.0, W, H) if px < 0.01 else (0.0, 0.0, px, H)
    else:                                               # a top or bottom band
        vis = (0.0, py + ph, W, H) if py < 0.01 else (0.0, 0.0, W, py)
    fx, fy, fw, fh = (v / 914400.0 for v in (pic.left, pic.top, pic.width, pic.height))
    cl, cr, ct, cb = pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom
    x0, x1 = max(vis[0], fx), min(vis[2], fx + fw)
    y0, y1 = max(vis[1], fy), min(vis[3], fy + fh)
    sx = lambda x: cl + (1 - cl - cr) * (x - fx) / fw
    sy = lambda y: ct + (1 - ct - cb) * (y - fy) / fh
    return (sx(x0), sy(y0), sx(x1), sy(y1)), vis


for pg in ("cover", "closing"):
    for tag, src in (("portrait", "photo/hall-repair.jpg"), ("landscape", "photo/tools-tray.jpg"),
                     ("square", "photo/bench-toaster.jpg")):
        path = str(vl.ASSETS / src)
        iw, ih = _PIL.open(path).size
        for g in vl.VARIANTS["editorial"]:
            prs, k = _deck_at("editorial", g, 13.333, 7.5)
            sl = k.new_slide()
            with contextlib.redirect_stdout(io.StringIO()):
                if pg == "cover":
                    k.cover(sl, title="Bring it broken, take it home working", kicker="Repair café", image=path)
                else:
                    k.closing(sl, title="Questions", line="Thank you.", image=path)
            (sx0, sy0, sx1, sy1), vis = _visible_source(sl)
            win_ar = ((sx1 - sx0) * iw) / ((sy1 - sy0) * ih)
            vis_ar = (vis[2] - vis[0]) / (vis[3] - vis[1])
            # the share of the picture's height a window of the visible rect's shape can show: all of it for a
            # picture wider than the window, vis_ar-limited (ih/iw * ...) for a taller one
            max_h = 1.0 if iw / ih >= vis_ar else (iw / ih) / vis_ar
            check(abs(win_ar - vis_ar) / vis_ar < 0.02,
                  "editorial {} ({} source, {}): the visible window shows the picture undistorted ({:.3f} vs {:.3f})".format(pg, tag, g, win_ar, vis_ar))
            check((sy1 - sy0) >= 0.95 * max_h,
                  "editorial {} ({} source, {}): the crop is made for the visible rect — it keeps {:.1%} of the picture's height "
                  "(the most that rect can show is {:.1%})".format(pg, tag, g, sy1 - sy0, max_h))
            if max_h < 0.999:                           # a taller picture is cropped top/bottom: the top is the side kept
                check(sy0 <= 0.12 and sy0 < (1 - sy1), "editorial {} ({} source, {}): the crop is biased to the top (top edge "
                      "{:.1%} cut, bottom {:.1%})".format(pg, tag, g, sy0, 1 - sy1))
dk.set_theme(None)

# ── Task 5 (A): soft: a quote card is sized to its content and the mark sits next to the quote ──
# The mark's 90pt line box (1.56in) is mostly blank under a glyph 0.38in tall, and the quote started under all of it:
# 1.16in of blank card between the mark and the first line, the top 58% of the card holding only the mark. Measured
# on the RENDER (the glyph's ink is not in the file): the longest run of blank rows inside the card, between its ink.
import subprocess
import render_deck as _rd


def _render_slides(prs, tag, dpi=60):
    """The slides of `prs` as PIL images, rendered by LibreOffice; None when it is not installed (the caller fails
    loudly — a render check that did not run must not read as one that passed)."""
    soffice = _rd.find_soffice()
    if not soffice:
        return None
    import fitz
    from PIL import Image
    d = WORK / ("r_" + tag); d.mkdir(exist_ok=True)
    f = d / (tag + ".pptx"); prs.save(str(f))
    for _try in range(3):                       # LibreOffice sometimes exits without a PDF when another instance is closing
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(d), str(f)], capture_output=True,
                       timeout=180)
        if (d / (tag + ".pdf")).exists():
            break
    if not (d / (tag + ".pdf")).exists():
        return None
    out = []
    for pg in fitz.open(str(d / (tag + ".pdf"))):
        pix = pg.get_pixmap(dpi=dpi)
        out.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
    return out


def _blank_runs(im, rect_in, dpi, bg, tol=60):
    """(first ink row, last ink row, longest blank run between them) in inches, for the ink inside `rect_in` — pixels
    further than `tol` (summed channel difference) from the colour `bg`."""
    x, y, w, h = (int(round(v * dpi)) for v in rect_in)
    rows = []
    for yy in range(y, y + h):
        rows.append(any(sum(abs(a - b) for a, b in zip(im.getpixel((xx, yy)), bg)) > tol for xx in range(x, x + w, 2)))
    inked = [i for i, v in enumerate(rows) if v]
    if not inked:
        return None
    gap = best = 0
    for v in rows[inked[0]:inked[-1] + 1]:
        gap = 0 if v else gap + 1
        best = max(best, gap)
    return (y + inked[0]) / dpi, (y + inked[-1] + 1) / dpi, best / dpi


for g in vl.VARIANTS["soft"]:
    prs, k = _deck_at("soft", g, 13.333, 7.5)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.quote(sl, quote="The visitor holds the screwdriver; the volunteer only guides.", attribution="How every repair begins",
                image=str(vl.ASSETS / "photo/hall-repair.jpg"))
    card = max((sh for sh in sl.shapes if sh.shape_type == 1), key=lambda sh: sh.width * sh.height)
    cx, cy, cw, ch = (v / 914400.0 for v in (card.left, card.top, card.width, card.height))
    ims = _render_slides(prs, "softq_" + g)
    if ims is None:
        check(False, "soft/{} quote card: LibreOffice not found — the render half of this check did not run".format(g)); continue
    bg = ims[0].getpixel((int((cx + 0.12) * 60), int((cy + 0.05) * 60)))      # the card's own colour, near its top-left corner
    first, last, blank = _blank_runs(ims[0], (cx + 0.2, cy + 0.12, cw - 0.4, ch - 0.24), 60, bg)
    check(blank <= 0.35, "soft/{} quote card: no blank band inside the card is taller than 0.35in — the mark sits next to the "
          "quote (longest blank run {:.2f}in)".format(g, blank))
    check(ch <= (last - first) + 1.0, "soft/{} quote card stays sized to its content: {:.2f}in of card for {:.2f}in of "
          "ink (<= 1.0in of padding in all)".format(g, ch, last - first))
dk.set_theme(None)

# ── Task 5 (A): soft: the figure's disc keeps its true colour and the digit takes the ground's ink (dusk) ──
# On dusk the sage disc was blended toward the ground (9DB8A0 -> 7C8D83) so the off-white digit read at 3.05:1 on it
# (the disc itself 3.98:1 on the ground). The disc stays the language's sage and the digit is set in the ground's
# dark ink — 6.5:1 — on the data page and on the section pill alike.
from pptx.enum.shapes import MSO_SHAPE as _MS


def _hexc(c):
    return str(c).upper()


for g in vl.VARIANTS["soft"]:
    acc_sage = vl.VARIANTS["soft"][g]["palette"]["accents"][1]
    for page, kw in (("data", dict(number="1", label="evening a month", note="Short enough to fit around work.")),
                     ("data", dict(number="87%", label="of repairs end well")),
                     ("section", dict(number="02", kicker="Repair café", title="Everything you need is on the table"))):
        for pic in (str(vl.ASSETS / "photo/jacket-mend.jpg"), None):
            for (W, H) in ((13.333, 7.5), (7.5, 10.0)):
                prs, k = _deck_at("soft", g, W, H)
                sl = k.new_slide()
                try:
                    with contextlib.redirect_stdout(io.StringIO()):
                        getattr(k, page)(sl, image=pic, **kw)
                except Exception as e:
                    check(False, "soft/{} {} builds ({}: {})".format(g, page, type(e).__name__, str(e)[:70])); continue
                num = next((sh for sh in _text_shapes(sl) if sh.text_frame.text.strip() == kw["number"]), None)
                if num is None:
                    check(False, "soft/{} {}: the figure {!r} is on the page".format(g, page, kw["number"])); continue
                # the disc is drawn on the figure's own line: it spans the figure's box from top to bottom (to 0.3in) and meets it across
                held = [sh for sh in sl.shapes if sh.shape_type == 1 and sh.auto_shape_type in (_MS.OVAL, _MS.ROUNDED_RECTANGLE)
                        and sh.top <= num.top + 0.3 * 914400 and sh.top + sh.height >= num.top + num.height - 0.3 * 914400
                        and sh.left < num.left + num.width and sh.left + sh.width > num.left]
                if not held:
                    check(False, "soft/{} {} '{}': a disc holds the figure".format(g, page, kw["number"])); continue
                disc = min(held, key=lambda sh: sh.width * sh.height)
                fill = _hexc(disc.fill.fore_color.rgb)
                ink = _hexc(_run_of(num).font.color.rgb)
                cr = vl._contrast(fill, ink)
                ground = vl.VARIANTS["soft"][g]["palette"]["ground"]
                check(fill == acc_sage and cr >= 4.5 and (g != "dusk" or ink == _hexc(ground)),
                      "soft/{} {} '{}' {}{}: the disc is the true sage {} (is {}) and its figure reads {:.2f}:1 ({} on it)".format(
                          g, page, kw["number"], "landscape" if W > H else "portrait", "" if pic else " (no picture)", acc_sage, fill, cr, ink))
                if g == "dusk":
                    # the lint reads dark ink against what is behind its text box: the disc must hold the box (it read the
                    # ground — '1' ink #2E2940 on fill #2E2940, INVISIBLE TEXT — until the box was the disc's)
                    mine = [f for f in _hard(prs, "sd_{}_{}_{}".format(page, int(W), int(bool(pic)))) if "'{}'".format(kw["number"]) in f]
                    check(not mine, "soft/dusk {} '{}' {}{}: the figure on its disc raises no lint finding ({})".format(
                        page, kw["number"], "landscape" if W > H else "portrait", "" if pic else " (no picture)", mine[:1]))
dk.set_theme(None)

# ── Task 5 (A): collage: on slate the prints and notes are cream paper with dark ink; kraft stays exactly as it is ──
# Every print on slate kept its FFFFFF border (14.4:1 on the ground — the brightest thing on a dark page) and the quote
# note was the ground's own dark panel 3B3934 (1.24:1 on the ground, a dark slab with tape the only paper cue). Both are
# cream now (E8DFC9), the note's words in the light ground's dark inks; kraft's white borders, white notes and ink are as before.
def _lumh(c):
    return vl._lum(c)


def _frames(slide):
    """(frame fill, picture) for each print on `slide`: the filled box drawn right under the picture."""
    sh = list(slide.shapes)
    return [(_hexc(sh[i - 1].fill.fore_color.rgb), s_) for i, s_ in enumerate(sh)
            if s_.shape_type == 13 and i and sh[i - 1].shape_type == 1 and sh[i - 1].fill.type == 1]


def _notes_text(slide):
    """[(card fill, run colour, text)] for every run of text that sits on a filled card (the smallest one holding the
    text box's centre), excluding the page's own ground."""
    out = []
    cards = [sh for sh in slide.shapes if sh.shape_type == 1 and sh.fill.type == 1
             and sh.width * sh.height < 0.5 * 13.333 * 7.5 * 914400 ** 2 and sh.height > 0.5 * 914400]
    for t in _text_shapes(slide):
        cx, cy = t.left + t.width / 2, t.top + t.height / 2
        held = [c for c in cards if c.left - 0.2 * 914400 <= cx <= c.left + c.width + 0.2 * 914400
                and c.top - 0.2 * 914400 <= cy <= c.top + c.height + 0.2 * 914400]
        if not held:
            continue
        card = min(held, key=lambda c: c.width * c.height)
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip() and r.font.color and r.font.color.type is not None:
                    out.append((_hexc(card.fill.fore_color.rgb), _hexc(r.font.color.rgb), r.text.strip()))
    return out


_CPTS = [("Your broken thing", "Any size that fits through the door."), ("An hour", "Most repairs take less."),
         ("Curiosity", "Volunteers explain every step.")]
for g in vl.VARIANTS["collage"]:
    ground = vl.VARIANTS["collage"][g]["palette"]["ground"]
    prs, k = _deck_at("collage", g, 13.333, 7.5)
    pic = str(vl.ASSETS / "photo/hall-repair.jpg")
    slides = {}
    with contextlib.redirect_stdout(io.StringIO()):
        slides["quote"] = k.new_slide(); k.quote(slides["quote"], quote="The visitor holds the screwdriver; the volunteer only guides.",
                                                 attribution="How every repair begins", image=pic)
        slides["points"] = k.new_slide(); k.points(slides["points"], kicker="A repair café", title="Three things to bring", items=_CPTS,
                                                   image=str(vl.ASSETS / "photo/tools-tray.jpg"))
        slides["data"] = k.new_slide(); k.data(slides["data"], number="1", label="evening a month", note="Short enough.",
                                               image=str(vl.ASSETS / "photo/jacket-mend.jpg"))
        slides["cover"] = k.new_slide(); k.cover(slides["cover"], kicker="Repair café", title="Bring it broken",
                                                 image=[str(vl.ASSETS / x) for x in vl.SAMPLE_IMAGES["photo"][:3]])
    fr = [f for sl in slides.values() for f in _frames(sl)]
    if g == "light":
        check(fr and all(f == "FFFFFF" for f, _p in fr), "collage/kraft: every print keeps its white border exactly ({} prints)".format(len(fr)))
        nt = _notes_text(slides["quote"]) + _notes_text(slides["points"])
        check(nt and all(c == "FFFFFF" and i != "FFFFFF" and vl._contrast(c, i) >= 4.5 for c, i, _t in nt),
              "collage/kraft: notes stay white with the language's own ink ({} runs)".format(len(nt)))
        tapes = [_hexc(sh.fill.fore_color.rgb) for sl in slides.values() for sh in sl.shapes
                 if sh.shape_type == 5 and sh.fill.type == 1]
        check(tapes and set(tapes) == {"EDE3C8"}, "collage/kraft: the tape stays washi EDE3C8 ({} strips: {})".format(len(tapes), sorted(set(tapes))))
    else:
        check(fr and all(0.55 <= _lumh(f) <= 0.80 and vl._contrast(f, ground) >= 4.5 for f, _p in fr),
              "collage/slate: every print has a cream border, not white: {} prints, luminance {} (white is 1.0), {:.1f}:1 on the ground".format(
                  len(fr), sorted({round(_lumh(f), 2) for f, _p in fr}), min([vl._contrast(f, ground) for f, _p in fr] or [0])))
        for pg in ("quote", "points"):
            nt = _notes_text(slides[pg])
            worst = min([vl._contrast(c, i) for c, i, _t in nt] or [0])
            check(nt and all(_lumh(c) >= 0.55 and _lumh(i) <= 0.2 and vl._contrast(c, i) >= 4.5 for c, i, _t in nt),
                  "collage/slate {}: every note is cream paper with dark ink on it — {} runs, card luminance {}, ink luminance <= {:.2f}, "
                  "worst {:.1f}:1".format(pg, len(nt), sorted({round(_lumh(c), 2) for c, _i, _t in nt}),
                                          max([_lumh(i) for _c, i, _t in nt] or [0]), worst))
dk.set_theme(None)

# ── Task 5 (A): storybook: the data page's watercolour is a real share of the page, and its figure is not a hairline ──
# The harvest-basket watercolour sat in a 5.07x2.85in slot on the data page: 10.9% of the page was picture where the other
# image-led languages gave 22-26%, and the figure beside it was Times New Roman 150pt regular — a stem 12% of its height.
# The figure is bold now (21%), and the slot follows the picture's shape (a wide picture gets a wide slot).
import numpy as np
from PIL import Image as _PILI


def _art_share(slide, W, H):
    """Share of the page that is picture ink: for each picture, the fraction of its pixels at alpha >= 128 times the area
    it is drawn at — a feathered watercolour fades to the paper, and the faded edge is not picture."""
    tot = 0.0
    for sh in slide.shapes:
        if sh.shape_type == 13:
            a = np.asarray(_PILI.open(io.BytesIO(sh.image.blob)).convert("RGBA"))[..., 3]
            tot += float((a >= 128).mean()) * (sh.width / 914400.0) * (sh.height / 914400.0) / (W * H)
    return tot


def _stem_share(face, bold, size=150):
    """Stem width / ink height of a "1" in `face`, from the glyph itself (PIL): the widest ink run across the middle row."""
    from PIL import ImageDraw
    f = dk._pil_font(face, size, bool(bold))
    px = size * dk._MEAS_PREC
    im = _PILI.new("L", (int(px * 1.5), int(px * 1.6)), 0)
    ImageDraw.Draw(im).text((int(px * 0.2), int(px * 0.1)), "1", font=f, fill=255)
    bb = im.getbbox()
    mid, best, run = (bb[1] + bb[3]) // 2, 0, 0
    for x in range(im.width):
        run = run + 1 if im.getpixel((x, mid)) > 128 else 0
        best = max(best, run)
    return best / float(bb[3] - bb[1])


WC = {n: str(vl.ASSETS / "watercolour" / n) for n in ("harvest-basket.jpg", "rooftop-garden.jpg", "garden-tools.jpg", "balcony-watering.jpg")}
for g in vl.VARIANTS["storybook"]:
    for (W, H, tag) in ((13.333, 7.5, "landscape"), (7.5, 10.0, "portrait")):
        for n_, path in WC.items():
            prs, k = _deck_at("storybook", g, W, H)
            sl = k.new_slide()
            with contextlib.redirect_stdout(io.StringIO()):
                k.data(sl, number="1", label="pot is enough to begin", note="A window sill will do.", image=path)
            share = _art_share(sl, W, H)
            check(share >= 0.20, "storybook/{} data {} {}: the watercolour is {:.1%} of the page (>= 20%)".format(g, tag, n_, share))
    prs, k = _deck_at("storybook", g, 13.333, 7.5)
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.data(sl, number="1", label="pot is enough to begin", note="A window sill will do.", image=WC["harvest-basket.jpg"])
    num = next(sh for sh in _text_shapes(sl) if sh.text_frame.text.strip() == "1")
    r = _run_of(num)
    stem = _stem_share(r.font.name, r.font.bold)
    ground = vl.VARIANTS["storybook"][g]["palette"]["ground"]
    check(stem >= 0.15 and vl._contrast(_hexc(r.font.color.rgb), ground) >= 4.5,
          "storybook/{} data: the figure's stem is >= 15% of its height ({:.1%}, {} {}) and it reads {:.2f}:1 on the ground".format(
              g, stem, r.font.name, "bold" if r.font.bold else "regular", vl._contrast(_hexc(r.font.color.rgb), ground)))
dk.set_theme(None)

# ── Task 5 (A): storybook: the point dots are one colour that reads on the ground ──
# The dots cycled orange / green / yellow (the fourth repeating the first): colour carried no meaning, and the yellow dot
# was 1.05:1 on meadow. They are one colour — the language's accent ink — at >= 3:1 on either ground.
for g in vl.VARIANTS["storybook"]:
    ground = vl.VARIANTS["storybook"][g]["palette"]["ground"]
    for n in (2, 3, 4):
        for pic in (str(vl.ASSETS / "watercolour/garden-tools.jpg"), None):
            prs, k = _deck_at("storybook", g, 13.333, 7.5)
            sl = k.new_slide()
            with contextlib.redirect_stdout(io.StringIO()):
                k.points(sl, kicker="A small city garden", title="Things to bring", items=_PTS[:n], image=pic)
            dots = [sh for sh in sl.shapes if sh.shape_type == 1 and sh.fill.type == 1 and sh.width < 0.4 * 914400
                    and abs(sh.width - sh.height) < 2000]            # a painted dot: a small filled shape as wide as it is tall
            fills = [_hexc(sh.fill.fore_color.rgb) for sh in dots]
            worst = min([vl._contrast(f, ground) for f in fills] or [0])
            check(len(dots) == n and len(set(fills)) == 1 and worst >= 3.0,
                  "storybook/{} points x{}{}: {} dots, {} colour(s) {}, the faintest {:.2f}:1 on the ground (>= 3)".format(
                      g, n, "" if pic else " (no picture)", len(dots), len(set(fills)), sorted(set(fills)), worst))
dk.set_theme(None)

# ══ Task 9: the whole-branch review's findings, each pinned (2026-10-16) ══
import native_art as na, vl_contract as vc, vl_conformance as vcf, vl_specs, copy as _copy
_KO_Q = "우리는 함께 고칠 수 있습니다 그리고 이웃을 데려오세요 모든 것은 작은 시작에서 비롯됩니다"
_KO_T = "왜 모든 거리에는 수리의 밤이 필요한가"


def _criticals(prs):
    return sorted({f[2] for f in dk.lint_layout(prs, verbose=False) if f[1] == "CRITICAL"})


# (1) a Korean line the glyphs set on one line but lint reads as two is not imposed by the native stack either: the
# imposed paragraphs ran off the canvas / into the headline (clean before Phase 3)
for L, cv, page, kw in (("tally", (13.333, 7.5), "quote", dict(quote=_KO_Q, attribution="Repair cafe volunteer")),
                        ("interface", (10, 7.5), "quote", dict(quote=_KO_Q, attribution="Repair cafe volunteer")),
                        ("wayfinding", (10, 7.5), "quote", dict(quote=_KO_Q, attribution="Repair cafe volunteer")),
                        ("chalkboard", (13.333, 7.5), "cover", dict(title=_KO_T, kicker="Repair café"))):
    prs = dk.blank_deck(*cv)
    with contextlib.redirect_stdout(io.StringIO()):
        k = vl.use(L, prs)
        getattr(k, page)(k.new_slide(), **kw)
    found = _criticals(prs)
    check(not found, "{} {} {:.0f}x{}: a Korean {} is lint-clean ({})".format(L, page, cv[0], cv[1], page, found))
    dk.set_theme(None)

# (2) native_art.chip is a documented primitive: an ordinary deck's pill keeps its 1.9em height at every size; only a
# language's own card-sized pill asks for the roomier padding
prs = dk.blank_deck(13.333, 7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
for z in (12, 14, 20):
    h = na.chip(s, 0.8, 0.6, "Q3 REVIEW", size=z, fill="C6F432", ink="0E0E10", face="Arial")[3]
    check(abs(h - z / 72.0 * 1.9) < 1e-9, "na.chip at {}pt keeps its 1.9em height ({:.3f}in)".format(z, h))
h = na.chip(s, 0.8, 2.0, "Q3 REVIEW", size=20, fill="C6F432", ink="0E0E10", face="Arial", roomy=True)[3]
check(abs(h - 20 / 72.0 * 2.1) < 1e-9, "a roomy chip at 20pt gets half an em above and below ({:.3f}in)".format(h))

# (3) ink's column rule runs as long as the LONGER of the two columns it separates (a long middle column's words ran
# 2.5in past the rule on its left)
_LONG = ("敞开大门欢迎每一位邻居今晚到来", "不用预约也不收费每个人都可以带着坏掉的东西来带着修好的东西走")
for tag, items in (("short,LONG,short", [("共享", "一抽屉"), _LONG, ("就近", "十分钟")]),
                   ("LONG,short,short", [_LONG, ("共享", "一抽屉"), ("就近", "十分钟")]),
                   ("short,short,LONG", [("共享", "一抽屉"), ("就近", "十分钟"), _LONG])):
    prs, k = _deck("ink")
    sl = k.new_slide()
    with contextlib.redirect_stdout(io.StringIO()):
        k.points(sl, kicker="年度回顾", title="为什么每条街都需要修理之夜", items=items)
    # the vertical columns' own boxes (vcol sizes each to its ink; _ink_box reads vertical text as horizontal and
    # ends it 0.7in low), on the page — the off-canvas copy of the title is for the outline, not the reader
    texts = [(x_, y_, w_, h_) for x_, y_, w_, h_ in ((v / 914400.0 for v in (sh.left, sh.top, sh.width, sh.height))
             for sh in sl.shapes if getattr(sh, "has_text_frame", False) and len(sh.text_frame.text.strip()) > 1)
             if y_ >= 0]
    rules = [(x_, y_ + h_) for x_, y_, w_, h_ in ((v / 914400.0 for v in (sh.left, sh.top, sh.width, sh.height))
             for sh in sl.shapes if not (getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()))
             if w_ <= 0.06 and h_ >= 1.0]
    short = []
    for x_, bot in rules:
        near = [t for t in texts if abs(t[0] + t[2] / 2.0 - x_) < 1.4]          # the columns on either side of it
        low = max((t[1] + t[3] for t in near), default=0)
        if bot < low - 0.05:
            short.append("x={:.2f} ends {:.2f}in above its words".format(x_, low - bot))
    check(rules and not short, "ink points ({}): every column rule reaches the longer of its two columns ({})".format(
        tag, short or ("none short" if rules else "no rules found")))
dk.set_theme(None)

# (4) the author-facing CLI shows a language's distinctiveness problem (the message names "lang/ground", unquoted)
import check_vl_contract as cvc
_dm = vc.DISTINCT_MIN
vc.DISTINCT_MIN = 1000.0
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    cvc.main(["editorial"])
vc.DISTINCT_MIN = _dm
check("look alike" in buf.getvalue(), "check_vl_contract editorial prints editorial's look-alike pairs")

# (5) every Phase 3 check BITES on a fixture made to fail it, and passes the twin made to pass it
_S = vl_specs.load()
_twin = _copy.deepcopy(_S["ink"]); _twin["name"] = "inkcopy"
for g_ in _twin["grounds"]:
    _twin["grounds"][g_]["palette"]["ground"] = _S["editorial"]["grounds"]["light"]["palette"]["ground"]
    _twin["grounds"][g_]["palette"]["accents"] = list(_S["editorial"]["grounds"]["light"]["palette"]["accents"])
_twin["look"]["fonts"] = _copy.deepcopy(_S["editorial"]["look"]["fonts"])
_dp = vc.distinct_pairs({"editorial": _S["editorial"], "inkcopy": _twin})
check(any("editorial/light" in p and "inkcopy" in p for p in _dp), "distinct_pairs flags a language set in another's "
      "ground, accent and faces ({})".format(_dp[:1]))
check(vc.distinct_pairs(_S) == [], "distinct_pairs passes the fifteen shipped languages, starlit and tally-night "
      "among them ({})".format(vc.distinct_pairs(_S)[:1]))
_small = _copy.deepcopy(_S["journal"])
_small["type"]["kicker"] = (10,) + tuple(_small["type"]["kicker"][1:])
check(any("kicker" in p and "smallest tier" in p for p in vc.validate(_small)), "the type floor flags a 10pt kicker")
_low = _copy.deepcopy(_S["journal"])
_low["type"]["attribution"] = tuple(_low["type"]["attribution"][:5]) + (9,)
check(any("attribution" in p and "shrink" in p for p in vc.validate(_low)),
      "the type floor flags an attribution allowed to SHRINK to 9pt")
check(not vc.validate(_S["journal"]), "journal's own manifest passes ({})".format(vc.validate(_S["journal"])[:1]))


def _fixture(make):
    prs = dk.blank_deck(13.333, 7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    dk.box(s, 0, 0, 13.333, 7.5, fill="101830")
    make(s)
    return [("points", s)], prs


_real = vcf._built
try:
    from PIL import Image as _Im
    _white = str(WORK / "white.jpg"); _Im.new("RGB", (800, 600), (255, 255, 255)).save(_white)
    vcf._built = lambda *a, **kw: _fixture(lambda s: dk.box(s, 1.0, 1.0, 6.0, 4.0, fill="FFFFFF"))
    check(bool(vcf.bright_slabs("starlit")), "bright_slabs flags a drawn white slab on a dark ground")
    vcf._built = lambda *a, **kw: _fixture(lambda s: dk.picture(s, _white, 1.0, 1.0, 6.0, 4.0))
    check(not vcf.bright_slabs("starlit"), "bright_slabs leaves a user's white photograph alone")

    def _col(over):
        def make(s):                         # a column of eight lines, and a rule beside it ending `over` past its ink
            tb = dk.text(s, 1.0, 1.0, 3.0, 2.6, [[("A line of a column", 16, dk.RGBColor(255, 255, 255), False, False)]
                                                 for _ in range(8)])
            ink = vcf._ink_box(tb)
            dk.box(s, 4.2, 1.0, 0.02, ink[1] + ink[3] + over - 1.0, fill="FFFFFF")
        return make
    vcf._built = lambda *a, **kw: _fixture(_col(3.0))
    check(bool(vcf.rules_past_content("starlit")), "rules_past_content flags a rule running 3in past its words")
    vcf._built = lambda *a, **kw: _fixture(_col(0.1))
    check(not vcf.rules_past_content("starlit"), "rules_past_content passes a rule that ends with its words")

    # under a STAND-IN face the two line models can disagree by a line (broadsheet on CI's Linux faces: the layout's
    # box held a line lint's model did not) — one line of slack there, and a dangling rule is still caught
    def _sub(over):
        def make(s):
            tb = dk.text(s, 1.0, 1.0, 3.0, 2.6, [[("A line of a column", 16, dk.RGBColor(255, 255, 255), False, False,
                                                   "Zz Probe Nowhere 9")] for _ in range(8)])
            ink = vcf._ink_box(tb)
            dk.box(s, 4.2, 1.0, 0.02, ink[1] + ink[3] + over - 1.0, fill="FFFFFF")
        return make
    vcf._built = lambda *a, **kw: _fixture(_sub(0.55))
    check(not vcf.rules_past_content("starlit"), "rules_past_content allows a stand-in face one line of slack")
    vcf._built = lambda *a, **kw: _fixture(_sub(3.0))
    check(bool(vcf.rules_past_content("starlit")), "rules_past_content still flags a dangling rule under a stand-in face")

    # a page's own frame is not a column rule: topo's thin side bands of the neatline (0.052in, the page's height, at
    # the page edge) read as "a rule at x=0.21 runs 2.05in past its text" (2026-10-22); a full-height rule BETWEEN
    # columns is still a rule and still caught
    def _edge(x):
        def make(s):
            dk.text(s, 0.6, 1.0, 3.0, 2.6, [[("A line of a column", 16, dk.RGBColor(255, 255, 255), False, False)]
                                             for _ in range(4)])
            dk.box(s, x, 0.25, 0.05, 7.0, fill="FFFFFF")
        return make
    vcf._built = lambda *a, **kw: _fixture(_edge(0.21))
    check(not vcf.rules_past_content("starlit"), "rules_past_content leaves a page-edge frame band alone: {}".format(
        vcf.rules_past_content("starlit")))
    vcf._built = lambda *a, **kw: _fixture(_edge(3.9))
    check(bool(vcf.rules_past_content("starlit")), "rules_past_content still flags a full-height rule beside a column")
finally:
    vcf._built = _real
dk.set_theme(None)

# (6) a very long quote on soft keeps its tucked mark clear of the first line (90 and 120 words were clean before)
_W = ("Every repair we made that evening was a small argument for keeping things, and the room kept making it until "
      "the last lamp came on and the tables were bare again.").split()
for n in (90, 120):
    prs, k = _deck("soft")
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            k.quote(k.new_slide(), quote=" ".join((_W * 5)[:n]) + ".", attribution="Mara Jensen, volunteer lead")
        found = _criticals(prs)
    except vl.VLTextOverflow:
        found = []                                          # refused loudly is an allowed answer
    check(not found, "soft: a {}-word quote is lint-clean or refused ({})".format(n, found))
dk.set_theme(None)

# (7) the smallest tier holds when a page SHRINKS: a 60-word quote's attribution stays >= 12pt (it fell to 10.3pt)
for L in vl.LANGS:
    prs, k = _deck(L)
    sl = k.new_slide()
    pic = str(vl.ASSETS / "photo/hall-repair.jpg") if "quote" in vl.REGISTRY[L]["image_pages"] else None
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            k.quote(sl, quote=" ".join((_W * 5)[:60]) + ".", attribution="Mara Jensen, volunteer lead",
                    **({"image": pic} if pic else {}))
    except vl.VLTextOverflow:
        continue
    z = min((r.font.size.pt for sh in sl.shapes if getattr(sh, "has_text_frame", False)
             and "jensen" in sh.text_frame.text.lower() for p in sh.text_frame.paragraphs for r in p.runs
             if r.font.size and r.text.strip()), default=0)
    check(z >= 12 - 0.05, "{}: a 60-word quote's attribution stays >= 12pt ({:.1f}pt)".format(L, z))
dk.set_theme(None)

# (8) lint's ORPHANED PUNCTUATION is the 避头尾 rule: a renderer breaks Latin only at spaces, so a full stop attached to
# its word cannot open a line. The char-level estimate put a lone "." on its own line for Latin columns whenever the
# widths happened to fill up to it (cutpaper and broadsheet points on CI's Linux faces; a fixer side-stepped it with
# "very long copy" before — 2026-10-16). In CJK a closing mark cannot open a line either — LibreOffice hangs one past a
# full line or takes the character before it down (measured 2026-10-22, lint_deck._HANG_MARKS) — but an ellipsis can,
# and the render leaves it alone: that is the CJK orphan the lint must still flag.
import lint_deck as _ld


def _orphan_width(text, size, want, face):
    """A box width at which lint's char-level estimate leaves exactly `want` on the last line, or None."""
    for i in range(4000, 300, -1):
        w = i / 1000.0
        if (_ld._est_lines([[(text, size)]], w, face, False) >= 2
                and _ld._last_line([[(text, size)]], w, face, False).strip() == want):
            return w
    return None


def _orphan_findings(text, size, w, face):
    prs = dk.blank_deck(13.333, 7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    dk.text(sl, 1.0, 1.0, w, 3.0, [[(text, size, dk.RGBColor(0x22, 0x22, 0x22), False, False, face)]])
    p = WORK / "orphan.pptx"; prs.save(str(p)); buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
        _ld.lint(str(p), static_ok=True)
    return [l for l in buf.getvalue().splitlines() if "ORPHANED PUNCTUATION" in l]


for _txt, _mark, _want, _face in (("Anything that switches on, unzips or wobbles is welcome.", ".", False, "Arial"),
                                  ("我们一起把坏掉的东西修好，然后带着它们回家…", "…", True, "PingFang SC")):
    _w = _orphan_width(_txt, 16, _mark, _face)
    if _w is None:
        check(False, "lint fixture: a width that orphans {!r} exists for {!r}".format(_mark, _txt[:16]))
        continue
    _f = _orphan_findings(_txt, 16, _w, _face)
    check(bool(_f) == _want, "lint {} a lone {!r} on {}'s last line ({})".format(
        "flags" if _want else "does not flag", _mark, "CJK" if _want else "a Latin column", _f[:1]))
_w = _orphan_width("我们一起把坏掉的东西修好，然后带着它们回家。", 16, "。", "PingFang SC")
check(_w is None, "lint's line model never leaves a CJK '。' alone on a line — it hangs or takes a character down, "
                  "as the render does (alone at {}in)".format(_w))
dk.set_theme(None)

# (9) broadsheet's cover plans its headline in what the standfirst and the Inside list leave it: a headline that WRAPS
# before it shrinks (Task 2) kept a mid-length title tall enough on a 10in page to leave the Inside list a 0.72in column
# (refused), and CI's wider faces refused the sigs.py scaffold's own cover (smoke, GitHub run 38890026577, 2026-10-16)
for _t in ("Every street needs a night for fixing things",
           "Every street needs a night for fixing the things we were told to throw away",
           "Every street in the city needs one evening a month for fixing the things we were told to throw away"):
    for _cv in ((10, 5.625), (13.333, 7.5), (7.5, 10)):
        _prs = dk.blank_deck(*_cv)
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                _k = vl.use("broadsheet", _prs)
                _k.cover(_k.new_slide(), title=_t, subtitle="How a repair evening works, and why it spreads.",
                         masthead="The Repair Weekly", edition="No. 12",
                         inside=["The idea", "The evening", "The numbers", "Next month"])
            _found = _hard(_prs, "bscov")
        except vl.VLTextOverflow as _e:
            _found = ["refused: " + str(_e)[:90]]
        check(not _found, "broadsheet cover ({} chars, {:g}x{:g}): the headline leaves the Inside list its room ({})".format(
            len(_t), _cv[0], _cv[1], _found[:1]))
        dk.set_theme(None)

# (10) the number gutter's 10% holds AFTER the pptx stores it: a width is written in whole EMU, truncated, so a box sized
# to exactly 1.10x its digits came back 1.0999x (GitHub run 38892848519, Linux faces, where the slack term beats the
# floor) — the floor is 0 here so the slack term is what is measured, on every machine
import display_type as _dtx
for L in ("editorial", "soft", "collage"):
    prs, k = _deck(L)
    face, bold = k.face("numeral"), bool(vl.TYPE[L]["item_no"][2])
    short = []
    for z10 in range(200, 420, 7):
        z = z10 / 10.0
        need = max(dk._natural_width_in([("01", bold)], z, face), _dtx._glyph_width("01", z, face, bold) or 0.0)
        g = vl._number_gutter(k, ["01"], z, 0.0)
        stored = int(g * 914400) / 914400.0                # what python-pptx writes: Inches() truncates to whole EMU
        if (stored - dk.TEXT_INSET_LR) / need < vl.NUMBER_SLACK:
            short.append(z)
    check(not short, "{}: the number gutter keeps its {:.0%} once stored in whole EMU (short at {})".format(
        L, vl.NUMBER_SLACK - 1, short[:4]))
    dk.set_theme(None)

# ── a NO-BREAK SPACE holds its two words together (iPhone 18 explainer, 2026-10-17) ─────────────────────────────
# The cover title "iPhone<NBSP>18, explained" was set "iPhone / 18, explained": the widow balancer split words with
# str.split(), which treats U+00A0 as whitespace, chose a break INSIDE the glued pair and joined the lines back with
# an ordinary space. Every language's cover, and the image-led ones, must keep the pair on one line.
def _title_paras(prs):
    out = []
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame and "explained" in sh.text_frame.text and sh.top >= 0:
            out.append([p.text for p in sh.text_frame.paragraphs])
    return out

_nb_bad, _nb_built = [], 0
for L in vl.LANGS:
    import formats as _fmt
    _p = _fmt.blank_deck("wide13")
    _k = vl.use(L, _p, ground="light")
    _img = str(ROOT / "assets" / "vl" / "photo" / "hall-repair.jpg")
    try:
        _k.cover(_k.new_slide(), kicker="An explainer", title="iPhone 18, explained",
                 subtitle="What you can buy now", image=_img if L in vl.IMAGE_LED else None)
    except Exception as _e:                      # a refusal is not this test's question
        continue
    finally:
        dk.set_theme(None)
    _nb_built += 1
    for paras in _title_paras(_p):
        if any(t.strip() in ("iPhone", "iPhone,") or t.strip().startswith("18") for t in paras):
            _nb_bad.append("{}: {}".format(L, paras))
check(_nb_built >= 12 and not _nb_bad, "a no-break space keeps 'iPhone 18' on one line on every language's "
      "cover ({} built; {})".format(_nb_built, _nb_bad[:3]))
import random as _rnd
_rng = _rnd.Random(7)
_alpha = "abc XYZ 12,.\t\n　 —é"
_diff = [t for t in ("".join(_rng.choice(_alpha) for _ in range(_rng.randint(0, 40))) for _ in range(500))
         if vl._words(t) != t.split()]
check(not _diff, "without a no-break space, _words() splits exactly as str.split() did (500 generated strings; "
                 "differs on {})".format([repr(t) for t in _diff[:2]]))
check(vl._words("a b c d e f") == ["a b", "c d", "e f"],
      "_words() keeps U+00A0, U+202F and U+2007 inside the word they glue")

# ── credit= on every page that carries a picture (iPhone 18 explainer, 2026-10-17) ─────────────────────────────
# Only image_text had a caption; a CC BY photo on a cover, a points page or a data page had nowhere lawful to carry
# its credit, so a real deck squeezed source lines into page corners by hand (and one landed in the footer band of a
# full-bleed cover). credit= is set by the page itself, on the picture it credits, at the language's caption size.
_photo = str(ROOT / "assets" / "vl" / "photo" / "hall-repair.jpg")
_cred_bad, _cred_cases = [], 0
# a Chinese credit too: its chip was measured in the Latin face (~3.1in for ~4.7in of 1-em characters), so the line
# wrapped and hung out of a one-line chip (review of 475f816)
for L, _CREDIT in [(L_, c_) for L_ in vl.LANGS for c_ in ("Photo: A. Example, CC BY 4.0",
                                                           "摄影：张三，CC BY 4.0 授权，来源：维基共享资源")]:
    _pages = vl.PAGE_FIELDS if L in vl.IMAGE_LED else vl.NATIVE_IMAGE_PAGES.get(L, ())
    import formats as _fmt
    _p = _fmt.blank_deck("wide13")
    _k = vl.use(L, _p, ground="light")
    _built = []
    _copy = {"cover": dict(kicker="Kicker", title="A short title", subtitle="A subtitle"),
             "section": dict(number="02", kicker="Kicker", title="A section"),
             "image_text": dict(kicker="Kicker", title="A title", body="One sentence of body copy."),
             "quote": dict(quote="A short quote that fits.", attribution="Someone"),
             "data": dict(number="42", label="a label", note="A note."),
             "closing": dict(title="Thank you", line="One line."),
             "points": dict(kicker="Kicker", title="Three points", items=[("One", "a line"), ("Two", "a line")])}
    try:
        for pg in _pages:
            s = _k.new_slide()
            try:
                getattr(_k, pg)(s, image=_photo, credit=_CREDIT, **_copy[pg])
            except (vl.VLTextOverflow, ValueError) as e:
                if "credit" in str(e).lower():
                    _cred_bad.append("{}.{}: {}".format(L, pg, str(e)[:90]))
                continue
            except TypeError as e:
                _cred_bad.append("{}.{}: {}".format(L, pg, str(e)[:90]))
                continue
            _cred_cases += 1
            _built.append((len(_p.slides), pg))
            W_, H_ = _p.slide_width, _p.slide_height
            hits = [sh for sh in s.shapes if sh.has_text_frame and _CREDIT in sh.text_frame.text]
            if not hits:
                _cred_bad.append("{}.{}: no credit text on the page".format(L, pg))
                continue
            c = hits[0]
            _sz = max((r.font.size.pt for p_ in c.text_frame.paragraphs for r in p_.runs if r.font.size), default=12)
            _need = dk.measure_text([(_CREDIT, False)], c.width / 914400.0, _sz,   # it subtracts the insets itself
                                    font=_k.ea_face("body", _CREDIT) if dk._has_cjk(_CREDIT) else _k.face("body"))
            if _need > c.height / 914400.0 + 0.02:
                _cred_bad.append("{}.{}: the credit needs {:.2f}in and its box is {:.2f}in ({!r})".format(
                    L, pg, _need, c.height / 914400.0, _CREDIT[:12]))
            if c.left < 0 or c.top < 0 or c.left + c.width > W_ or c.top + c.height > H_:
                _cred_bad.append("{}.{}: the credit leaves the canvas".format(L, pg))
            sizes = [r.font.size.pt for p_ in c.text_frame.paragraphs for r in p_.runs if r.font.size is not None]
            if not sizes or min(sizes) < 12:
                _cred_bad.append("{}.{}: the credit is set under 12pt ({})".format(L, pg, sizes))
            pics = [sh for sh in s.shapes if sh.shape_type == 13]
            inside = any(c.left >= p_.left - 9144 and c.top >= p_.top - 9144 and c.left + c.width <= p_.left + p_.width + 9144
                         and c.top + c.height <= p_.top + p_.height + 9144 for p_ in pics)
            if not inside:
                _cred_bad.append("{}.{}: the credit is not set on the picture it credits".format(L, pg))
        if _built:
            for f in _hard(_p, "credit_{}_{}".format(L, "cjk" if dk._has_cjk(_CREDIT) else "lat")):
                _cred_bad.append("{}: {}".format(L, f))
    finally:
        dk.set_theme(None)
check(_cred_cases >= 60 and not _cred_bad,
      "credit= sets a 12pt+ credit on the picture, inside the canvas, lint-clean, on {} picture page(s) in every "
      "language ({})".format(_cred_cases, _cred_bad[:4]))
_p = _fmt.blank_deck("square")
_k = vl.use("storybook", _p, ground="light")
try:
    _k.points(_k.new_slide(), kicker="Kicker", title="Points", items=[("One", "a"), ("Two", "b")], image=_photo,
              credit="Photo: " + "a very long attribution line " * 5)
    check(False, "a credit too long for its picture is refused (it was set across the page)")
except ValueError as e:
    check("credit" in str(e).lower() and "shorten" in str(e).lower(),
          "a credit too long for its picture is refused, asking to shorten it ({})".format(str(e)[:80]))
finally:
    dk.set_theme(None)
_p = _fmt.blank_deck("wide13")
_k = vl.use("editorial", _p, ground="light")
try:
    _k.data(_k.new_slide(), number="42", label="a label", credit="Photo: A. Example, CC BY 4.0")
    check(False, "credit= without a picture is refused, naming the reason (it was accepted)")
except (ValueError, TypeError) as e:
    check("unknown field" not in str(e) and "picture" in str(e).lower(),
          "credit= without a picture is refused, naming the reason ({})".format(str(e)[:80]))
finally:
    dk.set_theme(None)

# a balanced (narrower) measure must still hold its widest word: a greedy break lets an over-wide word overflow its
# line, so the line count still matched, and a 4.33in column set "Everything starts at the table" in 3.12in —
# LibreOffice broke it "Everythin / g" (the herbarium review, 2026-10-22). Georgia ships on every machine.
_kb = vl.use("ink", dk.blank_deck(13.333, 7.5))
for _tt, _ww in (("Everything starts at the table", 4.33), ("Neighbourhood repair evenings work", 4.6),
                 ("Extraordinary things happen here", 4.4)):
    _wb2 = vl._balanced_width(_kb, "title", _tt, 50, _ww)
    check(_wb2 is None or vl._widest_word_fits(_kb, "title", _tt, 50, _wb2),
          "a balanced measure holds the widest word of {!r} ({})".format(_tt, None if _wb2 is None else round(_wb2, 2)))

# align="r" sets the words to the right of their column (receipt's total, 2026-10-22): the box ends at the column's
# right edge and the paragraph is right-aligned — before, "r" raised a KeyError at draw time
_kr = vl.use("tally", dk.blank_deck(13.333, 7.5))
_sr = _kr.new_slide()
_rr, _dr = vl._flow(_kr, _sr, "data", (2.0, 1.0, 6.0, 3.0), [("label", "a total")], anchor="top", align="r")
try:
    _dr()
    _tb = [sh for sh in _sr.shapes if sh.has_text_frame and sh.text_frame.text == "a total"]
    from pptx.enum.text import PP_ALIGN as _PA
    check(_tb and _tb[0].text_frame.paragraphs[0].alignment == _PA.RIGHT
          and abs(_rr["label"][0] + _rr["label"][2] - 8.0) < 0.02,
          "align='r' right-aligns the words and ends the box at the column's right edge ({})".format(_rr.get("label")))
except KeyError as _e:
    check(False, "align='r' is accepted by the flow (KeyError {})".format(_e))

for line in ok:
    print("  ok   " + line)
for line in skipped:
    print("  SKIP " + line + " (the languages' real faces are not installed here)")
for line in bad:
    print("  FAIL " + line)
print("\n{} passed, {} failed".format(len(ok), len(bad)))
sys.exit(1 if bad else 0)
