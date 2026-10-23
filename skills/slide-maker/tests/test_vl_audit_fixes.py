#!/usr/bin/env python3
"""The 13 visual languages after the 2026-10-09 audit: every finding the audit reproduced is pinned here — text measured
with room to spare and by script (Latin words never split), no caller's words or pictures dropped, photos upright,
the kit's own pages clear of its own gates, and the image-led languages level with the drawn ones."""
from __future__ import annotations
import contextlib, io, re, sys, tempfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import deckkit as dk
import visual_languages as vl

ok, bad = [], []
def check(cond, why):
    (ok if cond else bad).append(why)
EMU = 914400.0
td = Path(tempfile.mkdtemp())
PHOTO = str(vl.ASSETS / "photo" / "hall-repair.jpg")


def use(name, W=13.333, H=7.5, ground="light"):
    prs = dk.blank_deck(W, H)
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        k = vl.use(name, prs, ground=ground)
    return prs, k


def texts(s):
    return [sh for sh in s.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip() and sh.top >= 0]


def img_for(name):
    return str(vl.ASSETS / "watercolour" / "rooftop-garden.jpg") if name == "storybook" else PHOTO


# ── A1: measured with room to spare, by script ──
# A line measured to fit its box exactly was set one line longer by LibreOffice: collage on 9:16 measured "fix what it
# already owns" at 4.894in in a 4.894in box and rendered "owns" on a fourth line, over the subtitle (audit sweep,
# 2026-10-09). Every language now measures its words in HEADROOM of the column and draws the box at the full width.
check(0.95 <= getattr(vl, "HEADROOM", 1.0) < 1.0, "the kit measures display text with headroom (vl.HEADROOM)")
TITLE = "Why every street deserves a place to fix what it already owns"
for name in vl.LANGS:
    for W, H in ((5.625, 10.0), (7.5, 10.0), (10.0, 5.625)):
        prs, k = use(name, W, H)
        s = k.new_slide()
        k.cover(s, kicker="Annual review of the repair network", title=TITLE,
                subtitle="What twelve months of evenings, benches and borrowed tools taught us",
                image=img_for(name) if name in vl.IMAGE_LED else None)
        hd = [sh for sh in texts(s) if sh.text_frame.text.replace("\n", " ").replace("\v", " ").lower().startswith("why every")]
        if not hd:
            check(False, "{} {}x{}: the cover title is drawn".format(name, W, H)); continue
        sh = hd[0]
        tf = sh.text_frame
        inner = sh.width / EMU - (tf.margin_left or 91440) / EMU - (tf.margin_right or 91440) / EMU
        sz = tf.paragraphs[0].runs[0].font.size.pt
        if len(tf.paragraphs) > 1:        # set in its measured lines: no line may come near the box's inner width
            widest = max(dk._natural_width_in([(p_.text, bool(p_.runs[0].font.bold))], sz, p_.runs[0].font.name)
                         for p_ in tf.paragraphs if p_.runs)
            check(widest <= inner * 0.985, "{} {}x{}: every title line leaves room in its box ({:.3f}in of {:.3f}in)".format(
                name, W, H, widest, inner))
        else:                             # one paragraph: the box holds the lines the title needs at 98.5% of its width
            need = len(vl._break_lines(k, "title", tf.text, sz, inner * 0.985 + vl._INSET))
            held = round((sh.height / EMU - 0.06) / (sz / 72.0 * dk._LINT_LINE_H))
            check(need <= held, "{} {}x{}: the title box holds its lines with room to spare (needs {}, holds {})".format(
                name, W, H, need, held))

# mixed CJK + Latin: a Latin word is one unit, measured in the Latin face; breaks fall between CJK characters, at the
# script boundary or at a space — never inside "report" (journal set "…进展re / port v2", the robustness audit)
MIXED = ("2026年第3季度MRI重建项目进展report v2", "GPT-4o与Claude在代码生成上的对比", "用ResNet50做图像分类的三个教训")
def word_cuts(t, ls):
    """Line boundaries that fall between two Latin letters/digits that were ADJACENT in `t` (a break at a space is a
    break between words)."""
    cuts, pos = [], 0
    for a, b in zip(ls, ls[1:]):
        pos = t.index(a.strip(), pos) + len(a.strip())
        if a and b and re.match(r"[A-Za-z0-9]", a.strip()[-1]) and re.match(r"[A-Za-z0-9]", b.strip()[0]) and t[pos:pos + 1] != " ":
            cuts.append((a, b))
    return cuts


for name in ("journal", "tally", "starlit", "editorial", "ink"):
    prs, k = use(name, 10.0, 5.625)
    for t in MIXED:
        for wi in (2.0, 2.6, 3.3, 4.1, 5.0):
            ls = vl._break_lines(k, "title", t, 30, wi)
            split = word_cuts(t, ls)
            check(not split and "".join(ls).replace(" ", "") == t.replace(" ", ""),
                  "{}: {!r} at {}in breaks only between words ({})".format(name, t, wi, ls))
for name in ("journal", "tally", "starlit", "broadsheet", "chalkboard"):
    prs, k = use(name, 10.0, 5.625)
    s = k.new_slide()
    k.cover(s, title=MIXED[0], subtitle="面向临床的加速方案")
    paras = [p_.text for sh in texts(s) for p_ in sh.text_frame.paragraphs if "MRI" in sh.text_frame.text or "port" in sh.text_frame.text]
    cut = word_cuts(MIXED[0], paras) if paras else []
    check(not cut, "{}: a mixed title is never split inside a Latin word ({})".format(name, paras))
# the Latin parts of a mixed line are measured in the Latin face they render in: a mixed line measures at least as wide
# as its Latin and CJK parts measured apart
prs, k = use("journal")
t = "Claude与GPT-4o"
w_mix = vl._line_width(k, "title", t, 30) if hasattr(vl, "_line_width") else 0.0
w_lat = dk._natural_width_in([("Claude", False)], 30, k.face("display")) + dk._natural_width_in([("GPT-4o", False)], 30, k.face("display"))
check(w_mix >= w_lat + 0.9 * 30 / 72.0, "a mixed line is measured per script ({:.2f}in for Latin {:.2f}in + one CJK em)".format(w_mix, w_lat))

# ── A2: a caller's words are never dropped because a sibling field is absent ──
# poster drew a data note, a quote's attribution and a closing line only inside the primary field's `if`: number+note
# lost the note, attribution alone drew a lone “, line alone a blank colour block (robustness audit, 2026-10-09)
SUBSETS = [("data", dict(number="12", note="Weighed at the door.")), ("data", dict(label="evenings a month", note="Since spring.")),
           ("quote", dict(quote="Fix it together.")), ("quote", dict(quote="Fix it together.", attribution="A volunteer")),
           ("quote", dict(attribution="A volunteer")),
           ("closing", dict(line="Bring a neighbour.")), ("closing", dict(title="See you next week")),
           ("section", dict(title="The evening")), ("section", dict(number="02", title="The evening")),
           ("cover", dict(title="Bring it broken")), ("cover", dict(subtitle="How an evening works", title="Bring it broken"))]
for name in vl.LANGS:
    for page, kw in SUBSETS:
        prs, k = use(name)
        s = k.new_slide()
        try:
            getattr(k, page)(s, **kw)
        except (vl.VLTextOverflow, ValueError) as e:
            check("needs" in str(e) or "give" in str(e) or "empty" in str(e),
                  "{} {} {}: a refusal says what is missing ({})".format(name, page, sorted(kw), str(e)[:80]))
            continue
        on = " ".join(sh.text_frame.text for sh in texts(s)).replace("\n", " ").replace("\v", " ").lower()
        on = re.sub(r"\s+", " ", on)
        lost = [v for v in kw.values() if v.lower() not in on and v.upper() not in on.upper()
                and not (name == "collage" and vl._outlinable(k, v))]      # collage draws a short figure as an outline
        check(not lost, "{} {} {}: every given word is on the page (lost {})".format(name, page, sorted(kw), lost))

# ── A3: a drawn language never drops a picture silently ──
# image= on a drawn language's cover/section/quote/data/closing/points built with no picture and no error — even a path
# to no file (robustness + docs audits, 2026-10-09). It is refused there, naming the pages that do draw one.
F = dict(cover=dict(title="T"), section=dict(title="T"), points=dict(title="T", items=[("a", "b"), ("c", "d")]),
         quote=dict(quote="Q"), data=dict(number="3", label="l"), closing=dict(title="T"))
check(set(getattr(vl, "NATIVE_IMAGE_PAGES", {})) == set(vl.NATIVE), "vl.NATIVE_IMAGE_PAGES names the picture pages of every drawn language")
for name in vl.NATIVE:
    pages = getattr(vl, "NATIVE_IMAGE_PAGES", {}).get(name, ("image_text",))
    for page, kw in F.items():
        kw = dict(kw)
        if page == "points" and name in ("tally", "broadsheet"):
            kw["tags"] = ["a", "b"]
        for im in (PHOTO, "/no/such/photo.jpg"):
            prs, k = use(name)
            s = k.new_slide()
            try:
                getattr(k, page)(s, image=im, **kw)
                drew = any(sh.shape_type == 13 for sh in s.shapes)
                check(page in pages and drew, "{} {}: image={} is drawn or refused (built, picture drawn: {})".format(
                    name, page, Path(im).name, drew))
            except (ValueError, FileNotFoundError) as e:
                check(page not in pages or im != PHOTO, "{} {}: a picture it draws is accepted ({})".format(name, page, str(e)[:60]))
                if page not in pages:
                    check("image_text" in str(e), "{} {}: the refusal names the pages that draw a picture ({})".format(
                        name, page, str(e)[:90]))
    for im in ([], [PHOTO, PHOTO]):                      # one picture per page: an empty or long list is refused by name
        prs, k = use(name)
        try:
            k.image_text(k.new_slide(), title="T", body="B", image=im)
            check(False, "{} image_text: image={} is refused".format(name, "[]" if not im else "two pictures"))
        except ValueError as e:
            check("image" in str(e), "{} image_text: image={} is refused by name ({})".format(name, len(im), str(e)[:60]))
        except Exception as e:
            check(False, "{} image_text: image={} raises a plain error, not {}".format(name, len(im), type(e).__name__))
    prs, k = use(name)
    s = k.new_slide(); k.image_text(s, title="T", body="B", image=[PHOTO])
    check(any(sh.shape_type == 13 for sh in s.shapes), "{} image_text: a one-picture list is drawn".format(name))

# ── A4 + A7: every picture arrives upright, readable and in a format PowerPoint embeds ──
from PIL import Image
import numpy as np
src = Image.open(PHOTO).convert("RGB")
portrait = src.crop((0, 0, int(src.height * 0.66), src.height))            # an upright portrait photo …
stored = portrait.rotate(90, expand=True)                                    # … stored on its side, as phones do
ex = Image.Exif(); ex[0x0112] = 6
EXIF6 = td / "phone.jpg"; stored.save(EXIF6, exif=ex.tobytes())
G16 = td / "figure16.png"
Image.fromarray((np.linspace(0, 65535, 500 * 400).reshape(400, 500)).astype(np.uint16)).save(G16)
WEBP = td / "photo.webp"; src.save(WEBP)
TRUNC = td / "cut.jpg"; TRUNC.write_bytes(Path(PHOTO).read_bytes()[:4000])
NOTIMG = td / "notes.jpg"; NOTIMG.write_text("not a picture")


def placed(s, before):
    """The pictures a page added (blob decoded), the ground's own picture(s) left out."""
    out = []
    for sh in s.shapes:
        if sh.shape_type == 13 and sh.image.sha1 not in before:
            out.append(Image.open(io.BytesIO(sh.image.blob)))
    return out


for name in vl.LANGS:
    prs, k = use(name)
    s0 = k.new_slide(); k.image_text(s0, title="Before", body="B", image=PHOTO)
    ground = {sh.image.sha1 for sh in s0.shapes if sh.shape_type == 13 and Image.open(io.BytesIO(sh.image.blob)).size != src.size}
    s = k.new_slide(); k.image_text(s, title="A photo from a phone", body="Taken upright.", image=str(EXIF6))
    pics = [im for im in placed(s, ground) if im.size[0] > 50]
    up = [im for im in pics if im.size[1] > im.size[0] and int(im.getexif().get(274, 1) or 1) == 1]
    check(bool(up), "{}: a phone photo (EXIF orientation 6) is placed upright ({})".format(name, [im.size for im in pics]))
    s = k.new_slide(); k.image_text(s, title="A 16-bit figure", body="From the scanner.", image=str(G16))
    flat = [np.asarray(im.convert("L")).std() for im in placed(s, ground) if im.size[0] > 50]
    check(flat and max(flat) > 30, "{}: a 16-bit PNG keeps its tones (std {})".format(name, [round(x) for x in flat]))
    s = k.new_slide()
    try:
        k.image_text(s, title="A WebP", body="From the web.", image=str(WEBP))
        check(any(sh.shape_type == 13 for sh in s.shapes), "{}: a WebP photo is placed".format(name))
    except Exception as e:
        check(False, "{}: a WebP photo is placed, not refused ({}: {})".format(name, type(e).__name__, str(e)[:60]))
    for bad_, why in ((TRUNC, "truncated"), (NOTIMG, "not an image")):
        try:
            k.image_text(k.new_slide(), title="Bad", body="B", image=str(bad_))
            check(False, "{}: a {} file is refused".format(name, why))
        except ValueError as e:
            check("image_text" in str(e) and bad_.name in str(e), "{}: a {} file is refused naming the page and file ({})".format(
                name, why, str(e)[:90]))
        except Exception as e:
            check(False, "{}: a {} file raises a plain refusal, not {} ({})".format(name, why, type(e).__name__, str(e)[:60]))
check(int(Image.open(EXIF6).getexif().get(274, 1)) == 6, "the caller's own file is never rewritten")

# ── A5: a deck built from a language's own pages passes that language's own prohibitions ──
# starlit's data page drew two oversized primitives (the glow behind the figure, the horizon) and its own CONFETTI guard
# blocked every landscape starlit deck at hand-off (docs audit, 2026-10-09) — its declared atmosphere is not a hero shape
import check_visual_language as cvl
FULL = [("cover", dict(kicker="Annual review", title="Every street needs a night for fixing things", subtitle="How it works")),
        ("section", dict(number="02", title="The evening")),
        ("image_text", dict(title="Tools on every bench", body="Volunteers sit beside you.", caption="The hall")),
        ("points", dict(title="How it works", items=[("Bring it", "Anything that switches on."), ("Fix it", "Your hands do the work.")])),
        ("quote", dict(quote="The visitor holds the screwdriver.", attribution="A volunteer")),
        ("data", dict(number="12", label="evenings this year", note="Counted at the door.")),
        ("closing", dict(title="See you next week", line="Bring a neighbour."))]
# one deck of each language's own pages per canvas, the two grounds taking turns (both meet a canvas), with the words
# only the caller can give — built once, read by the prohibitions here and by the delivery lint (B2)
DX = {"tally": {"points": {"tags": ["A", "B"]}, "data": {"total": "20"}}, "broadsheet": {"points": {"tags": ["A", "B"]}},
      "chalkboard": {"cover": {"doodle": "lucide:wrench"}, "closing": {"doodle": "lucide:wrench"}, "points": {"ordered": True}},
      "ink": {"cover": {"seal": "修"}}}
OWN = []
for name in vl.LANGS:
    grounds = list(vl.VARIANTS[name])
    for ci, (W, H) in enumerate(((13.333, 7.5), (7.5, 10.0))):
        ground = grounds[ci % len(grounds)]
        prs, k = use(name, W, H, ground)
        for page, kw in FULL:
            kw = dict(kw, **DX.get(name, {}).get(page, {}))
            if page == "image_text" or (name in vl.IMAGE_LED and page in ("cover", "closing")):
                kw["image"] = img_for(name)
            getattr(k, page)(k.new_slide(), **kw)
        path = td / "own_{}_{}_{}.pptx".format(name, ground, int(W))
        prs.save(str(path))
        OWN.append((name, ground, W, H, path))
        with contextlib.redirect_stdout(io.StringIO()):
            found, _f = cvl.check(str(path), {"name": name, "fonts": "both", "ground": ground})
        forb = [x for x in found if "FORBIDDEN" in str(x) or "forbids" in str(x)]
        check(not forb, "{} {} {}x{}: a deck of its own pages passes its own prohibitions ({})".format(
            name, ground, W, H, [str(x)[:90] for x in forb[:1]]))

# ── A6: a picture stays visible on a dark ground ──
# ink's night ground multiplied the picture by the dark/light paper ratio (~0.13) and the photo all but vanished (both
# audits, 2026-10-09); the paper tint is for moving a watercolour's paper to a LIGHT or mid ground (storybook meadow)
src_mean = np.asarray(Image.open(PHOTO).convert("L")).mean()
for name in ("ink", "storybook"):
    for ground in vl.VARIANTS[name]:
        prs, k = use(name, ground=ground)
        s = k.new_slide(); k.image_text(s, title="A photo", body="B", image=PHOTO)
        pics = [Image.open(io.BytesIO(sh.image.blob)).convert("RGBA") for sh in s.shapes if sh.shape_type == 13]
        pics = [im for im in pics if im.size[0] > 200]
        if not pics:
            check(False, "{} {}: the photo is placed".format(name, ground)); continue
        a = np.asarray(pics[0]).astype(float)
        core = a[a.shape[0] // 4: 3 * a.shape[0] // 4, a.shape[1] // 4: 3 * a.shape[1] // 4]
        lum = (0.299 * core[..., 0] + 0.587 * core[..., 1] + 0.114 * core[..., 2]).mean()
        check(lum >= 0.55 * src_mean, "{} {}: the photo keeps its light ({:.0f} of the source's {:.0f})".format(name, ground, lum, src_mean))

# ── B4: the palette --gates prints is one the register-pixels gate can find on the pages ──
# it declared the TEXT accents (a kicker's red, a hairline's gold): broadsheet, soft dusk, ink night and cutpaper night
# failed DECLARED HUES ABSENT on the bundled samples themselves (docs audit, 2026-10-09). The hues a language really
# paints are measured on its rendered sample pages (assets/vl/samples/hues.json, written by --sample-sheet).
import json as _json
import check_register_pixels as crp
HF = vl.ASSETS / "samples" / "hues.json"
hues = _json.loads(HF.read_text(encoding="utf-8")) if HF.exists() else {}
for name in vl.LANGS:
    for g in vl.VARIANTS[name]:
        stem = vl._sample_stem(name, g)
        P = vl.VARIANTS[name][g]["palette"]
        mine = {c.upper() for c in list(P["accents"]) + list(P["text_accents"]) + [P["ground"], P["ink"]]}
        rec = hues.get(stem)
        check(isinstance(rec, list) and all(h.lstrip("#").upper() in mine for h in rec),
              "{}: its painted hues are recorded from its sample pages, each one of its palette ({})".format(stem, rec))
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            vl._print_gates(name, str(td / "gates_{}".format(stem)), "a topic", "both", ground=g)
        line = [l_ for l_ in buf.getvalue().splitlines() if "design_plan.palette" in l_]
        got = {h.upper() for h in re.findall(r"#([0-9A-Fa-f]{6})", line[0])} if line else set()
        want = {P["ground"].upper(), P["ink"].upper()} | {h.lstrip("#").upper() for h in (rec or [])}
        check(got == want, "{}: --gates declares the ground, the ink and the hues it paints ({} vs {})".format(stem, sorted(got), sorted(want)))

# ── B1: the screen readers' title, parked off the page, is not text in a platform's safe zone ──
# every visual-language page on 小红书 3:4 / 9:16 failed SAFE ZONE "text sits -1.00in from the top" — the a11y title
# (non-Claude agent run, 2026-10-09)
import check_surface as csf
for name in vl.LANGS:
    for fmtname, (W, H) in (("red", (7.5, 10.0)), ("story", (5.625, 10.0))):
        prs, k = use(name, W, H)
        k.cover(k.new_slide(), kicker="A repair café", title="Bring it broken, take it home working", subtitle="How it works",
                image=img_for(name) if name in vl.IMAGE_LED else None)
        k.closing(k.new_slide(), title="See you next week", line="Bring a neighbour.",
                  image=img_for(name) if name in vl.IMAGE_LED else None)
        path = td / "safe_{}_{}.pptx".format(name, fmtname)
        prs.save(str(path))
        with contextlib.redirect_stdout(io.StringIO()):
            res = csf.check(str(path), fmtname)
        found = res[0] if isinstance(res, tuple) else res
        off = [x for x in found if x[0] == "SAFE ZONE" and "-1.00in" in x[1]]
        check(not off, "{} {}: the off-page a11y title is not SAFE ZONE text ({})".format(name, fmtname, [x[1][:60] for x in off[:1]]))

# ── B2: a deck of a language's own pages is clean under the delivery lint (lint_deck) ──
# the kit's own furniture failed it: soft's blob behind the photo, ink's sun over a ridge, cutpaper's cards tucked into
# the hills, chalkboard's doodle (OVERLAP), ink's figure inside its ensō (INVISIBLE TEXT: the ring's bounding box read as
# the figure's backing) — an agent told "fix every finding" could not (both audits, 2026-10-09)
import lint_deck as _ld
for name, ground, W, H, path in OWN:
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
        hard = _ld.lint(str(path), static_ok=True)
    why = [l_.strip()[:110] for l_ in buf.getvalue().splitlines()
           if re.match(r"\s+slide \d+: [A-Z]", l_) and "[warn]" not in l_ and "[stats]" not in l_]
    check(hard == 0, "{} {} {}x{}: a deck of its own pages has no hard lint finding ({})".format(name, ground, W, H, why[:2]))

# ── B3: a near-uniform texture ground is a solid ground for contrast ──
# tally's grid, chalkboard's slate, collage's grain, the drafting grid and the night sky are pictures painted as the
# slide background; lint called their colour unknowable and estimated it from the render, where the glyphs' own
# antialiasing scored white-on-navy at 2.8:1 — a TEXT-ON-IMAGE warning on every page (audits, 2026-10-09). The picture
# is measured instead: 99%+ of its pixels sit on one colour, which is the backing (a photo stays unknowable).
from pptx import Presentation as _Prs
for name, ground in (("tally", "night"), ("chalkboard", "light"), ("collage", "slate"), ("drafting", "cyanotype"), ("starlit", "light")):
    prs, k = use(name, ground=ground)
    k.cover(k.new_slide(), title="A title on the ground", subtitle="And a line")
    path = td / "bg_{}.pptx".format(name)
    prs.save(str(path))
    p2 = _Prs(str(path))
    rec = _ld._slide_bg_box(p2.slides[0], 13.333, 7.5)
    want = vl.VARIANTS[name][ground]["palette"]["ground"].upper()
    got = (rec or {}).get("fill")
    check(rec is not None and not rec.get("unk") and bool(got) and _ld._contrast(got, want) < 1.15,
          "{} {}: its texture ground reads as the solid {} (got {})".format(name, ground, want, got))
prs = dk.blank_deck(13.333, 7.5)
s = dk.add_slide(prs)
import native_art as _na
_part, _rid = s.part.get_or_add_image_part(PHOTO)
_na._bg(s, '<a:blipFill><a:blip r:embed="{}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'.format(_rid))
path = td / "bg_photo.pptx"; prs.save(str(path))
rec = _ld._slide_bg_box(_Prs(str(path)).slides[0], 13.333, 7.5)
check(rec is not None and rec.get("unk") and not rec.get("fill"), "a photo background stays unknowable ({})".format(rec and rec.get("fill")))

# ── a printed board takes a language's LIGHTEST ground: poster's default is a saturated blue, its paper ground is light
# (it said poster "has none" — docs agent, 2026-10-09); starlit and chalkboard are dark on both and say so
for name in vl.LANGS:
    prs = dk.blank_deck(8.27, 11.69)
    g, why = vl._auto_ground(name, prs)
    lums = {k_: vl._lum(vl.VARIANTS[name][k_]["palette"]["ground"]) for k_ in vl.VARIANTS[name]}
    best = max(lums, key=lums.get)
    if lums[best] < 0.2:                  # dark on both (starlit, chalkboard): the designed default, and it says so
        check(g == "light" and "dark" in why, "{}: a printed board keeps its default and says it prints dark ({}, {})".format(name, g, why[:60]))
    else:
        check(g == best, "{}: a printed board takes its lightest ground ({} picked, {} is lightest)".format(name, g, best))

# ── C2: a refused page leaves its slide as it found it ──
# collage left three furniture shapes and storybook its picture on a slide whose cover was refused; an agent that caught
# the refusal and retried shorter copy on the same slide got the art twice (audit sweep, 2026-10-09)
LONG = "Why every street deserves a place to fix what it already owns, " * 30
for name in vl.LANGS:
    prs, k = use(name, 5.625, 10.0)
    s = k.new_slide()
    n0, rels0 = len(s.shapes), len(s.part.rels)
    for kw in (dict(title=LONG, subtitle=LONG, image=img_for(name)), dict(title="Fine", image=str(NOTIMG))):
        if name in vl.NATIVE:
            kw.pop("image")
            if "subtitle" not in kw:
                continue
        try:
            k.cover(s, **kw)
            check(False, "{}: an overlong cover is refused on 9:16".format(name))
        except (ValueError, FileNotFoundError):
            check(len(s.shapes) == n0 and len(s.part.rels) == rels0,
                  "{} cover: a refusal leaves the slide as it was ({} -> {} shapes, {} -> {} rels)".format(
                      name, n0, len(s.shapes), rels0, len(s.part.rels)))

# ── C1: the image-led languages have a points page too ──
# "three practical tips" on a photo deck had nowhere to go — the points page belonged to the drawn languages only and an
# agent hand-rolled one from rs.ground + rs.card with a tiny label and no real title (non-Claude agent run, 2026-10-09)
import ooxml_safety as ox
PTS = {"en": [("Share the tools", "One drawer of screwdrivers feeds three tables."), ("Open the door", "No booking, no fee, no questions."),
              ("Keep it walkable", "Ten minutes on foot is the limit we keep."), ("Write every repair down", "A notebook per bench becomes next year's manual.")],
       "zh": [("共享工具", "一抽屉螺丝刀可以供三张桌子同时使用。"), ("敞开大门", "不用预约，不收费，不问是什么东西。"),
              ("步行可达", "步行十分钟是我们坚持的距离。"), ("记下每次修理", "每张桌一本笔记，就是明年的培训手册。")]}
PT_TITLE = {"en": "How a repair evening works", "zh": "修理之夜是怎么运作的"}
for name in vl.IMAGE_LED:
    for W, H in ((13.333, 7.5), (5.625, 10.0)):       # the long-copy corpus (test_native_generality) covers the others
        for lang in ("en", "zh"):
            for n, im in ((4, None), (3, img_for(name))):
                    prs, k = use(name, W, H)
                    s = k.new_slide()
                    tag = "{} {}x{} {} n={}{}".format(name, W, H, lang, n, " +image" if im else "")
                    try:
                        k.points(s, kicker="Practical tips" if lang == "en" else "实用建议", title=PT_TITLE[lang],
                                 items=PTS[lang][:n], image=im)
                    except Exception as e:
                        check(False, "{}: an ordinary points page builds ({}: {})".format(tag, type(e).__name__, str(e)[:90]))
                        continue
                    on = re.sub(r"\s+", "", " ".join(sh.text_frame.text for sh in texts(s)))
                    lost = [w_ for hl in PTS[lang][:n] for w_ in hl if re.sub(r"\s+", "", w_) not in on]
                    check(not lost, "{}: every point's words are on the page ({})".format(tag, lost[:2]))
                    with contextlib.redirect_stdout(io.StringIO()):
                        crit = [f_ for f_ in dk.lint_layout(prs, verbose=False) if f_[1] == "CRITICAL"]
                    check(not crit, "{}: no critical layout fault ({})".format(tag, [(c[2], c[3][:60]) for c in crit[:2]]))
                    check(ox.beyond_page(prs) == [], "{}: nothing past the page".format(tag))
                    if im:
                        check(any(sh.shape_type == 13 for sh in s.shapes), "{}: the picture is placed".format(tag))

# ── C4: what a caller passes is either used as given or refused by name — never str()'d or half-dropped ──
# title=["Bring it", "broken"] shipped the literal "['Bring it', 'broken']"; number=0.1+0.2 shipped 0.30000000000000004;
# a point given as a 3-tuple or a dict with a third key silently lost it; a dict keyed title/text was refused as "an
# empty one" (robustness audit, 2026-10-09)
for name in vl.LANGS:
    prs, k = use(name)
    for page, kw, field in (("cover", dict(title=["Bring it", "broken"]), "title"), ("data", dict(number=0.1 + 0.2, label="l"), "number"),
                            ("quote", dict(quote="Q", attribution=("A", "B")), "attribution")):
        try:
            getattr(k, page)(k.new_slide(), **kw)
            check(False, "{} {}: a non-text {}= is refused".format(name, page, field))
        except (ValueError, TypeError) as e:
            check(field in str(e), "{} {}: a non-text {}= is refused by name ({})".format(name, page, field, str(e)[:70]))
    s = k.new_slide(); k.data(s, number=12, label="evenings a month")
    check(any(sh.text_frame.text.strip() == "12" for sh in texts(s)) or (name == "collage" and any(sh.shape_type == 13 for sh in s.shapes)),
          "{} data: a whole number is shown as written".format(name))
    for items, why in (([("Bring it", "Anything.", "extra"), ("Fix it", "Together.")], "a 3-tuple"),
                       ([{"title": "Bring it", "text": "Anything."}, {"head": "Fix it"}], "a dict with other keys")):
        kw = dict(title="How", items=items)
        if name in ("tally", "broadsheet"):
            kw["tags"] = ["A", "B"]
        try:
            k.points(k.new_slide(), **kw)
            check(False, "{} points: {} is refused".format(name, why))
        except ValueError as e:
            check("head" in str(e) and "line" in str(e), "{} points: {} is refused naming head and line ({})".format(name, why, str(e)[:80]))
# journal: a margin note on 4:3 with three Japanese points (refused "shorten the copy", which did not help: the margin took
# the room, not the words — the non-Claude agent run)
JA = [("同じデータ", "三つの手法に同じ撮像データを与える。"), ("同じ指標", "誤差は一つの指標で測り、変えない。"), ("同じ条件", "加速率は四段階で揃える。")]
for W, H in ((10.0, 7.5), (7.5, 7.5), (10.0, 5.625)):
    prs, k = use("journal", W, H)
    try:
        k.points(k.new_slide(), kicker="実験の設計", title="三つの手法を、同じ条件で比べる", items=JA, margin="指標の具体的な定義は、次回までに決める。")
        check(True, "journal {}x{}: three points and a margin note fit".format(W, H))
    except vl.VLTextOverflow as e:
        check(False, "journal {}x{}: three points and a margin note fit ({})".format(W, H, str(e)[:100]))

# ── C3: cutpaper's words stay above its front hills ──
# on 3:4 a long section title's last line sat on the green hill (audit sweep, 2026-10-09): the column ran to 0.86H, the
# crest rises to 0.78H
LONGT = {"en": "Why every street deserves a place to fix what it already owns, not throw away",
         "zh": "为什么每条街道都值得拥有一个修理自己物品的地方，而不是随手扔掉"}
for W, H in ((7.5, 10.0), (5.625, 10.0), (10.0, 5.625), (13.333, 7.5), (7.5, 7.5)):
    for lang in ("en", "zh"):
        for page in ("section", "closing"):
            prs, k = use("cutpaper", W, H)
            s = k.new_slide()
            kw = dict(number="02", kicker="Annual review", title=LONGT[lang]) if page == "section" else dict(title=LONGT[lang], line="Bring a neighbour.")
            getattr(k, page)(s, **kw)
            hills, shs = [], list(s.shapes)
            cards = [i for i, sh in enumerate(shs) if "paper card" in (sh.name or "")]
            for i, sh in enumerate(shs):      # only the hills IN FRONT: back hills sit behind the card the words are on
                if "hill" in (sh.name or "") and (not cards or i > max(cards)):
                    hills += _ld._cust_polys(sh, sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU) or []
            hit = []
            for sh in texts(s):
                x, y, w, h = sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU
                for px in (x + 0.1, x + w / 2, x + w - 0.1):
                    if hills and _ld._in_polys(hills, px, y + h - 0.02):
                        hit.append(sh.text_frame.text[:16])
            check(hills and not hit, "cutpaper {} {}x{} {}: the words stay above the hills ({})".format(page, W, H, lang, hit[:2]))

# ── final review: the lint changes hold on ORDINARY decks too ──
# reading a custom shape's painted outline must never silence a real backing: a filled path plus an outline-only path,
# two overlapping filled paths, a flipped shape (dark text on the dark card was INVISIBLE TEXT on main, nothing on the
# branch); a background picture is a solid colour only where the text actually sits on that colour, and never when
# PowerPoint washes it (alphaModFix / lum) — final review, 2026-10-09
from pptx import Presentation as _P2
from pptx.util import Inches as _In, Pt as _Pt2
from pptx.dml.color import RGBColor as _RGB
from pptx.oxml.ns import qn as _qn
from lxml import etree as _et
from pptx.enum.shapes import MSO_SHAPE as _MS
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _deck():
    prs = _P2(); prs.slide_width, prs.slide_height = _In(13.333), _In(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _tb(sl, x, y, w, h, txt, ink, size=24):
    tb = sl.shapes.add_textbox(_In(x), _In(y), _In(w), _In(h)); tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run(); r.text = txt; r.font.size = _Pt2(size); r.font.color.rgb = _RGB.from_string(ink)


def _cust(sl, x, y, w, h, paths, flipH=False):
    sh = sl.shapes.add_shape(_MS.RECTANGLE, _In(x), _In(y), _In(w), _In(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = _RGB.from_string("1F2A44"); sh.line.fill.background()
    spPr = sh._element.spPr; prst = spPr.find(_qn("a:prstGeom"))
    prst.addprevious(_et.fromstring('<a:custGeom xmlns:a="%s"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" '
                                    'r="r" b="b"/><a:pathLst>%s</a:pathLst></a:custGeom>' % (_A, paths)))
    spPr.remove(prst)
    if flipH:
        spPr.find(_qn("a:xfrm")).set("flipH", "1")


def _lint_text(prs, tag):
    path = td / ("rv_" + tag + ".pptx"); prs.save(str(path))
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
        _ld.lint(str(path), static_ok=True)
    return buf.getvalue()


_R = ('<a:moveTo><a:pt x="0" y="0"/></a:moveTo><a:lnTo><a:pt x="1000" y="0"/></a:lnTo><a:lnTo><a:pt x="1000" y="1000"/>'
      '</a:lnTo><a:lnTo><a:pt x="0" y="1000"/></a:lnTo><a:close/>')
_P1 = _R.replace('x="1000"', 'x="700"')
_P2b = _R.replace('x="0"', 'x="300"')
_TRI = ('<a:moveTo><a:pt x="0" y="0"/></a:moveTo><a:lnTo><a:pt x="0" y="1000"/></a:lnTo><a:lnTo><a:pt x="1000" y="1000"/>'
        '</a:lnTo><a:close/>')
for tag, paths, flip, tx in (("fill_plus_outline", '<a:path w="1000" h="1000">%s</a:path><a:path w="1000" h="1000" fill="none">%s</a:path>' % (_R, _R), False, (3, 3)),
                             ("two_overlapping", '<a:path w="1000" h="1000">%s</a:path><a:path w="1000" h="1000">%s</a:path>' % (_P1, _P2b), False, (3, 3)),
                             ("one_path", '<a:path w="1000" h="1000">%s</a:path>' % _R, False, (3, 3)),
                             ("flipped_wedge", '<a:path w="1000" h="1000">%s</a:path>' % _TRI, True, (8.2, 4.6))):
    prs, sl = _deck()
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = _RGB.from_string("FFFFFF")
    if tag == "flipped_wedge":
        _cust(sl, 1, 1, 11, 6, paths, flip)
    else:
        _cust(sl, 2, 2, 6, 3, paths, flip)
    _tb(sl, tx[0], tx[1], 3.2, 0.8, "Dark words on the dark card", "2B2B2B")
    check("INVISIBLE TEXT" in _lint_text(prs, tag), "lint: dark text on a dark custom shape ({}) is still INVISIBLE TEXT".format(tag))
from PIL import ImageDraw as _ID


def _bgpic(sl, im, extra=""):
    buf = io.BytesIO(); im.save(buf, "PNG")
    _ip, rid = sl.part.get_or_add_image_part(io.BytesIO(buf.getvalue()))
    c = sl._element.find(_qn("p:cSld"))
    c.insert(0, _et.fromstring('<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="%s" '
                               'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:bgPr><a:blipFill '
                               'dpi="0" rotWithShape="1"><a:blip r:embed="%s">%s</a:blip><a:srcRect/><a:stretch><a:fillRect/>'
                               '</a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>' % (_A, rid, extra)))


for frac in (0.035, 0.08):                          # a white template picture with a navy brand strip, a white label on it
    prs, sl = _deck()
    im = Image.new("RGB", (1333, 750), (255, 255, 255)); _ID.Draw(im).rectangle([0, 0, 1333, int(750 * frac)], fill=(16, 42, 92))
    _bgpic(sl, im)
    _tb(sl, 0.4, 0.02, 9, 7.5 * frac - 0.04, "Clinical update", "FFFFFF", 14)
    out = _lint_text(prs, "band_{}".format(frac))
    check("INVISIBLE TEXT: 'Clinical" not in out, "lint: a white label on a {:.1%} brand strip of a background picture is not INVISIBLE".format(frac))
prs, sl = _deck()                                   # a dark patch on light paper, a white caption on the patch
im = Image.new("RGB", (1333, 750), (240, 236, 228)); _ID.Draw(im).rectangle([60, 600, 580, 690], fill=(30, 30, 30))
_bgpic(sl, im); _tb(sl, 0.6, 6.0, 5.2, 0.9, "A caption on the dark patch", "FFFFFF", 20)
check("INVISIBLE TEXT: 'A caption" not in _lint_text(prs, "patch"), "lint: a white caption on a dark patch of a light background picture is not INVISIBLE")
for extra, tag in (('<a:alphaModFix amt="15000"/>', "alpha"), ('<a:lum bright="70000" contrast="-70000"/>', "washout")):
    prs, sl = _deck()
    _bgpic(sl, Image.new("RGB", (400, 225), (22, 30, 58)), extra)
    rec = _ld._slide_bg_box(sl, 13.333, 7.5)
    check(rec is not None and not rec.get("fill"), "lint: a background picture PowerPoint washes ({}) is not read as its pixels' colour ({})".format(
        tag, rec and rec.get("fill")))

# ── final review: the upright copy stays a JPEG and a broken cache entry is never reused ──
# every rewritten picture became a PNG (a 1.2MB phone JPEG grew to 7.5MB); a build killed mid-write left a truncated
# copy under a stable key, embedded by every later build as a black rectangle
up = vl._usable(EXIF6)
check(Image.open(up).format == "JPEG" and Path(up).stat().st_size <= 2.5 * EXIF6.stat().st_size,
      "an upright copy of a phone JPEG is a JPEG of similar size ({} {} vs {} bytes)".format(
          Image.open(up).format, Path(up).stat().st_size, EXIF6.stat().st_size))
Path(up).write_bytes(Path(up).read_bytes()[:2000])                      # a write that was cut short
up2 = vl._usable(EXIF6)
try:
    Image.open(up2).load(); _ok = True
except Exception:
    _ok = False
check(_ok, "a truncated upright copy in the cache is made again, not reused")
fe = vl._feathered(PHOTO)
Path(fe).write_bytes(Path(fe).read_bytes()[:2000])
try:
    Image.open(vl._feathered(PHOTO)).load(); _ok = True
except Exception:
    _ok = False
check(_ok, "a truncated feathered copy in the cache is made again, not reused")

# ── final review: undoing a refused page never breaks what was already on the slide ──
# a refused image_text whose picture the caller had already placed on the slide dropped the SHARED relationship (the
# reopened file failed KeyError rId3); drafting's project= deleted the ground's old title block during an attempt that
# was then undone, leaving two stacked blocks or a sheet showing the new project the kit never kept
HUGE = " ".join(["Why every street deserves a place to fix what it already owns"] * 6)
for name in ("editorial", "ink", "poster"):
    prs, k = use(name)
    s = k.new_slide()
    mine = dk.picture(s, PHOTO, 9.5, 5.2, 3.0, 2.0, fit="cover", alt="my own thumbnail")
    rid = mine._element.find(".//" + _qn("a:blip")).get(_qn("r:embed"))
    try:
        k.image_text(s, image=PHOTO, title=HUGE, body=HUGE, caption=HUGE)
    except vl.VLTextOverflow:
        pass
    path = td / "shared_{}.pptx".format(name); prs.save(str(path))
    try:
        for sh in _P2(str(path)).slides[0].shapes:
            if sh.shape_type == 13:
                sh.image.blob
        reopened = True
    except Exception:
        reopened = False
    check(rid in s.part.rels and reopened, "{}: a refused page keeps the caller's picture that shares its file".format(name))


def _blocks(sl):
    return [[p_.text for p_ in sh.text_frame.paragraphs][:2] for sh in sl.shapes
            if getattr(sh, "has_text_frame", False) and sh.text_frame.paragraphs and sh.text_frame.paragraphs[0].text == "PROJECT"]


LONGDR = dict(kicker="Annual review of the repair network", title="Why every street deserves a place to fix what it already owns, not throw away",
              body="What twelve months of evenings, benches and borrowed tools taught us", caption="Weighed at the door, item by item.")
for W, H in ((7.5, 7.5), (10.0, 7.5)):
    prs, k = use("drafting", W, H)
    k.cover(k.new_slide(), title="Repair network", project="Alpha phase")
    s = k.new_slide()
    k.image_text(s, image=PHOTO, project="Beta phase", **LONGDR)
    b = _blocks(s)
    check(len(b) == 1 and b[0][1] == "Beta phase", "drafting {}x{}: a new project on a fallback layout leaves one title block ({})".format(W, H, b))
    s2 = k.new_slide(); before = _blocks(s2); kept = k.project
    try:
        k.image_text(s2, image=PHOTO, project="Gamma phase", kicker=LONGDR["kicker"], title=" ".join([LONGDR["title"]] * 4),
                     body=LONGDR["body"] * 3, caption=LONGDR["caption"])
    except vl.VLTextOverflow:
        pass
    check(_blocks(s2) == before and k.project == kept, "drafting {}x{}: a refused page with a new project leaves the sheet and the kit as they were ({} -> {}, {})".format(
        W, H, before, _blocks(s2), k.project))

# ── final review: headroom never brings a widow back ──
# words measured and widow-checked at 97% of the column were drawn in a box at 100%, and LibreOffice re-wrapped
# "…what you already / own" (soft), "…more than ever / before" (collage, ink): a wrapping display field is now SET in
# its measured lines, so the box width cannot re-wrap it
WT = ["Why small repair cafés matter more than ever before", "The quiet economics of fixing what you already own",
      "How to run a repair night with five volunteers", "A year of evenings spent fixing other people's toasters",
      "Mending as a habit, not a hobby, for everyone", "Our first hundred repairs and what went wrong"]
for name in ("editorial", "soft", "collage", "storybook", "ink", "poster", "cutpaper", "drafting"):
    for W, H in ((13.333, 7.5), (7.5, 7.5)):
        prs, k = use(name, W, H)
        for t in WT:
            s = k.new_slide()
            try:
                if name in vl.IMAGE_LED:
                    k.cover(s, title=t, image=img_for(name))
                else:
                    k.section(s, number="02", title=t)
            except vl.VLTextOverflow:
                continue
            sh = [x for x in texts(s) if " ".join(p_.text for p_ in x.text_frame.paragraphs).strip().lower() == t.lower()]
            if not sh or len(sh[0].text_frame.paragraphs) > 1:
                continue                    # set in its lines: nothing to re-wrap
            sz = sh[0].text_frame.paragraphs[0].runs[0].font.size.pt
            full = vl._break_lines(k, "title", sh[0].text_frame.text, sz, sh[0].width / EMU)   # the words as drawn (poster: capitals)
            stack = name == "poster" and len(full) >= 2 and len(full[-2].split()) == 1   # poster's word stack, by design
            check(len(full) < 2 or len(full[-1].split()) > 1 or stack,
                  "{} {}x{}: {!r} has no lone last word at the width it is drawn ({})".format(name, W, H, t[:30], full[-2:]))

# ── boardgame's quote: the rule between the quote and its attribution stays between them on any face metrics ──
# GitHub's Linux runner (stand-in faces) measured the quote's foot lower than macOS; the attribution was only ever
# pulled UP to it, never pushed down, so the rule set halfway between them crossed the attribution (felt/4:3/zh,
# RULE_THROUGH_TEXT, 2026-10-22). Simulated here by moving the quote's foot 0.5in lower.
import vl_langs as _vll
_bg = _vll.load()["boardgame"]
_foot0 = _bg._lint_foot
_bg._lint_foot = lambda sh: _foot0(sh) + 0.5
try:
    for _g in ("light", "felt"):
        _prs, _k = use("boardgame", 10.0, 7.5, _g)
        _k.quote(_k.new_slide(), quote="来访者握着螺丝刀，志愿者只在旁边指导，这就是整个夜晚的意义所在。", attribution="志愿者手册")
        with contextlib.redirect_stdout(io.StringIO()):
            _crit = [f_ for f_ in dk.lint_layout(_prs, verbose=False) if f_[1] == "CRITICAL"]
        check(not _crit, "boardgame/{} quote: the rule clears the attribution when the quote measures taller ({})".format(
            _g, [(c_[2], c_[3][:60]) for c_ in _crit[:2]]))
finally:
    _bg._lint_foot = _foot0

# ── C1 (2026-10-22): a line that mixes CJK and Latin is measured as LibreOffice sets it ──
# calendar on 10x5.625 measured "用 GPT-5 做 3 件事" (31.5pt bold, Times New Roman + Songti SC) at 3.67in in a 3.79in
# measure and planned one line; LibreOffice set 4.04in and put 事 alone on a second line. Measured (a PDF read back with
# PyMuPDF): a space or a mark of no script is set in the face of the script BEFORE it, and every change between CJK and
# Latin adds 0.24 em of autospace — four changes here, 0.41in.
import display_type as _dt
check(_dt.script_runs("用 GPT-5 做 3 件事") == [("用 ", "a"), ("GPT-5 ", "l"), ("做 ", "a"), ("3 ", "l"), ("件事", "a")],
      "a space takes the script before it; digits and '-' are Latin ({})".format(_dt.script_runs("用 GPT-5 做 3 件事")))
check(_dt.script_runs(" 用“中”") == [(" 用“中”", "a")] and _dt.script_runs("“A”中") == [("“A”", "l"), ("中", "a")],
      "a leading weak mark takes the first script; quotation marks take their neighbour's")
for _name in ("calendar", "journal", "chalkboard"):
    _p, _k = use(_name, 10.0, 5.625)
    for _f in ("title", "body"):
        if _f not in vl.TYPE[_name]:
            continue
        for _a, _b in (("GPT-5", "做三件事"), ("Claude", "与模型"), ("2026", "年计划")):
            for _sz in (14.0, 31.5):
                _w_ab = vl._line_width(_k, _f, _a + _b, _sz)
                _w_a, _w_b = vl._line_width(_k, _f, _a, _sz), vl._line_width(_k, _f, _b, _sz)
                check(_w_ab >= _w_a + _w_b + 0.235 * _sz / 72.0,
                      "{} {} {}pt: {!r} measures its parts plus the autospace between them ({:.3f} < {:.3f} + {:.3f} + "
                      "{:.3f})".format(_name, _f, _sz, _a + _b, _w_ab, _w_a, _w_b, 0.24 * _sz / 72.0))
# a run that carries CJK is set UPRIGHT (Kit.run), so a Latin part of it is measured upright too: measured in the quote's
# italic, "v2 2026" read 6% narrow in Didot Italic and "v2 2026 道具" was planned as two lines that render as three
_n_up = 0
for _name in vl.LANGS:
    for _f, _spec in vl.TYPE[_name].items():
        if not _spec[4]:
            continue
        for _fonts in ("both", "mac"):
            _p = dk.blank_deck(13.333, 7.5)
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                try:
                    _k = vl.use(_name, _p, fonts=_fonts)
                except ValueError:
                    continue                                   # a mac face not installed here
            _sz, _t = _spec[0] * 0.8, "v2 2026 道具"
            _role, _b = _spec[1], vl._weight(_k, _f, _t)
            for _w in (1.0, 1.2, 1.4, 1.8):
                for _l in vl._break_lines(_k, _f, _t, _sz, _w):
                    if " " not in _l.strip():
                        continue                               # one unit: nothing to break
                    _n_up += 1
                    _up = max(_dt.line_width(_l, _sz, _k.face(_role), _k.ea_face(_role, _t), _b),
                              _dt.em_width(_l, _sz, _k.face(_role), _b))
                    check(_up <= _w - vl._INSET + 1e-6, "{} {} {}: {!r} is planned on one line of {:.2f}in and measures "
                                                        "{:.3f}in upright, as it renders".format(
                                                            _name, _fonts, _f, _l, _w - vl._INSET, _up))
check(_n_up > 0, "the upright check ran ({} lines)".format(_n_up))
_p, _k = use("calendar", 10.0, 5.625)
_t = "用 GPT-5 做 3 件事"
if not (dk._font_substituted("Times New Roman") or dk._font_substituted("Songti SC")):
    _w = vl._line_width(_k, "title", _t, 31.5)
    check(4.035 <= _w <= 4.04 * 1.04, "calendar's title is measured as LibreOffice set it (4.04in, rendered "
                                      "2026-10-22): {:.3f}in".format(_w))
_lines = vl._break_lines(_k, "title", _t, 31.5, 3.96 * vl.HEADROOM)
check(len(_lines) == 2, "calendar's mixed title is planned on the two lines LibreOffice set it in ({})".format(_lines))
check(vl._field_height(_k, "title", _t, 31.5, 3.96 * vl.HEADROOM) >= 2 * 31.5 / 72.0 * dk._LINT_LINE_H,
      "…and its box is two lines tall")

# ── C2 (2026-10-22): one word or figure gets the same headroom as every line of words ──
# _widest_word_fits let a single word fill the whole column while _field_height measured in HEADROOM of it: periodic
# planned "1,250" at 108.1pt measuring 3.75in in a 3.81in column (a 3.754in inner width) and LibreOffice set "1,25 / 0";
# where the word fell between the two widths, the box was planned a line taller than the figure (calendar: "1,250" at the
# top of a box twice its height). LibreOffice keeps a word whole only once the box holds its full advance width
# (measured 2026-10-22: at -0.1% "1,25 / 0", at +0.1% whole), so the word is held to the measure the lines are.
_n_flow = 0
for _name in vl.LANGS:
    if "number" not in vl.TYPE[_name]:
        continue
    for _txt, _start, _cw in (("1,250", 147.0, 3.81), ("1,250,000", 120.0, 3.81), ("42", 200.0, 0.9)):
        _p, _k = use(_name)
        _s = _k.new_slide()
        try:
            _r, _d = vl._flow(_k, _s, "data", (0.5, 0.5, _cw, 6.0), [("number", _txt)], anchor="top", align="l",
                              start={"number": _start})
        except vl.VLTextOverflow:
            continue
        if vl._outlinable(_k, _txt):
            continue                                  # drawn as a picture of its own size
        _n_flow += 1
        _sz = _d.sizes["number"]
        _gw = _dt._glyph_width(_txt, _sz, _k.face("numeral"), vl._weight(_k, "number", _txt)) or 0.0
        _floor = max(9.0, vl.TYPE[_name]["number"][5] * 1.0)
        # above its floor a figure keeps HEADROOM of the column; AT its floor one word may take the whole column, less
        # WORD_HEADROOM, in a box exactly the column wide — never more
        _meas = _cw * (vl.HEADROOM if _sz > _floor + 0.1 else vl.WORD_HEADROOM) - vl._INSET
        check(_gw <= _meas + 1e-6 and _r["number"][2] <= _cw + 1e-6,
              "{}: {!r} planned at {}pt measures {:.3f}in — within the {:.3f}in measure of its {}in column, in a box no "
              "wider ({:.3f}in)".format(_name, _txt, _sz, _gw, _meas, _cw, _r["number"][2]))
        _one = _sz / 72.0 * dk._LINT_LINE_H + 0.06                 # _field_height's one Latin line
        check(_r["number"][3] <= _one + 1e-6, "{}: {!r}'s box is one line tall ({:.3f} for {:.3f}in)".format(
            _name, _txt, _r["number"][3], _one))
check(_n_flow >= 20, "the one-figure flow ran on the languages ({})".format(_n_flow))
# …and at its floor a figure that misses HEADROOM but fits the column (with WORD_HEADROOM to spare) is SET, on one line,
# not refused: "1,250,000" at 56pt in herbarium's 2.60in cabinet column (the generality corpus)
_n_floor = 0
for _name in ("herbarium", "periodic", "editorial", "calendar"):
    _p, _k = use(_name)
    _fl = max(9.0, vl.TYPE[_name]["number"][5])
    if vl._outlinable(_k, "1,250,000"):
        continue
    _gw = _dt._glyph_width("1,250,000", _fl, _k.face("numeral"), vl._weight(_k, "number", "1,250,000"))
    _cw = (_gw + vl._INSET) / 0.98                     # the figure fills 98% of the column at its floor
    try:
        _r, _d = vl._flow(_k, _k.new_slide(), "data", (0.5, 0.5, _cw, 6.0), [("number", "1,250,000")], anchor="top",
                          align="l", start={"number": _fl})
    except vl.VLTextOverflow as e:
        check(False, "{}: a figure that fits its column at the floor is set, not refused ({})".format(_name, str(e)[:90]))
        continue
    _n_floor += 1
    check(abs(_r["number"][2] - _cw) < 1e-6 and _r["number"][3] <= _fl / 72.0 * dk._LINT_LINE_H + 0.06 + 1e-6,
          "{}: at its floor '1,250,000' takes the whole {:.2f}in column on one line ({})".format(_name, _cw, _r["number"]))
check(_n_floor >= 3, "the floor case ran ({})".format(_n_floor))

# ── C3 (2026-10-22): a picture credit's chip holds its words by the render AND by the lint ──
# the chip was 1.3 em a CJK line; LibreOffice sets 1.344 em (1.2 x CJK_LS 1.12, every face measured) and the delivery
# lint reads 1.4 em — a two-line Chinese credit at 13pt+ (chalkboard image_text, 10x7.5) and a one-line one tall enough to
# be a card (A4, 15.4pt) were each "cramped against the card bottom" of their own chip.
_CJK_CREDIT = "摄影：张三，CC BY 4.0 授权，来源：维基共享资源"
_LAT_CREDIT = "Photo: Jane Doe, CC BY 4.0, via Wikimedia Commons"
for _name, _W, _H in (("chalkboard", 10.0, 7.5), ("chalkboard", 8.27, 11.69), ("chalkboard", 13.333, 7.5),
                      ("editorial", 13.333, 7.5), ("editorial", 7.5, 10.0), ("journal", 10.0, 7.5)):
    for _cr in (_CJK_CREDIT, _LAT_CREDIT):
        _p, _k = use(_name, _W, _H)
        _s = _k.new_slide()
        try:
            _k.image_text(_s, title="每张桌上都有工具", body="志愿者坐在你身边。", image=img_for(_name), credit=_cr)
        except ValueError as e:
            check("shorten" in str(e), "{} {}x{}: a credit is placed or refused asking to shorten it ({})".format(
                _name, _W, _H, str(e)[:60]))
            continue
        _path = td / "credit_{}_{}_{}.pptx".format(_name, _W, "zh" if dk._has_cjk(_cr) else "en")
        _p.save(str(_path))
        _buf = io.StringIO()
        with contextlib.redirect_stdout(_buf), contextlib.redirect_stderr(io.StringIO()):
            _ld.lint(str(_path), static_ok=True)
        _pad = [l_.strip()[:120] for l_ in _buf.getvalue().splitlines() if "TEXT PADDING" in l_ and _cr[:8] in l_]
        check(not _pad, "{} {}x{}: the credit is not cramped in its own chip ({})".format(_name, _W, _H, _pad[:1]))
        _tb = [x for x in texts(_s) if x.text_frame.text == _cr][0]
        _z = _tb.text_frame.paragraphs[0].runs[0].font.size.pt
        _pitch = _z / 72.0 * dk._LINT_LINE_H * (dk.CJK_LS if dk._has_cjk(_cr) else 1.0)
        _n = round((_tb.height / EMU) / _pitch) if _tb.height / EMU >= 1.5 * _pitch else 1
        check(_tb.height / EMU >= _n * _pitch - 1e-3, "{} {}x{}: the credit's box holds the lines LibreOffice sets "
                                                        "({:.3f}in)".format(_name, _W, _H, _tb.height / EMU))

# ── D1 (2026-10-23): figures are measured in the face they are DRAWN in ──
# Kit.runs sets a figure run of an old-style face (Didot, Hoefler Text, Georgia …) in a lining one (Times New Roman), but
# the flow measured those figures in the role's own face: Didot's "214" (73.7pt) read 1.29in for the 1.54in Times New
# Roman draws, in a 1.457in box — LibreOffice set "21 / 4" (stainedglass, fonts="mac"); airmail's Didot numeral "1:250"
# (130pt, 13.333x7.5) rendered "1:25 / 0" (rendered 2026-10-23). Every measure now takes the face of each run Kit.runs
# draws — never under the render — floored by the run's own face's reading, which is what both lints read (one face a
# paragraph): Georgia's figures, wider than Times New Roman's, keep that width where the lints read Georgia.
def _drawn_w(k, field, t, sz):
    """The width (in) of `t` as Kit.runs draws it in `field`: each run in its own face (Latin text)."""
    _, role, _b, _c, it, _f = vl.TYPE[k.name][field]
    b = vl._weight(k, field, t)
    return sum(_dt._glyph_width(r[0], sz, r[5], r[3], r[4]) or 0.0 for r in k.runs(t, sz, None, b, role, it))
def _need_w(k, field, t, sz):
    """…or, when wider, `t` read whole in its run's own face (Kit.run's face: the lining one for a run of figures)."""
    _, role, _b, _c, it, _f = vl.TYPE[k.name][field]
    r = k.run(t, sz, None, vl._weight(k, field, t), role, it)
    return max(_drawn_w(k, field, t, sz), _dt._glyph_width(t, sz, r[5], r[3], r[4]) or 0.0)
_n_d1 = 0
for _name in vl.LANGS:
    for _fonts in ("both", "mac"):
        _p = dk.blank_deck(13.333, 7.5)
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            try:
                _k = vl.use(_name, _p, fonts=_fonts)
            except ValueError:
                continue                                       # a mac face not installed here
        for _f, _spec in vl.TYPE[_name].items():
            if not isinstance(_spec, tuple) or not dk.has_oldstyle_figures(_k.face(_spec[1])):
                continue
            for _t in ("214", "1:250", "$4.2M", "The 2026 plan in 12 steps", "Q3 2025: 1,250 visitors"):
                _sz = 40.0
                _n_d1 += 1
                _want, _need, _got = _drawn_w(_k, _f, _t, _sz), _need_w(_k, _f, _t, _sz), vl._line_width(_k, _f, _t, _sz)
                check(abs(_got - _need) < 1e-6 and _got >= _want - 1e-9,
                      "{} {} {}: {!r} measures {:.3f}in — never under the {:.3f}in it is drawn in ({}), and the run's "
                      "own face's {:.3f}in".format(_name, _fonts, _f, _t, _got, _want,
                                                   "+".join(sorted({r[5] for r in _k.runs(_t, _sz, None, False, _spec[1])})),
                                                   _need))
                for _word in vl._words(_t):        # the widest-word test and the line breaks use the same widths
                    _dw, _nw = _drawn_w(_k, _f, _word, _sz), _need_w(_k, _f, _word, _sz)
                    check(vl._widest_word_fits(_k, _f, _word, _sz, _nw + vl._INSET + 1e-4)
                          and not vl._widest_word_fits(_k, _f, _word, _sz, _dw + vl._INSET - 1e-3),
                          "{} {} {}: {!r} fits its measure, never one narrower than it is drawn ({:.3f}in)".format(
                              _name, _fonts, _f, _word, _dw))
                _one = vl._break_lines(_k, _f, _t, _sz, _need + vl._INSET + 1e-4)
                check(len(_one) == 1, "{} {} {}: {!r} is one line in a measure as wide as it is read ({})".format(
                    _name, _fonts, _f, _t, _one))
                if len(vl._words(_t)) > 1:          # (one word is never broken by the plan: it is held to the measure)
                    _less = vl._break_lines(_k, _f, _t, _sz, _want + vl._INSET - 1e-3)
                    check(len(_less) >= 2, "{} {} {}: …and two in one a hair narrower ({})".format(
                        _name, _fonts, _f, _less))
check(_n_d1 >= 40, "the drawn-face check ran ({} cases)".format(_n_d1))
# the reported page: a figure in a Didot numeral is held to its box as Times New Roman draws it (fonts="mac")
_n_d1p = 0
for _name, _num in (("airmail", "1:250"), ("airmail", "214"), ("calendar", "1:250"), ("calendar", "214"),
                    ("calendar", "1,250,000")):
    for _W, _H in ((13.333, 7.5), (7.5, 7.5), (5.625, 10.0)):
        _p = dk.blank_deck(_W, _H)
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            try:
                _k = vl.use(_name, _p, fonts="mac")
            except ValueError:
                break
        if not dk.has_oldstyle_figures(_k.face("numeral")):
            continue
        for _cw, _start in ((1.6, 200.0), (2.4, 220.0), (3.81, 220.0)):
            _s = _k.new_slide()
            try:
                _r, _d = vl._flow(_k, _s, "data", (0.5, 0.5, _cw, 6.0), [("number", _num)], anchor="top", align="l",
                                  start={"number": _start})
            except vl.VLTextOverflow:
                continue
            _d()
            _tb = [x for x in _s.shapes if getattr(x, "has_text_frame", False) and x.text_frame.text == _num]
            if not _tb:
                continue
            _n_d1p += 1
            _tb = _tb[0]
            _drawn = sum(_dt._glyph_width(r.text, r.font.size.pt, r.font.name, bool(r.font.bold)) or 0.0
                         for r in _tb.text_frame.paragraphs[0].runs)
            check(_drawn <= _tb.width / EMU - vl._INSET + 1e-6,
                  "{} mac {}x{} {}in column: {!r} as drawn ({:.3f}in in {}) fits its box's {:.3f}in".format(
                      _name, _W, _H, _cw, _num, _drawn, _tb.text_frame.paragraphs[0].runs[0].font.name,
                      _tb.width / EMU - vl._INSET))
check(_n_d1p >= 10 or not any(dk.has_oldstyle_figures(f) for f in ("Didot",)),
      "the Didot figure flow ran ({})".format(_n_d1p))

# ── D2 (2026-10-23): CJK line ends as LibreOffice sets them, measured on renders ──
# Rendered 2,256 boxes (calendar body/title, journal body; 20/36pt; Songti SC, Hiragino Sans GB, Hiragino Sans with
# Arial / Times New Roman / Georgia), each at the width that puts one mark — or the space before it — at the line end,
# read back with PyMuPDF:
#  * a SPACE that itself overflows is a break, whatever follows: "螺丝刀握 / ，在来访" (the mark starts the line);
#    when the space fits, a mark after an ideograph and a space hangs ("螺丝刀握 ，"), a closing bracket there starts the
#    next line ("螺丝刀握 / 」在来访"), and after a Latin word and a space the word goes down with the mark ("螺丝刀握 /
#    GPT，在来访") — the model hung the mark, or carried one letter ("螺丝刀握 GP / T，");
#  * an opening quote or bracket never ends a line — “ ‘ ( as well as （「 — even with a space after it, unless that
#    space is what overflows;
#  * a closing mark (or ” ’) that cannot start a line takes the WHOLE word before it down: "螺丝刀握 / WikiGPT-5，在来访"
#    (the carry took one character: "螺丝刀握WikiGPT- / 5，"); with nothing on the line to move back to, the line
#    breaks at the measure: "「Claude / 」在来访";
#  * one mark hangs after an ideograph even when the line or the text has Latin elsewhere ("AB螺丝刀，" on one line,
#    "在来访者。 / AB螺丝"), and a half-width , . ! ; : ? after an ideograph hangs as the full-width ones do; nothing
#    hangs right after a Latin letter or digit.
_p, _k = use("calendar")
_role = vl.TYPE["calendar"]["body"][1]
for _t, _prefix, _slack, _want in (
        ("螺丝刀握 ，在来访者", "螺丝刀握", 0.1, ["螺丝刀握", "，在来访", "者"]),
        ("螺丝刀握 GPT ，在来访者", "螺丝刀握 GPT", 0.1, ["螺丝刀握 GPT", "，在来访者"]),
        ("螺丝刀握在来访“者手里”修好", "螺丝刀握在来访“", 0.1, ["螺丝刀握在来访", "“者手里”修好"]),
        ("螺丝刀握(在来访)者手", "螺丝刀握(", 0.1, ["螺丝刀握", "(在来访)", "者手"]),
        ("螺丝刀握WikiGPT-5，在来访者", "螺丝刀握WikiGPT-5", 0.1, ["螺丝刀握", "WikiGPT-5，在来访", "者"]),
        ("螺丝刀握Claude。在来访者", "螺丝刀握Claude", 0.1, ["螺丝刀握", "Claude。在来访", "者"]),
        ("ドライバーは来Claude、訪者の手", "ドライバーは来Claude", 0.1, ["ドライバーは来", "Claude、訪者の手"]),
        ("螺丝刀握在”来访者", "螺丝刀握", 1.3, ["螺丝刀握", "在”来访者"]),
        ("螺丝刀握「Claude」在来访", "螺丝刀握「", 0.1, ["螺丝刀握", "「Claude", "」在来访"]),
        ("螺丝刀握在（ 握在来", "螺丝刀握在（", 0.05, ["螺丝刀握在（", "握在来"]),
        ("螺丝刀握 ，在来访者", "螺丝刀握 ", 0.15, ["螺丝刀握 ，", "在来访者"]),
        ("螺丝刀握 GPT ，在来访者", "螺丝刀握 GPT ", 0.15, ["螺丝刀握", "GPT ，在来访", "者"]),
        ("螺丝刀握 」在来访者", "螺丝刀握 ", 0.15, ["螺丝刀握", "」在来访", "者"]),
        ("螺丝刀握（ 在来访者", "螺丝刀握（ ", 0.15, ["螺丝刀握", "（ 在来访者"]),
        ("螺。", None, 1.02, ["螺。"]),
        ("螺」。", None, 1.02, ["螺", "」。"]),
        ("AB螺丝刀，握在来访", "AB螺丝刀", 0.1, ["AB螺丝刀，", "握在来访"]),
        ("螺丝刀握在,握在来访", "螺丝刀握在", 0.1, ["螺丝刀握在,", "握在来访"]),
        ("螺丝刀握在来访者手里修好带走。AB螺丝", None, 7.02, ["螺丝刀握在来访", "者手里修好带走。", "AB螺丝"]),
        ("螺丝刀握在A。", None, 3.02, ["螺丝刀", "握在", "A。"])):
    for _sz in (20.0, 36.0):
        _em = _sz / 72.0
        _b = vl._weight(_k, "body", _t)
        _w = (_slack * _em if _prefix is None else
              _dt.line_width(_prefix, _sz, _k.face(_role), _k.ea_face(_role, _t) or _k.face(_role), _b) + _slack * _em)
        _got = [l_.replace(" ", "") for l_ in vl._break_lines(_k, "body", _t, _sz, _w + vl._INSET)]
        check(_got == [l_.replace(" ", "") for l_ in _want], "{}pt {!r}: lines as LibreOffice set them ({} — rendered "
                                                             "{})".format(_sz, _t, " / ".join(_got), " / ".join(_want)))

for line in ok:
    print("  ok   " + line)
for line in bad:
    print("  FAIL " + line)
print("\n{} passed, {} failed".format(len(ok), len(bad)))
sys.exit(1 if bad else 0)
